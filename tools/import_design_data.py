"""Extract the design team's example text and embedded pictures for the web UI.

Usage: python tools/import_design_data.py [source_directory]
The source workbooks and presentation are read only. Derived assets are
written to data/design-data and served through authenticated API routes.
"""

from __future__ import annotations

import bisect
import difflib
import hashlib
import io
import json
import os
import posixpath
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from PIL import Image, ImageOps, UnidentifiedImageError


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Desktop" / "设计数据"
OUTPUT = ROOT / "data" / "design-data"
IMAGES = OUTPUT / "images"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
DRAW = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def relationship_targets(archive: zipfile.ZipFile, path: str, base: str) -> dict[str, str]:
    if path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(path))
    return {
        item.attrib["Id"]: posixpath.normpath(posixpath.join(base, item.attrib["Target"]))
        for item in root.findall(f"{REL}Relationship")
        if item.attrib.get("Target")
    }


def image_url(archive: zipfile.ZipFile, member: str) -> str | None:
    if member not in archive.namelist():
        return None
    raw = archive.read(member)
    digest = hashlib.sha256(raw).hexdigest()[:20]
    target = IMAGES / f"{digest}.webp"
    if not target.exists():
        try:
            with Image.open(io.BytesIO(raw)) as opened:
                image = ImageOps.exif_transpose(opened)
                image.thumbnail((1400, 1400), Image.Resampling.LANCZOS)
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
                image.save(target, "WEBP", quality=82, method=5)
        except (UnidentifiedImageError, OSError):
            return None
    return f"/api/design/image/{target.name}"


def clean_text(value: object) -> str:
    return re.sub(r"\n{3,}", "\n\n", str(value or "").strip())


def title_for(text: str, fallback: str) -> str:
    first = re.sub(r"[#*`]+", "", text).strip().splitlines()[0].strip() if text else ""
    return (first[:72] + "…") if len(first) > 72 else (first or fallback)


def product_for(text: str) -> str:
    for label, words in [
        ("背包/箱包", ("背包", "双肩包", "手提包", "挎包", "包袋")),
        ("行李箱", ("行李箱", "拉杆箱", "旅行箱")),
        ("户外炊具", ("卡式炉", "烧烤炉", "锅具", "月光炉", "收纳箱")),
        ("服饰穿搭", ("服装", "上衣", "外套", "针织", "羽绒服", "裙装")),
        ("鞋履", ("运动鞋", "老爹鞋", "靴子")),
    ]:
        if any(word in text[:700] for word in words):
            return label
    return "其他"


def scene_for(text: str) -> str:
    for label, words in [
        ("户外/露营", ("露营", "山地", "草原", "徒步", "户外", "湖边")),
        ("城市/通勤", ("都市", "城市", "街头", "地铁", "通勤", "建筑")),
        ("旅行/度假", ("旅行", "度假", "机场", "海边", "车站")),
        ("室内/影棚", ("室内", "影棚", "厨房", "家居", "纯色背景")),
    ]:
        if any(word in text[:700] for word in words):
            return label
    return "其他场景"


def workbook_text_rows(path: Path, sheet: str, columns: tuple[int, ...]) -> list[tuple[int, int, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook[sheet]
    result: list[tuple[int, int, str]] = []
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
        for column in columns:
            value = row[column - 1] if column <= len(row) else None
            if isinstance(value, str) and value.strip() and not value.startswith("="):
                result.append((row_number, column, clean_text(value)))
    workbook.close()
    return result


def drawing_images(archive: zipfile.ZipFile) -> list[tuple[int, int, str]]:
    drawing = "xl/drawings/drawing1.xml"
    if drawing not in archive.namelist():
        return []
    rels = relationship_targets(archive, "xl/drawings/_rels/drawing1.xml.rels", "xl/drawings")
    root = ET.fromstring(archive.read(drawing))
    result = []
    for anchor in root:
        origin = anchor.find(f"{DRAW}from")
        if origin is None:
            continue
        row = origin.findtext(f"{DRAW}row")
        col = origin.findtext(f"{DRAW}col")
        blip = anchor.find(f".//{A}blip")
        member = rels.get(blip.attrib.get(f"{R}embed", "")) if blip is not None else None
        if row is None or col is None or not member:
            continue
        url = image_url(archive, member)
        if url:
            result.append((int(row) + 1, int(col) + 1, url))
    return result


def model_cases(path: Path) -> tuple[list[dict], list[dict]]:
    rows = workbook_text_rows(path, "Sheet1", (1, 22))
    cases = [
        {"id": f"model-{index:03d}", "title": title_for(text, f"模特案例 {index}"),
         "text": text, "images": [], "product": product_for(text), "scene": scene_for(text),
         "source": f"{path.name} · Sheet1 · {openpyxl.utils.get_column_letter(col)}{row}"}
        for index, (row, col, text) in enumerate(rows, 1)
    ]
    by_column: dict[int, tuple[list[int], list[int]]] = {}
    for column in (1, 22):
        pairs = [(row, index) for index, (row, col, _) in enumerate(rows) if col == column]
        by_column[column] = ([row for row, _ in pairs], [index for _, index in pairs])
    with zipfile.ZipFile(path) as archive:
        for image_row, column, url in drawing_images(archive):
            text_column = 22 if column >= 22 else 1
            positions, indices = by_column[text_column]
            offset = bisect.bisect_right(positions, image_row) - 1
            if offset >= 0:
                cases[indices[offset]]["images"].append(url)
    # Two source rows repeat an existing case description, while WPS stores
    # their picture block on the other copy of the text. Reuse the verified
    # picture block for the near-duplicate record so the web library does not
    # show a false missing-image state.
    for case in cases:
        if case["images"]:
            continue
        candidates = [item for item in cases if item["images"] and item["product"] == case["product"]]
        match = max(
            candidates,
            key=lambda item: difflib.SequenceMatcher(None, case["text"], item["text"]).ratio(),
            default=None,
        )
        if match and difflib.SequenceMatcher(None, case["text"], match["text"]).ratio() >= 0.97:
            case["images"] = list(match["images"])
    # The second sheet contains four useful image-replacement instructions.
    extra = [
        {"id": f"replace-{index:02d}", "title": title_for(text, f"替换提示词 {index}"),
         "text": text, "images": [], "product": product_for(text), "scene": scene_for(text),
         "source": f"{path.name} · Sheet2 · A{row}"}
        for index, (row, _, text) in enumerate(workbook_text_rows(path, "Sheet2", (1,)), 1)
    ]
    return cases, extra


def prompt_cases(path: Path) -> list[dict]:
    rows = [(r, t) for r, col, t in workbook_text_rows(path, "Sheet1", (1,)) if r >= 3]
    cases = [
        {"id": f"prompt-{index:03d}", "title": title_for(text, f"生图提示词 {index}"),
         "text": text, "images": [], "product": product_for(text), "scene": scene_for(text),
         "source": f"{path.name} · Sheet1 · A{row}"}
        for index, (row, text) in enumerate(rows, 1)
    ]
    by_row = {row: index for index, (row, _) in enumerate(rows)}
    # In this source workbook, several pictures are anchored to the previous
    # prompt's row. The visual content was checked against each prompt before
    # assigning these exceptions; A3's two pictures are genuine variants.
    target_rows_by_anchor = {
        8: (8, 9, 9),
        12: (12, 13),
        20: (20, 21),
        21: (22,),
        24: (24, 25),
        25: (26,),
        29: (29, 30),
    }
    with zipfile.ZipFile(path) as archive:
        anchored_images: dict[int, list[str]] = {}
        for image_row, _, url in drawing_images(archive):
            anchored_images.setdefault(image_row, []).append(url)
        for image_row, urls in anchored_images.items():
            targets = target_rows_by_anchor.get(image_row, (image_row,) * len(urls))
            if len(targets) != len(urls):
                raise ValueError(f"图片数量与已核对的行关联不符：{path.name} 第 {image_row} 行")
            for target_row, url in zip(targets, urls):
                if target_row not in by_row:
                    raise ValueError(f"图片关联的提示词行不存在：{path.name} 第 {target_row} 行")
                cases[by_row[target_row]]["images"].append(url)
    return cases


def cell_image_map(archive: zipfile.ZipFile) -> dict[str, str]:
    rels = relationship_targets(archive, "xl/_rels/cellimages.xml.rels", "xl")
    root = ET.fromstring(archive.read("xl/cellimages.xml"))
    result = {}
    for picture in root.findall(f".//{DRAW}pic"):
        props = picture.find(f".//{DRAW}cNvPr")
        blip = picture.find(f".//{A}blip")
        if props is None or blip is None:
            continue
        member = rels.get(blip.attrib.get(f"{R}embed", ""))
        if member:
            url = image_url(archive, member)
            if url:
                result[props.attrib.get("name", "")] = url
    return result


def gold_cases(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    rows = []
    for row in workbook["Sheet1"].iter_rows(values_only=True):
        if not isinstance(row[0], str) or row[0] == "描述":
            continue
        rows.append((len(rows) + 1, clean_text(row[0]), row[7], row[16]))
    workbook.close()
    with zipfile.ZipFile(path) as archive:
        images = cell_image_map(archive)
    result = []
    for index, text, output_formula, input_formula in rows:
        output_id = re.search(r'ID_[A-F0-9]+', str(output_formula))
        input_id = re.search(r'ID_[A-F0-9]+', str(input_formula))
        result.append({
            "id": f"gold-{index:02d}", "title": title_for(text, f"金宇案例 {index}"),
            "text": text, "generated": images.get(output_id.group(0)) if output_id else None,
            "original": images.get(input_id.group(0)) if input_id else None,
            "product": product_for(text), "scene": scene_for(text),
            "source": f"{path.name} · Sheet1 · 第 {index} 组",
        })
    return result


def style_slides(path: Path) -> list[dict]:
    groups = [(2, 3, "身材"), (5, 7, "拍摄角度"), (9, 19, "背景与构图"),
              (21, 29, "穿搭与姿态"), (31, 38, "发型与面容")]
    result = []
    with zipfile.ZipFile(path) as archive:
        for start, end, category in groups:
            for number in range(start, end + 1):
                member = f"ppt/slides/slide{number}.xml"
                root = ET.fromstring(archive.read(member))
                text = clean_text(" ".join(node.text or "" for node in root.iter(f"{A}t")))
                rels = relationship_targets(
                    archive, f"ppt/slides/_rels/slide{number}.xml.rels", "ppt/slides")
                urls = []
                for blip in root.iter(f"{A}blip"):
                    image_member = rels.get(blip.attrib.get(f"{R}embed", ""))
                    url = image_url(archive, image_member) if image_member else None
                    if url and url not in urls:
                        urls.append(url)
                result.append({
                    "id": f"style-{number:02d}", "title": f"{category}参考 · 第 {number} 页",
                    "category": category, "text": text, "images": urls,
                    "source": f"{path.name} · 第 {number} 页",
                })
    return result


def main() -> None:
    if not SOURCE.is_dir():
        raise SystemExit(f"Source directory not found: {SOURCE}")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    IMAGES.mkdir(parents=True, exist_ok=True)
    models, replacement_prompts = model_cases(SOURCE / "AI模特关键词.xlsx")
    prompts = prompt_cases(SOURCE / "AI生图提示词xlsx.xlsx") + replacement_prompts
    comparisons = gold_cases(SOURCE / "金宇生图.xlsx")
    styles = style_slides(SOURCE / "生图 -分类.pptx")
    payload = {
        "version": 1,
        "models": models,
        "prompts": prompts,
        "comparisons": comparisons,
        "styles": styles,
        "sources": ["AI模特关键词.xlsx", "AI生图提示词xlsx.xlsx", "金宇生图.xlsx", "生图 -分类.pptx"],
    }
    (OUTPUT / "catalog.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("records", {key: len(payload[key]) for key in ("models", "prompts", "comparisons", "styles")})
    print("image_files", len(list(IMAGES.glob("*.webp"))))
    print("model_image_distribution", dict(Counter(len(item["images"]) for item in models)))
    print("prompt_image_distribution", dict(Counter(len(item["images"]) for item in prompts)))
    print("comparison_complete", sum(bool(item["original"] and item["generated"]) for item in comparisons))


if __name__ == "__main__":
    main()
