"""Entirely generated fixtures for the original application, never copied from company data."""
from __future__ import annotations

import json
import hashlib
import math
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy import text

ROOT = Path(__file__).resolve().parents[1]
DEPARTMENTS = ["电商运营部", "客户服务部", "供应链部", "产品技术部", "设计部", "人力行政部", "财务部", "管理部"]
BRANDS = ["测试箱包", "测试生活", "测试数码", "测试户外", "测试旅行", "测试安防"]
CITIES = ["北京", "上海", "广州", "成都", "武汉", "沈阳", "西安", "德州"]


def sku(i):
    return f"TEST-SKU-{i:04d}"


def product(i):
    return ["测试通勤双肩包", "测试旅行收纳箱", "测试便携键盘", "测试户外水杯", "测试轻量拉杆箱", "测试五金锁具"][(i - 1) % 6] + f" · 款式{i:03d}"


def synthetic_identity(birthday, index):
    # Region 000000 is deliberately invalid: preserves masking/checksum demonstrations without a real ID.
    first = f"000000{birthday:%Y%m%d}{index:03d}"
    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    return first + "10X98765432"[sum(int(n) * weight for n, weight in zip(first, weights)) % 11]


def seed_people(engine, today, put, count):
    employees = []
    for i in range(1, 81):
        hire = today - timedelta(days=40 + (i * 29) % 1400)
        employees.append({"id": i, "source_key": f"synthetic:employee:{i}", "emp_no": f"TEST-{i:04d}",
                          "name": f"测试员工{i:02d}", "company": "示例传媒" if i % 4 == 0 else "示例科技",
                          "contract_type": "劳动合同", "dept_name": DEPARTMENTS[(i - 1) % 8],
                          "group_name": f"测试小组{(i - 1) % 3 + 1}", "position": ["专员", "主管", "工程师", "经理"][i % 4],
                          "hire_date": hire, "regular_date": hire + timedelta(days=90),
                          "birthday": date(1989 + i % 13, i % 12 + 1, i % 27 + 1), "gender": "男" if i % 2 else "女",
                          "phone": f"TEST-CONTACT-{i:03d}", "email": f"employee{i:03d}@example.invalid",
                          "education": ["本科", "大专", "硕士"][i % 3], "is_fulltime": "是", "school": "测试示范大学",
                          "major": ["计算机", "电子商务", "工商管理"][i % 3], "status": "resigned" if i > 72 else "active",
                          "remark": "完全生成的虚构人员；不包含有效身份证或银行账号。", "source_sheet": "测试生成器", "source_row": i})
    put(engine, "hr_employee", employees)
    with engine.begin() as connection:
        for employee in employees:
            identity = synthetic_identity(employee["birthday"], employee["id"])
            connection.execute(text("UPDATE hr_employee SET id_card=:identity,id_card_hash=:digest,bank_card=:bank WHERE source_key=:source AND id_card=''"),
                               {"identity": identity, "digest": hashlib.sha256(identity.encode()).hexdigest(), "bank": f"000000000000{employee['id']:06d}", "source": employee["source_key"]})
    put(engine, "hr_contract", [{"employee_id": e["id"], "contract_type": "劳动合同", "seq_no": 1,
                                "start_date": e["hire_date"], "end_date": today + timedelta(days=12 + e["id"] * 7),
                                "source_key": f"synthetic:contract:{e['id']}", "note": "测试合同"} for e in employees])
    put(engine, "hr_resign", [{"source_key": f"synthetic:resign:{e['id']}", "name": e["name"], "company": e["company"],
                              "dept_name": e["dept_name"], "position": e["position"], "hire_date": e["hire_date"],
                              "resign_date": today - timedelta(days=(e["id"] - 72) * 17), "resign_reason": ["个人发展", "岗位调整", "合同到期"][e["id"] % 3],
                              "education": e["education"], "source_row": e["id"]} for e in employees if e["id"] > 72])
    attendance, details, late, leaves, makeups = [], [], [], [], []
    for offset in range(3):
        first = (today.replace(day=1) - timedelta(days=offset * 28)).replace(day=1)
        month = first.strftime("%Y-%m")
        for e in employees[:72]:
            i = e["id"]
            attendance.append({"employee_id": i, "month": month, "name": e["name"], "dept_name": e["dept_name"],
                               "company": e["company"], "dept1": e["dept_name"], "position": e["position"], "hire_date": e["hire_date"],
                               "work_days": 22, "actual_days": 22 - (i % 4) * .5, "paid_days": 22, "actual_paid_days": 22 - (i % 4) * .5,
                               "late_count": int(i % 9 == 0), "leave_days": (i % 4) * .5, "sick_days": .5 if i % 11 == 0 else 0,
                               "overtime_hours": i % 7 * 1.5, "abnormal_count": int(i % 9 == 0), "source_key": f"synthetic:sum:{month}:{i}"})
            for day in range(1, min(today.day if offset == 0 else 22, 22) + 1):
                work_date = first + timedelta(days=day - 1)
                if work_date.weekday() > 4:
                    continue
                details.append({"employee_id": i, "month": month, "name": e["name"], "dept_name": e["dept_name"], "position": e["position"],
                                "work_date": work_date, "weekday": f"星期{'一二三四五六日'[work_date.weekday()]}", "on_time": "09:12" if i % 9 == 0 and day == 8 else "08:55",
                                "on_result": "迟到" if i % 9 == 0 and day == 8 else "正常", "off_time": "18:10", "off_result": "正常",
                                "work_days": 1, "work_hours": 8, "actual_days": 1, "actual_hours": 8, "shift_hours": 8,
                                "late_count": int(i % 9 == 0 and day == 8), "late_minutes": 12 if i % 9 == 0 and day == 8 else 0,
                                "source_key": f"synthetic:day:{work_date}:{i}"})
            if i % 9 == 0:
                late.append({"employee_id": i, "month": month, "name": e["name"], "late_date": first + timedelta(days=7), "late_time": "09:12", "penalty_amount": 0, "source_key": f"synthetic:late:{month}:{i}"})
            if i % 8 == 0:
                leaves.append({"employee_id": i, "month": month, "name": e["name"], "leave_type": "年假", "reason": "测试请假记录", "start_time": f"{month}-12 09:00", "end_time": f"{month}-12 18:00", "hours": 8, "days": 1, "source_key": f"synthetic:leave:{month}:{i}"})
            if i % 13 == 0:
                makeups.append({"employee_id": i, "month": month, "name": e["name"], "makeup_date": first + timedelta(days=5), "makeup_time": "09:00", "reason": "测试补卡", "source_key": f"synthetic:makeup:{month}:{i}"})
    for name, rows in [("hr_attendance_sum", attendance), ("hr_attendance_detail", details), ("hr_attendance_late", late), ("hr_attendance_leave", leaves), ("hr_attendance_makeup", makeups)]:
        put(engine, name, rows)
    put(engine, "hr_job_template", [{"job_name": title, "job_info": f"【虚构测试岗位】{title}，用于面试流程功能演示。"} for title in ["运营专员", "后端开发工程师", "客服主管", "采购专员", "数据分析师"]])
    put(engine, "interview_register", [{"inviter": "测试人事", "name": f"测试候选人{i:02d}", "position": ["运营专员", "后端开发工程师", "客服主管"][i % 3],
                                      "phone": f"TEST-APPLICANT-{i:02d}", "first_interview_date": (today + timedelta(days=i % 8 - 3)).isoformat(),
                                      "first_interview_time": "14:00", "is_present": ["是", "否", "待定"][i % 3],
                                      "interview_status": ["待面试", "初试通过", "复试通过", "已入职"][i % 4], "hire_status": ["待定", "已录用", "未录用"][i % 3],
                                      "remark": "合成测试候选人记录", "archived": 0} for i in range(1, 25)])
    put(engine, "resume_history", [{"job_title": "后端开发工程师", "jd_text": "测试岗位：Python、SQL、FastAPI、工程化实践。", "filename": f"test-resume-{i:02d}.txt",
                                  "name": f"测试候选人{i:02d}", "education": "测试示范大学 / 计算机 / 本科", "skills": json.dumps(["Python", "FastAPI", "SQL"], ensure_ascii=False),
                                  "experience": "合成履历：开发测试电商平台并设计查询接口。", "match_score": 60 + i % 5 * 8,
                                  "ai_review": "此为预置测试初筛记录，仅用于界面演示，不代表模型对真实人员的评价。", "created_at": datetime.combine(today - timedelta(days=i % 10), datetime.min.time())} for i in range(1, 16)])
    put(engine, "points_rule", [{"id": i, "name": name, "points": value, "description": "演示积分规则", "is_active": 1} for i, (name, value) in enumerate([("知识分享", 20), ("跨部门协作", 15), ("流程改进", 30), ("项目交付", 50), ("记录纠正", -5)], 1)])
    put(engine, "points_record", [{"employee_id": i % 72 + 1, "employee_name": f"测试员工{i % 72 + 1:02d}", "rule_id": i % 4 + 1,
                                 "rule_name": ["知识分享", "跨部门协作", "流程改进", "项目交付"][i % 4], "points": [20, 15, 30, 50][i % 4],
                                 "note": "合成积分流水", "operator": "测试管理员", "created_at": datetime.combine(today - timedelta(days=i % 60), datetime.min.time())} for i in range(216)])


def seed_inventory(engine, today, put, count):
    from utils import inventory_service, business_forecast_service
    put(engine, "inv_warehouse", [{"warehouse_code": city, "warehouse_name": city + "仓", "region": ["华北", "华东", "华南", "西南", "华中", "东北", "西北", "华北"][j],
                                  "warehouse_type": "JD_WAREHOUSE", "is_supply_hub": 0, "is_jd_warehouse": 1, "base_city": city,
                                  "latitude": [39.9, 31.2, 23.1, 30.6, 30.5, 41.8, 34.3, 37.4][j], "longitude": [116.4, 121.5, 113.3, 104.1, 114.3, 123.4, 108.9, 116.3][j]} for j, city in enumerate(CITIES)], append=True)
    put(engine, "inv_sku", [{"sku_code": sku(i), "product_name": product(i), "brand": BRANDS[(i - 1) % 6],
                            "category": "箱包" if i % 6 in (1, 2, 5) else "五金", "shop_name": f"测试店铺{(i - 1) % 6 + 1:02d}",
                            "supplier_name": f"测试供应商{i % 4 + 1:02d}", "barcode": f"TEST-BC-{i:05d}", "case_pack": 6} for i in range(1, 97)])
    snapshots, monitor, batches = [], [], []
    for days_ago in [2, 1, 0]:
        day = today - timedelta(days=days_ago)
        batch = f"TEST-INV-{day:%Y%m%d}"
        batches.append({"batch_no": batch, "file_name": f"synthetic-inventory-{day}.xlsx", "snapshot_date": day, "status": "SUCCESS", "rows_read": 96, "sku_rows": 96, "snapshot_rows": 768})
        for i in range(1, 97):
            metrics = {}
            for j, city in enumerate(CITIES):
                daily = 2 + (i * 3 + j * 5) % 17
                stock = 0 if (i + j) % 11 == 0 else int(daily * [4, 12, 28, 60, 90][(i + j) % 5])
                inbound = 24 if (i + j) % 4 == 0 else 0
                metrics[city] = {"stock_qty": stock + days_ago * 3, "orderable_qty": stock + inbound,
                                 "sales_1d": daily, "sales_7d": daily * 7, "sales_30d": daily * 30}
                snapshots.append({"batch_no": batch, "snapshot_date": day, "sku_code": sku(i), "warehouse_code": city,
                                  "product_status": "下柜" if i % 13 == 0 else "上柜", "is_gift": 0, "is_decision_active": int(i % 13 != 0),
                                  "stock_qty": stock + days_ago * 3, "inbound_qty": inbound, "available_qty": stock,
                                  "orderable_qty": stock + inbound, "sales_1d": daily, "sales_7d": daily * 7, "sales_14d": daily * 14, "sales_30d": daily * 30})
            monitor.append({"snapshot_date": day, "sku_code": sku(i), "product_name": product(i), "brand": BRANDS[(i - 1) % 6],
                            "category_l1": "箱包" if i % 6 in (1, 2, 5) else "五金", "category_l2": "测试二级类目", "category_l3": "测试三级类目",
                            "shop_name": f"测试店铺{(i - 1) % 6 + 1:02d}", "product_status": "下柜" if i % 13 == 0 else "上柜",
                            "national_stock_qty": sum(m["stock_qty"] for m in metrics.values()), "national_orderable_qty": sum(m["orderable_qty"] for m in metrics.values()),
                            "national_sales_1d": sum(m["sales_1d"] for m in metrics.values()), "national_sales_7d": sum(m["sales_7d"] for m in metrics.values()),
                            "national_sales_30d": sum(m["sales_30d"] for m in metrics.values()), "city_metrics": metrics})
    put(engine, "inv_import_batch", batches)
    put(engine, "inv_stock_snapshot", snapshots)
    put(engine, "inv_monitor_snapshot", monitor)
    put(engine, "inv_supply_stock_snapshot", [{"snapshot_time": datetime.combine(today, datetime.min.time()), "snapshot_date": today, "warehouse_code": "花东仓", "sku_code": sku(i),
                                              "stock_qty": 300 + i * 7, "available_qty": 300 + i * 7, "inbound_qty": 60, "source": "SYNTHETIC", "raw_payload": {"synthetic": True}} for i in range(1, 97)])
    if not count(engine, "inv_decision_run"):
        result = inventory_service.recalculate_inventory_decisions(snapshot_date=today, engine=engine)
        print("Real inventory decision engine:", str(result)[:180])
    business = []
    for k in range(90):
        day = today - timedelta(days=89 - k)
        for i in range(1, 49):
            units = 25 + i % 17 + int(k * (0.15 if i % 3 else -.09)) + k % 7
            visitors = 400 + i * 11 + k * 3
            business.append({"biz_date": day, "sku_code": sku(i), "product_name": product(i), "brand": BRANDS[(i - 1) % 6], "category_l1": "箱包", "category_l2": "测试配件", "category_l3": "测试商品",
                             "shop_name": f"测试店铺{(i - 1) % 6 + 1:02d}", "business_mode": "测试自营", "page_views": visitors * 2.4, "visitors": visitors,
                             "views_per_visitor": 2.4, "avg_stay_seconds": 45 + i % 20, "buyers": units * .8, "conversion_rate": units * .8 / visitors,
                             "order_count": int(units * .9), "sold_units": units, "gmv_amount": units * (79 + i * 3), "avg_order_value": units * (79 + i * 3) / int(units * .9),
                             "cart_users": visitors * .13, "cart_conversion_rate": .13, "cart_units": visitors * .18})
    put(engine, "inv_business_import_batch", [{"batch_no": "TEST-BUSINESS-001", "file_name": "synthetic-business.xlsx", "status": "SUCCESS", "min_biz_date": today - timedelta(days=89), "max_biz_date": today, "rows_read": len(business), "rows_imported": len(business), "unique_sku_count": 48, "message": "合成经营数据"}])
    put(engine, "inv_business_daily", business)
    if not count(engine, "inv_business_forecast"):
        result = business_forecast_service.recalculate_business_forecasts(forecast_date=today, engine=engine)
        print("Real business forecast engine:", str(result)[:160])


def seed_operating(engine, today, put, count):
    rows = []
    for k in range(180):
        day = today - timedelta(days=179 - k)
        for i in range(1, 37):
            quantity = 8 + i % 15 + k % 7 + (k // 30) * (1 + i % 3)
            price = 59 + i * 7
            visitors = 170 + i * 9 + k * 2
            rows.append({"day": day.isoformat(), "sku": sku(i), "name": product(i), "code": f"TEST-MODEL-{i:03d}",
                         "category": ["箱包/通勤/双肩包", "数码/电脑配件/键盘", "户外/日用/水杯"][i % 3], "filename": "synthetic-operating.xlsx", "source_row": k * 36 + i + 1,
                         "revenue": quantity * price * 1000000, "quantity": quantity, "orders": int(quantity * .85), "customers": int(quantity * .8), "views": int(visitors * 2.1), "visitors": visitors,
                         "cart_quantity": int(quantity * 1.5), "cart_customers": int(quantity * 1.2), "cart_amount": int(quantity * 1.5) * price * 1000000,
                         "raw": json.dumps({"synthetic": True, "时间": day.isoformat(), "SKU": sku(i), "SKU名称": product(i),
                                            "成交金额": quantity * price, "加购金额": int(quantity * 1.5) * price}, ensure_ascii=False),
                         "source": "示例科技" if i % 3 else "示例户外", "brand": BRANDS[(i - 1) % 6]})
    put(engine, "jy_batches", [{"filename": "synthetic-operating.xlsx", "digest": hashlib.sha256(b"synthetic-operating-v1").hexdigest(), "imported_at": datetime.now(), "row_count": len(rows),
                              "start": rows[0]["day"], "end": today.isoformat(), "control": "{}", "headers": json.dumps(["时间", "SKU", "成交金额"], ensure_ascii=False)}])
    put(engine, "jy_records", rows)
    put(engine, "jy_operating_brands", [{"name": product(i), "brand": BRANDS[(i - 1) % 6]} for i in range(1, 37)])
    put(engine, "jy_operating_model_mapping", [{"sku": sku(i), "code": f"TEST-MODEL-{i:03d}"} for i in range(1, 37)])
    # These are the real import service's schema-version markers. They prevent its legacy
    # pre-scale migration from treating already scaled synthetic values as old raw data.
    put(engine, "jy_operating_settings", [{"setting_key": "money_scale", "value": "1000000"}, {"setting_key": "brand_rules_v2", "value": "1"}], append=True)
    with engine.begin() as connection:
        connection.execute(text("UPDATE jy_records SET orders=FLOOR(orders),customers=FLOOR(customers),views=FLOOR(views),cart_quantity=FLOOR(cart_quantity),cart_customers=FLOOR(cart_customers) WHERE filename='synthetic-operating.xlsx'"))
        broken = connection.execute(text("SELECT COUNT(*) FROM jy_records WHERE filename='synthetic-operating.xlsx' AND revenue=0")).scalar()
        if broken:
            connection.execute(text("UPDATE jy_records SET revenue=:revenue,cart_amount=:cart_amount,raw=:raw WHERE day=:day AND sku=:sku AND filename='synthetic-operating.xlsx' AND revenue=0"),
                               [{key: row[key] for key in ("day", "sku", "revenue", "cart_amount", "raw")} for row in rows])
        for i in range(1, 37):
            connection.execute(text("UPDATE jy_operating_brands SET brand=:brand WHERE name=:name AND brand='待识别'"), {"name": product(i), "brand": BRANDS[(i - 1) % 6]})


def seed_supply_dashboards(engine, today, put, count):
    put(engine, "inventory_product_import_batch", [{"batch_id": "test-inventory-products", "report_date": today, "period_start": today.replace(day=1), "period_end": today,
                                                    "source_name": "synthetic-inventory-products.xlsx", "row_count": 96, "image_count": 0, "imported_by": "测试管理员"}])
    put(engine, "inventory_product_dashboard_item", [{"batch_id": "test-inventory-products", "source_row": i + 1, "category": ["箱包", "数码", "户外"][i % 3], "style_code": f"TEST-MODEL-{i:03d}",
                                                      "product_code": sku(i), "product_name": product(i), "color_spec": ["测试黑色 / 标准款", "测试蓝色 / 大号"][i % 2],
                                                      "brand": BRANDS[(i - 1) % 6], "opening_quantity": 250 + i * 7, "closing_quantity": 180 + i * 6} for i in range(1, 97)])
    put(engine, "jst_product_import_batch", [{"batch_id": "test-erp-products", "source_name": "synthetic-erp.xlsx", "row_count": 96, "image_count": 0, "imported_by": "测试管理员"}])
    put(engine, "jst_product_stock", [{"batch_id": "test-erp-products", "source_row": i + 1, "style_code": f"TEST-MODEL-{i:03d}", "product_code": sku(i), "product_name": product(i),
                                      "brand": BRANDS[(i - 1) % 6], "color_spec": "测试标准款", "color": ["黑色", "蓝色", "灰色"][i % 3], "spec": "标准款",
                                      "actual_stock": 400 + i * 9, "order_occupied": i * 2, "available_stock": 400 + i * 7, "stock_updated_at": datetime.now()} for i in range(1, 97)])
    put(engine, "sales_dashboard_import_batch", [{"batch_id": "test-sales-products", "source_name": "synthetic-sales.xlsx", "row_count": 96, "imported_by": "测试管理员"}])
    put(engine, "sales_dashboard_item", [{"batch_id": "test-sales-products", "source_row": i + 1, "product_code": sku(i), "product_name": product(i), "category": ["箱包", "数码", "户外"][i % 3],
                                        "brand": BRANDS[(i - 1) % 6], "sales_quantity": 250 + i * 7, "shipped_quantity": 240 + i * 7, "returned_quantity": i % 8 + 2} for i in range(1, 97)])
    put(engine, "purchase_order_dashboard_batch", [{"batch_id": "test-purchases", "source_name": "synthetic-purchases.xlsx", "date_start": today - timedelta(days=59), "date_end": today, "row_count": 80, "order_count": 40, "image_count": 0, "imported_by": "测试管理员"}])
    put(engine, "purchase_order_dashboard_item", [{"batch_id": "test-purchases", "source_row": i + 1, "order_no": f"TEST-PO-{(i - 1) // 2 + 1:04d}", "purchase_date": datetime.combine(today - timedelta(days=i % 60), datetime.min.time()),
                                                 "status": ["待审核", "已审核", "已完成"][i % 3], "receiving_status": ["未收货", "部分收货", "全部收货"][i % 3], "purchase_remark": "合成采购数据",
                                                 "receiving_address": "测试园区 / 虚构地址", "purchase_type": ["常规采购", "补货采购"][i % 2], "warehouse_party": CITIES[i % 8] + "测试仓", "purchaser": f"测试采购员{i % 4 + 1:02d}",
                                                 "review_date": datetime.combine(today - timedelta(days=max(0, i % 60 - 1)), datetime.min.time()), "last_modified_date": datetime.now(),
                                                 "product_code": sku(i), "style_code": f"TEST-MODEL-{i:03d}", "product_name": product(i), "color_spec": "测试标准款", "pending_ship_quantity": 20 + i % 10,
                                                 "available_quantity": 100 + i * 3, "quantity": 120 + i * 2, "received_quantity": 80 + i * 2, "variance_quantity": 40,
                                                 "supplier_product_code": f"TEST-SUP-{i:03d}", "supplier_style_no": f"TEST-STYLE-{i:03d}", "detail_no": f"TEST-DETAIL-{i:04d}"} for i in range(1, 81)])


def seed_shared_workspaces(engine, today, put, count):
    from utils import ecom_dashboard_service, cs_dashboard_service, ai_daily_board_service
    if not count(engine, "ecom_menu"):
        payload = json.loads((ROOT / "tools/ecom_seed.json").read_text(encoding="utf8"))
        ecom_dashboard_service.save_board_state(payload, engine)
    if not count(engine, "cs_dashboard_state"):
        rows = []
        for week in range(4):
            first = today - timedelta(days=today.weekday() + 7 * (3 - week))
            for store in range(1, 7):
                for agent in range(1, 6):
                    questions = 180 + store * 31 + agent * 27 + week * 17
                    robot = int(questions * (.5 + agent * .02))
                    human = questions * 2 - robot
                    rows.append({"store": f"测试店铺{store:02d}", "agentName": f"测试客服{agent:02d}", "account": f"test-agent-{agent:02d}-shop-{store:02d}",
                                 "agentRole": "售前" if agent < 3 else "售后", "reportWeekKey": first.isoformat(),
                                 "reportWeekLabel": f"{first:%m/%d}—{first + timedelta(days=6):%m/%d}", "visitors": int(questions * .8), "questions": questions,
                                 "consults": int(questions * .75), "robotMsgs": robot, "humanMsgs": human, "qaRatio": 1.86,
                                 "robotQaRatio": 1.93, "humanQaRatio": 1.8, "humanRespSec": round(18.5 - week * 1.2 + agent * 1.3, 2),
                                 "accountRespSec": round(12.4 - week * .7 + agent * .9, 2), "speedImprove": round(.28 + week * .03 + agent * .01, 3)})
        cs_dashboard_service.save_state({"rows": rows, "sourceName": "合成客服测试报表（4周）", "loaded": True}, "测试管理员", engine)
    workbook = ROOT / "uploads/demo-fixtures/customer-service-demo.xlsx"
    if not workbook.exists():
        from openpyxl import Workbook
        workbook.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "测试客服周报"
        fields = {"store": "店铺", "agentName": "客服名字", "account": "客服", "visitors": "接待人数", "questions": "接待问题数", "consults": "咨询人数",
                  "robotMsgs": "机器人发出消息数", "humanMsgs": "人工发出消息数", "qaRatio": "答问比", "robotQaRatio": "机器人答问比", "humanQaRatio": "人工答问比",
                  "humanRespSec": "人工平均响应时长（秒）", "accountRespSec": "账号平均响应时长（秒）", "speedImprove": "响应速度提升"}
        sheet.append(list(fields.values()))
        current = cs_dashboard_service.get_state(engine)["rows"]
        latest = max(row["reportWeekKey"] for row in current)
        for row in current:
            if row["reportWeekKey"] == latest:
                sheet.append([row.get(key, "") for key in fields])
        book.save(workbook)
    if not count(engine, "ai_daily_board_project"):
        names = ["知识库问答助手", "经营数据分析智能体", "库存预测与补货", "客服质检流程", "简历结构化解析", "设计素材管理", "采购协同助手", "权限与账号治理", "门店费用核算"]
        projects = [{"id": f"test-project-{i:02d}", "name": name, "status": ["待开启", "开发中", "已上线"][i % 3],
                     "focus": ["测试数据接入与字段核对", "联调查询、追溯与导出", "持续验证响应与权限"][i % 3], "owner": f"测试员工{i + 1:02d}",
                     "startDate": (today - timedelta(days=20 + i * 4)).isoformat(), "endDate": (today + timedelta(days=10 + i * 3)).isoformat()} for i, name in enumerate(names)]
        for offset in range(7, -1, -1):
            day = today - timedelta(days=offset)
            payload = {"settings": {"boardTitle": "AI 项目进度看板", "currentDate": day.isoformat(), "showCharts": True, "showStandup": True, "showTimeline": True, "showTable": True},
                       "projects": projects, "dailyLogs": [{"projectId": p["id"], "projectName": p["name"], "work": f"第{8 - offset}轮测试：完成字段联调、权限验证和记录核对。",
                                                          "tomorrow": "补齐边界场景并复核导出结果。", "risks": "等待测试评审。" if i % 4 == 0 else ""} for i, p in enumerate(projects[:5])],
                       "risks": [{"id": "test-risk-01", "risk": "测试接口字段仍需核对", "projectId": projects[2]["id"], "projectName": projects[2]["name"], "impact": "中", "owner": "测试员工03", "status": "处理中", "blockedSince": (today - timedelta(days=3)).isoformat()}]}
            ai_daily_board_service.save_board_state(payload, engine)
        for i in range(10):
            ai_daily_board_service.add_activity(engine, operator="测试管理员", action="测试进度更新", target=projects[i % len(projects)]["name"], detail="虚构的演示操作记录")


def write_shop_snapshot(today):
    from tools.export_shop_report import FEE_KEYS, FEE_LABELS
    output = ROOT / "data/shop_report.json"
    if output.exists():
        old = json.loads(output.read_text(encoding="utf-8-sig"))
        if old.get("source", {}).get("demo_version") == "full-application-v1":
            return
    shops = [
        ("self", "demo_bags", "测试箱包自营"), ("self", "demo_life", "测试生活自营"),
        ("self", "home", "测试家装自营"), ("self", "demo_apparel", "测试服饰自营"),
        ("self", "fly", "测试五金自营"), ("self", "demo_security", "测试安防自营"),
        ("pop", "demo_bags", "测试箱包旗舰店"), ("pop", "demo_apparel", "测试服饰旗舰店"),
        ("store", "demo_store", "测试体验门店"),
    ]
    months = [f"{today.year}-{m:02d}" for m in range(1, today.month + 1)]
    rows = []
    for j, (channel, shop_id, label) in enumerate(shops):
        for m, month in enumerate(months, 1):
            revenue = 95000 + j * 16800 + m * (3800 + j * 220) + (m * 7 + j) % 6 * 1350
            cost = round(revenue * (.48 + j * .018), 2)
            row = {key: 0 for key in FEE_KEYS}
            row.update(promo=round(revenue * .075, 2), maobaoFee=round(revenue * .012, 2), storage=620 + j * 75,
                       storageHd=180 + j * 20, express=1700 + m * 45, expressHd=420 + j * 35,
                       tcFwd=round(revenue * .004, 2), tcRev=round(revenue * .002, 2), platform=round(revenue * .018, 2),
                       commission=round(revenue * .007, 2), transFee=round(revenue * .006, 2), shipInsure=95 + j * 8,
                       jdBeans=120 + m * 12, afterSale=80 + j * 25, priceRebate=-round(revenue * .003, 2),
                       brand_rebate=-round(revenue * .002, 2), rent=2100 if channel == "store" else 0)
            fees = round(sum(row[key] for key in FEE_KEYS if key != "labor"), 2)
            row.update(channel=channel, shop_id=shop_id, month=month, revenue=revenue, costG=cost, rentAlloc=0,
                       gross_profit=round(revenue - cost, 2), total_fees=fees, product_profit=round(revenue - cost - fees, 2), net_profit=round(revenue - cost - fees, 2),
                       has_data=True, data_status="complete", data_note="完全虚构的公式生成测试数据。", quantity=850 + j * 53 + m * 37,
                       missing_cost=False, missing_fees=False, source_conflict=False, profit_provisional=False,
                       storageGen=round(row["storage"] * .7, 2), storageCons=round(row["storage"] * .3, 2), known_fee_keys=list(FEE_KEYS))
            rows.append(row)
    payload = {"source": {"name": "合成测试店铺报表", "year": today.year, "currency": "CNY", "labor_included": False, "synthetic": True, "demo_version": "full-application-v1",
                          "exported_at": datetime.combine(today, datetime.min.time()).isoformat(), "profit_policy": "收入减商品成本与已录费用；人工及税费未计入。"},
               "months": months, "shops": [{"channel": ch, "shop_id": sid, "shop_name": label, "months": months} for ch, sid, label in shops],
               "fee_keys": list(FEE_KEYS), "fee_labels": dict(FEE_LABELS), "rows": rows}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf8", newline="\n")


def seed_return_sources(engine, today, put, count):
    if count(engine, "sku_return_aggregate"):
        return
    from openpyxl import Workbook
    from utils import sku_return_service
    source = ROOT / "data/synthetic-return-sources"
    headers = ["供应商名称", "应付账单据类型", "SKU编码", "SKU名称", "数量", "单据ID", "单据编号"]
    for month in range(1, today.month + 1):
        folder = source / f"{today.year}年示例科技结算单" / f"{month}月" / "示例科技"
        folder.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "测试数据"
        sheet.append(headers)
        for i in range(1, 37):
            for kind, quantity in [("实销实结销售单", 110 + i * 3 + month * 7), ("实销实结退货入库单", -(i % 7 + 2)), ("售后退货", -(i % 4 + 1))]:
                sheet.append([f"测试供应商{i % 4 + 1:02d}", kind, sku(i), ("测试旅行" if i % 2 else "测试箱包") + f" 合成商品{i:03d}", quantity,
                              f"TEST-DOC-{month}-{i}-{kind}", f"TEST-NO-{month}-{i}-{kind}"])
        book.save(folder / f"测试箱包-{month:02d}.xlsx")
    result = sku_return_service.import_source_tree(source, apply=True, engine=engine)
    print("Real SKU-return importer:", result.get("status"), result.get("included_rows"))


def seed_permissions(engine):
    from utils import permission_service as permissions
    with engine.connect() as connection:
        done = connection.execute(text("SELECT meta_value FROM demo_runtime_meta WHERE meta_key='permissions_seeded'")).scalar()
    if done:
        return
    entries = []
    for key in permissions.all_page_defs():
        roles = ["admin", "boss"]
        if key.startswith("hr-"):
            roles += ["hr"]
        elif key.startswith("design-"):
            roles += ["designer"]
        elif key in {"inventory-decision", "inventory-product-dashboard", "jushuitan-dashboard", "purchase-order-dashboard", "sales-dashboard"}:
            roles += ["supply", "operator"]
        else:
            roles += ["hr", "designer", "operator", "supply", "finance", "user"]
        entries.append({"page_key": key, "roles": roles})
    permissions.save_page_permissions(engine, entries)
    subpages = [{"page_key": page, "sub_key": sub, "roles": ["admin", "boss", "operator", "supply"]}
                for page, definitions in permissions.SUB_PERM_DEFS.items() for sub in definitions]
    permissions.save_subpage_permissions(engine, subpages)
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO demo_runtime_meta VALUES ('permissions_seeded','1')"))


def seed_asset_libraries(engine, today, put, count):
    """Create sample visuals, PDF and knowledge files locally from fixture definitions."""
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject
    from utils import design_library_service as design, product_quality_report_service as reports, ecom_knowledge_service as knowledge
    if not count(engine, "design_library_item"):
        visual_urls = []
        for i in range(1, 13):
            colors = ["#557b98", "#c08659", "#626f59", "#7a6f98", "#5d8d89", "#8c6473"]
            image = Image.new("RGB", (720, 900), "#f1f2f3")
            draw = ImageDraw.Draw(image)
            draw.ellipse((120, 740, 600, 790), fill="#d8dcdf")
            color = colors[(i - 1) % len(colors)]
            draw.rounded_rectangle((220, 255, 500, 725), radius=70, fill=color)
            draw.rounded_rectangle((245, 490, 475, 670), radius=30, outline="#e0e7ed", width=5)
            draw.arc((285, 160, 435, 340), 180, 360, fill=color, width=20)
            draw.line((250, 380, 470, 380), fill="#e0e7ed", width=5)
            draw.text((38, 38), f"SYNTHETIC / PRODUCT STUDY {i:02d}", fill="#303c47", font_size=22)
            draw.text((38, 825), "TEST FIXTURE - NOT A REAL PRODUCT", fill="#66717a", font_size=20)
            buf = BytesIO(); image.save(buf, format="WEBP")
            saved = design.save_image(buf.getvalue())
            visual_urls.append(saved)
        for kind in design.ITEM_TYPES:
            for i in range(1, 7):
                before = visual_urls[(i - 1) % len(visual_urls)]
                after = visual_urls[(i + 2) % len(visual_urls)]
                payload = {"item_type": kind, "title": f"测试设计案例{i:02d} · {dict(models='商品与场景',prompts='提示词',comparisons='原图对比',styles='视觉风格')[kind]}",
                           "text": "## 合成测试素材\n\n此内容仅展示原设计资料库的分类、筛选、详情、编辑与回收流程。\n\n- 商品：虚构测试背包\n- 场景：简洁陈列\n- 光线：柔和漫射\n\n样例图片由绘图程序生成，不对应实际商品或商业作品。",
                           "product": ["测试箱包", "测试户外"][i % 2], "scene": ["电商主图", "场景陈列", "社媒封面"][i % 3], "category": "测试素材",
                           "images": [before], "original": before if kind == "comparisons" else None, "generated": after if kind == "comparisons" else None, "source": "程序生成测试素材"}
                design.create_item(payload, user_name="测试管理员", engine=engine)
    if not count(engine, "product_quality_report"):
        for i in range(1, 5):
            writer = PdfWriter()
            page = writer.add_blank_page(width=595, height=842)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            lines = ["SYNTHETIC PRODUCT INSPECTION REPORT", f"TEST REPORT {i:03d}", "Fixture only - no real product or certification.",
                     "Sample: fictional demonstration backpack", "Checks: visual inspection / dimensions / material", "Result: synthetic example, not a certification", f"Reference date: {today}"]
            commands = ["BT /F1 16 Tf 50 770 Td"]
            for j, line in enumerate(lines):
                if j: commands.append("0 -45 Td")
                commands.append(f"({line}) Tj")
            commands.append("ET")
            stream = DecodedStreamObject(); stream.set_data("\n".join(commands).encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
            raw = BytesIO(); writer.write(raw)
            reports.save_report(raw.getvalue(), f"TS1_0000{i}-TEST-INSPECTION.pdf", "测试质检", "测试管理员", engine)
    if not count(engine, "ecom_knowledge_file"):
        from fastapi import UploadFile
        from starlette.datastructures import Headers
        for library in knowledge.LIBRARIES:
            for i, title in enumerate(["测试业务指南", "测试操作流程"], 1):
                body = f"# {title}\n\n本文件是虚构测试知识资料。所属知识库：{library}。\n\n## 测试业务范围\n本演示使用独立数据库、测试店铺、测试员工和生成金额。\n\n## 报表核对\n先选择相同期间、店铺和渠道，再核对收入、成本、费用与净利。净利不等于回款。\n\n## 操作与追溯\n导入先预览，查看字段映射和质量报告，保留来源行。库存建议需要考虑仓库、在途、近30日出库和安全周期。\n"
                upload = UploadFile(BytesIO(body.encode("utf8")), filename=f"test-guide-{i:02d}.md", headers=Headers({"content-type": "text/markdown"}))
                knowledge.upload_file(library, "测试资料", upload, engine)
    after_sales = ROOT / "data/after_sales/2026-08-11-18-35-08--105874739770.xlsx"
    if not after_sales.exists():
        from openpyxl import Workbook
        after_sales.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook(); sheet = book.active
        sheet.append([f"测试字段{i}" for i in range(46)])
        sheet.append([f"测试子字段{i}" for i in range(46)])
        reasons = ["拉链损坏", "商品质量问题", "商品与页面描述不符", "七天无理由", "商品破损/包装问题", "配件脱落"]
        for i in range(1, 241):
            row = [""] * 46
            row[0], row[8], row[14] = f"TEST-AS-{i:04d}", f"TEST-ORDER-{i:05d}", sku(i % 36 + 1)
            row[15] = ("测试旅行" if i % 2 else "测试箱包") + f" 虚构箱包{i % 12:02d}"
            row[3], row[4], row[5] = reasons[i % 6], "测试说明", ["拉链卡顿", "肩带走线测试", "不需要了", "测试外观划痕"][i % 4]
            row[7] = datetime.combine(today - timedelta(days=i % 90), datetime.min.time())
            row[16], row[37], row[45] = 159 + i % 12 * 17, 99 + i % 9 * 15, 5 + i % 48
            row[21], row[23], row[24], row[33] = "测试审核通过", "产品质量" if i % 3 else "个人原因", "虚构描述", "测试收货原因"
            row[34], row[35] = f"测试店铺{i % 6 + 1:02d}", "已完成"
            sheet.append(row)
        book.save(after_sales)


def populate(engine, today, insert_rows, count):
    from utils.account_password import hash_password
    insert_rows(engine, "sys_user", [
        {"id": i, "username": username, "password": hash_password("Demo@2026"), "real_name": label,
         "department": department, "job_title": title, "role": role}
        for i, (username, label, role, department, title) in enumerate([
            ("demo_admin", "测试管理员", "admin", "产品技术部", "系统管理员"),
            ("demo_boss", "测试负责人", "boss", "管理部", "负责人"),
            ("demo_hr", "测试人事", "hr", "人力行政部", "人事专员"),
            ("demo_operator", "测试运营", "operator", "电商运营部", "运营专员"),
            ("demo_designer", "测试设计", "designer", "设计部", "设计师"),
            ("demo_supply", "测试供应链", "supply", "供应链部", "采购专员"),
        ], 1)
    ])
    seed_people(engine, today, insert_rows, count)
    seed_inventory(engine, today, insert_rows, count)
    seed_operating(engine, today, insert_rows, count)
    seed_supply_dashboards(engine, today, insert_rows, count)
    seed_shared_workspaces(engine, today, insert_rows, count)
    write_shop_snapshot(today)
    seed_return_sources(engine, today, insert_rows, count)
    seed_asset_libraries(engine, today, insert_rows, count)
    seed_permissions(engine)


if __name__ == "__main__":
    raise SystemExit("Run tools/full_demo_seed.py")
