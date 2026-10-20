"""Import and query the weekly inventory-by-product dashboard."""

from __future__ import annotations

from datetime import date, datetime, timedelta
import hashlib
from pathlib import Path
import re
from typing import Any
from urllib.parse import urlparse
import uuid

import httpx
from openpyxl import load_workbook
from sqlalchemy import Engine, bindparam, text

from utils import import_service


UPLOAD_ROOT = Path(__file__).resolve().parent.parent / "uploads" / "inventory_product"
IMAGE_MAX_BYTES = 15 * 1024 * 1024
IMAGE_TIMEOUT_SECONDS = 20.0
IMAGE_EXTENSIONS = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/webp": ".webp",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
}

_BATCH_DDL = """
CREATE TABLE IF NOT EXISTS inventory_product_import_batch (
    batch_id VARCHAR(32) NOT NULL,
    report_date DATE NOT NULL,
    period_start DATE NOT NULL,
    period_end DATE NOT NULL,
    source_name VARCHAR(255) NOT NULL DEFAULT '',
    row_count INT NOT NULL DEFAULT 0,
    image_count INT NOT NULL DEFAULT 0,
    imported_by VARCHAR(64) NOT NULL DEFAULT '',
    imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (batch_id),
    UNIQUE KEY uk_inventory_product_report_date (report_date),
    KEY idx_inventory_product_period (period_start, period_end)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='进销商品看板导入批次'
"""

_ITEM_DDL = """
CREATE TABLE IF NOT EXISTS inventory_product_dashboard_item (
    id BIGINT NOT NULL AUTO_INCREMENT,
    batch_id VARCHAR(32) NOT NULL,
    source_row INT NOT NULL,
    image_url TEXT,
    category VARCHAR(128) NOT NULL DEFAULT '',
    style_code VARCHAR(128) NOT NULL DEFAULT '',
    product_code VARCHAR(128) NOT NULL,
    product_name TEXT NOT NULL,
    color_spec VARCHAR(255) NOT NULL DEFAULT '',
    brand VARCHAR(128) NOT NULL DEFAULT '',
    opening_quantity DECIMAL(18,2) NOT NULL DEFAULT 0,
    closing_quantity DECIMAL(18,2) NOT NULL DEFAULT 0,
    imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (id),
    UNIQUE KEY uk_inventory_product_source (batch_id, source_row),
    KEY idx_inventory_product_code (product_code),
    KEY idx_inventory_product_category (category),
    KEY idx_inventory_product_brand (brand)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='进销商品看板明细快照'
"""

EXPECTED_HEADERS = (
    "图片",
    "分类",
    "款式编码",
    "商品编码",
    "商品名称",
    "颜色规格",
    "品牌",
    "期初数量",
    "期末数量",
)


def get_engine(engine: Engine | None = None) -> Engine:
    return engine or import_service.get_engine()


def ensure_schema(engine: Engine | None = None) -> None:
    db_engine = get_engine(engine)
    with db_engine.begin() as conn:
        conn.execute(text(_BATCH_DDL))
        conn.execute(text(_ITEM_DDL))


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _normalise_headers(values: tuple[Any, ...]) -> tuple[str, ...]:
    return tuple(_clean(value).replace("\ufeff", "").replace(" ", "") for value in values)


def parse_report_date(source_name: str | Path) -> date:
    """Extract the first YYYYMMDD date from the report filename."""

    stem = Path(source_name).stem
    match = re.search(r"(20\d{6})", stem)
    if not match:
        raise ValueError("文件名中未找到报告日期，请使用包含 YYYYMMDD 的文件名")
    try:
        return datetime.strptime(match.group(1), "%Y%m%d").date()
    except ValueError as exc:
        raise ValueError(f"文件名中的报告日期无效：{match.group(1)}") from exc


def report_period(report_date: date) -> tuple[date, date]:
    # Follow the requested Monday example: 2026-09-14 -> 2026-09-06 ~ 2026-09-13.
    return report_date - timedelta(days=8), report_date - timedelta(days=1)


def _read_products(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        headers = _normalise_headers(next(rows, ()))
        header_index = {header: index for index, header in enumerate(headers)}
        missing = [header for header in EXPECTED_HEADERS if header not in header_index]
        if missing:
            raise ValueError(f"不是可识别的进销存按商品表，缺少字段：{'、'.join(missing)}")

        products: list[dict[str, Any]] = []
        for source_row, values in enumerate(rows, start=2):
            if not any(value is not None and _clean(value) for value in values):
                continue

            def get_value(header: str) -> Any:
                index = header_index[header]
                return values[index] if index < len(values) else None

            product_code = _clean(get_value("商品编码"))
            if not product_code:
                continue
            products.append(
                {
                    "source_row": source_row,
                    "image_url": _clean(get_value("图片")),
                    "category": _clean(get_value("分类")),
                    "style_code": _clean(get_value("款式编码")),
                    "product_code": product_code,
                    "product_name": _clean(get_value("商品名称")),
                    "color_spec": _clean(get_value("颜色规格")),
                    "brand": _clean(get_value("品牌")),
                    "opening_quantity": _number(get_value("期初数量")),
                    "closing_quantity": _number(get_value("期末数量")),
                }
            )
    finally:
        workbook.close()

    if not products:
        raise ValueError("进销存按商品表未读取到有效商品")
    return products


def import_excel(
    path: str | Path,
    source_name: str = "",
    imported_by: str = "",
    engine: Engine | None = None,
) -> dict[str, Any]:
    source_path = Path(path)
    products = _read_products(source_path)
    report_date = parse_report_date(source_name or source_path.name)
    period_start, period_end = report_period(report_date)
    batch_id = uuid.uuid4().hex
    imported_at = datetime.now()
    image_count = sum(1 for product in products if product["image_url"])
    payload = [
        {
            **product,
            "batch_id": batch_id,
            "imported_at": imported_at,
        }
        for product in products
    ]

    db_engine = get_engine(engine)
    ensure_schema(db_engine)
    with db_engine.begin() as conn:
        old_batch_ids = conn.execute(
            text("SELECT batch_id FROM inventory_product_import_batch WHERE report_date = :report_date"),
            {"report_date": report_date},
        ).scalars().all()
        for old_batch_id in old_batch_ids:
            conn.execute(
                text("DELETE FROM inventory_product_dashboard_item WHERE batch_id = :batch_id"),
                {"batch_id": old_batch_id},
            )
        conn.execute(
            text("DELETE FROM inventory_product_import_batch WHERE report_date = :report_date"),
            {"report_date": report_date},
        )
        conn.execute(
            text(
                """
                INSERT INTO inventory_product_import_batch
                    (batch_id, report_date, period_start, period_end, source_name,
                     row_count, image_count, imported_by, imported_at)
                VALUES
                    (:batch_id, :report_date, :period_start, :period_end, :source_name,
                     :row_count, :image_count, :imported_by, :imported_at)
                """
            ),
            {
                "batch_id": batch_id,
                "report_date": report_date,
                "period_start": period_start,
                "period_end": period_end,
                "source_name": (source_name or source_path.name)[:255],
                "row_count": len(products),
                "image_count": image_count,
                "imported_by": _clean(imported_by)[:64],
                "imported_at": imported_at,
            },
        )
        for start in range(0, len(payload), 500):
            conn.execute(
                text(
                    """
                    INSERT INTO inventory_product_dashboard_item (
                        batch_id, source_row, image_url, category, style_code,
                        product_code, product_name, color_spec, brand,
                        opening_quantity, closing_quantity, imported_at
                    ) VALUES (
                        :batch_id, :source_row, :image_url, :category, :style_code,
                        :product_code, :product_name, :color_spec, :brand,
                        :opening_quantity, :closing_quantity, :imported_at
                    )
                    """
                ),
                payload[start : start + 500],
            )

    return {
        "batch_id": batch_id,
        "report_date": report_date.isoformat(),
        "period_start": period_start.isoformat(),
        "period_end": period_end.isoformat(),
        "period_label": f"{period_start:%Y-%m-%d} ~ {period_end:%Y-%m-%d}",
        "source_name": source_name or source_path.name,
        "row_count": len(products),
        "image_count": image_count,
        "imported_at": imported_at.strftime("%Y-%m-%d %H:%M:%S"),
    }


def _parse_report_date_arg(value: str | None) -> date | None:
    raw = _clean(value)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError("report_date 必须是 YYYY-MM-DD 或 YYYYMMDD")


def _where_clause(
    keyword: str,
    brands: list[str] | None,
    categories: list[str] | None,
    stagnant: str,
) -> tuple[str, dict[str, Any], list[Any]]:
    clauses = ["batch_id = :batch_id"]
    params: dict[str, Any] = {}
    bind_params: list[Any] = []
    clean_keyword = _clean(keyword)
    if clean_keyword:
        clauses.append(
            "(product_code LIKE :keyword OR style_code LIKE :keyword "
            "OR product_name LIKE :keyword OR color_spec LIKE :keyword)"
        )
        params["keyword"] = f"%{clean_keyword}%"
    clean_brands = [item.strip() for item in (brands or []) if item and item.strip()]
    if clean_brands:
        clauses.append("brand IN :brands")
        params["brands"] = clean_brands
        bind_params.append(bindparam("brands", expanding=True))
    clean_categories = [item.strip() for item in (categories or []) if item and item.strip()]
    if clean_categories:
        clauses.append("category IN :categories")
        params["categories"] = clean_categories
        bind_params.append(bindparam("categories", expanding=True))
    if stagnant == "yes":
        clauses.append("opening_quantity = closing_quantity")
    elif stagnant == "no":
        clauses.append("opening_quantity <> closing_quantity")
    return " WHERE " + " AND ".join(clauses), params, bind_params


def list_dashboard(
    keyword: str = "",
    brands: list[str] | None = None,
    categories: list[str] | None = None,
    stagnant: str = "",
    report_date: str | None = None,
    page: int = 1,
    page_size: int = 50,
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_schema(db_engine)
    page = max(1, int(page))
    page_size = min(200, max(10, int(page_size)))
    selected_date = _parse_report_date_arg(report_date)

    with db_engine.connect() as conn:
        batches = conn.execute(
            text(
                """
                SELECT batch_id, report_date, period_start, period_end, source_name,
                       row_count, image_count, imported_by, imported_at
                FROM inventory_product_import_batch
                ORDER BY report_date DESC, imported_at DESC
                """
            )
        ).mappings().all()
        selected_batch = next(
            (row for row in batches if selected_date is not None and row["report_date"] == selected_date),
            None,
        )
        if selected_date is None and batches:
            selected_batch = batches[0]

        periods = [
            {
                "report_date": row["report_date"].isoformat(),
                "period_start": row["period_start"].isoformat(),
                "period_end": row["period_end"].isoformat(),
                "period_label": f"{row['period_start']:%Y-%m-%d} ~ {row['period_end']:%Y-%m-%d}",
                "source_name": row["source_name"],
                "row_count": int(row["row_count"] or 0),
                "image_count": int(row["image_count"] or 0),
                "imported_at": row["imported_at"].strftime("%Y-%m-%d %H:%M:%S") if row["imported_at"] else "",
            }
            for row in batches
        ]
        if selected_batch is None:
            return {
                "batch": {},
                "periods": periods,
                "summary": {
                    "row_count": 0,
                    "sku_count": 0,
                    "category_count": 0,
                    "brand_count": 0,
                    "opening_quantity": 0,
                    "closing_quantity": 0,
                    "stagnant_count": 0,
                    "image_count": 0,
                },
                "filter_options": {"brands": [], "categories": []},
                "items": [],
                "total": 0,
                "page": page,
                "page_size": page_size,
            }

        where_sql, params, bind_params = _where_clause(keyword, brands, categories, stagnant)
        params["batch_id"] = selected_batch["batch_id"]
        offset = (page - 1) * page_size

        def stmt(sql: str):
            statement = text(sql)
            return statement.bindparams(*bind_params) if bind_params else statement

        summary = conn.execute(
            stmt(
                """
                SELECT COUNT(*) AS row_count,
                       COUNT(DISTINCT product_code) AS sku_count,
                       COUNT(DISTINCT NULLIF(category, '')) AS category_count,
                       COUNT(DISTINCT NULLIF(brand, '')) AS brand_count,
                       COALESCE(SUM(opening_quantity), 0) AS opening_quantity,
                       COALESCE(SUM(closing_quantity), 0) AS closing_quantity,
                       COALESCE(SUM(CASE WHEN opening_quantity = closing_quantity THEN 1 ELSE 0 END), 0) AS stagnant_count,
                       COALESCE(SUM(CASE WHEN image_url IS NOT NULL AND image_url <> '' THEN 1 ELSE 0 END), 0) AS image_count
                FROM inventory_product_dashboard_item
                """ + where_sql
            ),
            params,
        ).mappings().one()
        total = int(summary["row_count"] or 0)
        items = conn.execute(
            stmt(
                """
                SELECT id, batch_id, source_row, image_url, category, style_code,
                       product_code, product_name, color_spec, brand,
                       opening_quantity, closing_quantity
                FROM inventory_product_dashboard_item
                """ + where_sql + " ORDER BY closing_quantity DESC, product_code LIMIT :limit OFFSET :offset"
            ),
            {**params, "limit": page_size, "offset": offset},
        ).mappings().all()
        brand_options = [
            str(value)
            for value in conn.execute(
                text(
                    "SELECT DISTINCT brand FROM inventory_product_dashboard_item "
                    "WHERE batch_id = :batch_id AND brand <> '' ORDER BY brand"
                ),
                {"batch_id": selected_batch["batch_id"]},
            ).scalars()
        ]
        category_options = [
            str(value)
            for value in conn.execute(
                text(
                    "SELECT DISTINCT category FROM inventory_product_dashboard_item "
                    "WHERE batch_id = :batch_id AND category <> '' ORDER BY category"
                ),
                {"batch_id": selected_batch["batch_id"]},
            ).scalars()
        ]

    def serialise_row(row: Any) -> dict[str, Any]:
        item = dict(row)
        item["opening_quantity"] = float(item["opening_quantity"] or 0)
        item["closing_quantity"] = float(item["closing_quantity"] or 0)
        item["is_stagnant"] = item["opening_quantity"] == item["closing_quantity"]
        image_url = _clean(item.get("image_url"))
        if image_url.startswith(("http://", "https://")):
            item["image_url"] = (
                f"/api/inventory-product-dashboard/images/{item['batch_id']}/{item['source_row']}"
            )
        return item

    batch_payload = dict(selected_batch)
    for key in ("report_date", "period_start", "period_end"):
        if batch_payload.get(key):
            batch_payload[key] = batch_payload[key].isoformat()
    if batch_payload.get("imported_at"):
        batch_payload["imported_at"] = batch_payload["imported_at"].strftime("%Y-%m-%d %H:%M:%S")
    batch_payload["period_label"] = f"{batch_payload['period_start']} ~ {batch_payload['period_end']}"

    serialised_summary = {
        "row_count": int(summary["row_count"] or 0),
        "sku_count": int(summary["sku_count"] or 0),
        "category_count": int(summary["category_count"] or 0),
        "brand_count": int(summary["brand_count"] or 0),
        "opening_quantity": float(summary["opening_quantity"] or 0),
        "closing_quantity": float(summary["closing_quantity"] or 0),
        "stagnant_count": int(summary["stagnant_count"] or 0),
        "image_count": int(summary["image_count"] or 0),
    }
    return {
        "batch": batch_payload,
        "periods": periods,
        "summary": serialised_summary,
        "filter_options": {"brands": brand_options, "categories": category_options},
        "items": [serialise_row(row) for row in items],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def _image_source_url(batch_id: str, source_row: int, engine: Engine | None = None) -> str:
    db_engine = get_engine(engine)
    ensure_schema(db_engine)
    with db_engine.connect() as conn:
        value = conn.execute(
            text(
                "SELECT image_url FROM inventory_product_dashboard_item "
                "WHERE batch_id = :batch_id AND source_row = :source_row LIMIT 1"
            ),
            {"batch_id": batch_id, "source_row": source_row},
        ).scalar()
    return _clean(value)


def _cached_image_path(batch_id: str, source_url: str) -> Path:
    cache_key = hashlib.sha256(source_url.encode("utf-8")).hexdigest()
    return UPLOAD_ROOT / batch_id / cache_key


def _is_allowed_image_url(source_url: str) -> bool:
    parsed = urlparse(source_url)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def resolve_image(
    batch_id: str,
    source_row: int,
    engine: Engine | None = None,
) -> tuple[Path, str] | None:
    if not re.fullmatch(r"[0-9a-f]{32}", batch_id.lower()):
        return None
    try:
        source_row = int(source_row)
    except (TypeError, ValueError):
        return None
    if source_row < 1:
        return None

    source_url = _image_source_url(batch_id, source_row, engine)
    if not _is_allowed_image_url(source_url):
        return None

    cache_path = _cached_image_path(batch_id, source_url)
    metadata_path = cache_path.with_suffix(".mime")
    if metadata_path.is_file():
        media_type = metadata_path.read_text(encoding="ascii").strip() or "image/*"
        final_path = cache_path.with_suffix(IMAGE_EXTENSIONS.get(media_type, ".bin"))
        if final_path.is_file():
            return final_path, media_type

    try:
        with httpx.Client(
            follow_redirects=True,
            timeout=httpx.Timeout(IMAGE_TIMEOUT_SECONDS),
            headers={
                "Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8",
                "User-Agent": "EnterpriseAgent/1.0 image-proxy",
            },
        ) as client:
            response = client.get(source_url)
            response.raise_for_status()
            media_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            if not media_type.startswith("image/") or len(response.content) > IMAGE_MAX_BYTES:
                return None
    except (httpx.HTTPError, OSError):
        return None

    final_path = cache_path.with_suffix(IMAGE_EXTENSIONS.get(media_type, ".bin"))
    final_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = final_path.with_suffix(final_path.suffix + ".tmp")
    try:
        temp_path.write_bytes(response.content)
        temp_path.replace(final_path)
        metadata_path.write_text(media_type, encoding="ascii")
    except OSError:
        temp_path.unlink(missing_ok=True)
        return None
    return final_path, media_type
