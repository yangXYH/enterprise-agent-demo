"""Inventory import, risk scoring, and transfer decision engine.

This module intentionally keeps LLM usage out of inventory math.  It owns the
deterministic data pipeline:

Excel wide table -> inv_stock_snapshot -> inv_risk_score -> inv_transfer_plan
"""

from __future__ import annotations

import logging
import hashlib
import json
import math
import os
import re
import time
import uuid
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import Engine, bindparam, text

from utils import import_service
import requests


LOGGER = logging.getLogger(__name__)

XML_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

TARGET_DAYS = 30.0
DONOR_RESERVE_DAYS = 15.0
RISK_MID_DAYS = 40.0
MAX_IMPORT_ERRORS_PER_BATCH = 5000
HISTORY_LOOKBACK_DAYS = 7
INVENTORY_HISTORY_KEEP_DAYS = 3
INBOUND_ETA_DAYS = 7
SUPPLIER_TO_HUB_ETA_DAYS = 7
SUPPLY_HUB_WAREHOUSE = "花东仓"
NANSHA_SUPPLY_WAREHOUSE = "南沙仓"
SUPPLIER_WAREHOUSE = "供应商"
FOSHAN_SUPPLIER_WAREHOUSE = "供应商（佛山）"

HUADONG_ROUTE_ETA_DAYS = {
    "广州": 2,
    "上海": 4,
    "北京": 4,
    "成都": 5,
    "武汉": 6,
    "西安": 5,
    "德州": 5,
    "沈阳": 6,
}

BOX_CATEGORY = "箱包"
WAREHOUSE_TYPE_SUPPLY_HUB = "OWN_HUB"
WAREHOUSE_TYPE_JD = "JD_WAREHOUSE"
WAREHOUSE_TYPE_SUPPLIER = "SUPPLIER"
SOURCE_TYPE_JD_TRANSFER = "JD_TRANSFER"
SOURCE_TYPE_HUB_TRANSFER = "HUB_TRANSFER"
SOURCE_TYPE_SUPPLIER_REPLENISH = "SUPPLIER_REPLENISH"
SOURCE_TYPE_SUPPLIER_DIRECT = "SUPPLIER_DIRECT"


@dataclass(frozen=True)
class JushuitanConfig:
    endpoint: str
    app_key: str
    app_secret: str
    access_token: str
    refresh_token: str | None = None

BASE_COLUMNS = {
    "snapshot_date": "时间",
    "product_name": "商品名称",
    "sku_code": "SKU",
    "brand": "品牌",
    "supplier_name": "供应商名称",
    "shop_name": "店铺名称",
    "barcode": "条形码",
    "product_status": "上下柜状态",
    "gift_flag": "是否赠品",
}

SUPPLIER_NAME_COLUMN_ALIASES = ("供应商名称", "供应商")

METRIC_SUFFIXES = {
    "stock_qty": ("库存件数",),
    "inbound_qty": ("采购未到货",),
    "sales_1d": ("昨日出库商品件数",),
    "sales_7d": ("近7日出库商品件数",),
    "sales_14d": ("近14日出库商品件数",),
    "sales_30d": ("近30日出库商品件数",),
    "available_qty": ("可用库存",),
    "orderable_qty": ("可订购库存", "可订购"),
}
REQUIRED_METRIC_FIELDS = ("stock_qty", "inbound_qty", "sales_1d", "sales_7d", "sales_30d")

NATIONAL_PREFIX = "全国"
MONITOR_WAREHOUSES = ("北京", "上海", "广州", "成都", "武汉", "沈阳", "西安", "德州")
MONITOR_METRICS = (
    ("stock_qty", "现货库存"),
    ("orderable_qty", "可订购库存"),
    ("sales_1d", "昨日出库商品件数"),
    ("sales_7d", "近7日出库商品件数"),
    ("sales_30d", "近30日出库商品件数"),
)
ACTIVE_STATUSES = {"上柜", "可上柜"}
MONITOR_ON_SHELF_STATUSES = {"上柜", "可上柜", "上架", "可上架", "在售"}
MONITOR_OFF_SHELF_STATUSES = {"下柜", "不可上柜", "下架", "不可上架", "停售", "下线"}
MONITOR_METRIC_LABELS = {
    "stockout": "断货",
    "shortage": "缺货",
    "listed": "上柜",
    "delisted": "下柜",
}
ENTERPRISE_PURCHASE_SUPPLIER_NAME = "广州示例户外网络科技有限公司"
SAMSONITE_BRAND = "测试旅行"
SAMSONITE_ENTERPRISE_BRAND = "测试旅行企业购"
AMERICAN_TOURISTER_BRAND = "美旅（AMERICAN TOURISTER）"
AMERICAN_TOURISTER_ENTERPRISE_BRAND = "美旅企业购"

# iGear 品牌按一级类目拆分：工业品 -> iGear工业品；家装建材 -> iGear；其余类目不导入。
# 用于区分全品牌全集表(表1)与 iGear/五金专表(表3)导入后的两套 iGear 商品。
IGEAR_BRAND = "iGear"
IGEAR_INDUSTRIAL_BRAND = "iGear工业品"
IGEAR_CATEGORY_INDUSTRIAL = "工业品"
IGEAR_CATEGORY_BUILDING_MATERIAL = "家装建材"
IGEAR_CATEGORY_COLUMN = "一级类目"

BRAND_CATEGORY_MAP = {
    "campingmoon": "户外",
    "demo_bags": "箱包",
    "demo_apparel": "箱包",
    "demo_travel": "箱包",
    "demo_life": "五金",
    "demo_home": "咖啡机",
    "demo_rain": "箱包",
    "venque": "箱包",
    "京塑": "箱包",
    "测试运动(aonijie)": "户外",
    "测试旅行": "箱包",
    "测试旅行企业购": "箱包",
    "美旅(americantourister)": "箱包",
    "美旅企业购": "箱包",
    "测试家居": "户外",
    "联想(lenovo)": "箱包",
    "金宇(jinyu)": "户外",
    "测试五金": "五金",
    "高山(demo_outdoor)": "箱包",
    "魅族(meizu)": "箱包",
}


REGION_BY_CITY = {
    "上海": "华东",
    "南京": "华东",
    "苏州": "华东",
    "杭州": "华东",
    "宁波": "华东",
    "金华": "华东",
    "温州": "华东",
    "福州": "华东",
    "厦门": "华东",
    "合肥": "华东",
    "南昌": "华东",
    "济南": "华东",
    "青岛": "华东",
    "临沂": "华东",
    "德州": "华东",
    "潍坊": "华东",
    "济宁": "华东",
    "徐州": "华东",
    "北京": "华北",
    "天津": "华北",
    "石家庄": "华北",
    "唐山": "华北",
    "沧州": "华北",
    "邢台": "华北",
    "太原": "华北",
    "大同": "华北",
    "呼和浩特": "华北",
    "郑州": "华中",
    "武汉": "华中",
    "长沙": "华中",
    "衡阳": "华中",
    "襄阳": "华中",
    "荆州": "华中",
    "驻马店": "华中",
    "鹤壁": "华中",
    "广州": "华南",
    "佛山": "华南",
    "中山": "华南",
    "茂名": "华南",
    "南宁": "华南",
    "柳州": "华南",
    "海口": "华南",
    "成都": "西南",
    "重庆": "西南",
    "昆明": "西南",
    "贵阳": "西南",
    "宜宾": "西南",
    "拉萨": "西南",
    "西安": "西北",
    "兰州": "西北",
    "西宁": "西北",
    "银川": "西北",
    "汉中": "西北",
    "乌鲁木齐": "西北",
    "沈阳": "东北",
    "大连": "东北",
    "长春": "东北",
    "哈尔滨": "东北",
    "宿迁": "华东",
    "阜阳": "华东",
    "赣州": "华东",
    "泰州": "华东",
}

PROVINCE_BASE_CITY = {
    "安徽": "合肥",
    "北京": "北京",
    "重庆": "重庆",
    "福建": "福州",
    "甘肃": "兰州",
    "广东": "广州",
    "广西": "南宁",
    "贵州": "贵阳",
    "海南": "海口",
    "河北": "石家庄",
    "河南": "郑州",
    "黑龙江": "哈尔滨",
    "湖北": "武汉",
    "湖南": "长沙",
    "吉林": "长春",
    "江苏": "南京",
    "江西": "南昌",
    "辽宁": "沈阳",
    "内蒙古": "呼和浩特",
    "宁夏": "银川",
    "青海": "西宁",
    "山东": "济南",
    "山西": "太原",
    "陕西": "西安",
    "上海": "上海",
    "四川": "成都",
    "天津": "天津",
    "西藏": "拉萨",
    "新疆": "乌鲁木齐",
    "云南": "昆明",
    "浙江": "杭州",
}

CITY_COORDINATES = {
    "上海": (31.2304, 121.4737),
    "成都": (30.5728, 104.0668),
    "武汉": (30.5928, 114.3055),
    "北京": (39.9042, 116.4074),
    "南京": (32.0603, 118.7969),
    "沈阳": (41.8057, 123.4315),
    "广州": (23.1291, 113.2644),
    "阜阳": (32.8901, 115.8142),
    "临沂": (35.1047, 118.3564),
    "宜宾": (28.7513, 104.6417),
    "西安": (34.3416, 108.9398),
    "福州": (26.0745, 119.2965),
    "沧州": (38.3044, 116.8387),
    "杭州": (30.2741, 120.1551),
    "重庆": (29.5630, 106.5516),
    "宁波": (29.8683, 121.5440),
    "郑州": (34.7466, 113.6254),
    "厦门": (24.4798, 118.0894),
    "青岛": (36.0671, 120.3826),
    "石家庄": (38.0428, 114.5149),
    "太原": (37.8706, 112.5489),
    "南宁": (22.8170, 108.3669),
    "哈尔滨": (45.8038, 126.5349),
    "大连": (38.9140, 121.6147),
    "长沙": (28.2282, 112.9388),
    "兰州": (36.0611, 103.8343),
    "乌鲁木齐": (43.8256, 87.6168),
    "合肥": (31.8206, 117.2272),
    "宿迁": (33.9630, 118.2752),
    "昆明": (24.8801, 102.8329),
    "贵阳": (26.6470, 106.6302),
    "南昌": (28.6829, 115.8582),
    "长春": (43.8171, 125.3235),
    "襄阳": (32.0089, 112.1224),
    "金华": (29.0791, 119.6474),
    "海口": (20.0442, 110.1999),
    "银川": (38.4872, 106.2309),
    "驻马店": (32.9802, 114.0259),
    "潍坊": (36.7068, 119.1618),
    "唐山": (39.6305, 118.1802),
    "呼和浩特": (40.8426, 111.7492),
    "温州": (27.9938, 120.6994),
    "茂名": (21.6630, 110.9255),
    "柳州": (24.3264, 109.4281),
    "济宁": (35.4149, 116.5872),
    "衡阳": (26.8932, 112.5719),
    "德州": (37.4355, 116.3592),
    "荆州": (30.3352, 112.2397),
    "拉萨": (29.6500, 91.1000),
    "大同": (40.0768, 113.3001),
    "泰州": (32.4558, 119.9231),
    "鹤壁": (35.7470, 114.2975),
    "西宁": (36.6171, 101.7782),
    "赣州": (25.8311, 114.9348),
    "中山": (22.5176, 113.3928),
    "汉中": (33.0675, 107.0238),
    "邢台": (37.0706, 114.5048),
    "苏州": (31.2989, 120.5853),
    "济南": (36.6512, 117.1201),
    "佛山": (23.0215, 113.1214),
    "徐州": (34.2044, 117.2858),
}


@dataclass(frozen=True)
class ImportBatchResult:
    batch_no: str
    file_name: str
    snapshot_date: str | None
    rows_read: int
    sku_rows: int
    snapshot_rows: int
    warehouses: int
    warnings_count: int


@dataclass(frozen=True)
class ReplenishmentPolicy:
    policy_type: str
    brand: str
    category: str
    target_days: float
    reserve_days: float
    supplier_eta_days: int
    allow_jd_transfer: bool
    preferred_source_warehouse: str


SKU_CODE_TABLES = (
    "inv_sku",
    "inv_stock_snapshot",
    "inv_risk_score",
    "inv_transfer_plan",
    "inv_import_error",
)


def get_engine(engine: Engine | None = None) -> Engine:
    return engine or import_service.get_engine()


def ensure_inventory_schema(engine: Engine | None = None) -> None:
    """Create inventory tables with table and column comments."""

    db_engine = get_engine(engine)
    ddl_statements = [
        """
        CREATE TABLE IF NOT EXISTS inv_sku (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU唯一编码',
            product_name VARCHAR(512) NOT NULL DEFAULT '' COMMENT '商品名称',
            brand VARCHAR(128) NOT NULL DEFAULT '' COMMENT '品牌',
            category VARCHAR(64) NOT NULL DEFAULT '' COMMENT '商品类目',
            shop_name VARCHAR(255) NOT NULL DEFAULT '' COMMENT '店铺名称',
            supplier_name VARCHAR(255) NOT NULL DEFAULT '' COMMENT '供应商名称',
            barcode VARCHAR(64) NOT NULL DEFAULT '' COMMENT '条形码/聚水潭库存查询编码',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_sku_code (sku_code),
            KEY idx_inv_sku_brand (brand),
            KEY idx_inv_sku_category (category),
            KEY idx_inv_sku_shop (shop_name),
            KEY idx_inv_sku_supplier (supplier_name),
            KEY idx_inv_sku_barcode (barcode)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存SKU基础表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_warehouse (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            warehouse_code VARCHAR(128) NOT NULL COMMENT '仓库编码，V1使用Excel城市/仓库名前缀',
            warehouse_name VARCHAR(128) NOT NULL COMMENT '仓库/城市名称',
            region VARCHAR(32) NOT NULL DEFAULT '其他' COMMENT '大区，例如华东、华北、华南',
            warehouse_type VARCHAR(32) NOT NULL DEFAULT 'JD_WAREHOUSE' COMMENT '仓库类型 OWN_HUB/JD_WAREHOUSE/SUPPLIER',
            is_supply_hub TINYINT NOT NULL DEFAULT 0 COMMENT '是否自有供给仓，1是0否',
            is_jd_warehouse TINYINT NOT NULL DEFAULT 1 COMMENT '是否京东仓/外部仓，1是0否',
            base_city VARCHAR(64) NOT NULL DEFAULT '' COMMENT '用于距离估算的基础城市名',
            latitude DECIMAL(10,6) NULL COMMENT '基础城市纬度',
            longitude DECIMAL(10,6) NULL COMMENT '基础城市经度',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_warehouse_code (warehouse_code),
            KEY idx_inv_warehouse_region (region),
            KEY idx_inv_warehouse_type (warehouse_type)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存城市/仓库表，城市即仓库'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_import_batch (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            batch_no VARCHAR(64) NOT NULL COMMENT '导入批次号',
            file_name VARCHAR(512) NOT NULL DEFAULT '' COMMENT '导入文件名',
            file_size BIGINT NOT NULL DEFAULT 0 COMMENT '文件大小字节数',
            snapshot_date DATE NULL COMMENT '快照日期',
            status VARCHAR(32) NOT NULL DEFAULT 'RUNNING' COMMENT '批次状态 RUNNING/SUCCESS/FAILED',
            rows_read INT NOT NULL DEFAULT 0 COMMENT 'Excel读取商品行数',
            sku_rows INT NOT NULL DEFAULT 0 COMMENT '入库SKU数量',
            snapshot_rows INT NOT NULL DEFAULT 0 COMMENT '入库SKU×仓库快照数量',
            warnings_count INT NOT NULL DEFAULT 0 COMMENT '数据校验警告数量',
            error_message TEXT NULL COMMENT '失败错误信息',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            finished_at DATETIME NULL COMMENT '结束时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_import_batch_no (batch_no),
            KEY idx_inv_import_batch_date (snapshot_date)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存Excel导入批次表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_import_error (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            batch_no VARCHAR(64) NOT NULL COMMENT '导入批次号',
            snapshot_date DATE NULL COMMENT '快照日期',
            sku_code VARCHAR(64) NOT NULL DEFAULT '' COMMENT 'SKU编码',
            warehouse_code VARCHAR(128) NOT NULL DEFAULT '' COMMENT '仓库编码',
            error_type VARCHAR(64) NOT NULL COMMENT '异常类型',
            severity VARCHAR(16) NOT NULL DEFAULT 'WARN' COMMENT '严重程度 WARN/ERROR',
            message VARCHAR(1000) NOT NULL DEFAULT '' COMMENT '异常说明',
            raw_value VARCHAR(1000) NOT NULL DEFAULT '' COMMENT '原始值或校验上下文',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            PRIMARY KEY (id),
            KEY idx_inv_import_error_batch (batch_no),
            KEY idx_inv_import_error_type (error_type),
            KEY idx_inv_import_error_sku (sku_code)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存导入数据校验异常表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_stock_snapshot (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            batch_no VARCHAR(64) NOT NULL COMMENT '导入批次号',
            snapshot_date DATE NOT NULL COMMENT '库存快照日期',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU编码',
            warehouse_code VARCHAR(128) NOT NULL COMMENT '仓库编码',
            product_status VARCHAR(32) NOT NULL DEFAULT '' COMMENT '商品上下柜状态',
            is_gift TINYINT NOT NULL DEFAULT 0 COMMENT '是否赠品，1是0否',
            is_decision_active TINYINT NOT NULL DEFAULT 1 COMMENT '是否纳入主决策，1是0否',
            stock_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '库存件数',
            inbound_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '采购未到货件数',
            available_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '可用库存件数',
            orderable_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '可订购库存件数（运营表40周转口径）',
            sales_1d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '昨日出库商品件数',
            sales_7d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '近7日出库商品件数',
            sales_14d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '近14日出库商品件数',
            sales_30d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '近30日出库商品件数',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_snapshot_date_sku_wh (snapshot_date, sku_code, warehouse_code),
            KEY idx_inv_snapshot_sku_date (sku_code, snapshot_date),
            KEY idx_inv_snapshot_wh_date (warehouse_code, snapshot_date),
            KEY idx_inv_snapshot_active_date (is_decision_active, snapshot_date)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存与销量快照核心表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_monitor_snapshot (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            snapshot_date DATE NOT NULL COMMENT '快照日期',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU',
            product_name VARCHAR(512) NOT NULL DEFAULT '' COMMENT '商品名称',
            brand VARCHAR(128) NOT NULL DEFAULT '' COMMENT '品牌',
            category_l1 VARCHAR(128) NOT NULL DEFAULT '' COMMENT '一级类目',
            category_l2 VARCHAR(128) NOT NULL DEFAULT '' COMMENT '二级类目',
            category_l3 VARCHAR(128) NOT NULL DEFAULT '' COMMENT '三级类目',
            shop_name VARCHAR(255) NOT NULL DEFAULT '' COMMENT '店铺名称',
            product_status VARCHAR(32) NOT NULL DEFAULT '' COMMENT '商品上下柜状态',
            national_stock_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '全国现货库存',
            national_orderable_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '全国可订购库存',
            national_sales_1d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '全国昨日出库商品件数',
            national_sales_7d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '全国近7日出库商品件数',
            national_sales_30d DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '全国近30日出库商品件数',
            city_metrics JSON NOT NULL COMMENT '八大仓库存及出库指标',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_monitor_date_sku (snapshot_date, sku_code),
            KEY idx_inv_monitor_date (snapshot_date),
            KEY idx_inv_monitor_brand (brand),
            KEY idx_inv_monitor_shop (shop_name)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存看板监控快照表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_decision_run (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            run_no VARCHAR(64) NOT NULL COMMENT '计算任务编号',
            snapshot_date DATE NOT NULL COMMENT '计算使用的快照日期',
            status VARCHAR(32) NOT NULL DEFAULT 'RUNNING' COMMENT '计算状态 RUNNING/SUCCESS/FAILED',
            risk_count INT NOT NULL DEFAULT 0 COMMENT '风险评分数量',
            plan_count INT NOT NULL DEFAULT 0 COMMENT '调拨方案数量',
            high_count INT NOT NULL DEFAULT 0 COMMENT '高风险数量',
            mid_count INT NOT NULL DEFAULT 0 COMMENT '中风险数量',
            low_count INT NOT NULL DEFAULT 0 COMMENT '低风险数量',
            message TEXT NULL COMMENT '任务说明或失败原因',
            started_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '开始时间',
            finished_at DATETIME NULL COMMENT '结束时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_decision_run_no (run_no),
            KEY idx_inv_decision_run_date (snapshot_date)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存决策计算任务表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_risk_score (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            run_id BIGINT NOT NULL COMMENT '决策任务ID',
            snapshot_date DATE NOT NULL COMMENT '快照日期',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU编码',
            warehouse_code VARCHAR(128) NOT NULL COMMENT '仓库编码',
            daily_sales DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '日均销量，V1按近7日出库/7计算',
            target_days DECIMAL(10,2) NOT NULL DEFAULT 30 COMMENT '本次计算采用的目标库存周转天数',
            days_of_supply DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '可售天数',
            demand_30d DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '未来30天需求预测',
            shortage_days DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '未来30天内预计缺货天数',
            risk_level VARCHAR(16) NOT NULL DEFAULT 'LOW' COMMENT '风险等级 HIGH/MID/LOW',
            excess_qty DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '富余库存，库存-30天需求',
            need_qty DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '补足30天目标需要数量',
            available_qty DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT 'V2有效可用库存，库存+可信采购未到货',
            inbound_credit_qty DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT 'V2可信采购未到货抵扣件数',
            inbound_credit_rate DECIMAL(10,4) NOT NULL DEFAULT 0 COMMENT 'V2采购未到货可信抵扣比例',
            trend_factor DECIMAL(10,4) NOT NULL DEFAULT 1 COMMENT 'V2历史趋势修正系数',
            volatility_factor DECIMAL(10,4) NOT NULL DEFAULT 1 COMMENT 'V2销量波动修正系数',
            confidence_score DECIMAL(10,4) NOT NULL DEFAULT 0 COMMENT 'V2趋势可信度0到1',
            consecutive_high_days INT NOT NULL DEFAULT 0 COMMENT 'V2连续高风险天数',
            is_shortage_risk TINYINT NOT NULL DEFAULT 0 COMMENT '是否缺货风险',
            is_excess_risk TINYINT NOT NULL DEFAULT 0 COMMENT '是否积压/可调出风险',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_risk_date_sku_wh (snapshot_date, sku_code, warehouse_code),
            KEY idx_inv_risk_run (run_id),
            KEY idx_inv_risk_level (snapshot_date, risk_level),
            KEY idx_inv_risk_sku (sku_code)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存风险评分计算结果表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_transfer_route (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            from_warehouse VARCHAR(128) NOT NULL COMMENT '调出仓库编码',
            to_warehouse VARCHAR(128) NOT NULL COMMENT '调入仓库编码',
            distance_km DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '估算城市距离公里',
            eta_days INT NOT NULL DEFAULT 1 COMMENT '调拨预计到达天数',
            route_source VARCHAR(32) NOT NULL DEFAULT 'ESTIMATED' COMMENT '时效来源 ESTIMATED/MANUAL',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_route_pair (from_warehouse, to_warehouse)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存调拨城市时效估算表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_replenishment_policy (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            policy_type VARCHAR(32) NOT NULL DEFAULT 'CATEGORY' COMMENT '策略类型 SYSTEM/CATEGORY/BRAND',
            category VARCHAR(64) NOT NULL DEFAULT '' COMMENT '适用类目',
            brand VARCHAR(128) NOT NULL DEFAULT '' COMMENT '适用品牌',
            target_days DECIMAL(10,2) NOT NULL DEFAULT 30 COMMENT '目标库存周转天数',
            reserve_days DECIMAL(10,2) NOT NULL DEFAULT 15 COMMENT '调出仓保留安全天数',
            supplier_eta_days INT NOT NULL DEFAULT 7 COMMENT '供应商到花东暂定到货天数',
            allow_jd_transfer TINYINT NOT NULL DEFAULT 1 COMMENT '是否允许京东仓之间互调，1是0否',
            preferred_source_warehouse VARCHAR(128) NOT NULL DEFAULT '花东仓' COMMENT '默认供给仓',
            remark VARCHAR(512) NOT NULL DEFAULT '' COMMENT '策略备注',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_policy_scope (policy_type, category, brand),
            KEY idx_inv_policy_brand (brand),
            KEY idx_inv_policy_category (category)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存补货策略表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_supply_stock_snapshot (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            snapshot_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '供给仓库存快照时间',
            snapshot_date DATE NOT NULL COMMENT '供给仓库存快照日期',
            warehouse_code VARCHAR(128) NOT NULL DEFAULT '花东仓' COMMENT '供给仓编码',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU编码',
            stock_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '供给仓库存件数',
            available_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '供给仓可用库存件数',
            inbound_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '供给仓在途/采购未到货件数',
            source VARCHAR(64) NOT NULL DEFAULT 'JUSHUITAN' COMMENT '数据来源',
            raw_payload JSON NULL COMMENT '原始数据',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_supply_stock_time (warehouse_code, sku_code, snapshot_time),
            KEY idx_inv_supply_stock_latest (warehouse_code, sku_code, snapshot_date)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='花东等自有供给仓库存快照表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_huadong_sync_setting (
            id TINYINT NOT NULL DEFAULT 1 COMMENT '固定主键，花东库存同步配置仅一行',
            enabled TINYINT NOT NULL DEFAULT 0 COMMENT '是否启用自动同步，1启用0停用',
            interval_minutes INT NOT NULL DEFAULT 10 COMMENT '自动拉取聚水潭花东库存的分钟间隔',
            recalculate TINYINT NOT NULL DEFAULT 0 COMMENT '同步后是否自动重算库存决策，1是0否',
            last_run_at DATETIME NULL COMMENT '最近一次自动同步时间',
            last_status VARCHAR(32) NOT NULL DEFAULT 'IDLE' COMMENT '最近同步状态 IDLE/RUNNING/SUCCESS/FAILED',
            last_message VARCHAR(1000) NOT NULL DEFAULT '' COMMENT '最近同步结果或失败原因',
            last_result JSON NULL COMMENT '最近同步结构化结果',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='花东聚水潭库存自动同步配置表'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_sku_jst_lookup (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT 'Primary key',
            sku_code VARCHAR(64) NOT NULL COMMENT 'Internal SKU code',
            lookup_code VARCHAR(128) NOT NULL COMMENT 'Jushuitan inventory query code',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT 'Created at',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT 'Updated at',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_sku_jst_lookup_sku (sku_code),
            KEY idx_inv_sku_jst_lookup_code (lookup_code)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='SKU to Jushuitan lookup code mapping'
        """,
        """
        CREATE TABLE IF NOT EXISTS inv_transfer_plan (
            id BIGINT NOT NULL AUTO_INCREMENT COMMENT '主键ID',
            run_id BIGINT NOT NULL COMMENT '决策任务ID',
            plan_date DATE NOT NULL COMMENT '方案日期',
            sku_code VARCHAR(64) NOT NULL COMMENT 'SKU编码',
            from_warehouse VARCHAR(128) NOT NULL COMMENT '调出仓库编码',
            to_warehouse VARCHAR(128) NOT NULL COMMENT '调入仓库编码',
            transfer_qty DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '建议调拨数量',
            eta_days INT NOT NULL DEFAULT 1 COMMENT '调拨预计到达时间，单位天',
            priority VARCHAR(16) NOT NULL DEFAULT 'LOW' COMMENT '优先级 HIGH/MID/LOW',
            risk_level VARCHAR(16) NOT NULL DEFAULT 'LOW' COMMENT '调入仓风险等级',
            risk_days DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '调入仓当前可售天数',
            surplus_days_after DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '调出仓调拨后可售天数',
            eta_late TINYINT NOT NULL DEFAULT 0 COMMENT '预计到货是否晚于调入仓当前可售天数，1是0否',
            source_type VARCHAR(32) NOT NULL DEFAULT 'JD_TRANSFER' COMMENT '建议来源类型 JD_TRANSFER/HUB_TRANSFER/SUPPLIER_REPLENISH/SUPPLIER_DIRECT',
            route_path VARCHAR(255) NOT NULL DEFAULT '' COMMENT '调拨/补货路径展示文案',
            target_days DECIMAL(10,2) NOT NULL DEFAULT 30 COMMENT '本方案采用的目标库存周转天数',
            reserve_days DECIMAL(10,2) NOT NULL DEFAULT 15 COMMENT '本方案采用的调出仓保留天数',
            supplier_eta_days INT NOT NULL DEFAULT 0 COMMENT '供应商到供给仓预计天数',
            hub_eta_days INT NOT NULL DEFAULT 0 COMMENT '供给仓到调入仓预计天数',
            score DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT '调拨匹配评分',
            reason TEXT NULL COMMENT 'AI或规则生成的中文调拨解释',
            status VARCHAR(32) NOT NULL DEFAULT 'PENDING' COMMENT '方案状态 PENDING/CONFIRMED/EXECUTED/CANCELLED',
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
            updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
            PRIMARY KEY (id),
            UNIQUE KEY uk_inv_plan_run_pair (run_id, sku_code, from_warehouse, to_warehouse),
            KEY idx_inv_plan_date_priority (plan_date, priority),
            KEY idx_inv_plan_sku (sku_code),
            KEY idx_inv_plan_to_wh (to_warehouse),
            KEY idx_inv_plan_status (status)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='库存跨城市调拨方案表'
        """,
    ]

    with db_engine.begin() as conn:
        for ddl in ddl_statements:
            conn.execute(text(ddl))
        _ensure_inventory_columns(conn)


def _ensure_inventory_columns(conn: Any) -> None:
    """Apply additive migrations for inventory tables created by earlier builds."""

    database_name = conn.execute(text("SELECT DATABASE()")).scalar()
    if not database_name:
        return

    def has_column(table: str, column: str) -> bool:
        return bool(
            conn.execute(
                text(
                    """
                    SELECT COUNT(*)
                    FROM information_schema.COLUMNS
                    WHERE TABLE_SCHEMA=:schema
                      AND TABLE_NAME=:table_name
                      AND COLUMN_NAME=:column_name
                    """
                ),
                {"schema": database_name, "table_name": table, "column_name": column},
            ).scalar()
        )

    def has_index(table: str, index_name: str) -> bool:
        return bool(
            conn.execute(
                text(
                    """
                    SELECT COUNT(*)
                    FROM information_schema.STATISTICS
                    WHERE TABLE_SCHEMA=:schema
                      AND TABLE_NAME=:table_name
                      AND INDEX_NAME=:index_name
                    """
                ),
                {"schema": database_name, "table_name": table, "index_name": index_name},
            ).scalar()
        )

    if not has_column("inv_sku", "category"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_sku
                ADD COLUMN category VARCHAR(64) NOT NULL DEFAULT ''
                COMMENT '商品类目'
                AFTER brand
                """
            )
        )

    if not has_index("inv_sku", "idx_inv_sku_category"):
        conn.execute(text("ALTER TABLE inv_sku ADD INDEX idx_inv_sku_category (category)"))

    if not has_column("inv_sku", "barcode"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_sku
                ADD COLUMN barcode VARCHAR(64) NOT NULL DEFAULT ''
                COMMENT '条形码/聚水潭库存查询编码'
                AFTER shop_name
                """
            )
        )

    if not has_column("inv_sku", "supplier_name"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_sku
                ADD COLUMN supplier_name VARCHAR(255) NOT NULL DEFAULT ''
                COMMENT '供应商名称'
                AFTER shop_name
                """
            )
        )

    if not has_index("inv_sku", "idx_inv_sku_supplier"):
        conn.execute(text("ALTER TABLE inv_sku ADD INDEX idx_inv_sku_supplier (supplier_name)"))

    if not has_index("inv_sku", "idx_inv_sku_barcode"):
        conn.execute(text("ALTER TABLE inv_sku ADD INDEX idx_inv_sku_barcode (barcode)"))

    if not has_column("inv_sku", "case_pack"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_sku
                ADD COLUMN case_pack DECIMAL(12,2) NOT NULL DEFAULT 0
                COMMENT '箱规/整箱件数（iGear补货按箱取整用，0=未知）'
                AFTER barcode
                """
            )
        )

    if not has_column("inv_stock_snapshot", "sales_14d"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_stock_snapshot
                ADD COLUMN sales_14d DECIMAL(18,2) NOT NULL DEFAULT 0
                COMMENT '近14日出库商品件数'
                AFTER sales_7d
                """
            )
        )

    snapshot_orderable_columns = [
        (
            "available_qty",
            "DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '可用库存件数'",
            "inbound_qty",
        ),
        (
            "orderable_qty",
            "DECIMAL(18,2) NOT NULL DEFAULT 0 COMMENT '可订购库存件数（运营表40周转口径）'",
            "available_qty",
        ),
    ]
    for column, definition, after_column in snapshot_orderable_columns:
        if not has_column("inv_stock_snapshot", column):
            conn.execute(
                text(
                    f"""
                    ALTER TABLE inv_stock_snapshot
                    ADD COLUMN {column} {definition}
                    AFTER {after_column}
                    """
                )
            )

    warehouse_columns = [
        ("warehouse_type", "VARCHAR(32) NOT NULL DEFAULT 'JD_WAREHOUSE' COMMENT '仓库类型 OWN_HUB/JD_WAREHOUSE/SUPPLIER'", "region"),
        ("is_supply_hub", "TINYINT NOT NULL DEFAULT 0 COMMENT '是否自有供给仓，1是0否'", "warehouse_type"),
        ("is_jd_warehouse", "TINYINT NOT NULL DEFAULT 1 COMMENT '是否京东仓/外部仓，1是0否'", "is_supply_hub"),
    ]
    for column, definition, after_column in warehouse_columns:
        if not has_column("inv_warehouse", column):
            conn.execute(
                text(
                    f"""
                    ALTER TABLE inv_warehouse
                    ADD COLUMN {column} {definition}
                    AFTER {after_column}
                    """
                )
            )

    if not has_index("inv_warehouse", "idx_inv_warehouse_type"):
        conn.execute(text("ALTER TABLE inv_warehouse ADD INDEX idx_inv_warehouse_type (warehouse_type)"))

    if not has_column("inv_transfer_plan", "eta_late"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_transfer_plan
                ADD COLUMN eta_late TINYINT NOT NULL DEFAULT 0
                COMMENT '预计到货是否晚于调入仓当前可售天数，1是0否'
                AFTER surplus_days_after
                """
            )
        )
    if not has_column("inv_monitor_snapshot", "product_status"):
        conn.execute(
            text(
                """
                ALTER TABLE inv_monitor_snapshot
                ADD COLUMN product_status VARCHAR(32) NOT NULL DEFAULT ''
                COMMENT '商品上下柜状态'
                AFTER shop_name
                """
            )
        )
    transfer_columns = [
        ("source_type", "VARCHAR(32) NOT NULL DEFAULT 'JD_TRANSFER' COMMENT '建议来源类型 JD_TRANSFER/HUB_TRANSFER/SUPPLIER_REPLENISH/SUPPLIER_DIRECT'", "eta_late"),
        ("route_path", "VARCHAR(255) NOT NULL DEFAULT '' COMMENT '调拨/补货路径展示文案'", "source_type"),
        ("target_days", "DECIMAL(10,2) NOT NULL DEFAULT 30 COMMENT '本方案采用的目标库存周转天数'", "route_path"),
        ("reserve_days", "DECIMAL(10,2) NOT NULL DEFAULT 15 COMMENT '本方案采用的调出仓保留天数'", "target_days"),
        ("supplier_eta_days", "INT NOT NULL DEFAULT 0 COMMENT '供应商到供给仓预计天数'", "reserve_days"),
        ("hub_eta_days", "INT NOT NULL DEFAULT 0 COMMENT '供给仓到调入仓预计天数'", "supplier_eta_days"),
    ]
    for column, definition, after_column in transfer_columns:
        if not has_column("inv_transfer_plan", column):
            conn.execute(
                text(
                    f"""
                    ALTER TABLE inv_transfer_plan
                    ADD COLUMN {column} {definition}
                    AFTER {after_column}
                    """
                )
            )

    risk_columns = [
        ("target_days", "DECIMAL(10,2) NOT NULL DEFAULT 30 COMMENT '本次计算采用的目标库存周转天数'", "daily_sales"),
        ("available_qty", "DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT 'V2有效可用库存，库存+可信采购未到货'", "need_qty"),
        ("inbound_credit_qty", "DECIMAL(18,4) NOT NULL DEFAULT 0 COMMENT 'V2可信采购未到货抵扣件数'", "available_qty"),
        ("inbound_credit_rate", "DECIMAL(10,4) NOT NULL DEFAULT 0 COMMENT 'V2采购未到货可信抵扣比例'", "inbound_credit_qty"),
        ("trend_factor", "DECIMAL(10,4) NOT NULL DEFAULT 1 COMMENT 'V2历史趋势修正系数'", "inbound_credit_rate"),
        ("volatility_factor", "DECIMAL(10,4) NOT NULL DEFAULT 1 COMMENT 'V2销量波动修正系数'", "trend_factor"),
        ("confidence_score", "DECIMAL(10,4) NOT NULL DEFAULT 0 COMMENT 'V2趋势可信度0到1'", "volatility_factor"),
        ("consecutive_high_days", "INT NOT NULL DEFAULT 0 COMMENT 'V2连续高风险天数'", "confidence_score"),
    ]
    for column, definition, after_column in risk_columns:
        if has_column("inv_risk_score", column):
            continue
        conn.execute(
            text(
                f"""
                ALTER TABLE inv_risk_score
                ADD COLUMN {column} {definition}
                AFTER {after_column}
                """
            )
        )

    supply_columns = [
        ("raw_payload", "JSON NULL COMMENT '原始数据'", "source"),
    ]
    for column, definition, after_column in supply_columns:
        if has_column("inv_supply_stock_snapshot", column):
            continue
        conn.execute(
            text(
                f"""
                ALTER TABLE inv_supply_stock_snapshot
                ADD COLUMN {column} {definition}
                AFTER {after_column}
                """
            )
        )

    _seed_inventory_defaults(conn)


def _seed_inventory_defaults(conn: Any) -> None:
    supply_warehouse_rows = []
    for warehouse_code in (SUPPLY_HUB_WAREHOUSE, NANSHA_SUPPLY_WAREHOUSE):
        record = _warehouse_record(warehouse_code)
        supply_warehouse_rows.append(
            {
                **record,
                "warehouse_type": WAREHOUSE_TYPE_SUPPLY_HUB,
                "is_supply_hub": 1,
                "is_jd_warehouse": 0,
            }
        )
    conn.execute(
        text(
            """
            INSERT INTO inv_warehouse
                (warehouse_code, warehouse_name, region, warehouse_type, is_supply_hub, is_jd_warehouse,
                 base_city, latitude, longitude, created_at, updated_at)
            VALUES
                (:warehouse_code, :warehouse_name, :region, :warehouse_type, :is_supply_hub, :is_jd_warehouse,
                 :base_city, :latitude, :longitude, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                warehouse_name=VALUES(warehouse_name),
                region=VALUES(region),
                warehouse_type=VALUES(warehouse_type),
                is_supply_hub=VALUES(is_supply_hub),
                is_jd_warehouse=VALUES(is_jd_warehouse),
                base_city=VALUES(base_city),
                latitude=VALUES(latitude),
                longitude=VALUES(longitude),
                updated_at=NOW()
            """
        ),
        supply_warehouse_rows,
    )
    supplier_record = _warehouse_record(FOSHAN_SUPPLIER_WAREHOUSE)
    conn.execute(
        text(
            """
            INSERT INTO inv_warehouse
                (warehouse_code, warehouse_name, region, warehouse_type, is_supply_hub, is_jd_warehouse,
                 base_city, latitude, longitude, created_at, updated_at)
            VALUES
                (:warehouse_code, :warehouse_name, :region, :warehouse_type, :is_supply_hub, :is_jd_warehouse,
                 :base_city, :latitude, :longitude, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                warehouse_name=VALUES(warehouse_name),
                region=VALUES(region),
                warehouse_type=VALUES(warehouse_type),
                is_supply_hub=VALUES(is_supply_hub),
                is_jd_warehouse=VALUES(is_jd_warehouse),
                base_city=VALUES(base_city),
                latitude=VALUES(latitude),
                longitude=VALUES(longitude),
                updated_at=NOW()
            """
        ),
        {
            **supplier_record,
            "warehouse_type": WAREHOUSE_TYPE_SUPPLIER,
            "is_supply_hub": 0,
            "is_jd_warehouse": 0,
        },
    )

    policies = [
        {
            "policy_type": "SYSTEM",
            "category": "",
            "brand": "",
            "target_days": TARGET_DAYS,
            "reserve_days": DONOR_RESERVE_DAYS,
            "supplier_eta_days": SUPPLIER_TO_HUB_ETA_DAYS,
            "allow_jd_transfer": 1,
            "preferred_source_warehouse": SUPPLY_HUB_WAREHOUSE,
            "remark": "系统默认策略",
        },
        {
            "policy_type": "CATEGORY",
            "category": BOX_CATEGORY,
            "brand": "",
            "target_days": TARGET_DAYS,
            "reserve_days": DONOR_RESERVE_DAYS,
            "supplier_eta_days": SUPPLIER_TO_HUB_ETA_DAYS,
            "allow_jd_transfer": 0,
            "preferred_source_warehouse": SUPPLY_HUB_WAREHOUSE,
            "remark": "箱包类目只能由花东仓补货，不允许京东仓之间互调",
        },
        {
            "policy_type": "BRAND",
            "category": "",
            "brand": "iGear",
            "target_days": TARGET_DAYS,
            "reserve_days": DONOR_RESERVE_DAYS,
            "supplier_eta_days": SUPPLIER_TO_HUB_ETA_DAYS,
            "allow_jd_transfer": 1,
            "preferred_source_warehouse": NANSHA_SUPPLY_WAREHOUSE,
            "remark": "iGear品牌供给仓为南沙仓，南沙到各仓时效按花东标准估算",
        },
    ]
    conn.execute(
        text(
            """
            INSERT INTO inv_replenishment_policy
                (policy_type, category, brand, target_days, reserve_days, supplier_eta_days,
                 allow_jd_transfer, preferred_source_warehouse, remark, created_at, updated_at)
            VALUES
                (:policy_type, :category, :brand, :target_days, :reserve_days, :supplier_eta_days,
                 :allow_jd_transfer, :preferred_source_warehouse, :remark, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                id=id
            """
        ),
        policies,
    )
    conn.execute(
        text(
            """
            UPDATE inv_replenishment_policy
            SET
                allow_jd_transfer = 1,
                preferred_source_warehouse = :preferred_source_warehouse,
                remark = :remark,
                updated_at = NOW()
            WHERE policy_type = 'BRAND' AND brand = :brand
            """
        ),
        {
            "brand": "iGear",
            "preferred_source_warehouse": NANSHA_SUPPLY_WAREHOUSE,
            "remark": "iGear品牌供给仓为南沙仓，南沙到各仓时效按花东标准估算",
        },
    )

    route_rows = []
    for source_warehouse in (SUPPLY_HUB_WAREHOUSE, NANSHA_SUPPLY_WAREHOUSE):
        route_rows.extend(
            {
                "from_warehouse": source_warehouse,
                "to_warehouse": warehouse,
                "distance_km": 0.0,
                "eta_days": eta_days,
                "route_source": "MANUAL",
            }
            for warehouse, eta_days in HUADONG_ROUTE_ETA_DAYS.items()
        )
    conn.execute(
        text(
            """
            INSERT INTO inv_transfer_route
                (from_warehouse, to_warehouse, distance_km, eta_days, route_source, created_at, updated_at)
            VALUES
                (:from_warehouse, :to_warehouse, :distance_km, :eta_days, :route_source, NOW(), NOW())
            ON DUPLICATE KEY UPDATE
                eta_days=VALUES(eta_days),
                route_source=VALUES(route_source),
                updated_at=NOW()
            """
        ),
        route_rows,
    )

    default_enabled_text = _env_first("JUSHUITAN_SYNC_ENABLED", "JST_SYNC_ENABLED").lower()
    if default_enabled_text in {"1", "true", "yes", "on"}:
        default_enabled = 1
    elif default_enabled_text in {"0", "false", "no", "off"}:
        default_enabled = 0
    else:
        default_enabled = 1 if _env_first("JUSHUITAN_APP_KEY", "JST_APP_KEY") and _env_first("JUSHUITAN_ACCESS_TOKEN", "JST_ACCESS_TOKEN") else 0
    try:
        default_interval = max(1, min(1440, int(_env_first("JUSHUITAN_SYNC_INTERVAL_MINUTES", "JST_SYNC_INTERVAL_MINUTES") or "10")))
    except ValueError:
        default_interval = 10
    default_recalculate_text = _env_first("JUSHUITAN_SYNC_RECALCULATE", "JST_SYNC_RECALCULATE").lower()
    default_recalculate = 0 if default_recalculate_text in {"0", "false", "no", "off"} else 1
    conn.execute(
        text(
            """
            INSERT IGNORE INTO inv_huadong_sync_setting
                (id, enabled, interval_minutes, recalculate, last_status, last_message, created_at, updated_at)
            VALUES
                (1, :enabled, :interval_minutes, :recalculate, 'IDLE', '', NOW(), NOW())
            """
        ),
        {
            "enabled": default_enabled,
            "interval_minutes": default_interval,
            "recalculate": default_recalculate,
        },
    )
    if default_recalculate:
        conn.execute(text("UPDATE inv_huadong_sync_setting SET recalculate=1 WHERE id=1"))


def normalize_inventory_sku_codes(engine: Engine | None = None) -> dict[str, Any]:
    """Normalize scientific-notation SKU codes already persisted in inventory tables."""

    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.begin() as conn:
        mappings = _collect_sku_normalization_mappings(conn)
        if not mappings:
            return {
                "normalized_sku_count": 0,
                "sample_mappings": [],
                "tables": {},
                "remaining_normalizable_rows": _count_normalizable_sku_rows(conn),
            }

        table_stats = {
            "inv_sku": _normalize_inv_sku_table(conn, mappings),
            "inv_stock_snapshot": _normalize_unique_sku_table(
                conn,
                "inv_stock_snapshot",
                key_columns=("snapshot_date", "warehouse_code"),
                update_columns=(
                    "batch_no",
                    "product_status",
                    "is_gift",
                    "is_decision_active",
                    "stock_qty",
                    "inbound_qty",
                    "sales_1d",
                    "sales_7d",
                    "sales_14d",
                    "sales_30d",
                ),
                mappings=mappings,
            ),
            "inv_risk_score": _normalize_unique_sku_table(
                conn,
                "inv_risk_score",
                key_columns=("snapshot_date", "warehouse_code"),
                update_columns=(
                    "run_id",
                    "daily_sales",
                    "days_of_supply",
                    "demand_30d",
                    "shortage_days",
                    "risk_level",
                    "excess_qty",
                    "need_qty",
                    "available_qty",
                    "inbound_credit_qty",
                    "inbound_credit_rate",
                    "trend_factor",
                    "volatility_factor",
                    "confidence_score",
                    "consecutive_high_days",
                    "is_shortage_risk",
                    "is_excess_risk",
                ),
                mappings=mappings,
            ),
            "inv_transfer_plan": _normalize_unique_sku_table(
                conn,
                "inv_transfer_plan",
                key_columns=("run_id", "from_warehouse", "to_warehouse"),
                update_columns=(
                    "plan_date",
                    "transfer_qty",
                    "eta_days",
                    "priority",
                    "risk_level",
                    "risk_days",
                    "surplus_days_after",
                    "eta_late",
                    "score",
                    "reason",
                    "status",
                ),
                mappings=mappings,
            ),
            "inv_import_error": _normalize_plain_sku_table(conn, "inv_import_error", mappings),
        }

        return {
            "normalized_sku_count": len(mappings),
            "sample_mappings": [
                {"from": old, "to": new}
                for old, new in list(mappings.items())[:20]
            ],
            "tables": table_stats,
            "remaining_normalizable_rows": _count_normalizable_sku_rows(conn),
        }


def refresh_inventory_sku_categories(engine: Engine | None = None) -> dict[str, Any]:
    """Backfill SKU category by brand using the current category mapping."""

    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    updated = 0
    unmatched: set[str] = set()
    category_counts: dict[str, int] = defaultdict(int)

    with db_engine.begin() as conn:
        rows = conn.execute(text("SELECT id, brand, category FROM inv_sku")).mappings().all()
        for row in rows:
            brand = str(row["brand"] or "").strip()
            category = _brand_category(brand)
            if category:
                category_counts[category] += 1
            elif brand:
                unmatched.add(brand)
            if category and category != (row["category"] or ""):
                result = conn.execute(
                    text("UPDATE inv_sku SET category=:category, updated_at=NOW() WHERE id=:id"),
                    {"id": row["id"], "category": category},
                )
                updated += int(result.rowcount or 0)

    return {
        "updated": updated,
        "category_counts": dict(sorted(category_counts.items())),
        "unmatched_brands": sorted(unmatched),
    }


def _collect_sku_normalization_mappings(conn: Any) -> dict[str, str]:
    mappings: dict[str, str] = {}
    for table in SKU_CODE_TABLES:
        rows = conn.execute(
            text(f"SELECT DISTINCT sku_code FROM {table} WHERE sku_code IS NOT NULL AND sku_code<>''")
        ).scalars()
        for raw_value in rows:
            old_sku = str(raw_value or "").strip()
            new_sku = _clean_sku(old_sku)
            if not new_sku or old_sku == new_sku or len(new_sku) > 64:
                continue
            mappings[old_sku] = new_sku
    return dict(sorted(mappings.items(), key=lambda item: (item[1], item[0])))


def _count_normalizable_sku_rows(conn: Any) -> dict[str, int]:
    counts: dict[str, int] = {}
    for table in SKU_CODE_TABLES:
        total = 0
        rows = conn.execute(
            text(f"SELECT sku_code, COUNT(*) AS row_count FROM {table} WHERE sku_code IS NOT NULL AND sku_code<>'' GROUP BY sku_code")
        ).mappings()
        for row in rows:
            old_sku = str(row["sku_code"] or "").strip()
            if old_sku and _clean_sku(old_sku) != old_sku:
                total += int(row["row_count"] or 0)
        counts[table] = total
    return counts


def _normalize_inv_sku_table(conn: Any, mappings: dict[str, str]) -> dict[str, int]:
    stats = {"updated": 0, "merged": 0, "deleted": 0}
    for old_sku, new_sku in mappings.items():
        source = conn.execute(
            text(
                """
                SELECT id, product_name, brand, category, shop_name, supplier_name
                FROM inv_sku
                WHERE sku_code=:old_sku
                LIMIT 1
                """
            ),
            {"old_sku": old_sku},
        ).mappings().fetchone()
        if not source:
            continue

        target = conn.execute(
            text("SELECT id FROM inv_sku WHERE sku_code=:new_sku LIMIT 1"),
            {"new_sku": new_sku},
        ).mappings().fetchone()
        if target:
            conn.execute(
                text(
                    """
                    UPDATE inv_sku
                    SET product_name=CASE
                            WHEN product_name='' AND :product_name<>'' THEN :product_name
                            ELSE product_name
                        END,
                        brand=CASE
                            WHEN brand='' AND :brand<>'' THEN :brand
                            ELSE brand
                        END,
                        category=CASE
                            WHEN category='' AND :category<>'' THEN :category
                            ELSE category
                        END,
                        shop_name=CASE
                            WHEN shop_name='' AND :shop_name<>'' THEN :shop_name
                            ELSE shop_name
                        END,
                        supplier_name=CASE
                            WHEN supplier_name='' AND :supplier_name<>'' THEN :supplier_name
                            ELSE supplier_name
                        END,
                        updated_at=NOW()
                    WHERE id=:target_id
                    """
                ),
                {
                    "target_id": target["id"],
                    "product_name": source["product_name"] or "",
                    "brand": source["brand"] or "",
                    "category": source["category"] or _brand_category(source["brand"]),
                    "shop_name": source["shop_name"] or "",
                    "supplier_name": source["supplier_name"] or "",
                },
            )
            conn.execute(text("DELETE FROM inv_sku WHERE id=:source_id"), {"source_id": source["id"]})
            stats["merged"] += 1
            stats["deleted"] += 1
        else:
            result = conn.execute(
                text("UPDATE inv_sku SET sku_code=:new_sku, updated_at=NOW() WHERE id=:source_id"),
                {"new_sku": new_sku, "source_id": source["id"]},
            )
            stats["updated"] += int(result.rowcount or 0)
    return stats


def _normalize_unique_sku_table(
    conn: Any,
    table: str,
    key_columns: tuple[str, ...],
    update_columns: tuple[str, ...],
    mappings: dict[str, str],
) -> dict[str, int]:
    stats = {"updated": 0, "merged": 0, "deleted": 0}
    key_where = " AND ".join(f"{column}=:{column}" for column in key_columns)
    update_set = ", ".join(f"{column}=:{column}" for column in update_columns)

    for old_sku, new_sku in mappings.items():
        source_rows = conn.execute(
            text(f"SELECT * FROM {table} WHERE sku_code=:old_sku"),
            {"old_sku": old_sku},
        ).mappings().all()
        for source in source_rows:
            lookup_params = {column: source[column] for column in key_columns}
            lookup_params["new_sku"] = new_sku
            lookup_params["source_id"] = source["id"]
            target = conn.execute(
                text(
                    f"""
                    SELECT id
                    FROM {table}
                    WHERE sku_code=:new_sku
                      AND {key_where}
                      AND id<>:source_id
                    LIMIT 1
                    """
                ),
                lookup_params,
            ).mappings().fetchone()
            if target:
                update_params = {column: source[column] for column in update_columns}
                update_params["target_id"] = target["id"]
                conn.execute(
                    text(f"UPDATE {table} SET {update_set} WHERE id=:target_id"),
                    update_params,
                )
                conn.execute(text(f"DELETE FROM {table} WHERE id=:source_id"), {"source_id": source["id"]})
                stats["merged"] += 1
                stats["deleted"] += 1
            else:
                result = conn.execute(
                    text(f"UPDATE {table} SET sku_code=:new_sku WHERE id=:source_id"),
                    {"new_sku": new_sku, "source_id": source["id"]},
                )
                stats["updated"] += int(result.rowcount or 0)
    return stats


def _normalize_plain_sku_table(conn: Any, table: str, mappings: dict[str, str]) -> dict[str, int]:
    stats = {"updated": 0, "merged": 0, "deleted": 0}
    for old_sku, new_sku in mappings.items():
        result = conn.execute(
            text(f"UPDATE {table} SET sku_code=:new_sku WHERE sku_code=:old_sku"),
            {"new_sku": new_sku, "old_sku": old_sku},
        )
        stats["updated"] += int(result.rowcount or 0)
    return stats


def import_inventory_excel(
    file_path: str | os.PathLike[str],
    engine: Engine | None = None,
    original_file_name: str | None = None,
) -> dict[str, Any]:
    """Import one RDC inventory wide workbook into normalized inventory tables."""

    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)

    excel_path = Path(file_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel文件不存在: {excel_path}")

    batch_no = f"INV{datetime.now():%Y%m%d%H%M%S}{uuid.uuid4().hex[:8].upper()}"
    file_name = original_file_name or excel_path.name
    file_size = excel_path.stat().st_size

    with db_engine.begin() as conn:
        conn.execute(
            text(
                """
                INSERT INTO inv_import_batch
                    (batch_no, file_name, file_size, status, created_at)
                VALUES
                    (:batch_no, :file_name, :file_size, 'RUNNING', NOW())
                """
            ),
            {"batch_no": batch_no, "file_name": file_name, "file_size": file_size},
        )

    try:
        parsed = _parse_inventory_workbook(excel_path, batch_no=batch_no)
        with db_engine.begin() as conn:
            _upsert_skus(conn, parsed["skus"].values())
            _upsert_warehouses(conn, parsed["warehouses"])
            _upsert_stock_snapshots(conn, parsed["snapshots"])
            _purge_old_inventory_history(conn)
            _insert_import_errors(conn, parsed["errors"])
            conn.execute(
                text(
                    """
                    UPDATE inv_import_batch
                    SET status='SUCCESS',
                        snapshot_date=:snapshot_date,
                        rows_read=:rows_read,
                        sku_rows=:sku_rows,
                        snapshot_rows=:snapshot_rows,
                        warnings_count=:warnings_count,
                        finished_at=NOW()
                    WHERE batch_no=:batch_no
                    """
                ),
                {
                    "batch_no": batch_no,
                    "snapshot_date": parsed["snapshot_date"],
                    "rows_read": parsed["rows_read"],
                    "sku_rows": len(parsed["skus"]),
                    "snapshot_rows": len(parsed["snapshots"]),
                    "warnings_count": len(parsed["errors"]),
                },
            )

        result = ImportBatchResult(
            batch_no=batch_no,
            file_name=file_name,
            snapshot_date=str(parsed["snapshot_date"]) if parsed["snapshot_date"] else None,
            rows_read=parsed["rows_read"],
            sku_rows=len(parsed["skus"]),
            snapshot_rows=len(parsed["snapshots"]),
            warehouses=len(parsed["warehouses"]),
            warnings_count=len(parsed["errors"]),
        )
        return result.__dict__
    except Exception as exc:
        LOGGER.exception("Inventory import failed for %s", excel_path)
        with db_engine.begin() as conn:
            conn.execute(
                text(
                    """
                    UPDATE inv_import_batch
                    SET status='FAILED', error_message=:error_message, finished_at=NOW()
                    WHERE batch_no=:batch_no
                    """
                ),
                {"batch_no": batch_no, "error_message": str(exc)[:4000]},
            )
        raise


def recalculate_inventory_decisions(
    snapshot_date: str | date | None = None,
    engine: Engine | None = None,
    brand: str | None = None,
    category: str | None = None,
    sku_codes: Iterable[str] | None = None,
) -> dict[str, Any]:
    """Compute risk scores and transfer plans for a snapshot date."""

    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date)
    clean_brand = _clean_text(brand)
    clean_category = _clean_text(category)
    clean_skus = sorted({_clean_sku(sku) for sku in (sku_codes or []) if _clean_sku(sku)})
    run_no = f"RUN{datetime.now():%Y%m%d%H%M%S}{uuid.uuid4().hex[:8].upper()}"

    with db_engine.begin() as conn:
        run_id = conn.execute(
            text(
                """
                INSERT INTO inv_decision_run (run_no, snapshot_date, status, started_at)
                VALUES (:run_no, :snapshot_date, 'RUNNING', NOW())
                """
            ),
            {"run_no": run_no, "snapshot_date": clean_date},
        ).lastrowid

    try:
        policies = _load_replenishment_policies(db_engine)
        supply_stocks = _latest_supply_stocks_by_warehouse(db_engine)
        snapshots = _load_decision_snapshots(
            db_engine,
            clean_date,
            clean_brand or None,
            clean_category or None,
            clean_skus or None,
        )
        risk_rows, risk_stats = _calculate_risk_rows(run_id, clean_date, snapshots, policies)
        plan_rows = _calculate_transfer_plans(run_id, clean_date, snapshots, risk_rows, policies, supply_stocks)

        with db_engine.begin() as conn:
            if clean_brand or clean_category or clean_skus:
                params: dict[str, Any] = {"d": clean_date}
                risk_filters = ["snapshot_date=:d"]
                plan_filters = ["plan_date=:d"]
                if clean_skus:
                    sku_placeholders = []
                    for index, sku_code in enumerate(clean_skus):
                        key = f"sku_{index}"
                        sku_placeholders.append(f":{key}")
                        params[key] = sku_code
                    sku_filter = f"sku_code IN ({', '.join(sku_placeholders)})"
                    risk_filters.append(sku_filter)
                    plan_filters.append(sku_filter)
                if clean_brand:
                    risk_filters.append("sku_code IN (SELECT sku_code FROM inv_sku WHERE brand=:brand)")
                    plan_filters.append("sku_code IN (SELECT sku_code FROM inv_sku WHERE brand=:brand)")
                    params["brand"] = clean_brand
                if clean_category:
                    risk_filters.append("sku_code IN (SELECT sku_code FROM inv_sku WHERE category=:category)")
                    plan_filters.append("sku_code IN (SELECT sku_code FROM inv_sku WHERE category=:category)")
                    params["category"] = clean_category
                conn.execute(text(f"DELETE FROM inv_risk_score WHERE {' AND '.join(risk_filters)}"), params)
                conn.execute(text(f"DELETE FROM inv_transfer_plan WHERE {' AND '.join(plan_filters)}"), params)
            else:
                conn.execute(text("DELETE FROM inv_risk_score WHERE snapshot_date=:d"), {"d": clean_date})
                conn.execute(text("DELETE FROM inv_transfer_plan WHERE plan_date=:d"), {"d": clean_date})
            _insert_risk_scores(conn, risk_rows)
            _insert_transfer_plans(conn, plan_rows)
            conn.execute(
                text(
                    """
                    UPDATE inv_decision_run
                    SET status='SUCCESS',
                        risk_count=:risk_count,
                        plan_count=:plan_count,
                        high_count=:high_count,
                        mid_count=:mid_count,
                        low_count=:low_count,
                        finished_at=NOW(),
                        message=:message
                    WHERE id=:run_id
                    """
                ),
                {
                    "run_id": run_id,
                    "risk_count": len(risk_rows),
                    "plan_count": len(plan_rows),
                    "high_count": risk_stats["HIGH"],
                    "mid_count": risk_stats["MID"],
                    "low_count": risk_stats["LOW"],
                    "message": "库存风险评分与调拨方案计算完成",
                },
            )

        return {
            "run_id": run_id,
            "run_no": run_no,
            "snapshot_date": str(clean_date),
            "brand": clean_brand,
            "category": clean_category,
            "sku_count": len(clean_skus),
            "sku_codes": clean_skus,
            "risk_count": len(risk_rows),
            "plan_count": len(plan_rows),
            "high_count": risk_stats["HIGH"],
            "mid_count": risk_stats["MID"],
            "low_count": risk_stats["LOW"],
        }
    except Exception as exc:
        LOGGER.exception("Inventory decision run failed")
        with db_engine.begin() as conn:
            conn.execute(
                text(
                    """
                    UPDATE inv_decision_run
                    SET status='FAILED', message=:message, finished_at=NOW()
                    WHERE id=:run_id
                    """
                ),
                {"run_id": run_id, "message": str(exc)[:4000]},
            )
        raise


def get_latest_snapshot_date(engine: Engine | None = None) -> str | None:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.connect() as conn:
        value = conn.execute(text("SELECT MAX(snapshot_date) FROM inv_stock_snapshot")).scalar()
    return str(value) if value else None


def get_inventory_summary(engine: Engine | None = None, snapshot_date: str | None = None) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return {
            "snapshot_date": None,
            "sku_count": 0,
            "warehouse_count": 0,
            "risk_high": 0,
            "risk_mid": 0,
            "risk_low": 0,
            "plan_count": 0,
            "pending_count": 0,
        }
    with db_engine.connect() as conn:
        row = conn.execute(
            text(
                """
                SELECT
                    COUNT(DISTINCT s.sku_code) AS sku_count,
                    COUNT(DISTINCT s.warehouse_code) AS warehouse_count,
                    COALESCE(SUM(CASE WHEN r.risk_level='HIGH' THEN 1 ELSE 0 END), 0) AS risk_high,
                    COALESCE(SUM(CASE WHEN r.risk_level='MID' THEN 1 ELSE 0 END), 0) AS risk_mid,
                    COALESCE(SUM(CASE WHEN r.risk_level='LOW' THEN 1 ELSE 0 END), 0) AS risk_low
                FROM inv_stock_snapshot s
                LEFT JOIN inv_risk_score r
                  ON r.snapshot_date=s.snapshot_date
                 AND r.sku_code=s.sku_code
                 AND r.warehouse_code=s.warehouse_code
                WHERE s.snapshot_date=:d AND s.is_decision_active=1
                """
            ),
            {"d": clean_date},
        ).mappings().one()
        plan_row = conn.execute(
            text(
                """
                SELECT
                    COUNT(*) AS plan_count,
                    COALESCE(SUM(CASE WHEN status='PENDING' THEN 1 ELSE 0 END), 0) AS pending_count,
                    COALESCE(SUM(CASE WHEN priority='HIGH' THEN 1 ELSE 0 END), 0) AS high_plan_count
                FROM inv_transfer_plan
                WHERE plan_date=:d
                """
            ),
            {"d": clean_date},
        ).mappings().one()
        batch = conn.execute(
            text(
                """
                SELECT batch_no, file_name, rows_read, snapshot_rows, warnings_count, created_at
                FROM inv_import_batch
                WHERE snapshot_date=:d AND status='SUCCESS'
                ORDER BY created_at DESC
                LIMIT 1
                """
            ),
            {"d": clean_date},
        ).mappings().fetchone()

    return {
        "snapshot_date": str(clean_date),
        "sku_count": int(row["sku_count"] or 0),
        "warehouse_count": int(row["warehouse_count"] or 0),
        "risk_high": int(row["risk_high"] or 0),
        "risk_mid": int(row["risk_mid"] or 0),
        "risk_low": int(row["risk_low"] or 0),
        "plan_count": int(plan_row["plan_count"] or 0),
        "pending_count": int(plan_row["pending_count"] or 0),
        "high_plan_count": int(plan_row["high_plan_count"] or 0),
        "latest_batch": dict(batch) if batch else None,
    }


def list_transfer_plans(
    engine: Engine | None = None,
    snapshot_date: str | None = None,
    priority: str | None = None,
    status: str | None = None,
    sku_code: str | None = None,
    product_name: str | None = None,
    brand: str | None = None,
    category: str | None = None,
    min_transfer_qty: float | None = None,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return {"items": [], "total": 0, "page": page, "page_size": page_size}

    where = ["p.plan_date=:d"]
    params: dict[str, Any] = {"d": clean_date}
    if priority:
        where.append("p.priority=:priority")
        params["priority"] = priority
    if status:
        where.append("p.status=:status")
        params["status"] = status
    if sku_code:
        clean_sku = _clean_sku(sku_code)
        where.append("p.sku_code LIKE :sku_code")
        params["sku_code"] = f"%{clean_sku}%"
    if product_name:
        where.append("sku.product_name LIKE :product_name")
        params["product_name"] = f"%{product_name.strip()}%"
    if brand:
        where.append("sku.brand LIKE :brand")
        params["brand"] = f"%{brand.strip()}%"
    if category:
        where.append("sku.category=:category")
        params["category"] = category.strip()
    if min_transfer_qty is not None:
        qty_value = _to_float(min_transfer_qty)
        if qty_value > 0:
            where.append("p.transfer_qty > :min_transfer_qty")
            params["min_transfer_qty"] = qty_value

    clean_page = max(int(page or 1), 1)
    clean_size = min(max(int(page_size or 20), 1), 100)
    params["limit"] = clean_size
    params["offset"] = (clean_page - 1) * clean_size
    where_sql = " AND ".join(where)

    with db_engine.connect() as conn:
        total = conn.execute(
            text(
                f"""
                SELECT COUNT(*)
                FROM inv_transfer_plan p
                LEFT JOIN inv_sku sku ON sku.sku_code=p.sku_code
                WHERE {where_sql}
                """
            ),
            params,
        ).scalar()
        rows = conn.execute(
            text(
                f"""
                SELECT
                    p.id, p.run_id, p.plan_date, p.sku_code, sku.product_name, sku.brand, sku.category, sku.shop_name,
                    p.from_warehouse, fw.warehouse_name AS from_warehouse_name,
                    p.to_warehouse, tw.warehouse_name AS to_warehouse_name,
                    p.transfer_qty, p.eta_days, p.priority, p.risk_level, p.risk_days,
                    p.surplus_days_after, p.eta_late, p.source_type, p.route_path,
                    p.target_days, p.reserve_days, p.supplier_eta_days, p.hub_eta_days,
                    p.score, p.reason, p.status, p.created_at,
                    from_snap.stock_qty AS from_stock_qty,
                    from_snap.inbound_qty AS from_inbound_qty,
                    from_snap.sales_1d AS from_sales_1d,
                    from_snap.sales_7d AS from_sales_7d,
                    from_snap.sales_30d AS from_sales_30d,
                    from_risk.daily_sales AS from_daily_sales,
                    from_risk.demand_30d AS from_demand_30d,
                    to_snap.stock_qty AS to_stock_qty,
                    to_snap.inbound_qty AS to_inbound_qty,
                    to_snap.sales_1d AS to_sales_1d,
                    to_snap.sales_7d AS to_sales_7d,
                    to_snap.sales_30d AS to_sales_30d,
                    to_risk.daily_sales AS to_daily_sales,
                    to_risk.demand_30d AS to_demand_30d
                FROM inv_transfer_plan p
                LEFT JOIN inv_sku sku ON sku.sku_code=p.sku_code
                LEFT JOIN inv_warehouse fw ON fw.warehouse_code=p.from_warehouse
                LEFT JOIN inv_warehouse tw ON tw.warehouse_code=p.to_warehouse
                LEFT JOIN inv_stock_snapshot from_snap
                  ON from_snap.snapshot_date=p.plan_date
                 AND from_snap.sku_code=p.sku_code
                 AND from_snap.warehouse_code=p.from_warehouse
                LEFT JOIN inv_stock_snapshot to_snap
                  ON to_snap.snapshot_date=p.plan_date
                 AND to_snap.sku_code=p.sku_code
                 AND to_snap.warehouse_code=p.to_warehouse
                LEFT JOIN inv_risk_score from_risk
                  ON from_risk.snapshot_date=p.plan_date
                 AND from_risk.sku_code=p.sku_code
                 AND from_risk.warehouse_code=p.from_warehouse
                LEFT JOIN inv_risk_score to_risk
                  ON to_risk.snapshot_date=p.plan_date
                 AND to_risk.sku_code=p.sku_code
                 AND to_risk.warehouse_code=p.to_warehouse
                WHERE {where_sql}
                ORDER BY
                    p.sku_code ASC,
                    sku.product_name ASC,
                    FIELD(p.priority, 'HIGH', 'MID', 'LOW'),
                    p.eta_late DESC,
                    p.risk_days ASC,
                    p.eta_days ASC,
                    p.transfer_qty DESC,
                    p.score DESC,
                    p.id ASC
                LIMIT :limit OFFSET :offset
                """
            ),
            params,
        ).mappings().all()

    return {
        "items": [_jsonify_mapping(row) for row in rows],
        "total": int(total or 0),
        "page": clean_page,
        "page_size": clean_size,
    }


def list_transfer_plan_brands(
    engine: Engine | None = None,
    snapshot_date: str | None = None,
    sku_code: str | None = None,
    product_name: str | None = None,
    category: str | None = None,
    min_transfer_qty: float | None = None,
    limit: int = 500,
) -> list[str]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return []

    where = [
        "p.plan_date=:d",
        "sku.brand IS NOT NULL",
        "TRIM(sku.brand)<>''",
    ]
    params: dict[str, Any] = {"d": clean_date}
    if sku_code:
        clean_sku = _clean_sku(sku_code)
        where.append("p.sku_code LIKE :sku_code")
        params["sku_code"] = f"%{clean_sku}%"
    if product_name:
        where.append("sku.product_name LIKE :product_name")
        params["product_name"] = f"%{product_name.strip()}%"
    if category:
        where.append("sku.category=:category")
        params["category"] = category.strip()
    if min_transfer_qty is not None:
        qty_value = _to_float(min_transfer_qty)
        if qty_value > 0:
            where.append("p.transfer_qty > :min_transfer_qty")
            params["min_transfer_qty"] = qty_value

    params["limit"] = min(max(int(limit or 500), 1), 1000)
    where_sql = " AND ".join(where)
    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                f"""
                SELECT DISTINCT sku.brand AS brand
                FROM inv_transfer_plan p
                LEFT JOIN inv_sku sku ON sku.sku_code=p.sku_code
                WHERE {where_sql}
                ORDER BY sku.brand
                LIMIT :limit
                """
            ),
            params,
        ).mappings().all()
    return [str(row["brand"]).strip() for row in rows if str(row["brand"] or "").strip()]


def list_risk_scores(
    engine: Engine | None = None,
    snapshot_date: str | None = None,
    risk_level: str | None = None,
    sku_code: str | None = None,
    product_name: str | None = None,
    brand: str | None = None,
    category: str | None = None,
    warehouse_code: str | None = None,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return {"items": [], "total": 0, "page": page, "page_size": page_size}

    where = ["r.snapshot_date=:d"]
    params: dict[str, Any] = {"d": clean_date}
    if risk_level:
        where.append("r.risk_level=:risk_level")
        params["risk_level"] = risk_level
    if sku_code:
        clean_sku = _clean_sku(sku_code)
        where.append("r.sku_code LIKE :sku_code")
        params["sku_code"] = f"%{clean_sku}%"
    if product_name:
        where.append("sku.product_name LIKE :product_name")
        params["product_name"] = f"%{product_name.strip()}%"
    if brand:
        where.append("sku.brand LIKE :brand")
        params["brand"] = f"%{brand.strip()}%"
    if category:
        where.append("sku.category=:category")
        params["category"] = category.strip()
    if warehouse_code:
        where.append("r.warehouse_code=:warehouse_code")
        params["warehouse_code"] = warehouse_code

    clean_page = max(int(page or 1), 1)
    clean_size = min(max(int(page_size or 20), 1), 100)
    params["limit"] = clean_size
    params["offset"] = (clean_page - 1) * clean_size
    where_sql = " AND ".join(where)

    with db_engine.connect() as conn:
        total = conn.execute(
            text(
                f"""
                SELECT COUNT(*)
                FROM inv_risk_score r
                LEFT JOIN inv_sku sku ON sku.sku_code=r.sku_code
                WHERE {where_sql}
                """
            ),
            params,
        ).scalar()
        rows = conn.execute(
            text(
                f"""
                SELECT
                    r.snapshot_date, r.sku_code, sku.product_name, sku.brand, sku.category, sku.shop_name,
                    r.warehouse_code, wh.warehouse_name, wh.region,
                    r.daily_sales, r.target_days, r.days_of_supply, r.demand_30d, r.shortage_days,
                    r.risk_level, r.excess_qty, r.need_qty, r.available_qty, r.inbound_credit_qty,
                    r.inbound_credit_rate, r.trend_factor, r.volatility_factor, r.confidence_score,
                    r.consecutive_high_days, r.is_shortage_risk, r.is_excess_risk
                FROM inv_risk_score r
                LEFT JOIN inv_sku sku ON sku.sku_code=r.sku_code
                LEFT JOIN inv_warehouse wh ON wh.warehouse_code=r.warehouse_code
                WHERE {where_sql}
                ORDER BY
                    FIELD(r.risk_level, 'HIGH', 'MID', 'LOW'),
                    r.days_of_supply ASC,
                    r.need_qty DESC
                LIMIT :limit OFFSET :offset
                """
            ),
            params,
        ).mappings().all()

    return {
        "items": [_jsonify_mapping(row) for row in rows],
        "total": int(total or 0),
        "page": clean_page,
        "page_size": clean_size,
    }


def list_city_health(engine: Engine | None = None, snapshot_date: str | None = None) -> list[dict[str, Any]]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return []

    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT
                    wh.warehouse_code,
                    wh.warehouse_name,
                    wh.region,
                    COUNT(*) AS active_sku_count,
                    COALESCE(SUM(CASE WHEN r.risk_level='HIGH' THEN 1 ELSE 0 END), 0) AS high_count,
                    COALESCE(SUM(CASE WHEN r.risk_level='MID' THEN 1 ELSE 0 END), 0) AS mid_count,
                    COALESCE(SUM(CASE WHEN r.risk_level='LOW' THEN 1 ELSE 0 END), 0) AS low_count,
                    COALESCE(SUM(CASE WHEN r.is_excess_risk=1 THEN 1 ELSE 0 END), 0) AS excess_count,
                    COALESCE(SUM(s.stock_qty), 0) AS stock_qty,
                    COALESCE(SUM(s.inbound_qty), 0) AS inbound_qty,
                    COALESCE(SUM(s.sales_7d), 0) AS sales_7d
                FROM inv_stock_snapshot s
                JOIN inv_warehouse wh ON wh.warehouse_code=s.warehouse_code
                LEFT JOIN inv_risk_score r
                  ON r.snapshot_date=s.snapshot_date
                 AND r.sku_code=s.sku_code
                 AND r.warehouse_code=s.warehouse_code
                WHERE s.snapshot_date=:d AND s.is_decision_active=1
                GROUP BY wh.warehouse_code, wh.warehouse_name, wh.region
                ORDER BY high_count DESC, mid_count DESC, active_sku_count DESC
                """
            ),
            {"d": clean_date},
        ).mappings().all()
    return [_jsonify_mapping(row) for row in rows]


def list_inventory_brands(engine: Engine | None = None, keyword: str | None = None, limit: int = 200) -> list[str]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_limit = min(max(int(limit or 200), 1), 500)
    where = ["brand IS NOT NULL", "TRIM(brand) <> ''"]
    params: dict[str, Any] = {"limit": clean_limit}
    if keyword:
        where.append("brand LIKE :keyword")
        params["keyword"] = f"%{keyword.strip()}%"

    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                f"""
                SELECT DISTINCT brand
                FROM inv_sku
                WHERE {' AND '.join(where)}
                ORDER BY brand
                LIMIT :limit
                """
            ),
            params,
        ).scalars().all()
    return [str(row) for row in rows if str(row).strip()]


def list_inventory_categories(engine: Engine | None = None) -> list[str]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT DISTINCT category
                FROM inv_sku
                WHERE category IS NOT NULL AND TRIM(category) <> ''
                ORDER BY category
                """
            )
        ).scalars().all()
    return [str(row) for row in rows if str(row).strip()]


def list_replenishment_policies(engine: Engine | None = None) -> list[dict[str, Any]]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT id, policy_type, category, brand, target_days, reserve_days,
                       supplier_eta_days, allow_jd_transfer, preferred_source_warehouse,
                       remark, updated_at
                FROM inv_replenishment_policy
                ORDER BY FIELD(policy_type, 'SYSTEM', 'CATEGORY', 'BRAND'), category, brand
                """
            )
        ).mappings().all()
    return [_jsonify_mapping(row) for row in rows]


def upsert_replenishment_policy(
    engine: Engine | None = None,
    policy_type: str = "BRAND",
    brand: str = "",
    category: str = "",
    target_days: float | None = None,
    reserve_days: float | None = None,
    supplier_eta_days: int | None = None,
    allow_jd_transfer: bool | None = None,
    preferred_source_warehouse: str | None = None,
    remark: str = "",
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_type = str(policy_type or "BRAND").strip().upper()
    if clean_type not in {"SYSTEM", "CATEGORY", "BRAND"}:
        raise ValueError("策略类型必须是 SYSTEM/CATEGORY/BRAND")
    clean_brand = _clean_text(brand) if clean_type == "BRAND" else ""
    clean_category = _clean_text(category) if clean_type == "CATEGORY" else ""
    if clean_type == "BRAND" and not clean_brand:
        raise ValueError("品牌策略必须选择品牌")
    if clean_type == "CATEGORY" and not clean_category:
        raise ValueError("类目策略必须选择类目")

    base_policy = _default_replenishment_policy()
    target = _clamp(_to_float(target_days or base_policy.target_days), 1.0, 365.0)
    reserve = _clamp(_to_float(reserve_days or base_policy.reserve_days), 1.0, target)
    supplier_eta = int(max(1, _to_float(supplier_eta_days or base_policy.supplier_eta_days)))
    hub = _clean_text(preferred_source_warehouse) or SUPPLY_HUB_WAREHOUSE
    jd_flag = base_policy.allow_jd_transfer if allow_jd_transfer is None else bool(allow_jd_transfer)
    if clean_category == BOX_CATEGORY:
        jd_flag = False

    with db_engine.begin() as conn:
        conn.execute(
            text(
                """
                INSERT INTO inv_replenishment_policy
                    (policy_type, category, brand, target_days, reserve_days, supplier_eta_days,
                     allow_jd_transfer, preferred_source_warehouse, remark, created_at, updated_at)
                VALUES
                    (:policy_type, :category, :brand, :target_days, :reserve_days, :supplier_eta_days,
                     :allow_jd_transfer, :preferred_source_warehouse, :remark, NOW(), NOW())
                ON DUPLICATE KEY UPDATE
                    target_days=VALUES(target_days),
                    reserve_days=VALUES(reserve_days),
                    supplier_eta_days=VALUES(supplier_eta_days),
                    allow_jd_transfer=VALUES(allow_jd_transfer),
                    preferred_source_warehouse=VALUES(preferred_source_warehouse),
                    remark=VALUES(remark),
                    updated_at=NOW()
                """
            ),
            {
                "policy_type": clean_type,
                "category": clean_category,
                "brand": clean_brand,
                "target_days": target,
                "reserve_days": reserve,
                "supplier_eta_days": supplier_eta,
                "allow_jd_transfer": 1 if jd_flag else 0,
                "preferred_source_warehouse": hub,
                "remark": remark.strip()[:512],
            },
        )

    return {
        "policy_type": clean_type,
        "brand": clean_brand,
        "category": clean_category,
        "target_days": target,
        "reserve_days": reserve,
        "supplier_eta_days": supplier_eta,
        "allow_jd_transfer": bool(jd_flag),
        "preferred_source_warehouse": hub,
    }


def upsert_supply_stock_snapshot(
    rows: list[dict[str, Any]],
    engine: Engine | None = None,
    warehouse_code: str = SUPPLY_HUB_WAREHOUSE,
    source: str = "JUSHUITAN",
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_warehouse = _clean_text(warehouse_code) or SUPPLY_HUB_WAREHOUSE
    now = datetime.now()
    payload_rows: list[dict[str, Any]] = []
    for row in rows:
        sku_code = _clean_sku(row.get("sku_code") or row.get("sku"))
        if not sku_code:
            continue
        stock_qty = _to_float(row.get("stock_qty"))
        available_qty = _to_float(row.get("available_qty"))
        if available_qty <= 0 and stock_qty > 0:
            available_qty = stock_qty
        snapshot_time = _parse_datetime_value(row.get("snapshot_time")) or now
        raw_payload = row.get("raw_payload")
        if raw_payload is not None and not isinstance(raw_payload, str):
            raw_payload = json.dumps(raw_payload, ensure_ascii=False, default=str)
        payload_rows.append(
            {
                "snapshot_time": snapshot_time,
                "snapshot_date": snapshot_time.date(),
                "warehouse_code": clean_warehouse,
                "sku_code": sku_code,
                "stock_qty": stock_qty,
                "available_qty": available_qty,
                "inbound_qty": _to_float(row.get("inbound_qty")),
                "source": _clean_text(source) or "JUSHUITAN",
                "raw_payload": raw_payload,
            }
        )
    if not payload_rows:
        return {"inserted": 0, "warehouse_code": clean_warehouse}

    with db_engine.begin() as conn:
        if clean_warehouse == SUPPLY_HUB_WAREHOUSE:
            sku_codes = sorted({row["sku_code"] for row in payload_rows})
            brand_rows = conn.execute(
                text(
                    """
                    SELECT sku_code, brand
                    FROM inv_sku
                    WHERE sku_code IN :sku_codes
                    """
                ).bindparams(bindparam("sku_codes", expanding=True)),
                {"sku_codes": sku_codes},
            ).mappings().all()
            brand_by_sku = {str(row["sku_code"]): row["brand"] for row in brand_rows}
            for payload_row in payload_rows:
                if _is_nansha_supply_brand(brand_by_sku.get(payload_row["sku_code"])):
                    payload_row["warehouse_code"] = NANSHA_SUPPLY_WAREHOUSE
        conn.execute(
            text(
                """
                INSERT INTO inv_supply_stock_snapshot
                    (snapshot_time, snapshot_date, warehouse_code, sku_code, stock_qty,
                     available_qty, inbound_qty, source, raw_payload, created_at)
                VALUES
                    (:snapshot_time, :snapshot_date, :warehouse_code, :sku_code, :stock_qty,
                     :available_qty, :inbound_qty, :source, :raw_payload, NOW())
                ON DUPLICATE KEY UPDATE
                    stock_qty=VALUES(stock_qty),
                    available_qty=VALUES(available_qty),
                    inbound_qty=VALUES(inbound_qty),
                    source=VALUES(source),
                    raw_payload=VALUES(raw_payload)
                """
            ),
            payload_rows,
        )
    warehouse_codes = sorted({str(row["warehouse_code"]) for row in payload_rows})
    return {
        "inserted": len(payload_rows),
        "warehouse_code": warehouse_codes[0] if len(warehouse_codes) == 1 else "、".join(warehouse_codes),
        "warehouse_codes": warehouse_codes,
    }


def upsert_jushuitan_stock_sync_payload(
    payload: Any,
    engine: Engine | None = None,
    warehouse_code: str = SUPPLY_HUB_WAREHOUSE,
) -> dict[str, Any]:
    """Persist 聚水潭 business_sku_syn inventory messages into the supply hub stock table."""
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    items = _extract_jushuitan_stock_items(payload)
    rows: list[dict[str, Any]] = []
    ignored = 0
    for item in items:
        if not isinstance(item, dict):
            ignored += 1
            continue
        sku_code = _clean_sku(
            item.get("skuId")
            or item.get("sku_id")
            or item.get("skuCode")
            or item.get("sku_code")
            or item.get("sku")
            or item.get("iId")
            or item.get("iid")
        )
        if not sku_code:
            ignored += 1
            continue

        stock_qty = _to_float(item.get("qty"))
        order_lock = _to_float(item.get("orderLock"))
        pick_lock = _to_float(item.get("pickLock"))
        available_qty = max(stock_qty - order_lock - pick_lock, 0.0)
        rows.append(
            {
                "sku_code": sku_code,
                "stock_qty": stock_qty,
                "available_qty": available_qty,
                "inbound_qty": _to_float(item.get("purchaseQty")),
                "snapshot_time": item.get("modified"),
                "raw_payload": item,
            }
        )

    code_map = _sku_code_map_for_external_codes(db_engine, [row["sku_code"] for row in rows])
    for row in rows:
        row["sku_code"] = code_map.get(row["sku_code"], row["sku_code"])

    result = upsert_supply_stock_snapshot(
        rows,
        db_engine,
        warehouse_code=warehouse_code,
        source="JUSHUITAN:business_sku_syn",
    )
    return {
        **result,
        "received": len(items),
        "accepted": len(rows),
        "ignored": ignored,
        "recalculated": False,
    }


def _env_first(*names: str) -> str:
    for name in names:
        value = os.getenv(name)
        if value:
            return value.strip()
    return ""


def _load_jushuitan_config() -> JushuitanConfig:
    app_key = _env_first("JUSHUITAN_APP_KEY", "JST_APP_KEY")
    app_secret = _env_first("JUSHUITAN_APP_SECRET", "JST_APP_SECRET")
    access_token = _env_first("JUSHUITAN_ACCESS_TOKEN", "JST_ACCESS_TOKEN")
    refresh_token = _env_first("JUSHUITAN_REFRESH_TOKEN", "JST_REFRESH_TOKEN") or None
    endpoint = _env_first("JUSHUITAN_ENDPOINT", "JST_ENDPOINT") or "https://openapi.jushuitan.com"
    missing = [
        label
        for label, value in (
            ("JUSHUITAN_APP_KEY", app_key),
            ("JUSHUITAN_APP_SECRET", app_secret),
            ("JUSHUITAN_ACCESS_TOKEN", access_token),
        )
        if not value
    ]
    if missing:
        raise ValueError(f"聚水潭配置缺失: {', '.join(missing)}")
    return JushuitanConfig(
        endpoint=endpoint.rstrip("/"),
        app_key=app_key,
        app_secret=app_secret,
        access_token=access_token,
        refresh_token=refresh_token,
    )


def _jushuitan_sign(config: JushuitanConfig, params: dict[str, Any]) -> str:
    source = config.app_secret + "".join(f"{key}{params[key]}" for key in sorted(params))
    return hashlib.md5(source.encode("utf-8")).hexdigest()


def _jushuitan_post(path: str, biz: dict[str, Any], config: JushuitanConfig | None = None) -> dict[str, Any]:
    api_config = config or _load_jushuitan_config()
    biz_text = json.dumps(biz, ensure_ascii=False, separators=(",", ":"))
    params: dict[str, Any] = {
        "access_token": api_config.access_token,
        "app_key": api_config.app_key,
        "biz": biz_text,
        "charset": "utf-8",
        "timestamp": str(int(datetime.now().timestamp())),
        "version": "2",
    }
    params["sign"] = _jushuitan_sign(api_config, params)

    session = requests.Session()
    session.trust_env = False
    try:
        response = session.post(
            f"{api_config.endpoint}{path}",
            data=params,
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            timeout=25,
        )
    finally:
        session.close()
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, dict):
        raise ValueError("聚水潭接口返回格式异常")
    if int(payload.get("code") or 0) != 0:
        raise ValueError(f"聚水潭接口调用失败: {payload.get('code')} {payload.get('msg') or payload.get('message') or ''}".strip())
    return payload


def _is_retryable_jushuitan_error(exc: Exception) -> bool:
    message = str(exc)
    return (
        isinstance(exc, (requests.Timeout, requests.ConnectionError, requests.HTTPError))
        or "199" in message
        or "too frequent" in message.lower()
        or "频繁" in message
        or "502" in message
        or "503" in message
        or "504" in message
    )


def _first_list(value: Any) -> list[Any]:
    if isinstance(value, list):
        return value
    if not isinstance(value, dict):
        return []
    for key in ("inventorys", "wms_partners", "wmsPartners", "partners", "items", "datas", "data", "rows", "list"):
        item = value.get(key)
        if isinstance(item, list):
            return item
    for item in value.values():
        nested = _first_list(item)
        if nested:
            return nested
    return []


def _jst_field(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] is not None:
            return row[name]
    return None


def _query_jushuitan_wms_partners(config: JushuitanConfig | None = None, page_size: int = 100) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    page_index = 1
    api_config = config or _load_jushuitan_config()
    while page_index <= 20:
        payload = _jushuitan_post(
            "/open/wms/partner/query",
            {"page_index": page_index, "page_size": page_size},
            api_config,
        )
        data = payload.get("data")
        page_rows = [row for row in _first_list(data) if isinstance(row, dict)]
        rows.extend(page_rows)
        has_next = bool(data.get("has_next") or data.get("hasNext")) if isinstance(data, dict) else False
        if not has_next or not page_rows:
            break
        page_index += 1
    return rows


def _resolve_huadong_wms_co_id(config: JushuitanConfig | None = None) -> int:
    configured = _env_first("JUSHUITAN_HUADONG_WMS_CO_ID", "JST_HUADONG_WMS_CO_ID")
    if configured:
        try:
            return int(configured)
        except ValueError as exc:
            raise ValueError("JUSHUITAN_HUADONG_WMS_CO_ID 必须是数字") from exc

    partners = _query_jushuitan_wms_partners(config)
    for partner in partners:
        raw_text = json.dumps(partner, ensure_ascii=False, default=str)
        if "花东" not in raw_text and "华东" not in raw_text:
            continue
        raw_id = _jst_field(partner, "wms_co_id", "wmsCoId", "wms_coid", "wmsCoid", "co_id", "coId", "coid", "id")
        try:
            return int(raw_id)
        except (TypeError, ValueError):
            continue
    raise ValueError("未识别到花东仓仓库编号，请配置 JUSHUITAN_HUADONG_WMS_CO_ID")


def _default_jushuitan_sync_skus(engine: Engine, limit: int | None = None) -> list[str]:
    return sorted(
        {
            sku_code
            for item in _jushuitan_sync_targets(engine, None, limit)
            for sku_code in item.get("sku_codes", [])
        }
    )


def _load_sku_jst_lookup_map(
    engine: Engine,
    sku_codes: Iterable[str] | None = None,
) -> dict[str, str]:
    clean_codes = sorted({_clean_sku(code) for code in (sku_codes or []) if _clean_sku(code)})
    with engine.connect() as conn:
        if clean_codes:
            rows = conn.execute(
                text(
                    """
                    SELECT sku_code, lookup_code
                    FROM inv_sku_jst_lookup
                    WHERE sku_code IN :codes
                    """
                ).bindparams(bindparam("codes", expanding=True)),
                {"codes": clean_codes},
            ).mappings().all()
        else:
            rows = conn.execute(
                text(
                    """
                    SELECT sku_code, lookup_code
                    FROM inv_sku_jst_lookup
                    """
                )
            ).mappings().all()
    lookup_map: dict[str, str] = {}
    for row in rows:
        sku_code = _clean_sku(row.get("sku_code"))
        lookup_code = _clean_sku(row.get("lookup_code"))
        if sku_code and lookup_code:
            lookup_map[sku_code] = lookup_code
    return lookup_map


def list_sku_jst_lookup_mappings(
    engine: Engine | None = None,
    sku_code: str | None = None,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_page = max(1, int(page or 1))
    clean_page_size = min(200, max(1, int(page_size or 20)))
    sku_text = str(sku_code or "").strip()
    where = ""
    params: dict[str, Any] = {}
    if sku_text:
        where = "WHERE m.sku_code LIKE :sku_code"
        params["sku_code"] = f"%{sku_text}%"
    with db_engine.connect() as conn:
        total = conn.execute(
            text(f"SELECT COUNT(*) FROM inv_sku_jst_lookup m {where}"),
            params,
        ).scalar()
        rows = conn.execute(
            text(
                f"""
                SELECT m.sku_code, m.lookup_code, m.updated_at,
                       s.product_name, s.brand
                FROM inv_sku_jst_lookup m
                LEFT JOIN inv_sku s ON s.sku_code=m.sku_code
                {where}
                ORDER BY m.updated_at DESC, m.id DESC
                LIMIT :limit OFFSET :offset
                """
            ),
            {**params, "limit": clean_page_size, "offset": (clean_page - 1) * clean_page_size},
        ).mappings().all()
    items = [
        {
            "sku_code": _clean_sku(row["sku_code"]),
            "lookup_code": _clean_sku(row["lookup_code"]),
            "product_name": str(row.get("product_name") or ""),
            "brand": str(row.get("brand") or ""),
            "updated_at": row["updated_at"].isoformat(sep=" ")
            if hasattr(row.get("updated_at"), "isoformat")
            else str(row.get("updated_at") or ""),
        }
        for row in rows
    ]
    return {
        "items": items,
        "total": int(total or 0),
        "page": clean_page,
        "page_size": clean_page_size,
    }


def upsert_sku_jst_lookup_mappings(
    rows: Iterable[dict[str, Any]],
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    rows_list = list(rows)
    clean_rows: list[dict[str, Any]] = []
    for raw in rows_list:
        sku_code = _clean_sku(raw.get("sku_code") or raw.get("sku"))
        lookup_code = _clean_sku(raw.get("lookup_code") or raw.get("jushuitan_lookup_code"))
        if sku_code and lookup_code:
            clean_rows.append({"sku_code": sku_code, "lookup_code": lookup_code})
    if not clean_rows:
        return {"imported": 0, "skipped": len(rows_list)}
    with db_engine.begin() as conn:
        _execute_chunks(
            conn,
            text(
                """
                INSERT INTO inv_sku_jst_lookup (sku_code, lookup_code, created_at, updated_at)
                VALUES (:sku_code, :lookup_code, NOW(), NOW())
                ON DUPLICATE KEY UPDATE
                    lookup_code=VALUES(lookup_code),
                    updated_at=NOW()
                """
            ),
            clean_rows,
        )
    return {"imported": len(clean_rows), "skipped": max(0, len(rows_list) - len(clean_rows))}


def build_sku_jst_lookup_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "聚水潭查找编码"
    headers = ["sku", "聚水潭查找编码"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(name="Microsoft YaHei", size=12, bold=True)
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "A2"
    io = BytesIO()
    wb.save(io)
    return io.getvalue()


def import_sku_jst_lookup_excel(
    excel_path: str | Path,
    engine: Engine | None = None,
) -> dict[str, Any]:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError("聚水潭查找编码模板没有表头")
    headers = [str(value or "").strip() if value is not None else "" for value in header_row]
    sku_index = next((idx for idx, header in enumerate(headers) if header.lower() == "sku"), None)
    lookup_index = next(
        (idx for idx, header in enumerate(headers) if header == "聚水潭查找编码"),
        None,
    )
    if sku_index is None or lookup_index is None:
        raise ValueError("模板必须包含 sku 和 聚水潭查找编码 两列")

    mappings: list[dict[str, Any]] = []
    rows_read = 0
    for values in rows_iter:
        rows_read += 1
        if values is None:
            continue
        sku_value = values[sku_index] if sku_index < len(values) else None
        lookup_value = values[lookup_index] if lookup_index < len(values) else None
        sku_code = _clean_sku(sku_value)
        lookup_code = _clean_sku(lookup_value)
        if sku_code and lookup_code:
            mappings.append({"sku_code": sku_code, "lookup_code": lookup_code})
    wb.close()
    result = upsert_sku_jst_lookup_mappings(mappings, engine)
    return {
        "rows_read": rows_read,
        "sku_codes": [row["sku_code"] for row in mappings],
        **result,
    }


def _jushuitan_sync_target_data(
    engine: Engine,
    codes: Iterable[str] | None = None,
    limit: int | None = None,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    clean_codes = list(dict.fromkeys(_clean_sku(code) for code in (codes or []) if _clean_sku(code)))
    with engine.connect() as conn:
        if clean_codes:
            params: dict[str, Any] = {"codes": clean_codes}
            barcode_like_filters = []
            for index, code in enumerate(clean_codes):
                key = f"barcode_like_{index}"
                barcode_like_filters.append(f"barcode LIKE :{key}")
                params[key] = f"%{code}%"
            rows = conn.execute(
                text(
                    f"""
                    SELECT sku_code, barcode
                    FROM inv_sku
                    WHERE sku_code IN :codes
                       OR barcode IN :codes
                       OR {' OR '.join(barcode_like_filters)}
                    ORDER BY id
                    """
                ).bindparams(bindparam("codes", expanding=True)),
                params,
            ).mappings().all()
        else:
            sql = "SELECT sku_code, barcode FROM inv_sku ORDER BY id"
            params = {}
            if limit is not None:
                sql += " LIMIT :limit"
                params["limit"] = max(1, int(limit))
            rows = conn.execute(text(sql), params).mappings().all()

    lookup_map = _load_sku_jst_lookup_map(
        engine,
        [row["sku_code"] for row in rows if row.get("sku_code")],
    )
    requested_codes = set(clean_codes)
    query_to_skus: dict[str, list[str]] = {}
    query_source: dict[str, str] = {}
    total_skus: set[str] = set()
    skipped_without_barcode: set[str] = set()
    barcode_links = 0
    lookup_code_links = 0

    for row in rows:
        sku_code = _clean_sku(row["sku_code"])
        if not sku_code:
            continue
        lookup_code = lookup_map.get(sku_code)
        all_barcode_tokens = _barcode_tokens(row.get("barcode"))
        requested_by_sku = sku_code in requested_codes
        requested_by_lookup = bool(lookup_code and lookup_code in requested_codes)

        if requested_codes and not requested_by_sku and not requested_by_lookup:
            matched_tokens = [token for token in all_barcode_tokens if token in requested_codes]
            if not matched_tokens:
                continue
            all_barcode_tokens = matched_tokens[:1]
        elif lookup_code:
            total_skus.add(sku_code)
            sku_codes = query_to_skus.setdefault(lookup_code, [])
            if sku_code not in sku_codes:
                sku_codes.append(sku_code)
            query_source[lookup_code] = "JST_LOOKUP"
            lookup_code_links += 1
            continue
        else:
            all_barcode_tokens = all_barcode_tokens[:1]

        total_skus.add(sku_code)
        if not all_barcode_tokens:
            skipped_without_barcode.add(sku_code)
            continue
        for query_code in all_barcode_tokens:
            barcode_links += 1
            sku_codes = query_to_skus.setdefault(query_code, [])
            if sku_code not in sku_codes:
                sku_codes.append(sku_code)
            query_source.setdefault(query_code, "BARCODE")

    targets = [
        {
            "sku_code": sku_codes[0],
            "sku_codes": sku_codes,
            "query_code": query_code,
            "query_code_source": query_source.get(query_code, "BARCODE"),
        }
        for query_code, sku_codes in query_to_skus.items()
    ]
    stats = {
        "total_skus": len(total_skus),
        "requested_skus": len(total_skus - skipped_without_barcode),
        "requested_query_codes": len(targets),
        "skipped_without_barcode": len(skipped_without_barcode),
        "barcode_links": barcode_links,
        "lookup_code_links": lookup_code_links,
        "duplicate_query_codes": max(barcode_links + lookup_code_links - len(targets), 0),
    }
    return targets, stats


def _jushuitan_sync_targets(
    engine: Engine,
    codes: Iterable[str] | None = None,
    limit: int | None = None,
) -> list[dict[str, Any]]:
    return _jushuitan_sync_target_data(engine, codes, limit)[0]


def _sku_code_map_for_external_codes(engine: Engine, codes: Iterable[str]) -> dict[str, str]:
    clean_codes = sorted({_clean_sku(code) for code in codes if _clean_sku(code)})
    if not clean_codes:
        return {}
    with engine.connect() as conn:
        params: dict[str, Any] = {"codes": clean_codes}
        barcode_like_filters = []
        for index, code in enumerate(clean_codes):
            key = f"barcode_like_{index}"
            barcode_like_filters.append(f"barcode LIKE :{key}")
            params[key] = f"%{code}%"
        rows = conn.execute(
            text(
                f"""
                SELECT sku_code, barcode
                FROM inv_sku
                WHERE sku_code IN :codes
                   OR barcode IN :codes
                   OR {' OR '.join(barcode_like_filters)}
                """
            ).bindparams(bindparam("codes", expanding=True)),
            params,
        ).mappings().all()
    code_map: dict[str, str] = {}
    for row in rows:
        sku_code = _clean_sku(row["sku_code"])
        barcode = _clean_barcode(row.get("barcode"))
        if sku_code:
            code_map[sku_code] = sku_code
        if sku_code:
            for barcode_token in _barcode_tokens(barcode):
                code_map[barcode_token] = sku_code
    lookup_map = _load_sku_jst_lookup_map(engine)
    for sku_code, lookup_code in lookup_map.items():
        code_map[lookup_code] = sku_code
        code_map[sku_code] = sku_code
    return code_map


def _jushuitan_inventory_item_to_supply_row(item: dict[str, Any]) -> dict[str, Any] | None:
    sku_code = _clean_sku(
        _jst_field(item, "sku_id", "skuId", "sku_code", "skuCode", "sku", "i_id", "iId", "iid")
    )
    if not sku_code:
        return None
    stock_qty = _to_float(_jst_field(item, "qty", "stock_qty", "stockQty"))
    order_lock = _to_float(_jst_field(item, "order_lock", "orderLock"))
    pick_lock = _to_float(_jst_field(item, "pick_lock", "pickLock"))
    lock_qty = _to_float(_jst_field(item, "lock_qty", "lockQty"))
    if order_lock > 0 or pick_lock > 0:
        available_qty = max(stock_qty - order_lock - pick_lock, 0.0)
    elif lock_qty > 0:
        available_qty = max(stock_qty - lock_qty, 0.0)
    else:
        available_qty = stock_qty
    return {
        "sku_code": sku_code,
        "stock_qty": stock_qty,
        "available_qty": available_qty,
        "inbound_qty": _to_float(_jst_field(item, "purchase_qty", "purchaseQty")),
        "snapshot_time": _jst_field(item, "modified", "modified_time", "modifiedTime"),
        "raw_payload": item,
    }


def _jushuitan_inventory_item_codes(item: dict[str, Any]) -> list[str]:
    codes: list[str] = []
    for field_name in ("sku_id", "skuId", "sku_code", "skuCode", "sku", "i_id", "iId", "iid"):
        code = _clean_sku(item.get(field_name))
        if code and code not in codes:
            codes.append(code)
    return codes


def sync_jushuitan_huadong_stock(
    sku_codes: list[str] | None = None,
    engine: Engine | None = None,
    wms_co_id: int | None = None,
    recalculate: bool = False,
    snapshot_date: str | None = None,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Pull 聚水潭库存查询 API and persist 花东仓 supply stock snapshots."""
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_skus = [_clean_sku(sku) for sku in (sku_codes or []) if _clean_sku(sku)]
    targets, target_stats = _jushuitan_sync_target_data(db_engine, clean_skus or None)
    if not targets:
        if target_stats.get("total_skus"):
            raise ValueError("没有可同步的条形码，请检查商品条形码数据")
        raise ValueError("没有可同步的 SKU")
    query_codes = [item["query_code"] for item in targets]
    query_to_skus = {item["query_code"]: item["sku_codes"] for item in targets}

    config = _load_jushuitan_config()
    clean_wms_co_id = int(wms_co_id) if wms_co_id else _resolve_huadong_wms_co_id(config)
    rows: list[dict[str, Any]] = []
    api_rows = 0
    try:
        query_sleep_seconds = max(
            0.0,
            float(_env_first("JUSHUITAN_QUERY_SLEEP_SECONDS", "JST_QUERY_SLEEP_SECONDS") or "0.35"),
        )
    except ValueError:
        query_sleep_seconds = 0.35
    try:
        query_max_workers = max(
            1,
            min(16, int(_env_first("JUSHUITAN_QUERY_MAX_WORKERS", "JST_QUERY_MAX_WORKERS") or "2")),
        )
    except ValueError:
        query_max_workers = 2
    try:
        query_start_gap_seconds = max(
            0.0,
            float(_env_first("JUSHUITAN_QUERY_START_GAP_SECONDS", "JST_QUERY_START_GAP_SECONDS") or "0.8"),
        )
    except ValueError:
        query_start_gap_seconds = 0.8
    try:
        query_max_retries = max(
            0,
            min(8, int(_env_first("JUSHUITAN_QUERY_MAX_RETRIES", "JST_QUERY_MAX_RETRIES") or "3")),
        )
    except ValueError:
        query_max_retries = 3
    try:
        query_retry_base_seconds = max(
            0.1,
            float(_env_first("JUSHUITAN_QUERY_RETRY_BASE_SECONDS", "JST_QUERY_RETRY_BASE_SECONDS") or "1.0"),
        )
    except ValueError:
        query_retry_base_seconds = 1.0
    try:
        query_batch_size = max(
            1,
            min(100, int(_env_first("JUSHUITAN_QUERY_BATCH_SIZE", "JST_QUERY_BATCH_SIZE") or "2")),
        )
    except ValueError:
        query_batch_size = 2
    sku_id_codes: list[str] = []
    i_id_codes: list[str] = []
    for item in targets:
        code = item["query_code"]
        sku_id_codes.append(code)
        if item.get("query_code_source") == "JST_LOOKUP":
            i_id_codes.append(code)
    chunks: list[tuple[list[str], str]] = []
    for query_field, codes in (("sku_ids", sku_id_codes), ("i_ids", i_id_codes)):
        if not codes:
            continue
        chunks.extend(
            (codes[index : index + query_batch_size], query_field)
            for index in range(0, len(codes), query_batch_size)
        )
    query_errors: list[dict[str, Any]] = []
    processed_chunks = 0
    last_progress_emit_at = 0.0

    def _progress_payload(phase: str) -> dict[str, Any]:
        percent = round((processed_chunks / len(chunks)) * 100, 1) if chunks else 100.0
        return {
            "progress": {
                "phase": phase,
                "processed": processed_chunks,
                "total": len(chunks),
                "percent": min(percent, 100.0),
                "accepted": len(rows),
                "api_rows": api_rows,
                "errors": len(query_errors),
                "total_skus": target_stats["total_skus"],
                "requested_skus": target_stats["requested_skus"],
                "requested_query_codes": target_stats["requested_query_codes"],
                "skipped_without_barcode": target_stats["skipped_without_barcode"],
                "query_max_workers": query_max_workers,
            }
        }

    def _emit_progress(phase: str = "querying", force: bool = False) -> None:
        nonlocal last_progress_emit_at
        if not progress_callback:
            return
        now = time.time()
        if not force and now - last_progress_emit_at < 2:
            return
        last_progress_emit_at = now
        progress_callback(_progress_payload(phase))

    _emit_progress("querying", force=True)

    def _post_inventory_page_with_retry(
        chunk: list[str],
        page_index: int,
        query_field: str,
    ) -> dict[str, Any]:
        last_error: Exception | None = None
        for attempt in range(query_max_retries + 1):
            try:
                return _jushuitan_post(
                    "/open/inventory/query",
                    {
                        "page_index": page_index,
                        "page_size": 100,
                        query_field: ",".join(chunk),
                        "wms_co_id": clean_wms_co_id,
                        "has_lock_qty": True,
                    },
                    config,
                )
            except Exception as exc:
                last_error = exc
                if attempt >= query_max_retries or not _is_retryable_jushuitan_error(exc):
                    raise
                time.sleep(query_retry_base_seconds * (2**attempt))
        raise last_error or RuntimeError("Jushuitan inventory query failed")

    def _query_inventory_chunk(
        chunk: list[str],
        query_field: str,
    ) -> tuple[list[dict[str, Any]], int]:
        chunk_rows: list[dict[str, Any]] = []
        chunk_api_rows = 0
        page_index = 1
        while True:
            payload = _post_inventory_page_with_retry(chunk, page_index, query_field)
            data = payload.get("data")
            items = [item for item in _first_list(data) if isinstance(item, dict)]
            chunk_api_rows += len(items)
            for item in items:
                row = _jushuitan_inventory_item_to_supply_row(item)
                if row:
                    mapped_skus: list[str] = []
                    for returned_code in _jushuitan_inventory_item_codes(item):
                        mapped_skus = query_to_skus.get(returned_code, [])
                        if mapped_skus:
                            break
                    if not mapped_skus and len(chunk) == 1:
                        mapped_skus = query_to_skus.get(chunk[0], [])
                    for mapped_sku in mapped_skus:
                        mapped_row = dict(row)
                        mapped_row["sku_code"] = mapped_sku
                        chunk_rows.append(mapped_row)
            has_next = bool(data.get("has_next") or data.get("hasNext")) if isinstance(data, dict) else False
            if not has_next or not items:
                break
            page_index += 1
        return chunk_rows, chunk_api_rows

    if query_max_workers <= 1 or len(chunks) <= 1:
        for index, (chunk, query_field) in enumerate(chunks):
            if index > 0 and query_sleep_seconds > 0:
                time.sleep(query_sleep_seconds)
            try:
                chunk_rows, chunk_api_rows = _query_inventory_chunk(chunk, query_field)
                rows.extend(chunk_rows)
                api_rows += chunk_api_rows
            except Exception as exc:
                query_errors.append(
                    {
                        "query_field": query_field,
                        "query_code": ",".join(chunk),
                        "error": str(exc)[:300],
                    }
                )
            processed_chunks += 1
            _emit_progress("querying", force=processed_chunks == len(chunks))
    else:
        with ThreadPoolExecutor(max_workers=min(query_max_workers, len(chunks))) as executor:
            future_to_chunk = {}
            next_chunk_index = 0

            def _submit_next_chunk() -> None:
                nonlocal next_chunk_index
                if next_chunk_index >= len(chunks):
                    return
                if next_chunk_index > 0 and query_start_gap_seconds > 0:
                    time.sleep(query_start_gap_seconds)
                chunk, query_field = chunks[next_chunk_index]
                next_chunk_index += 1
                future_to_chunk[executor.submit(_query_inventory_chunk, chunk, query_field)] = (
                    chunk,
                    query_field,
                )

            for _ in range(min(query_max_workers, len(chunks))):
                _submit_next_chunk()

            while future_to_chunk:
                future = next(as_completed(list(future_to_chunk)))
                chunk, query_field = future_to_chunk.pop(future)
                try:
                    chunk_rows, chunk_api_rows = future.result()
                    rows.extend(chunk_rows)
                    api_rows += chunk_api_rows
                except Exception as exc:
                    query_errors.append(
                        {
                            "query_field": query_field,
                            "query_code": ",".join(chunk),
                            "error": str(exc)[:300],
                        }
                    )
                processed_chunks += 1
                _emit_progress("querying", force=processed_chunks == len(chunks))
                _submit_next_chunk()

    if query_errors and not rows:
        raise ValueError(f"聚水潭库存同步全部失败: {query_errors[0]['error']}")

    _emit_progress("saving", force=True)
    result = upsert_supply_stock_snapshot(
        rows,
        db_engine,
        warehouse_code=SUPPLY_HUB_WAREHOUSE,
        source="JUSHUITAN:/open/inventory/query",
    )
    accepted_skus = sorted({row["sku_code"] for row in rows if row.get("sku_code")})
    recalculate_result: dict[str, Any] | None = None
    recalculated = False
    if recalculate and rows:
        _emit_progress("recalculating", force=True)
        recalculate_result = recalculate_inventory_decisions(
            snapshot_date=snapshot_date,
            engine=db_engine,
            sku_codes=accepted_skus,
        )
        recalculated = True
    return {
        **result,
        **target_stats,
        "requested_query_codes": len(query_codes),
        "requested_batch_size": query_batch_size,
        "requested_chunks": len(chunks),
        "query_max_workers": query_max_workers,
        "query_start_gap_seconds": query_start_gap_seconds,
        "query_max_retries": query_max_retries,
        "api_rows": api_rows,
        "accepted": len(rows),
        "accepted_skus": accepted_skus,
        "query_error_count": len(query_errors),
        "query_errors": query_errors[:20],
        "progress": {
            **_progress_payload("done")["progress"],
            "processed": len(chunks),
            "percent": 100.0,
            "accepted": len(rows),
            "api_rows": api_rows,
            "errors": len(query_errors),
        },
        "wms_co_id": clean_wms_co_id,
        "recalculated": recalculated,
        "recalculate_result": recalculate_result,
    }


def list_jushuitan_wms_partners() -> dict[str, Any]:
    config = _load_jushuitan_config()
    partners = _query_jushuitan_wms_partners(config)
    huadong_id: int | None = None
    try:
        huadong_id = _resolve_huadong_wms_co_id(config)
    except ValueError:
        huadong_id = None
    return {"items": partners, "huadong_wms_co_id": huadong_id}


def _bool_from_db(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    try:
        return int(value or 0) == 1
    except (TypeError, ValueError):
        return False


def _sync_setting_to_dict(row: Any) -> dict[str, Any]:
    if not row:
        return {
            "enabled": False,
            "interval_minutes": 10,
            "recalculate": False,
            "last_run_at": None,
            "last_status": "IDLE",
            "last_message": "",
            "last_result": None,
        }
    last_result = row.get("last_result")
    if isinstance(last_result, str) and last_result:
        try:
            last_result = json.loads(last_result)
        except json.JSONDecodeError:
            pass
    last_run_at = row.get("last_run_at")
    updated_at = row.get("updated_at")
    return {
        "enabled": _bool_from_db(row.get("enabled")),
        "interval_minutes": int(row.get("interval_minutes") or 10),
        "recalculate": _bool_from_db(row.get("recalculate")),
        "last_run_at": last_run_at.isoformat(sep=" ") if hasattr(last_run_at, "isoformat") else last_run_at,
        "last_status": row.get("last_status") or "IDLE",
        "last_message": row.get("last_message") or "",
        "last_result": last_result,
        "updated_at": updated_at.isoformat(sep=" ") if hasattr(updated_at, "isoformat") else updated_at,
    }


def get_huadong_sync_setting(engine: Engine | None = None) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.connect() as conn:
        row = conn.execute(
            text(
                """
                SELECT enabled, interval_minutes, recalculate, last_run_at,
                       last_status, last_message, last_result, updated_at
                FROM inv_huadong_sync_setting
                WHERE id=1
                """
            )
        ).mappings().fetchone()
    return _sync_setting_to_dict(row)


def save_huadong_sync_setting(
    enabled: bool,
    interval_minutes: int,
    recalculate: bool = True,
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_interval = max(1, min(1440, int(interval_minutes or 10)))
    with db_engine.begin() as conn:
        conn.execute(
            text(
                """
                INSERT INTO inv_huadong_sync_setting
                    (id, enabled, interval_minutes, recalculate, last_status, last_message, created_at, updated_at)
                VALUES
                    (1, :enabled, :interval_minutes, :recalculate, 'IDLE', '', NOW(), NOW())
                ON DUPLICATE KEY UPDATE
                    enabled=VALUES(enabled),
                    interval_minutes=VALUES(interval_minutes),
                    recalculate=VALUES(recalculate),
                    updated_at=NOW()
                """
            ),
            {
                "enabled": 1 if enabled else 0,
                "interval_minutes": clean_interval,
                "recalculate": 1 if recalculate else 0,
            },
        )
    return get_huadong_sync_setting(db_engine)


def _update_huadong_sync_status(
    engine: Engine,
    status: str,
    message: str = "",
    result: dict[str, Any] | None = None,
    touch_run_at: bool = False,
) -> None:
    with engine.begin() as conn:
        conn.execute(
            text(
                """
                UPDATE inv_huadong_sync_setting
                SET last_status=:status,
                    last_message=:message,
                    last_result=:last_result,
                    last_run_at=IF(:touch_run_at=1, NOW(), last_run_at),
                    updated_at=NOW()
                WHERE id=1
                """
            ),
            {
                "status": status,
                "message": message[:1000],
                "last_result": json.dumps(result, ensure_ascii=False, default=str) if result is not None else None,
                "touch_run_at": 1 if touch_run_at else 0,
            },
        )


def run_huadong_sync_once(engine: Engine | None = None) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    setting = get_huadong_sync_setting(db_engine)
    if not setting.get("enabled"):
        return {"skipped": True, "reason": "花东库存自动同步未启用", "setting": setting}

    _update_huadong_sync_status(db_engine, "RUNNING", "正在拉取聚水潭花东库存", touch_run_at=True)
    running_message = "正在拉取聚水潭花东库存"

    def _sync_progress_callback(progress_result: dict[str, Any]) -> None:
        _update_huadong_sync_status(db_engine, "RUNNING", running_message, progress_result)

    try:
        result = sync_jushuitan_huadong_stock(
            engine=db_engine,
            recalculate=bool(setting.get("recalculate")),
            progress_callback=_sync_progress_callback,
        )
    except Exception as exc:
        _update_huadong_sync_status(db_engine, "FAILED", str(exc), {"error": str(exc)})
        raise

    recalculated_sku_count = (result.get("recalculate_result") or {}).get("sku_count") or 0
    recalculate_note = f"，已局部重算 {recalculated_sku_count} 个 SKU" if result.get("recalculated") else ""
    message = (
        f"全量 SKU {result.get('total_skus', 0)} 个，"
        f"查询唯一条形码 {result.get('requested_query_codes', 0)} 个，"
        f"成功同步 {len(result.get('accepted_skus') or [])} 个 SKU"
    )
    skipped_without_barcode = int(result.get("skipped_without_barcode") or 0)
    if skipped_without_barcode:
        message = f"{message}，无条形码跳过 {skipped_without_barcode} 个"
    message = f"{message}{recalculate_note}"
    error_count = int(result.get("query_error_count") or 0)
    if error_count:
        message = f"{message}，{error_count} 个查询失败等待下轮重试"
    _update_huadong_sync_status(db_engine, "SUCCESS", message, result)
    return {**result, "setting": get_huadong_sync_setting(db_engine)}


def get_sku_replenishment_decision(
    sku_code: str,
    snapshot_date: str | None = None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_sku = _clean_sku(sku_code)
    if not clean_sku:
        raise ValueError("SKU不能为空")
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=False)
    warehouses = list_sku_warehouses(db_engine, clean_sku, str(clean_date))
    plans = list_transfer_plans(db_engine, str(clean_date), None, None, clean_sku, None, None, None, None, 1, 50)["items"]

    with db_engine.connect() as conn:
        sku_row = conn.execute(
            text(
                """
                SELECT sku_code, product_name, brand, category, shop_name, barcode
                FROM inv_sku
                WHERE sku_code=:sku_code
                """
            ),
            {"sku_code": clean_sku},
        ).mappings().fetchone()

    policy = _resolve_policy(_load_replenishment_policies(db_engine), sku_row["brand"] if sku_row else "", sku_row["category"] if sku_row else "")
    hub_stock = _latest_supply_stock(db_engine, clean_sku, policy.preferred_source_warehouse)
    high_risks = [row for row in warehouses if row.get("risk_level") == "HIGH"]
    return {
        "snapshot_date": str(clean_date),
        "sku": _jsonify_mapping(sku_row) if sku_row else {"sku_code": clean_sku},
        "policy": _policy_to_dict(policy),
        "supply_hub": {
            "warehouse_code": policy.preferred_source_warehouse,
            "stock": hub_stock,
        },
        "summary": {
            "warehouse_count": len(warehouses),
            "high_risk_count": len(high_risks),
            "plan_count": len(plans),
            "allow_jd_transfer": policy.allow_jd_transfer,
        },
        "warehouses": warehouses,
        "plans": plans,
    }


def list_sku_warehouses(
    engine: Engine | None = None,
    sku_code: str | None = None,
    snapshot_date: str | None = None,
) -> list[dict[str, Any]]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_sku = _clean_sku(sku_code)
    if not clean_sku:
        return []
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date, allow_none=True)
    if clean_date is None:
        return []

    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT
                    s.snapshot_date,
                    s.sku_code,
                    s.warehouse_code,
                    wh.warehouse_name,
                    wh.region,
                    s.stock_qty,
                    s.inbound_qty,
                    s.sales_1d,
                    s.sales_7d,
                    s.sales_30d,
                    r.daily_sales,
                    r.target_days,
                    r.demand_30d,
                    r.days_of_supply,
                    r.risk_level,
                    r.need_qty,
                    r.excess_qty
                FROM inv_stock_snapshot s
                LEFT JOIN inv_warehouse wh ON wh.warehouse_code=s.warehouse_code
                LEFT JOIN inv_risk_score r
                  ON r.snapshot_date=s.snapshot_date
                 AND r.sku_code=s.sku_code
                 AND r.warehouse_code=s.warehouse_code
                WHERE s.snapshot_date=:snapshot_date
                  AND s.sku_code=:sku_code
                  AND s.is_decision_active=1
                ORDER BY
                    FIELD(r.risk_level, 'HIGH', 'MID', 'LOW'),
                    s.stock_qty DESC,
                    s.sales_30d DESC,
                    s.warehouse_code ASC
                """
            ),
            {"snapshot_date": clean_date, "sku_code": clean_sku},
        ).mappings().all()
    return [_jsonify_mapping(row) for row in rows]


def simulate_transfer(
    sku_code: str,
    from_warehouse: str,
    to_warehouse: str,
    transfer_qty: float,
    snapshot_date: str | None = None,
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_date = _resolve_snapshot_date(db_engine, snapshot_date)
    clean_sku = _clean_sku(sku_code)
    qty = max(_to_float(transfer_qty), 0)
    if qty <= 0:
        raise ValueError("调拨数量必须大于0")

    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT s.*, sku.product_name, sku.brand, sku.category
                FROM inv_stock_snapshot s
                LEFT JOIN inv_sku sku ON sku.sku_code=s.sku_code
                WHERE s.snapshot_date=:d
                  AND s.sku_code=:sku
                  AND s.warehouse_code IN (:from_wh, :to_wh)
                  AND s.is_decision_active=1
                """
            ),
            {"d": clean_date, "sku": clean_sku, "from_wh": from_warehouse, "to_wh": to_warehouse},
        ).mappings().all()

    by_wh = {str(row["warehouse_code"]): row for row in rows}
    if from_warehouse not in by_wh:
        raise ValueError("调出仓没有该SKU的有效库存快照")
    if to_warehouse not in by_wh:
        raise ValueError("调入仓没有该SKU的有效库存快照")

    src = by_wh[from_warehouse]
    dst = by_wh[to_warehouse]
    src_daily = _daily_sales(src["sales_7d"])
    dst_daily = _daily_sales(dst["sales_7d"])
    src_stock_after = _to_float(src["stock_qty"]) - qty
    dst_stock_after = _to_float(dst["stock_qty"]) + qty
    src_days_after = _days_of_supply(src_stock_after, src_daily, allow_no_sales_stock=True)
    dst_days_after = _days_of_supply(dst_stock_after, dst_daily, allow_no_sales_stock=True)
    route = estimate_transfer_route(from_warehouse, to_warehouse)
    safe = src_stock_after >= 0 and src_days_after >= DONOR_RESERVE_DAYS
    eta_late = dst_daily > 0 and route["eta_days"] > _days_of_supply(_to_float(dst["stock_qty"]), dst_daily)

    return {
        "snapshot_date": str(clean_date),
        "sku_code": clean_sku,
        "product_name": src.get("product_name") or dst.get("product_name") or "",
        "brand": src.get("brand") or dst.get("brand") or "",
        "category": src.get("category") or dst.get("category") or "",
        "from_warehouse": from_warehouse,
        "to_warehouse": to_warehouse,
        "transfer_qty": qty,
        "eta_days": route["eta_days"],
        "distance_km": route["distance_km"],
        "from_stock_before": _to_float(src["stock_qty"]),
        "from_stock_after": src_stock_after,
        "from_days_after": round(src_days_after, 2),
        "to_stock_before": _to_float(dst["stock_qty"]),
        "to_stock_after": dst_stock_after,
        "to_days_after": round(dst_days_after, 2),
        "is_safe": bool(safe and not eta_late),
        "warnings": [
            msg
            for msg in [
                "调出后库存小于0" if src_stock_after < 0 else "",
                f"调出后低于{int(DONOR_RESERVE_DAYS)}天安全库存" if src_stock_after >= 0 and src_days_after < DONOR_RESERVE_DAYS else "",
                "预计到货晚于调入仓当前可售天数" if eta_late else "",
            ]
            if msg
        ],
    }


def update_plan_reason(plan_id: int, reason: str, engine: Engine | None = None) -> None:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.begin() as conn:
        conn.execute(
            text("UPDATE inv_transfer_plan SET reason=:reason, updated_at=NOW() WHERE id=:id"),
            {"id": plan_id, "reason": reason},
        )


def get_transfer_plan_for_explain(plan_id: int, engine: Engine | None = None) -> dict[str, Any] | None:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    with db_engine.connect() as conn:
        row = conn.execute(
            text(
                """
                SELECT
                    p.id, p.sku_code, sku.product_name, sku.brand, sku.category,
                    p.from_warehouse, p.to_warehouse, p.transfer_qty, p.eta_days,
                    p.risk_level, p.risk_days, p.surplus_days_after, p.reason,
                    p.eta_late, p.source_type, p.route_path, p.target_days, p.reserve_days,
                    p.supplier_eta_days, p.hub_eta_days, p.score, p.plan_date,
                    r.daily_sales, r.target_days AS risk_target_days, r.days_of_supply, r.demand_30d, r.shortage_days,
                    r.excess_qty, r.need_qty, r.available_qty, r.inbound_credit_qty,
                    r.inbound_credit_rate, r.trend_factor, r.volatility_factor,
                    r.confidence_score, r.consecutive_high_days,
                    from_snap.stock_qty AS from_stock_qty,
                    from_snap.inbound_qty AS from_inbound_qty,
                    from_snap.sales_1d AS from_sales_1d,
                    from_snap.sales_7d AS from_sales_7d,
                    from_snap.sales_30d AS from_sales_30d,
                    to_snap.stock_qty AS to_stock_qty,
                    to_snap.inbound_qty AS to_inbound_qty,
                    to_snap.sales_1d AS to_sales_1d,
                    to_snap.sales_7d AS to_sales_7d,
                    to_snap.sales_30d AS to_sales_30d
                FROM inv_transfer_plan p
                LEFT JOIN inv_sku sku ON sku.sku_code=p.sku_code
                LEFT JOIN inv_risk_score r
                  ON r.snapshot_date=p.plan_date
                 AND r.sku_code=p.sku_code
                 AND r.warehouse_code=p.to_warehouse
                LEFT JOIN inv_stock_snapshot from_snap
                  ON from_snap.snapshot_date=p.plan_date
                 AND from_snap.sku_code=p.sku_code
                 AND from_snap.warehouse_code=p.from_warehouse
                LEFT JOIN inv_stock_snapshot to_snap
                  ON to_snap.snapshot_date=p.plan_date
                 AND to_snap.sku_code=p.sku_code
                 AND to_snap.warehouse_code=p.to_warehouse
                WHERE p.id=:id
                """
            ),
            {"id": plan_id},
        ).mappings().fetchone()
    return _jsonify_mapping(row) if row else None


def estimate_transfer_route(from_warehouse: str, to_warehouse: str) -> dict[str, Any]:
    from_city = _base_city(from_warehouse)
    to_city = _base_city(to_warehouse)
    if from_warehouse in {SUPPLY_HUB_WAREHOUSE, NANSHA_SUPPLY_WAREHOUSE} and to_city in HUADONG_ROUTE_ETA_DAYS:
        return {
            "distance_km": 0.0,
            "eta_days": HUADONG_ROUTE_ETA_DAYS[to_city],
            "source": "MANUAL_SUPPLY_ROUTE",
        }
    if from_city == to_city:
        return {"distance_km": 0.0, "eta_days": 1, "source": "ESTIMATED"}

    from_coord = CITY_COORDINATES.get(from_city)
    to_coord = CITY_COORDINATES.get(to_city)
    if from_coord and to_coord:
        distance = _haversine_km(from_coord[0], from_coord[1], to_coord[0], to_coord[1])
    else:
        from_region = REGION_BY_CITY.get(from_city, "其他")
        to_region = REGION_BY_CITY.get(to_city, "其他")
        distance = 500.0 if from_region == to_region else 1400.0

    eta = _eta_days(distance, from_city, to_city)
    return {"distance_km": round(distance, 2), "eta_days": eta, "source": "ESTIMATED"}


def estimate_logistics_eta(origin: str, destination: str) -> dict[str, Any]:
    """Estimate point-to-point logistics ETA for an arbitrary supplier/city text."""

    clean_origin = str(origin or "").strip()
    clean_destination = str(destination or "").strip()
    from_city = _base_city(clean_origin)
    to_city = _base_city(clean_destination)
    route = estimate_transfer_route(clean_origin, clean_destination)
    has_from_coord = from_city in CITY_COORDINATES
    has_to_coord = to_city in CITY_COORDINATES
    confidence = "HIGH" if has_from_coord and has_to_coord else "MID"
    if not clean_origin or not clean_destination:
        confidence = "LOW"

    eta_days = int(route["eta_days"])
    eta_range = [eta_days, eta_days + 1] if confidence != "HIGH" else [eta_days, eta_days]
    return {
        "origin": clean_origin,
        "destination": clean_destination,
        "from_city": from_city,
        "to_city": to_city,
        "distance_km": route["distance_km"],
        "eta_days": eta_days,
        "eta_range_days": eta_range,
        "confidence": confidence,
        "source": "ESTIMATED_BY_CITY_DISTANCE",
        "note": "基于省份/城市基准点和距离规则预估，未接入真实承运商线路时效，可后续用物流数据覆盖。",
    }


def _parse_inventory_workbook(excel_path: Path, batch_no: str) -> dict[str, Any]:
    with ZipFile(excel_path) as zf:
        shared_strings = _load_shared_strings(zf)
        sheet_path = _find_first_sheet_path(zf)
        headers = _read_header_row(zf, sheet_path, shared_strings)

        header_to_col = {value: col for col, value in headers.items() if value}
        missing = [name for name in ("时间", "商品名称", "SKU", "品牌", "店铺名称") if name not in header_to_col]
        if missing:
            raise ValueError(f"库存Excel缺少必要列: {', '.join(missing)}")

        base_cols = {key: header_to_col.get(label) for key, label in BASE_COLUMNS.items()}
        if not base_cols.get("supplier_name"):
            base_cols["supplier_name"] = next(
                (header_to_col.get(label) for label in SUPPLIER_NAME_COLUMN_ALIASES if header_to_col.get(label)),
                None,
            )
        category_l1_col = header_to_col.get(IGEAR_CATEGORY_COLUMN)
        prefix_cols: dict[str, dict[str, int]] = defaultdict(dict)
        for col, header in headers.items():
            if not isinstance(header, str):
                continue
            for field, suffixes in METRIC_SUFFIXES.items():
                matched = next((suffix for suffix in suffixes if header.endswith(suffix)), None)
                if matched is None:
                    continue
                prefix = header[: -len(matched)].strip()
                if prefix:
                    prefix_cols[prefix][field] = col
                break

        warehouses = [
            prefix
            for prefix, cols in prefix_cols.items()
            if prefix != NATIONAL_PREFIX and all(field in cols for field in REQUIRED_METRIC_FIELDS)
        ]
        if not warehouses:
            raise ValueError("库存Excel未识别到城市/仓库库存字段")

        required_cols: set[int] = {col for col in base_cols.values() if col}
        if category_l1_col is not None:
            required_cols.add(category_l1_col)
        for prefix in [NATIONAL_PREFIX, *warehouses]:
            for col in prefix_cols.get(prefix, {}).values():
                required_cols.add(col)

        skus: dict[str, dict[str, Any]] = {}
        snapshots: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        rows_read = 0
        snapshot_date: date | None = None

        for row_values in _iter_sheet_rows(zf, sheet_path, shared_strings, required_cols):
            row_number = int(row_values.pop("__row_number__", 0))
            if row_number <= 1:
                continue
            rows_read += 1
            sku_code = _clean_sku(row_values.get(base_cols["sku_code"]))
            if not sku_code:
                _append_error(errors, batch_no, None, "", "", "MISSING_SKU", "ERROR", "SKU为空", row_number)
                continue

            row_date = _parse_snapshot_date(row_values.get(base_cols["snapshot_date"]), excel_path.name)
            if row_date:
                snapshot_date = snapshot_date or row_date
            else:
                _append_error(errors, batch_no, None, sku_code, "", "MISSING_DATE", "ERROR", "时间字段无法识别", row_number)
                continue

            product_name = _clean_text(row_values.get(base_cols["product_name"]))
            supplier_name = _clean_text(row_values.get(base_cols.get("supplier_name")))
            brand = _normalize_inventory_brand_by_supplier(row_values.get(base_cols["brand"]), supplier_name)
            # iGear 品牌按一级类目拆分：工业品 -> iGear工业品；家装建材 -> iGear；其余类目不导入。
            # 该规则仅在表含「一级类目」列时生效；旧格式表无此列则保持原样导入。
            if category_l1_col is not None and _normalize_brand_key(brand) == _normalize_brand_key(IGEAR_BRAND):
                category_l1 = _clean_text(row_values.get(category_l1_col))
                if category_l1 == IGEAR_CATEGORY_INDUSTRIAL:
                    brand = IGEAR_INDUSTRIAL_BRAND
                elif category_l1 == IGEAR_CATEGORY_BUILDING_MATERIAL:
                    brand = IGEAR_BRAND
                else:
                    _append_error(
                        errors,
                        batch_no,
                        row_date,
                        sku_code,
                        "",
                        "IGEAR_CATEGORY_DROPPED",
                        "WARN",
                        "iGear非工业品/家装建材类目，按规则不导入",
                        category_l1 or "(空)",
                    )
                    continue
            shop_name = _clean_text(row_values.get(base_cols["shop_name"]))
            barcode = _clean_barcode(row_values.get(base_cols.get("barcode")))
            product_status = _clean_text(row_values.get(base_cols.get("product_status")))
            gift_text = _clean_text(row_values.get(base_cols.get("gift_flag")))
            is_gift = 0 if gift_text in ("", "非赠品", "否", "0") else 1
            is_decision_active = 1 if product_status in ACTIVE_STATUSES and is_gift == 0 else 0

            skus[sku_code] = {
                "sku_code": sku_code,
                "product_name": product_name,
                "brand": brand,
                "category": _brand_category(brand),
                "shop_name": shop_name,
                "supplier_name": supplier_name,
                "barcode": barcode,
            }

            national_stock = _metric(row_values, prefix_cols, NATIONAL_PREFIX, "stock_qty")
            national_inbound = _metric(row_values, prefix_cols, NATIONAL_PREFIX, "inbound_qty")
            city_stock_sum = 0.0
            city_inbound_sum = 0.0

            for warehouse_code in warehouses:
                stock_qty = _metric(row_values, prefix_cols, warehouse_code, "stock_qty")
                inbound_qty = _metric(row_values, prefix_cols, warehouse_code, "inbound_qty")
                available_qty = _metric(row_values, prefix_cols, warehouse_code, "available_qty")
                orderable_qty = _metric(row_values, prefix_cols, warehouse_code, "orderable_qty")
                sales_1d = _metric(row_values, prefix_cols, warehouse_code, "sales_1d")
                sales_7d = _metric(row_values, prefix_cols, warehouse_code, "sales_7d")
                sales_14d = _metric(row_values, prefix_cols, warehouse_code, "sales_14d")
                sales_30d = _metric(row_values, prefix_cols, warehouse_code, "sales_30d")

                if (
                    stock_qty == 0
                    and inbound_qty == 0
                    and available_qty == 0
                    and orderable_qty == 0
                    and sales_1d == 0
                    and sales_7d == 0
                    and sales_14d == 0
                    and sales_30d == 0
                ):
                    continue

                city_stock_sum += stock_qty
                city_inbound_sum += inbound_qty

                if stock_qty < 0:
                    _append_error(errors, batch_no, row_date, sku_code, warehouse_code, "NEGATIVE_STOCK", "WARN", "存在负库存", stock_qty)
                if sales_1d > max(stock_qty + inbound_qty, 0):
                    _append_error(errors, batch_no, row_date, sku_code, warehouse_code, "SALES_GT_STOCK", "WARN", "昨日销量大于库存与未到货合计", f"sales_1d={sales_1d},stock={stock_qty},inbound={inbound_qty}")
                if stock_qty <= 0 and (sales_1d > 0 or sales_7d > 0 or sales_30d > 0):
                    _append_error(errors, batch_no, row_date, sku_code, warehouse_code, "ZERO_STOCK_WITH_SALES", "WARN", "无库存但存在出库销量", f"sales_7d={sales_7d},sales_30d={sales_30d}")

                snapshots.append(
                    {
                        "batch_no": batch_no,
                        "snapshot_date": row_date,
                        "sku_code": sku_code,
                        "warehouse_code": warehouse_code,
                        "product_status": product_status,
                        "is_gift": is_gift,
                        "is_decision_active": is_decision_active,
                        "stock_qty": stock_qty,
                        "inbound_qty": inbound_qty,
                        "available_qty": available_qty,
                        "orderable_qty": orderable_qty,
                        "sales_1d": sales_1d,
                        "sales_7d": sales_7d,
                        "sales_14d": sales_14d,
                        "sales_30d": sales_30d,
                    }
                )

            if national_stock and abs(national_stock - city_stock_sum) > max(5.0, abs(national_stock) * 0.05):
                _append_error(
                    errors,
                    batch_no,
                    row_date,
                    sku_code,
                    NATIONAL_PREFIX,
                    "NATIONAL_STOCK_MISMATCH",
                    "WARN",
                    "全国库存与各城市库存合计偏差超过5%",
                    f"national={national_stock},city_sum={round(city_stock_sum, 2)}",
                )
            if national_inbound and abs(national_inbound - city_inbound_sum) > max(5.0, abs(national_inbound) * 0.05):
                _append_error(
                    errors,
                    batch_no,
                    row_date,
                    sku_code,
                    NATIONAL_PREFIX,
                    "NATIONAL_INBOUND_MISMATCH",
                    "WARN",
                    "全国采购未到货与各城市合计偏差超过5%",
                    f"national={national_inbound},city_sum={round(city_inbound_sum, 2)}",
                )

    warehouse_rows = [_warehouse_record(code) for code in warehouses]
    return {
        "snapshot_date": snapshot_date,
        "rows_read": rows_read,
        "skus": skus,
        "warehouses": warehouse_rows,
        "snapshots": snapshots,
        "errors": errors[:MAX_IMPORT_ERRORS_PER_BATCH],
    }


def _load_shared_strings(zf: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    strings: list[str] = []
    with zf.open("xl/sharedStrings.xml") as fh:
        for event, elem in ET.iterparse(fh, events=("end",)):
            if elem.tag == f"{XML_MAIN_NS}si":
                strings.append("".join(t.text or "" for t in elem.iter(f"{XML_MAIN_NS}t")))
                elem.clear()
    return strings


def _find_first_sheet_path(zf: ZipFile) -> str:
    names = zf.namelist()
    if "xl/worksheets/sheet1.xml" in names:
        return "xl/worksheets/sheet1.xml"
    sheets = sorted(name for name in names if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
    if not sheets:
        raise ValueError("Excel文件未找到工作表")
    return sheets[0]


def _read_header_row(zf: ZipFile, sheet_path: str, shared_strings: list[str]) -> dict[int, Any]:
    with zf.open(sheet_path) as fh:
        for event, elem in ET.iterparse(fh, events=("end",)):
            if elem.tag == f"{XML_MAIN_NS}row":
                if elem.attrib.get("r") == "1":
                    headers = {}
                    for cell in elem.iter(f"{XML_MAIN_NS}c"):
                        col = _cell_col_index(cell.attrib.get("r", ""))
                        if col:
                            headers[col] = _cell_value(cell, shared_strings)
                    elem.clear()
                    return headers
                elem.clear()
    raise ValueError("Excel文件未找到表头行")


def import_inventory_monitor_excel(
    file_path: str | os.PathLike[str], engine: Engine | None = None
) -> dict[str, Any]:
    """Import the eight-warehouse operating monitor without touching decision data."""
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    excel_path = Path(file_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel文件不存在: {excel_path}")

    with ZipFile(excel_path) as zf:
        strings = _load_shared_strings(zf)
        sheet_path = _find_first_sheet_path(zf)
        headers = _read_header_row(zf, sheet_path, strings)
        columns = {str(value).strip(): col for col, value in headers.items() if value is not None}
        base_names = ("时间", "商品名称", "SKU", "品牌", "一级类目", "二级类目", "三级类目", "店铺名称")
        required = [*base_names, *(f"全国{suffix}" for _, suffix in MONITOR_METRICS)]
        missing = [name for name in required if name not in columns]
        if missing:
            raise ValueError(f"库存看板Excel缺少字段: {', '.join(missing)}")
        product_status_col = next(
            (columns[name] for name in ("上下柜状态", "上下架状态", "上下柜") if name in columns),
            None,
        )
        city_columns: dict[str, dict[str, int]] = {}
        for city in MONITOR_WAREHOUSES:
            city_columns[city] = {}
            for key, suffix in MONITOR_METRICS:
                header = f"{city}{suffix}"
                if header not in columns:
                    raise ValueError(f"库存看板Excel缺少八大仓字段: {header}")
                city_columns[city][key] = columns[header]

        wanted = {columns[name] for name in base_names}
        if product_status_col is not None:
            wanted.add(product_status_col)
        wanted.update(columns[f"全国{suffix}"] for _, suffix in MONITOR_METRICS)
        wanted.update(col for metric_cols in city_columns.values() for col in metric_cols.values())
        rows: list[dict[str, Any]] = []
        snapshot_date: date | None = None
        for values in _iter_sheet_rows(zf, sheet_path, strings, wanted):
            if values.get("__row_number__", 0) <= 1:
                continue
            sku_code = _clean_sku(values.get(columns["SKU"]))
            row_date = _parse_snapshot_date(values.get(columns["时间"]), excel_path.name)
            if not sku_code or not row_date:
                continue
            snapshot_date = snapshot_date or row_date
            city_metrics = {
                city: {key: _to_float(values.get(col)) for key, col in metric_cols.items()}
                for city, metric_cols in city_columns.items()
            }
            rows.append({
                "snapshot_date": row_date,
                "sku_code": sku_code,
                "product_name": _clean_text(values.get(columns["商品名称"])),
                "brand": _clean_text(values.get(columns["品牌"])),
                "category_l1": _clean_text(values.get(columns["一级类目"])),
                "category_l2": _clean_text(values.get(columns["二级类目"])),
                "category_l3": _clean_text(values.get(columns["三级类目"])),
                "shop_name": _clean_text(values.get(columns["店铺名称"])),
                "product_status": _clean_text(values.get(product_status_col)) if product_status_col else "",
                **{f"national_{key}": _to_float(values.get(columns[f"全国{suffix}"])) for key, suffix in MONITOR_METRICS},
                "city_metrics": json.dumps(city_metrics, ensure_ascii=False),
            })

    if not rows or snapshot_date is None:
        raise ValueError("库存看板Excel未读取到有效SKU数据")
    sql = text("""
        INSERT INTO inv_monitor_snapshot
          (snapshot_date, sku_code, product_name, brand, category_l1, category_l2, category_l3, shop_name, product_status,
           national_stock_qty, national_orderable_qty, national_sales_1d, national_sales_7d, national_sales_30d, city_metrics)
        VALUES
          (:snapshot_date, :sku_code, :product_name, :brand, :category_l1, :category_l2, :category_l3, :shop_name, :product_status,
           :national_stock_qty, :national_orderable_qty, :national_sales_1d, :national_sales_7d, :national_sales_30d, :city_metrics)
        ON DUPLICATE KEY UPDATE
          product_name=VALUES(product_name), brand=VALUES(brand), category_l1=VALUES(category_l1), category_l2=VALUES(category_l2),
          category_l3=VALUES(category_l3), shop_name=VALUES(shop_name), product_status=VALUES(product_status), national_stock_qty=VALUES(national_stock_qty),
          national_orderable_qty=VALUES(national_orderable_qty), national_sales_1d=VALUES(national_sales_1d),
          national_sales_7d=VALUES(national_sales_7d), national_sales_30d=VALUES(national_sales_30d), city_metrics=VALUES(city_metrics)
    """)
    with db_engine.begin() as conn:
        _execute_chunks(conn, sql, rows)
    return {
        "snapshot_date": str(snapshot_date),
        "rows": len(rows),
        "warehouses": list(MONITOR_WAREHOUSES),
        "status_field_detected": product_status_col is not None,
    }


def _monitor_filter_context(
    snapshot_date: list[str] | str | None = None,
    keyword: str = "",
    brand: list[str] | str | None = None,
    category_l1: list[str] | str | None = None,
    shop_name: list[str] | str | None = None,
    product_status: list[str] | str | None = None,
) -> tuple[str, dict[str, Any]]:
    filters: list[str] = []
    params: dict[str, Any] = {}
    def _norm(value: list[str] | str | None) -> list[str]:
        if value is None:
            return []
        if isinstance(value, str):
            return [value.strip()] if value.strip() else []
        return [str(v).strip() for v in value if v is not None and str(v).strip()]
    def _in_clause(field: str, values: list[str] | str | None, prefix: str) -> None:
        cleaned = _norm(values)
        if not cleaned:
            return
        placeholders = ", ".join(f":{prefix}{i}" for i in range(len(cleaned)))
        for i, value in enumerate(cleaned):
            params[f"{prefix}{i}"] = value
        filters.append(f"{field} IN ({placeholders})")
    dates = _norm(snapshot_date)
    if dates:
        placeholders = ", ".join(f":sd{i}" for i in range(len(dates)))
        for i, value in enumerate(dates):
            params[f"sd{i}"] = value
        filters.append(f"snapshot_date IN ({placeholders})")
    else:
        filters.append("snapshot_date=(SELECT MAX(snapshot_date) FROM inv_monitor_snapshot)")
    if keyword.strip():
        filters.append("(sku_code LIKE :keyword OR product_name LIKE :keyword)")
        params["keyword"] = f"%{keyword.strip()}%"
    _in_clause("brand", brand, "b")
    _in_clause("category_l1", category_l1, "c")
    _in_clause("shop_name", shop_name, "s")
    status_values = _norm(product_status)
    if status_values:
        expanded_statuses: set[str] = set()
        for value in status_values:
            if value in {"上柜", "listed"}:
                expanded_statuses.update(MONITOR_ON_SHELF_STATUSES)
            elif value in {"下柜", "delisted"}:
                expanded_statuses.update(MONITOR_OFF_SHELF_STATUSES)
            else:
                expanded_statuses.add(value)
        _in_clause("product_status", sorted(expanded_statuses), "ps")
    return " WHERE " + " AND ".join(filters), params

def _decode_monitor_city_metrics(value: Any) -> dict[str, Any]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except (TypeError, ValueError):
            return {}
    return value if isinstance(value, dict) else {}

def _monitor_metric_counts(rows: Iterable[dict[str, Any]], turnover_days: float = TARGET_DAYS) -> dict[str, int]:
    counts = {key: 0 for key in MONITOR_METRIC_LABELS}
    for row in rows:
        status = _clean_text(row.get("product_status"))
        if status in MONITOR_ON_SHELF_STATUSES:
            counts["listed"] += 1
        if status in MONITOR_OFF_SHELF_STATUSES:
            counts["delisted"] += 1
        city_metrics = _decode_monitor_city_metrics(row.get("city_metrics"))
        for city in MONITOR_WAREHOUSES:
            metrics = city_metrics.get(city)
            if not isinstance(metrics, dict) or not metrics:
                continue
            stock_qty = _to_float(metrics.get("stock_qty"))
            sales_30d = _to_float(metrics.get("sales_30d"))
            if stock_qty == 0:
                counts["stockout"] += 1
            if (sales_30d / 30.0) * turnover_days - stock_qty > 0:
                counts["shortage"] += 1
    return counts

def _monitor_metric_details(
    rows: Iterable[dict[str, Any]], metric: str, turnover_days: float = TARGET_DAYS
) -> list[dict[str, Any]]:
    details: list[dict[str, Any]] = []
    for row in rows:
        base = {
            "snapshot_date": str(row.get("snapshot_date")),
            "sku_code": str(row.get("sku_code") or ""),
            "product_name": str(row.get("product_name") or ""),
            "brand": str(row.get("brand") or ""),
            "category_l1": str(row.get("category_l1") or ""),
            "category_l2": str(row.get("category_l2") or ""),
            "category_l3": str(row.get("category_l3") or ""),
            "shop_name": str(row.get("shop_name") or ""),
            "product_status": str(row.get("product_status") or ""),
            "national_stock_qty": round(_to_float(row.get("national_stock_qty")), 2),
            "national_orderable_qty": round(_to_float(row.get("national_orderable_qty")), 2),
            "national_sales_1d": round(_to_float(row.get("national_sales_1d")), 2),
            "national_sales_7d": round(_to_float(row.get("national_sales_7d")), 2),
            "national_sales_30d": round(_to_float(row.get("national_sales_30d")), 2),
            "city_metrics": _decode_monitor_city_metrics(row.get("city_metrics")),
        }
        if metric in {"listed", "delisted"}:
            expected = MONITOR_ON_SHELF_STATUSES if metric == "listed" else MONITOR_OFF_SHELF_STATUSES
            if base["product_status"] in expected:
                details.append(base)
            continue
        city_metrics = _decode_monitor_city_metrics(row.get("city_metrics"))
        for city in MONITOR_WAREHOUSES:
            metrics = city_metrics.get(city)
            if not isinstance(metrics, dict) or not metrics:
                continue
            stock_qty = _to_float(metrics.get("stock_qty"))
            sales_30d = _to_float(metrics.get("sales_30d"))
            daily_sales = sales_30d / 30.0
            shortage_qty = daily_sales * turnover_days - stock_qty
            is_match = stock_qty == 0 if metric == "stockout" else shortage_qty > 0
            if not is_match:
                continue
            details.append(
                {
                    **base,
                    "warehouse": f"{city}仓",
                    "stock_qty": round(stock_qty, 2),
                    "sales_30d": round(sales_30d, 2),
                    "daily_sales": round(daily_sales, 4),
                    "turnover_days": round(turnover_days, 2),
                    "target_stock_qty": round(daily_sales * turnover_days, 2),
                    "shortage_qty": round(max(0.0, shortage_qty), 2),
                }
            )
    if metric == "shortage":
        details.sort(key=lambda row: (-float(row.get("shortage_qty") or 0), row["sku_code"], row["warehouse"]))
    else:
        details.sort(key=lambda row: (row.get("shop_name", ""), row["sku_code"], row.get("warehouse", "")))
    return details

def list_inventory_monitor(
    engine: Engine | None = None, snapshot_date: list[str] | None = None, keyword: str = "",
    brand: list[str] | None = None, category_l1: list[str] | None = None,
    shop_name: list[str] | None = None, product_status: list[str] | str | None = None,
    page: int = 1, page_size: int = 50,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    where, params = _monitor_filter_context(snapshot_date, keyword, brand, category_l1, shop_name, product_status)
    page_size_clean = max(1, min(page_size, 200))
    params["offset"] = max(page - 1, 0) * page_size_clean
    params["limit"] = page_size_clean
    with db_engine.connect() as conn:
        total = conn.execute(text(f"SELECT COUNT(*) FROM inv_monitor_snapshot{where}"), params).scalar_one()
        summary = conn.execute(text(f"""SELECT MAX(snapshot_date) snapshot_date, COALESCE(SUM(national_stock_qty),0) stock_qty,
              COALESCE(SUM(national_orderable_qty),0) orderable_qty, COALESCE(SUM(national_sales_1d),0) sales_1d,
              COALESCE(SUM(national_sales_7d),0) sales_7d, COALESCE(SUM(national_sales_30d),0) sales_30d
              FROM inv_monitor_snapshot{where}"""), params).mappings().one()
        metric_rows = conn.execute(text(f"""SELECT snapshot_date, sku_code, product_name, brand, shop_name, product_status,
              national_sales_30d, city_metrics FROM inv_monitor_snapshot{where}"""), params).mappings().all()
        items = conn.execute(text(f"""SELECT snapshot_date, sku_code, product_name, brand, category_l1, category_l2, category_l3, shop_name,
              product_status, national_stock_qty, national_orderable_qty, national_sales_1d, national_sales_7d, national_sales_30d, city_metrics
              FROM inv_monitor_snapshot{where} ORDER BY national_stock_qty DESC, sku_code LIMIT :limit OFFSET :offset"""), params).mappings().all()
        filter_options = {}
        for key, column in (("snapshot_dates", "snapshot_date"), ("brands", "brand"),
                            ("categories", "category_l1"), ("shops", "shop_name")):
            values = conn.execute(text(
                f"SELECT DISTINCT {column} FROM inv_monitor_snapshot "
                f"WHERE {column} IS NOT NULL AND TRIM({column}) <> :empty"
            ), {"empty": ""}).scalars().all()
            filter_options[key] = sorted({str(value) for value in values}, reverse=key == "snapshot_dates")
    output = []
    for item in items:
        row = dict(item)
        row["snapshot_date"] = str(row["snapshot_date"])
        row["city_metrics"] = _decode_monitor_city_metrics(row["city_metrics"])
        output.append(row)
    result_summary = dict(summary)
    result_summary["metrics"] = _monitor_metric_counts(metric_rows)
    result_summary["turnover_days"] = TARGET_DAYS
    return {"summary": result_summary, "items": output, "total": total,
            "warehouses": list(MONITOR_WAREHOUSES), "filter_options": filter_options}

def list_inventory_monitor_details(
    engine: Engine | None = None, metric: str = "stockout",
    snapshot_date: list[str] | None = None, keyword: str = "",
    brand: list[str] | None = None, category_l1: list[str] | None = None,
    shop_name: list[str] | None = None, product_status: list[str] | str | None = None,
    page: int = 1, page_size: int = 50,
    turnover_days: float = TARGET_DAYS,
) -> dict[str, Any]:
    if metric not in MONITOR_METRIC_LABELS:
        raise ValueError(f"不支持的库存指标: {metric}")
    db_engine = get_engine(engine)
    ensure_inventory_schema(db_engine)
    clean_turnover_days = max(1.0, min(float(turnover_days or TARGET_DAYS), 365.0))
    where, params = _monitor_filter_context(snapshot_date, keyword, brand, category_l1, shop_name, product_status)
    with db_engine.connect() as conn:
        rows = conn.execute(text(f"""SELECT snapshot_date, sku_code, product_name, brand, category_l1, category_l2, category_l3,
              shop_name, product_status, national_stock_qty, national_orderable_qty, national_sales_1d,
              national_sales_7d, national_sales_30d, city_metrics
              FROM inv_monitor_snapshot{where}"""), params).mappings().all()
    details = _monitor_metric_details(rows, metric, clean_turnover_days)
    page_size_clean = max(1, min(page_size, 200))
    start = max(page - 1, 0) * page_size_clean
    return {
        "metric": metric,
        "metric_label": MONITOR_METRIC_LABELS[metric],
        "turnover_days": clean_turnover_days,
        "items": details[start:start + page_size_clean],
        "total": len(details),
    }


def _iter_sheet_rows(
    zf: ZipFile,
    sheet_path: str,
    shared_strings: list[str],
    wanted_cols: set[int],
) -> Iterable[dict[int | str, Any]]:
    with zf.open(sheet_path) as fh:
        for event, elem in ET.iterparse(fh, events=("end",)):
            if elem.tag != f"{XML_MAIN_NS}row":
                continue
            row_number = int(elem.attrib.get("r", "0") or "0")
            row_values: dict[int | str, Any] = {"__row_number__": row_number}
            for cell in elem.iter(f"{XML_MAIN_NS}c"):
                col = _cell_col_index(cell.attrib.get("r", ""))
                if col in wanted_cols:
                    row_values[col] = _cell_value(cell, shared_strings)
            elem.clear()
            yield row_values


_CELL_COL_INDEX_CACHE: dict[str, int] = {}


def _cell_col_index(ref: str) -> int | None:
    """Resolve an Excel cell reference prefix (e.g. \"AB\") to a 1-based column index.

    Cell references repeat the same column letters across every row, so the
    letter->index mapping is cached once; the per-cell letter extraction uses a
    manual scan instead of a regex to keep the hot path cheap.
    """
    source = ref or ""
    end = 0
    length = len(source)
    while end < length and "A" <= source[end] <= "Z":
        end += 1
    if end == 0:
        return None
    key = source[:end]
    cached = _CELL_COL_INDEX_CACHE.get(key)
    if cached is not None:
        return cached
    total = 0
    for char in key:
        total = total * 26 + (ord(char) - 64)
    _CELL_COL_INDEX_CACHE[key] = total
    return total


def _cell_value(cell: ET.Element, shared_strings: list[str]) -> Any:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(f"{XML_MAIN_NS}is")
        return "".join(t.text or "" for t in inline.iter(f"{XML_MAIN_NS}t")) if inline is not None else ""
    value_elem = cell.find(f"{XML_MAIN_NS}v")
    if value_elem is None or value_elem.text is None:
        return None
    raw = value_elem.text
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except Exception:
            return ""
    if cell_type == "b":
        return raw == "1"
    return raw


def _upsert_skus(conn: Any, rows: Iterable[dict[str, Any]]) -> None:
    rows_list = list(rows)
    if not rows_list:
        return
    for row in rows_list:
        row["category"] = row.get("category") or _brand_category(row.get("brand"))
        row["barcode"] = _clean_barcode(row.get("barcode"))
        row["supplier_name"] = _clean_text(row.get("supplier_name"))
    sql = text(
        """
        INSERT INTO inv_sku (sku_code, product_name, brand, category, shop_name, supplier_name, barcode, created_at, updated_at)
        VALUES (:sku_code, :product_name, :brand, :category, :shop_name, :supplier_name, :barcode, NOW(), NOW())
        ON DUPLICATE KEY UPDATE
            product_name=VALUES(product_name),
            brand=VALUES(brand),
            category=VALUES(category),
            shop_name=VALUES(shop_name),
            supplier_name=VALUES(supplier_name),
            barcode=IF(VALUES(barcode)<>'', VALUES(barcode), barcode),
            updated_at=NOW()
        """
    )
    _execute_chunks(conn, sql, rows_list)


def _upsert_warehouses(conn: Any, rows: Iterable[dict[str, Any]]) -> None:
    rows_list = list(rows)
    if not rows_list:
        return
    sql = text(
        """
        INSERT INTO inv_warehouse
            (warehouse_code, warehouse_name, region, warehouse_type, is_supply_hub, is_jd_warehouse,
             base_city, latitude, longitude, created_at, updated_at)
        VALUES
            (:warehouse_code, :warehouse_name, :region, :warehouse_type, :is_supply_hub, :is_jd_warehouse,
             :base_city, :latitude, :longitude, NOW(), NOW())
        ON DUPLICATE KEY UPDATE
            warehouse_name=VALUES(warehouse_name),
            region=VALUES(region),
            warehouse_type=VALUES(warehouse_type),
            is_supply_hub=VALUES(is_supply_hub),
            is_jd_warehouse=VALUES(is_jd_warehouse),
            base_city=VALUES(base_city),
            latitude=VALUES(latitude),
            longitude=VALUES(longitude),
            updated_at=NOW()
        """
    )
    _execute_chunks(conn, sql, rows_list)


def _upsert_stock_snapshots(conn: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    sql = text(
        """
        INSERT INTO inv_stock_snapshot
            (batch_no, snapshot_date, sku_code, warehouse_code, product_status, is_gift,
             is_decision_active, stock_qty, inbound_qty, available_qty, orderable_qty,
             sales_1d, sales_7d, sales_14d, sales_30d,
             created_at, updated_at)
        VALUES
            (:batch_no, :snapshot_date, :sku_code, :warehouse_code, :product_status, :is_gift,
             :is_decision_active, :stock_qty, :inbound_qty, :available_qty, :orderable_qty,
             :sales_1d, :sales_7d, :sales_14d, :sales_30d,
             NOW(), NOW())
        ON DUPLICATE KEY UPDATE
            batch_no=VALUES(batch_no),
            product_status=VALUES(product_status),
            is_gift=VALUES(is_gift),
            is_decision_active=VALUES(is_decision_active),
            stock_qty=VALUES(stock_qty),
            inbound_qty=VALUES(inbound_qty),
            available_qty=VALUES(available_qty),
            orderable_qty=VALUES(orderable_qty),
            sales_1d=VALUES(sales_1d),
            sales_7d=VALUES(sales_7d),
            sales_14d=VALUES(sales_14d),
            sales_30d=VALUES(sales_30d),
            updated_at=NOW()
        """
    )
    _execute_chunks(conn, sql, rows)


def _purge_old_inventory_history(
    conn: Any,
    keep_days: int = INVENTORY_HISTORY_KEEP_DAYS,
) -> int:
    """Delete inventory history older than the newest keep_days snapshot dates."""

    keep_days = max(1, int(keep_days))
    cutoff = conn.execute(
        text(
            """
            SELECT DISTINCT snapshot_date
            FROM inv_stock_snapshot
            WHERE snapshot_date IS NOT NULL
            ORDER BY snapshot_date DESC
            LIMIT 1 OFFSET :offset
            """
        ),
        {"offset": keep_days - 1},
    ).scalar()
    if not cutoff:
        return 0

    deleted = 0
    for table, column in (
        ("inv_stock_snapshot", "snapshot_date"),
        ("inv_risk_score", "snapshot_date"),
        ("inv_transfer_plan", "plan_date"),
        ("inv_supply_stock_snapshot", "snapshot_date"),
        ("inv_decision_run", "snapshot_date"),
        ("inv_import_batch", "snapshot_date"),
        ("inv_import_error", "snapshot_date"),
    ):
        result = conn.execute(
            text(f"DELETE FROM {table} WHERE {column} < :cutoff"),
            {"cutoff": cutoff},
        )
        deleted += int(result.rowcount or 0)
    return deleted


def _insert_import_errors(conn: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    sql = text(
        """
        INSERT INTO inv_import_error
            (batch_no, snapshot_date, sku_code, warehouse_code, error_type, severity, message, raw_value, created_at)
        VALUES
            (:batch_no, :snapshot_date, :sku_code, :warehouse_code, :error_type, :severity, :message, :raw_value, NOW())
        """
    )
    _execute_chunks(conn, sql, rows)


def _insert_risk_scores(conn: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    sql = text(
        """
        INSERT INTO inv_risk_score
            (run_id, snapshot_date, sku_code, warehouse_code, daily_sales, target_days, days_of_supply,
             demand_30d, shortage_days, risk_level, excess_qty, need_qty,
             available_qty, inbound_credit_qty, inbound_credit_rate, trend_factor, volatility_factor,
             confidence_score, consecutive_high_days,
             is_shortage_risk, is_excess_risk, created_at)
        VALUES
            (:run_id, :snapshot_date, :sku_code, :warehouse_code, :daily_sales, :target_days, :days_of_supply,
             :demand_30d, :shortage_days, :risk_level, :excess_qty, :need_qty,
             :available_qty, :inbound_credit_qty, :inbound_credit_rate, :trend_factor, :volatility_factor,
             :confidence_score, :consecutive_high_days,
             :is_shortage_risk, :is_excess_risk, NOW())
        """
    )
    _execute_chunks(conn, sql, rows)


def _insert_transfer_plans(conn: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    sql = text(
        """
        INSERT INTO inv_transfer_plan
            (run_id, plan_date, sku_code, from_warehouse, to_warehouse, transfer_qty,
             eta_days, priority, risk_level, risk_days, surplus_days_after, eta_late,
             source_type, route_path, target_days, reserve_days, supplier_eta_days, hub_eta_days,
             score, reason, status,
             created_at, updated_at)
        VALUES
            (:run_id, :plan_date, :sku_code, :from_warehouse, :to_warehouse, :transfer_qty,
             :eta_days, :priority, :risk_level, :risk_days, :surplus_days_after, :eta_late,
             :source_type, :route_path, :target_days, :reserve_days, :supplier_eta_days, :hub_eta_days,
             :score, :reason, 'PENDING',
             NOW(), NOW())
        """
    )
    _execute_chunks(conn, sql, rows)


def _execute_chunks(conn: Any, sql: Any, rows: list[dict[str, Any]], chunk_size: int = 1000) -> None:
    for idx in range(0, len(rows), chunk_size):
        conn.execute(sql, rows[idx : idx + chunk_size])


def _resolve_snapshot_date(
    engine: Engine,
    snapshot_date: str | date | None,
    allow_none: bool = False,
) -> date | None:
    if isinstance(snapshot_date, date):
        return snapshot_date
    if snapshot_date:
        parsed = _parse_date_text(str(snapshot_date))
        if parsed:
            return parsed
        raise ValueError("snapshot_date必须是YYYY-MM-DD格式")
    with engine.connect() as conn:
        value = conn.execute(text("SELECT MAX(snapshot_date) FROM inv_stock_snapshot")).scalar()
    if value:
        return value
    if allow_none:
        return None
    raise ValueError("暂无库存快照数据，请先导入Excel")


def _load_decision_snapshots(
    engine: Engine,
    snapshot_date: date,
    brand: str | None = None,
    category: str | None = None,
    sku_codes: Iterable[str] | None = None,
) -> list[dict[str, Any]]:
    where = ["s.snapshot_date=:d", "s.is_decision_active=1"]
    params: dict[str, Any] = {"d": snapshot_date}
    clean_skus = sorted({_clean_sku(sku) for sku in (sku_codes or []) if _clean_sku(sku)})
    if clean_skus:
        sku_placeholders = []
        for index, sku_code in enumerate(clean_skus):
            key = f"sku_{index}"
            sku_placeholders.append(f":{key}")
            params[key] = sku_code
        where.append(f"s.sku_code IN ({', '.join(sku_placeholders)})")
    if brand:
        where.append("sku.brand=:brand")
        params["brand"] = brand
    if category:
        where.append("sku.category=:category")
        params["category"] = category
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                f"""
                SELECT
                    s.snapshot_date, s.sku_code, s.warehouse_code,
                    s.stock_qty, s.inbound_qty, s.sales_1d, s.sales_7d, s.sales_30d,
                    sku.product_name, sku.brand, sku.category, sku.shop_name
                FROM inv_stock_snapshot s
                LEFT JOIN inv_sku sku ON sku.sku_code=s.sku_code
                WHERE {' AND '.join(where)}
                """
            ),
            params,
        ).mappings().all()
    return [dict(row) for row in rows]


def _default_replenishment_policy() -> ReplenishmentPolicy:
    return ReplenishmentPolicy(
        policy_type="SYSTEM",
        brand="",
        category="",
        target_days=TARGET_DAYS,
        reserve_days=DONOR_RESERVE_DAYS,
        supplier_eta_days=SUPPLIER_TO_HUB_ETA_DAYS,
        allow_jd_transfer=True,
        preferred_source_warehouse=SUPPLY_HUB_WAREHOUSE,
    )


def _load_replenishment_policies(engine: Engine) -> dict[str, ReplenishmentPolicy]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT policy_type, brand, category, target_days, reserve_days, supplier_eta_days,
                       allow_jd_transfer, preferred_source_warehouse
                FROM inv_replenishment_policy
                """
            )
        ).mappings().all()
    policies: dict[str, ReplenishmentPolicy] = {}
    for row in rows:
        policy = ReplenishmentPolicy(
            policy_type=str(row["policy_type"] or "SYSTEM").upper(),
            brand=str(row["brand"] or ""),
            category=str(row["category"] or ""),
            target_days=max(1.0, _to_float(row["target_days"]) or TARGET_DAYS),
            reserve_days=max(1.0, _to_float(row["reserve_days"]) or DONOR_RESERVE_DAYS),
            supplier_eta_days=max(1, int(_to_float(row["supplier_eta_days"]) or SUPPLIER_TO_HUB_ETA_DAYS)),
            allow_jd_transfer=bool(int(row["allow_jd_transfer"] or 0)),
            preferred_source_warehouse=str(row["preferred_source_warehouse"] or SUPPLY_HUB_WAREHOUSE),
        )
        policies[_policy_key(policy.policy_type, policy.brand, policy.category)] = policy
    if "SYSTEM::" not in policies:
        policies["SYSTEM::"] = _default_replenishment_policy()
    return policies


def _policy_key(policy_type: str, brand: str = "", category: str = "") -> str:
    if policy_type == "BRAND":
        return f"BRAND:{_normalize_brand_key(brand)}:"
    if policy_type == "CATEGORY":
        return f"CATEGORY::{category}"
    return "SYSTEM::"


def _resolve_policy(
    policies: dict[str, ReplenishmentPolicy],
    brand: Any,
    category: Any,
) -> ReplenishmentPolicy:
    brand_text = _clean_text(brand)
    category_text = _clean_text(category)
    policy = (
        policies.get(_policy_key("BRAND", brand_text, ""))
        or policies.get(_policy_key("CATEGORY", "", category_text))
        or policies.get("SYSTEM::")
        or _default_replenishment_policy()
    )
    if category_text == BOX_CATEGORY and policy.allow_jd_transfer:
        return ReplenishmentPolicy(
            policy_type=policy.policy_type,
            brand=policy.brand,
            category=policy.category,
            target_days=policy.target_days,
            reserve_days=policy.reserve_days,
            supplier_eta_days=policy.supplier_eta_days,
            allow_jd_transfer=False,
            preferred_source_warehouse=policy.preferred_source_warehouse or SUPPLY_HUB_WAREHOUSE,
        )
    return policy


def _policy_to_dict(policy: ReplenishmentPolicy) -> dict[str, Any]:
    return {
        "policy_type": policy.policy_type,
        "brand": policy.brand,
        "category": policy.category,
        "target_days": policy.target_days,
        "reserve_days": policy.reserve_days,
        "supplier_eta_days": policy.supplier_eta_days,
        "allow_jd_transfer": policy.allow_jd_transfer,
        "preferred_source_warehouse": policy.preferred_source_warehouse,
    }


def _latest_supply_stock(engine: Engine, sku_code: str, warehouse_code: str = SUPPLY_HUB_WAREHOUSE) -> dict[str, Any] | None:
    with engine.connect() as conn:
        row = conn.execute(
            text(
                """
                SELECT snapshot_time, snapshot_date, warehouse_code, sku_code, stock_qty,
                       available_qty, inbound_qty, source
                FROM inv_supply_stock_snapshot
                WHERE sku_code=:sku_code
                  AND warehouse_code=:warehouse_code
                ORDER BY snapshot_time DESC, id DESC
                LIMIT 1
                """
            ),
            {"sku_code": sku_code, "warehouse_code": warehouse_code},
        ).mappings().fetchone()
    return _jsonify_mapping(row) if row else None


def _latest_supply_stocks(engine: Engine, warehouse_code: str = SUPPLY_HUB_WAREHOUSE) -> dict[str, dict[str, Any]]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT s.snapshot_time, s.snapshot_date, s.warehouse_code, s.sku_code,
                       s.stock_qty, s.available_qty, s.inbound_qty, s.source
                FROM inv_supply_stock_snapshot s
                JOIN (
                    SELECT sku_code, warehouse_code, MAX(snapshot_time) AS max_snapshot_time
                    FROM inv_supply_stock_snapshot
                    WHERE warehouse_code=:warehouse_code
                    GROUP BY sku_code, warehouse_code
                ) latest
                  ON latest.sku_code=s.sku_code
                 AND latest.warehouse_code=s.warehouse_code
                 AND latest.max_snapshot_time=s.snapshot_time
                WHERE s.warehouse_code=:warehouse_code
                """
            ),
            {"warehouse_code": warehouse_code},
        ).mappings().all()
    return {str(row["sku_code"]): _jsonify_mapping(row) for row in rows}


def _latest_supply_stocks_by_warehouse(engine: Engine) -> dict[tuple[str, str], dict[str, Any]]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT s.snapshot_time, s.snapshot_date, s.warehouse_code, s.sku_code,
                       s.stock_qty, s.available_qty, s.inbound_qty, s.source
                FROM inv_supply_stock_snapshot s
                JOIN (
                    SELECT warehouse_code, sku_code, MAX(snapshot_time) AS max_snapshot_time
                    FROM inv_supply_stock_snapshot
                    GROUP BY warehouse_code, sku_code
                ) latest
                  ON latest.warehouse_code=s.warehouse_code
                 AND latest.sku_code=s.sku_code
                 AND latest.max_snapshot_time=s.snapshot_time
                """
            )
        ).mappings().all()
    return {
        (str(row["warehouse_code"]), str(row["sku_code"])): _jsonify_mapping(row)
        for row in rows
    }


def _supply_stock_for_policy(
    supply_stocks: dict[tuple[str, str], dict[str, Any]],
    policy: ReplenishmentPolicy,
    sku_code: str,
) -> dict[str, Any] | None:
    preferred = policy.preferred_source_warehouse or SUPPLY_HUB_WAREHOUSE
    stock = supply_stocks.get((preferred, sku_code))
    if stock is not None or preferred != SUPPLY_HUB_WAREHOUSE:
        return stock
    return supply_stocks.get((SUPPLY_HUB_WAREHOUSE, sku_code))


def _is_supplier_direct_brand(brand: Any) -> bool:
    return _normalize_brand_key(brand) == _normalize_brand_key("测试五金")


def _is_nansha_supply_brand(brand: Any) -> bool:
    # iGear 及其拆分子品牌（iGear工业品）统一视为南沙仓供给品牌。
    return _normalize_brand_key(brand).startswith(_normalize_brand_key(IGEAR_BRAND))


def _calculate_risk_rows(
    run_id: int,
    snapshot_date: date,
    snapshots: list[dict[str, Any]],
    policies: dict[str, ReplenishmentPolicy] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    rows: list[dict[str, Any]] = []
    stats = {"HIGH": 0, "MID": 0, "LOW": 0}

    for snap in snapshots:
        stock_qty = _to_float(snap["stock_qty"])
        inbound_qty = _to_float(snap["inbound_qty"])
        sales_7d = _to_float(snap["sales_7d"])
        sales_30d = _to_float(snap["sales_30d"])
        policy = _resolve_policy(policies or {}, snap.get("brand"), snap.get("category"))
        target_days = float(policy.target_days)
        forecast = _forecast_daily_sales(snap)
        daily_sales = forecast["daily_sales"]
        inbound_credit_qty = _inbound_credit_qty(inbound_qty, daily_sales)
        inbound_credit_rate = inbound_credit_qty / inbound_qty if inbound_qty > 0 else 0.0
        available_qty = stock_qty + inbound_credit_qty
        has_sales = sales_7d > 0 or sales_30d > 0 or _to_float(snap["sales_1d"]) > 0
        days_of_supply = _days_of_supply(stock_qty, daily_sales, allow_no_sales_stock=not has_sales)
        demand_30d = daily_sales * target_days
        consecutive_high_days = 0

        if not has_sales and stock_qty > 0:
            risk_level = "LOW"
            shortage_days = 0.0
            need_qty = 0.0
            is_shortage_risk = 0
        else:
            if days_of_supply < target_days:
                risk_level = "HIGH"
            elif days_of_supply < target_days + max(10.0, target_days * 0.25):
                risk_level = "MID"
            else:
                risk_level = "LOW"
            shortage_days = max(0.0, target_days - days_of_supply) if daily_sales > 0 else 0.0
            need_qty = max(0.0, demand_30d - available_qty)
            is_shortage_risk = 1 if daily_sales > 0 and (demand_30d > available_qty or days_of_supply < target_days) else 0

        excess_qty = stock_qty - demand_30d
        is_excess_risk = 1 if excess_qty > 0 and days_of_supply > target_days else 0
        stats[risk_level] += 1
        rows.append(
            {
                "run_id": run_id,
                "snapshot_date": snapshot_date,
                "sku_code": snap["sku_code"],
                "warehouse_code": snap["warehouse_code"],
                "daily_sales": round(daily_sales, 4),
                "target_days": round(target_days, 2),
                "days_of_supply": round(days_of_supply, 4),
                "demand_30d": round(demand_30d, 4),
                "shortage_days": round(shortage_days, 4),
                "risk_level": risk_level,
                "excess_qty": round(excess_qty, 4),
                "need_qty": round(need_qty, 4),
                "available_qty": round(available_qty, 4),
                "inbound_credit_qty": round(inbound_credit_qty, 4),
                "inbound_credit_rate": round(inbound_credit_rate, 4),
                "trend_factor": round(forecast["trend_factor"], 4),
                "volatility_factor": round(forecast["volatility_factor"], 4),
                "confidence_score": round(forecast["confidence_score"], 4),
                "consecutive_high_days": consecutive_high_days,
                "is_shortage_risk": is_shortage_risk,
                "is_excess_risk": is_excess_risk,
            }
        )
    return rows, stats


def _forecast_daily_sales(
    snap: dict[str, Any],
    history: list[dict[str, Any]] | None = None,
) -> dict[str, float]:
    # Inventory snapshots are point-in-time: the latest date supersedes older
    # snapshots, so skipped days must not change trend/volatility.
    del history

    sales_7d = _to_float(snap["sales_7d"])
    sales_30d = _to_float(snap["sales_30d"])
    sales_1d = _to_float(snap["sales_1d"])
    daily_7d = sales_7d / 7.0
    daily_30d = sales_30d / 30.0

    if daily_7d > 0 and daily_30d > 0:
        base_daily = daily_7d * 0.65 + daily_30d * 0.25 + sales_1d * 0.10
    elif daily_7d > 0:
        base_daily = daily_7d * 0.85 + sales_1d * 0.15
    elif daily_30d > 0:
        base_daily = daily_30d
    else:
        base_daily = sales_1d

    confidence = 1.0 if (daily_7d > 0 or daily_30d > 0 or sales_1d > 0) else 0.0
    trend_factor = 1.0
    volatility_factor = 1.0

    daily_sales = base_daily * trend_factor * volatility_factor
    return {
        "daily_sales": daily_sales,
        "trend_factor": trend_factor,
        "volatility_factor": volatility_factor,
        "confidence_score": confidence,
    }


def _inbound_credit_qty(inbound_qty: float, daily_sales: float) -> float:
    """Treat in-transit stock as already arrived; no arrival-window deduction."""
    return max(0.0, _to_float(inbound_qty))


def _calculate_transfer_plans(
    run_id: int,
    snapshot_date: date,
    snapshots: list[dict[str, Any]],
    risk_rows: list[dict[str, Any]],
    policies: dict[str, ReplenishmentPolicy] | None = None,
    supply_stocks: dict[tuple[str, str], dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    snap_by_key = {(row["sku_code"], row["warehouse_code"]): row for row in snapshots}
    shortage_rows = [
        row
        for row in risk_rows
        if row["is_shortage_risk"] == 1 and row["risk_level"] in {"HIGH", "MID"}
    ]
    shortage_rows.sort(
        key=lambda r: (
            _risk_rank(r["risk_level"]),
            -int(r.get("consecutive_high_days") or 0),
            float(r["days_of_supply"]),
            -float(r["need_qty"]),
        )
    )

    donor_remaining: dict[tuple[str, str], float] = {}
    donor_after_daily: dict[tuple[str, str], float] = {}
    donor_candidates_by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for risk in risk_rows:
        if float(risk["excess_qty"]) <= 0:
            continue
        snap = snap_by_key.get((risk["sku_code"], risk["warehouse_code"]))
        if not snap:
            continue
        donor_policy = _resolve_policy(policies or {}, snap.get("brand"), snap.get("category"))
        daily_sales = float(risk["daily_sales"])
        reserve_qty = daily_sales * donor_policy.reserve_days
        transferable = max(0.0, min(float(risk["excess_qty"]), _to_float(snap["stock_qty"]) - reserve_qty))
        if transferable <= 0:
            continue
        key = (risk["sku_code"], risk["warehouse_code"])
        donor_remaining[key] = transferable
        donor_after_daily[key] = daily_sales
        donor_candidates_by_sku[risk["sku_code"]].append(risk)

    plans: list[dict[str, Any]] = []
    planned_pairs: set[tuple[str, str, str]] = set()

    for shortage in shortage_rows:
        sku_code = shortage["sku_code"]
        to_warehouse = shortage["warehouse_code"]
        need_qty = math.ceil(float(shortage["need_qty"]))
        if need_qty <= 0:
            continue

        to_snap = snap_by_key.get((sku_code, to_warehouse))
        policy = _resolve_policy(policies or {}, to_snap.get("brand") if to_snap else "", to_snap.get("category") if to_snap else "")

        if not policy.allow_jd_transfer:
            supply_plan = _build_supply_plan(
                run_id=run_id,
                snapshot_date=snapshot_date,
                sku_code=sku_code,
                shortage=shortage,
                snap_by_key=snap_by_key,
                policy=policy,
                supply_stock=_supply_stock_for_policy(supply_stocks or {}, policy, sku_code),
                reason_context="箱包类目策略不允许京东仓之间互调",
            )
            if supply_plan:
                plans.append(supply_plan)
            continue

        donors = []
        for donor in donor_candidates_by_sku.get(sku_code, []):
            from_warehouse = donor["warehouse_code"]
            if from_warehouse == to_warehouse:
                continue
            if (sku_code, to_warehouse, from_warehouse) in planned_pairs:
                continue
            remaining = donor_remaining.get((sku_code, from_warehouse), 0.0)
            if remaining <= 0:
                continue
            route = estimate_transfer_route(from_warehouse, to_warehouse)
            risk_penalty = 25.0 if donor["risk_level"] == "MID" else 0.0
            eta_late_penalty = 30.0 if route["eta_days"] > float(shortage["days_of_supply"]) else 0.0
            urgent_bonus = min(int(shortage.get("consecutive_high_days") or 0), 7) * 2.0
            score = remaining * 0.8 - route["distance_km"] * 0.015 - route["eta_days"] * 12.0 - risk_penalty - eta_late_penalty + urgent_bonus
            can_cover = math.floor(remaining) >= need_qty
            donors.append((route["eta_days"], -score, donor, route, remaining, score, can_cover))

        donors.sort(key=lambda item: (item[0], item[1]))
        selected = next((item for item in donors if item[6]), None)
        if selected is None:
            supply_plan = _build_supply_plan(
                run_id=run_id,
                snapshot_date=snapshot_date,
                sku_code=sku_code,
                shortage=shortage,
                snap_by_key=snap_by_key,
                policy=policy,
                supply_stock=_supply_stock_for_policy(supply_stocks or {}, policy, sku_code),
                reason_context="没有单个京东仓能在不低于安全线的前提下覆盖本次缺口",
            )
            if supply_plan:
                plans.append(supply_plan)
            continue

        eta_days, _neg_score, donor, route, remaining, score, can_cover = selected
        from_warehouse = donor["warehouse_code"]
        qty = need_qty
        daily_after = donor_after_daily.get((sku_code, from_warehouse), 0.0)
        donor_snap = snap_by_key[(sku_code, from_warehouse)]
        donor_policy = _resolve_policy(policies or {}, donor_snap.get("brand"), donor_snap.get("category"))
        stock_after = _to_float(donor_snap["stock_qty"]) - qty
        surplus_days_after = _days_of_supply(stock_after, daily_after, allow_no_sales_stock=True)
        if stock_after < 0 or surplus_days_after < donor_policy.reserve_days:
            continue
        eta_late = 1 if route["eta_days"] > float(shortage["days_of_supply"]) else 0
        reason = _build_rule_reason(
            from_warehouse=from_warehouse,
            to_warehouse=to_warehouse,
            qty=qty,
            eta_days=route["eta_days"],
            risk_level=shortage["risk_level"],
            risk_days=float(shortage["days_of_supply"]),
            target_days=float(shortage.get("target_days") or policy.target_days),
            reserve_days=donor_policy.reserve_days,
            surplus_days_after=surplus_days_after,
            donor_excess=float(donor["excess_qty"]),
            donor_stock_after=stock_after,
            inbound_qty=_to_float(snap_by_key[(sku_code, to_warehouse)]["inbound_qty"]),
            inbound_credit_qty=float(shortage.get("inbound_credit_qty") or 0),
            consecutive_high_days=int(shortage.get("consecutive_high_days") or 0),
            trend_factor=float(shortage.get("trend_factor") or 1),
            volatility_factor=float(shortage.get("volatility_factor") or 1),
            eta_late=bool(eta_late),
            source_type=SOURCE_TYPE_JD_TRANSFER,
        )
        donor_remaining[(sku_code, from_warehouse)] = remaining - qty

        plans.append(
            {
                "run_id": run_id,
                "plan_date": snapshot_date,
                "sku_code": sku_code,
                "from_warehouse": from_warehouse,
                "to_warehouse": to_warehouse,
                "transfer_qty": qty,
                "eta_days": route["eta_days"],
                "priority": shortage["risk_level"],
                "risk_level": shortage["risk_level"],
                "risk_days": round(float(shortage["days_of_supply"]), 4),
                "surplus_days_after": round(surplus_days_after, 4),
                "eta_late": eta_late,
                "source_type": SOURCE_TYPE_JD_TRANSFER,
                "route_path": f"{from_warehouse} → {to_warehouse}",
                "target_days": round(float(shortage.get("target_days") or policy.target_days), 2),
                "reserve_days": round(donor_policy.reserve_days, 2),
                "supplier_eta_days": 0,
                "hub_eta_days": 0,
                "score": round(score, 4),
                "reason": reason,
            }
        )
        planned_pairs.add((sku_code, from_warehouse, to_warehouse))

    return plans


def _build_supply_plan(
    run_id: int,
    snapshot_date: date,
    sku_code: str,
    shortage: dict[str, Any],
    snap_by_key: dict[tuple[str, str], dict[str, Any]],
    policy: ReplenishmentPolicy,
    supply_stock: dict[str, Any] | None,
    reason_context: str,
) -> dict[str, Any] | None:
    to_warehouse = shortage["warehouse_code"]
    need_qty = math.ceil(float(shortage["need_qty"]))
    if need_qty <= 0:
        return None

    hub_warehouse = policy.preferred_source_warehouse or SUPPLY_HUB_WAREHOUSE
    hub_available_qty = max(
        0.0,
        _to_float((supply_stock or {}).get("available_qty")) or _to_float((supply_stock or {}).get("stock_qty")),
    )
    to_snap = snap_by_key.get((sku_code, to_warehouse), {})
    if _is_supplier_direct_brand(to_snap.get("brand")):
        direct_route = estimate_transfer_route(FOSHAN_SUPPLIER_WAREHOUSE, to_warehouse)
        eta_days = int(direct_route["eta_days"])
        eta_late = 1 if eta_days > float(shortage["days_of_supply"]) else 0
        reason = _build_rule_reason(
            from_warehouse=FOSHAN_SUPPLIER_WAREHOUSE,
            to_warehouse=to_warehouse,
            qty=need_qty,
            eta_days=eta_days,
            risk_level=shortage["risk_level"],
            risk_days=float(shortage["days_of_supply"]),
            target_days=float(shortage.get("target_days") or policy.target_days),
            reserve_days=policy.reserve_days,
            surplus_days_after=0.0,
            donor_excess=0.0,
            donor_stock_after=0.0,
            inbound_qty=_to_float(to_snap.get("inbound_qty")),
            inbound_credit_qty=float(shortage.get("inbound_credit_qty") or 0),
            consecutive_high_days=int(shortage.get("consecutive_high_days") or 0),
            trend_factor=float(shortage.get("trend_factor") or 1),
            volatility_factor=float(shortage.get("volatility_factor") or 1),
            eta_late=bool(eta_late),
            source_type=SOURCE_TYPE_SUPPLIER_DIRECT,
            supplier_eta_days=eta_days,
            hub_eta_days=0,
            reason_context=reason_context,
        )
        return {
            "run_id": run_id,
            "plan_date": snapshot_date,
            "sku_code": sku_code,
            "from_warehouse": FOSHAN_SUPPLIER_WAREHOUSE,
            "to_warehouse": to_warehouse,
            "transfer_qty": need_qty,
            "eta_days": eta_days,
            "priority": shortage["risk_level"],
            "risk_level": shortage["risk_level"],
            "risk_days": round(float(shortage["days_of_supply"]), 4),
            "surplus_days_after": 0.0,
            "eta_late": eta_late,
            "source_type": SOURCE_TYPE_SUPPLIER_DIRECT,
            "route_path": f"{FOSHAN_SUPPLIER_WAREHOUSE} → {to_warehouse}",
            "target_days": round(float(shortage.get("target_days") or policy.target_days), 2),
            "reserve_days": round(policy.reserve_days, 2),
            "supplier_eta_days": eta_days,
            "hub_eta_days": 0,
            "score": round(-5000.0 - eta_days * 20.0, 4),
            "reason": reason,
        }

    hub_route = estimate_transfer_route(hub_warehouse, to_warehouse)
    hub_eta_days = int(hub_route["eta_days"])
    can_hub_cover = math.floor(hub_available_qty) >= need_qty
    if can_hub_cover:
        source_type = SOURCE_TYPE_HUB_TRANSFER
        supplier_eta_days = 0
        eta_days = hub_eta_days
        route_path = f"{hub_warehouse} → {to_warehouse}"
        surplus_days_after = 9999.0
        score = 2500.0 + hub_available_qty - hub_eta_days * 20.0
    else:
        source_type = SOURCE_TYPE_SUPPLIER_REPLENISH
        supplier_eta_days = policy.supplier_eta_days
        eta_days = supplier_eta_days + hub_eta_days
        route_path = f"{SUPPLIER_WAREHOUSE} → {hub_warehouse} → {to_warehouse}"
        surplus_days_after = 0.0
        score = -9999.0

    eta_late = 1 if eta_days > float(shortage["days_of_supply"]) else 0
    reason = _build_rule_reason(
        from_warehouse=hub_warehouse,
        to_warehouse=to_warehouse,
        qty=need_qty,
        eta_days=eta_days,
        risk_level=shortage["risk_level"],
        risk_days=float(shortage["days_of_supply"]),
        target_days=float(shortage.get("target_days") or policy.target_days),
        reserve_days=policy.reserve_days,
        surplus_days_after=surplus_days_after,
        donor_excess=max(0.0, hub_available_qty - need_qty),
        donor_stock_after=max(0.0, hub_available_qty - need_qty),
        inbound_qty=_to_float(to_snap.get("inbound_qty")),
        inbound_credit_qty=float(shortage.get("inbound_credit_qty") or 0),
        consecutive_high_days=int(shortage.get("consecutive_high_days") or 0),
        trend_factor=float(shortage.get("trend_factor") or 1),
        volatility_factor=float(shortage.get("volatility_factor") or 1),
        eta_late=bool(eta_late),
        source_type=source_type,
        source_available_qty=hub_available_qty,
        supplier_eta_days=supplier_eta_days,
        hub_eta_days=hub_eta_days,
        reason_context=reason_context,
    )

    return {
        "run_id": run_id,
        "plan_date": snapshot_date,
        "sku_code": sku_code,
        "from_warehouse": hub_warehouse,
        "to_warehouse": to_warehouse,
        "transfer_qty": need_qty,
        "eta_days": eta_days,
        "priority": shortage["risk_level"],
        "risk_level": shortage["risk_level"],
        "risk_days": round(float(shortage["days_of_supply"]), 4),
        "surplus_days_after": round(surplus_days_after, 4),
        "eta_late": eta_late,
        "source_type": source_type,
        "route_path": route_path,
        "target_days": round(float(shortage.get("target_days") or policy.target_days), 2),
        "reserve_days": round(policy.reserve_days, 2),
        "supplier_eta_days": supplier_eta_days,
        "hub_eta_days": hub_eta_days,
        "score": round(score, 4),
        "reason": reason,
    }


def _build_rule_reason(
    from_warehouse: str,
    to_warehouse: str,
    qty: float,
    eta_days: int,
    risk_level: str,
    risk_days: float,
    target_days: float,
    reserve_days: float,
    surplus_days_after: float,
    donor_excess: float,
    donor_stock_after: float,
    inbound_qty: float = 0.0,
    inbound_credit_qty: float = 0.0,
    consecutive_high_days: int = 0,
    trend_factor: float = 1.0,
    volatility_factor: float = 1.0,
    eta_late: bool = False,
    source_type: str = SOURCE_TYPE_JD_TRANSFER,
    source_available_qty: float = 0.0,
    supplier_eta_days: int = 0,
    hub_eta_days: int = 0,
    reason_context: str = "",
) -> str:
    risk_label = {"HIGH": "高风险", "MID": "中风险", "LOW": "低风险"}.get(risk_level, risk_level)
    eta_note = "预计到货晚于当前可售天数，到货前存在短缺窗口，建议运营同步安排临时采购或前端限流。" if eta_late else "预计到货不晚于当前可售天数。"
    donor_supply_note = (
        "近7日无出库"
        if surplus_days_after >= 9999
        else f"累计调出后剩余约 {donor_stock_after:.0f} 件、可维持约 {surplus_days_after:.1f} 天"
    )
    if inbound_qty > 0:
        inbound_note = f"采购未到货约 {inbound_qty:.0f} 件按已到仓计入可售。"
    else:
        inbound_note = ""
    history_notes = []
    if consecutive_high_days >= 3:
        history_notes.append(f"连续 {consecutive_high_days} 天高风险")
    if trend_factor >= 1.08:
        history_notes.append("近期需求上升")
    if volatility_factor >= 1.05:
        history_notes.append("销量波动偏大")
    history_note = f"{'、'.join(history_notes)}，优先级上调。" if history_notes else ""

    target_text = f"{target_days:.0f}" if float(target_days).is_integer() else f"{target_days:.1f}"
    context_note = f"{reason_context}。" if reason_context else ""
    if source_type == SOURCE_TYPE_SUPPLIER_DIRECT:
        return (
            f"{to_warehouse} 当前可售约 {risk_days:.1f} 天，低于 {target_text} 天目标，属于{risk_label}。"
            f"{inbound_note}{history_note}{context_note}测试五金品牌支持品牌方佛山仓直发，"
            f"建议由{from_warehouse}直接发往 {to_warehouse}，补足 {int(qty)} 件，"
            f"按佛山到目的仓距离估算预计 {eta_days} 天到货。{eta_note}"
        )

    if source_type == SOURCE_TYPE_SUPPLIER_REPLENISH:
        return (
            f"{to_warehouse} 当前可售约 {risk_days:.1f} 天，低于 {target_text} 天目标，属于{risk_label}。"
            f"{inbound_note}{history_note}{context_note}{from_warehouse} 当前可用约 {source_available_qty:.0f} 件，"
            f"无法覆盖需补足的 {int(qty)} 件，建议先向供应商采购补到{from_warehouse}，"
            f"暂按 {supplier_eta_days} 天到{from_warehouse}，再由{from_warehouse}发往 {to_warehouse}，"
            f"花东到仓预计 {hub_eta_days} 天，合计预计 {eta_days} 天。{eta_note}"
        )

    if source_type == SOURCE_TYPE_HUB_TRANSFER:
        return (
            f"{to_warehouse} 当前可售约 {risk_days:.1f} 天，低于 {target_text} 天目标，属于{risk_label}。"
            f"{inbound_note}{history_note}{context_note}{from_warehouse} 当前可用约 {source_available_qty:.0f} 件，"
            f"可覆盖需补足的 {int(qty)} 件，建议由{from_warehouse}发往 {to_warehouse}，"
            f"预计 {eta_days} 天到货。{eta_note}"
        )

    return (
        f"{to_warehouse} 当前可售约 {risk_days:.1f} 天，低于 {target_text} 天目标，属于{risk_label}。"
        f"{inbound_note}{history_note}建议从 {from_warehouse} 调拨 {int(qty)} 件，预计 {eta_days} 天到货；"
        f"{from_warehouse} 当前富余约 {donor_excess:.0f} 件，{donor_supply_note}，"
        f"不低于 {reserve_days:.0f} 天安全线。{eta_note}"
    )


def _daily_sales(sales_7d: Any) -> float:
    return _to_float(sales_7d) / 7.0


def _avg(values: Iterable[float]) -> float:
    clean = [float(value) for value in values]
    return sum(clean) / len(clean) if clean else 0.0


def _clamp(value: float, min_value: float, max_value: float) -> float:
    return max(min_value, min(max_value, value))


def _days_of_supply(stock_qty: float, daily_sales: float, allow_no_sales_stock: bool = False) -> float:
    if daily_sales <= 0:
        return 9999.0 if allow_no_sales_stock and stock_qty > 0 else 0.0
    return stock_qty / daily_sales


def _risk_rank(level: str) -> int:
    return {"HIGH": 0, "MID": 1, "LOW": 2}.get(level, 9)


def _warehouse_record(warehouse_code: str) -> dict[str, Any]:
    base_city = _base_city(warehouse_code)
    region = REGION_BY_CITY.get(base_city, "其他")
    coord = CITY_COORDINATES.get(base_city)
    is_hub = warehouse_code == SUPPLY_HUB_WAREHOUSE
    return {
        "warehouse_code": warehouse_code,
        "warehouse_name": warehouse_code,
        "region": region,
        "warehouse_type": WAREHOUSE_TYPE_SUPPLY_HUB if is_hub else WAREHOUSE_TYPE_JD,
        "is_supply_hub": 1 if is_hub else 0,
        "is_jd_warehouse": 0 if is_hub else 1,
        "base_city": base_city,
        "latitude": coord[0] if coord else None,
        "longitude": coord[1] if coord else None,
    }


def _base_city(name: str) -> str:
    text_value = str(name or "").strip()
    if text_value in {SUPPLY_HUB_WAREHOUSE, "花东"}:
        return "上海"
    if text_value in {NANSHA_SUPPLY_WAREHOUSE, "南沙"}:
        return "广州"
    if text_value in {FOSHAN_SUPPLIER_WAREHOUSE, "供应商佛山仓", "供应商（佛山）", "佛山供应商"}:
        return "佛山"
    text_value = re.sub(r"(供应商|工厂|厂家|发货地|所在地|配送中心|补货B|FDC|RDC|仓)$", "", text_value)
    text_value = re.sub(r"(供应商|工厂|厂家|发货地|所在地|配送中心|补货B|FDC|RDC)", "", text_value)
    for city in sorted(CITY_COORDINATES.keys(), key=len, reverse=True):
        if text_value.startswith(city) or city in text_value:
            return city
    for province, city in sorted(PROVINCE_BASE_CITY.items(), key=lambda item: len(item[0]), reverse=True):
        if text_value.startswith(province) or province in text_value:
            return city
    return text_value[:4] if text_value else ""


def _haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    radius = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * radius * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def _eta_days(distance_km: float, from_city: str, to_city: str) -> int:
    if distance_km <= 120:
        eta = 1
    elif distance_km <= 500:
        eta = 2
    elif distance_km <= 1000:
        eta = 3
    elif distance_km <= 1700:
        eta = 4
    elif distance_km <= 2600:
        eta = 5
    else:
        eta = 6
    if from_city in {"拉萨", "乌鲁木齐"} or to_city in {"拉萨", "乌鲁木齐"}:
        eta = min(7, eta + 1)
    return max(1, eta)


def _append_error(
    errors: list[dict[str, Any]],
    batch_no: str,
    snapshot_date: date | None,
    sku_code: str,
    warehouse_code: str,
    error_type: str,
    severity: str,
    message: str,
    raw_value: Any,
) -> None:
    if len(errors) >= MAX_IMPORT_ERRORS_PER_BATCH:
        return
    errors.append(
        {
            "batch_no": batch_no,
            "snapshot_date": snapshot_date,
            "sku_code": str(sku_code or "")[:64],
            "warehouse_code": str(warehouse_code or "")[:128],
            "error_type": error_type,
            "severity": severity,
            "message": message[:1000],
            "raw_value": str(raw_value or "")[:1000],
        }
    )


def _metric(row_values: dict[int | str, Any], prefix_cols: dict[str, dict[str, int]], prefix: str, field: str) -> float:
    col = prefix_cols.get(prefix, {}).get(field)
    return _to_float(row_values.get(col)) if col else 0.0


def _clean_sku(value: Any) -> str:
    text_value = _clean_text(value)
    expanded = _expand_scientific_identifier(text_value)
    if expanded:
        return expanded
    if re.fullmatch(r"\d+\.0", text_value):
        return text_value[:-2]
    return text_value


def _clean_barcode(value: Any) -> str:
    return _clean_sku(value)


def _barcode_tokens(value: Any) -> list[str]:
    text_value = _clean_barcode(value)
    if not text_value:
        return []
    tokens: list[str] = []
    for part in re.split(r"[;,，；、\s]+", text_value):
        token = _clean_barcode(part)
        if token and token not in tokens:
            tokens.append(token)
    return tokens


def _preferred_barcode_query(sku_code: Any, barcode: Any) -> str:
    clean_sku_code = _clean_sku(sku_code)
    tokens = _barcode_tokens(barcode)
    for token in tokens:
        if token != clean_sku_code:
            return token
    return tokens[0] if tokens else clean_sku_code


def _normalize_brand_key(value: Any) -> str:
    return (
        _clean_text(value)
        .replace("（", "(")
        .replace("）", ")")
        .replace(" ", "")
        .strip()
        .lower()
    )


def _brand_category(brand: Any) -> str:
    brand_key = _normalize_brand_key(brand)
    if not brand_key:
        return ""
    if brand_key in BRAND_CATEGORY_MAP:
        return BRAND_CATEGORY_MAP[brand_key]
    for mapped_brand, category in BRAND_CATEGORY_MAP.items():
        if mapped_brand in brand_key or brand_key in mapped_brand:
            return category
    return ""


def _expand_scientific_identifier(value: Any) -> str:
    text_value = str(value or "").strip().replace(",", "")
    match = re.fullmatch(r"([+]?)?(\d+)(?:\.(\d+))?[eE]([+-]?\d+)", text_value)
    if not match:
        return ""

    integer_part = match.group(2)
    fraction_part = match.group(3) or ""
    exponent = int(match.group(4))
    digits = integer_part + fraction_part
    decimal_index = len(integer_part) + exponent

    if decimal_index >= len(digits):
        return digits + ("0" * (decimal_index - len(digits)))
    if decimal_index <= 0:
        return ""

    fractional_tail = digits[decimal_index:]
    if any(char != "0" for char in fractional_tail):
        return ""
    return digits[:decimal_index]


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text_value = str(value).strip()
    if re.fullmatch(r"\d+\.0", text_value):
        return text_value[:-2]
    return text_value


def _normalize_inventory_brand_by_supplier(brand: Any, supplier_name: Any) -> str:
    brand_text = _clean_text(brand)
    if _normalize_brand_key(supplier_name) != _normalize_brand_key(ENTERPRISE_PURCHASE_SUPPLIER_NAME):
        return brand_text

    brand_key = _normalize_brand_key(brand_text)
    if _normalize_brand_key(SAMSONITE_BRAND) in brand_key or "samsonite" in brand_key:
        return SAMSONITE_ENTERPRISE_BRAND
    if _normalize_brand_key(AMERICAN_TOURISTER_BRAND) in brand_key or "americantourister" in brand_key:
        return AMERICAN_TOURISTER_ENTERPRISE_BRAND
    return brand_text


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except Exception:
        return 0.0


def _parse_datetime_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text_value = str(value or "").strip()
    if not text_value:
        return None
    text_value = text_value.replace("T", " ").replace("Z", "")
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            limit = 26 if "%f" in fmt else 19 if "%H" in fmt else 10
            return datetime.strptime(text_value[:limit], fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text_value)
    except ValueError:
        return None


def _json_candidate(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    clean_value = value.strip()
    if not clean_value or clean_value[0] not in "[{":
        return value
    try:
        return json.loads(clean_value)
    except json.JSONDecodeError:
        return value


def _extract_jushuitan_stock_items(payload: Any, depth: int = 0) -> list[dict[str, Any]]:
    if depth > 4:
        return []
    payload = _json_candidate(payload)
    if isinstance(payload, list):
        items: list[dict[str, Any]] = []
        for item in payload:
            items.extend(_extract_jushuitan_stock_items(item, depth + 1))
        return items
    if not isinstance(payload, dict):
        return []

    stock_keys = {"skuId", "sku_id", "skuCode", "sku_code", "iId", "iid", "qty", "purchaseQty"}
    if stock_keys.intersection(payload.keys()):
        return [payload]

    items: list[dict[str, Any]] = []
    for key in ("data", "datas", "items", "rows", "payload", "body", "content", "message", "msg"):
        if key in payload:
            items.extend(_extract_jushuitan_stock_items(payload[key], depth + 1))
    return items


def _parse_snapshot_date(value: Any, filename: str = "") -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text_value = _clean_text(value)
    parsed = _parse_date_text(text_value)
    if parsed:
        return parsed
    numeric = _to_float(text_value)
    if numeric > 20000:
        try:
            return date(1899, 12, 30) + timedelta(days=int(numeric))
        except Exception:
            pass
    match = re.search(r"(20\d{6})", filename)
    if match:
        return datetime.strptime(match.group(1), "%Y%m%d").date()
    return None


def _parse_date_text(text_value: str) -> date | None:
    value = str(text_value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(value[:10] if fmt != "%Y%m%d" else value[:8], fmt).date()
        except Exception:
            continue
    return None


def _jsonify_mapping(row: Any) -> dict[str, Any]:
    data = dict(row)
    for key, value in list(data.items()):
        if isinstance(value, (datetime, date)):
            data[key] = value.isoformat()
        elif hasattr(value, "__float__") and value.__class__.__name__ == "Decimal":
            data[key] = float(value)
    return data
