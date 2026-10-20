"""Deterministic SKU return import, versioned brand rules and read APIs.

The spreadsheet directory is only read by the CLI. HTTP handlers call the
query functions in this module and never open a user's desktop files.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
import uuid
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from sqlalchemy import Engine, bindparam, text
from sqlalchemy.exc import SQLAlchemyError

from utils import import_service


MIGRATION = Path(__file__).resolve().parents[1] / "database" / "migrations" / "20260925_create_sku_return_tables.sql"
TARGET_TYPES = {
    "实销实结退货入库单": "return_in",
    "售后退货": "after_sales",
    "实销实结销售单": "sale",
}
HEADERS = {
    "supplier_name": {"供应商名称"},
    "document_type": {"应付帐单据类型", "应付账单据类型"},
    "sku_code": {"SKU编码"},
    "sku_name": {"SKU名称"},
    "quantity": {"数量"},
    "technical_document_id": {"单据ID"},
    "technical_document_no": {"单据编号"},
}
MONTH_RE = re.compile(r"^(1[0-2]|[1-9])月$")
YEAR_RE = re.compile(r"^(20\d{2})年(?:示例科技|示例户外)结算单$")
MAIN_ENTITIES = {"B2B", "示例生活", "示例科技"}
CATEGORIES = ("工业品", "五金", "箱包", "户外", "居家", "家电", "联想")
# Bump when source path/name classification semantics change. Including this
# value in the manifest makes a parser-only update publish a fresh data run.
IMPORT_PARSER_VERSION = 2
QUANTITY_SCALE = Decimal("0.000001")
ZERO = Decimal("0")
MAX_EXPORT_ROWS = 100_000

# Initial mappings are explicit, editable prefixes. Descriptive names remain
# 待识别. The rule table is versioned; an update never mutates an active run.
_INITIAL_ALIASES = {
    "测试旅行": ("测试旅行", "Samsonite"),
    "金宇": ("金宇", "Jinyu"),
    "DEMO_BAGS": ("Delsey", "测试箱包", "法国大使"),
    "测试运动": ("测试运动", "Aonijie"),
    "阿迪达斯": ("阿迪达斯", "Adidas"),
    "高山": ("高山", "High Sierra"),
    "测试家居": ("测试家居", "Campingmoon"),
    "Jordan": ("Jordan",),
    "测试服饰": ("测试服饰",),
    "联想": ("联想", "Lenovo"),
    "魅族": ("魅族", "Meizu"),
    "Rains": ("Rains",),
    "GRANITEGEAR": ("Granitegear", "花岗岩"),
    "iGear": ("iGear",),
    "VENQUE": ("Venque", "测试户外"),
    "迪卡侬": ("迪卡侬", "Decathlon"),
    "美旅": ("美旅", "American Tourister"),
    "京塑": ("京塑",),
    "米小圈": ("米小圈",),
    "测试五金": ("测试五金", "测试五金（DEMO_TOOLS）", "测试五金(DEMO_TOOLS)", "DEMO_TOOLS"),
    "demo_home": ("demo_home",),
}
DEFAULT_BRAND_RULES = [
    {"match_type": "prefix", "pattern": alias, "canonical_brand": brand, "priority": 0, "enabled": True}
    for brand, aliases in _INITIAL_ALIASES.items()
    for alias in aliases
]
ENTITY_BRAND_OVERRIDES = {("示例户外", "测试旅行"): "测试旅行企业购"}


class StaleRunError(RuntimeError):
    """The version requested by the page is no longer retained."""


class RuleConflictError(StaleRunError):
    """Brand rules were updated by another request."""


def _engine(engine: Engine | None = None) -> Engine:
    return engine or import_service.get_engine()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value).strip().strip("\ufeff"))


def _header(value: Any) -> str:
    return _clean(value).replace("\ufeff", "").replace("\u3000", "").replace("\r", "").replace("\n", "").replace(" ", "")


def _identifier(value: Any, number_format: str | None = None) -> tuple[str, str | None]:
    """Preserve text IDs and restore only an explicit Excel zero-fill format."""
    if value is None:
        return "", None
    if isinstance(value, bool):
        return str(value), None
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if number == number.to_integral_value():
            result = format(number.quantize(Decimal(1)), "f")
            if number_format and re.fullmatch(r"0{2,30}", number_format):
                return result.zfill(len(number_format)), "excel_zero_format"
            return result, "numeric_precision_risk" if len(result.lstrip("-")) > 15 else None
        return format(number, "f"), "numeric_fractional_identifier"
    return _clean(value), None


def _decimal(value: Any) -> Decimal:
    if value is None or _clean(value) == "":
        raise ValueError("数量为空")
    try:
        number = Decimal(_clean(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("数量不可解析") from exc
    if not number.is_finite():
        raise ValueError("数量不是有限数")
    if number.as_tuple().exponent < -6 or abs(number) >= Decimal("1000000000000000000"):
        raise ValueError("数量超出 DECIMAL(24,6) 范围或精度")
    return number.quantize(QUANTITY_SCALE)


def _decimal_str(value: Any) -> str:
    value = ZERO if value is None else Decimal(str(value))
    result = format(value, "f")
    return result.rstrip("0").rstrip(".") if "." in result else result


def _rate(sales: Any, returned: Any) -> str | None:
    sales_value = Decimal(str(sales or 0))
    if sales_value <= 0:
        return None
    return _decimal_str(Decimal(str(returned or 0)) / sales_value)


def _normalize_rules(rules: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for raw in rules:
        match_type = _clean(raw.get("match_type"))
        pattern = _clean(raw.get("pattern"))
        brand = _clean(raw.get("canonical_brand"))
        if match_type not in {"exact", "prefix"} or not pattern or not brand:
            raise ValueError("品牌规则需要 match_type=exact/prefix、pattern 与 canonical_brand")
        if len(pattern) > 178 or len(brand) > 191:
            raise ValueError("品牌匹配文本最多 178 字符，标准品牌最多 191 字符")
        key = (match_type, pattern.casefold())
        if key in seen:
            raise ValueError(f"品牌规则重复：{pattern}")
        seen.add(key)
        output.append({
            "id": raw.get("id"),
            "match_type": match_type,
            "pattern": pattern,
            "pattern_key": pattern.casefold(),
            "canonical_brand": brand,
            "priority": int(raw.get("priority") or 0),
            "enabled": bool(raw.get("enabled", True)),
        })
    return output


def match_brand(name: str, rules: Iterable[dict[str, Any]]) -> tuple[str, int | None]:
    key = _clean(name).casefold()
    matches = []
    for rule in rules:
        if not rule.get("enabled", True):
            continue
        pattern = str(rule.get("pattern_key") or _clean(rule.get("pattern")).casefold())
        if (rule["match_type"] == "exact" and key == pattern) or (rule["match_type"] == "prefix" and key.startswith(pattern)):
            matches.append(rule)
    if not matches:
        return "待识别", None
    matches.sort(key=lambda rule: (
        0 if rule["match_type"] == "exact" else 1,
        -len(str(rule.get("pattern_key") or _clean(rule.get("pattern")).casefold())),
        -int(rule.get("priority") or 0),
        int(rule.get("id") or 0),
    ))
    winner = matches[0]
    return str(winner["canonical_brand"]), winner.get("id")


def match_brand_for_entity(name: str, entity: str, rules: Iterable[dict[str, Any]]) -> tuple[str, int | None]:
    """Apply an explicit entity-specific display brand after normal name matching."""
    brand, rule_id = match_brand(name, rules)
    return ENTITY_BRAND_OVERRIDES.get((entity, brand), brand), rule_id


def _sha_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source_files(root: Path) -> list[Path]:
    if not root.is_dir():
        raise ValueError(f"导入根目录不存在：{root}")
    return sorted(
        (path for path in root.rglob("*.xlsx")
         if len(path.relative_to(root).parts) > 1
         and YEAR_RE.fullmatch(path.relative_to(root).parts[0])
         and not path.name.startswith(("~$", "."))
         and not any(part.startswith(".") for part in path.relative_to(root).parts)),
        key=lambda path: path.relative_to(root).as_posix().casefold(),
    )


def _category(filename: str, entity: str) -> tuple[str, str | None]:
    stem = Path(filename).stem
    stem = re.sub(r"^\s*\d+\s*", "", stem)
    stem = re.sub(r"\s*20\d{2}[-年][01]?\d(?:[-月][0-3]?\d日?)?\s*$", "", stem).strip()
    if stem.startswith(entity):
        stem = stem[len(entity):].strip()
    if entity == "B2B" and not stem:
        return "喜来登", None
    if entity == "示例科技" and stem == "电脑事业":
        return "电脑事业", None
    found = [category for category in CATEGORIES if category in stem]
    if len(found) == 1:
        category = found[0]
        warning = "new_entity_category" if (entity == "B2B" and category == "户外") else None
        return category, warning
    return "未分类", stem or "(无类目标签)"


def _parse_source(root: Path, path: Path, sheet_name: str) -> tuple[str, str, str, str | None]:
    parts = path.relative_to(root).parts
    if len(parts) == 4 and YEAR_RE.fullmatch(parts[0]) and parts[1] and parts[2] in MAIN_ENTITIES:
        month_match = MONTH_RE.fullmatch(parts[1])
        if not month_match:
            raise ValueError(f"非法月份目录：{path.relative_to(root)}")
        year = int(YEAR_RE.fullmatch(parts[0]).group(1))
        period = f"{year}-{int(month_match.group(1)):02d}"
        category, warning = _category(path.name, parts[2])
        return period, parts[2], category, warning
    if len(parts) == 2 and YEAR_RE.fullmatch(parts[0]) and "示例户外" in parts[0] and "示例户外" in path.stem:
        month_match = MONTH_RE.fullmatch(sheet_name)
        if not month_match:
            raise ValueError(f"示例户外非月份工作表：{sheet_name}")
        year = int(YEAR_RE.fullmatch(parts[0]).group(1))
        return f"{year}-{int(month_match.group(1)):02d}", "示例户外", "示例户外箱包", None
    raise ValueError(f"不支持的目录结构：{path.relative_to(root)}")


def _column_map(values: tuple[Any, ...]) -> dict[str, int] | None:
    normalized = [_header(value) for value in values]
    result = {}
    for field, aliases in HEADERS.items():
        candidates = [index for index, value in enumerate(normalized) if value in aliases]
        if len(candidates) > 1:
            raise ValueError(f"表头重复或歧义：{field}")
        if not candidates:
            return None
        result[field] = candidates[0]
    return result


def _manifest(root: Path, paths: list[Path]) -> tuple[str, dict[Path, str]]:
    digests = {path: _sha_file(path) for path in paths}
    lines = [f"parser\t{IMPORT_PARSER_VERSION}"]
    lines.extend(f"{path.relative_to(root).as_posix()}\t{digests[path]}" for path in paths)
    return sha256("\n".join(lines).encode("utf-8")).hexdigest(), digests


def _discover_units(root: Path, paths: list[Path], digests: dict[Path, str]) -> tuple[list[dict[str, Any]], list[str]]:
    units: list[dict[str, Any]] = []
    auxiliary: list[str] = []
    for path in paths:
        relative_path = path.relative_to(root).as_posix()
        parts = path.relative_to(root).parts
        is_weige = len(parts) == 2 and "示例户外" in parts[0]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            candidates = []
            for worksheet in workbook.worksheets:
                first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                columns = _column_map(first)
                if is_weige:
                    if MONTH_RE.fullmatch(worksheet.title):
                        if columns is None:
                            raise ValueError(f"示例户外月表缺少目标字段：{relative_path}#{worksheet.title}")
                        candidates.append((worksheet.title, columns))
                    else:
                        auxiliary.append(f"{relative_path}#{worksheet.title}")
                elif columns is not None:
                    candidates.append((worksheet.title, columns))
                else:
                    auxiliary.append(f"{relative_path}#{worksheet.title}")
            if is_weige:
                if not candidates:
                    raise ValueError(f"示例户外工作簿没有有效月份表：{relative_path}")
            elif len(candidates) != 1:
                raise ValueError(f"主目录工作簿应有且仅有一张业务表，实际 {len(candidates)} 张：{relative_path}")
            for sheet_name, columns in candidates:
                period, entity, category, warning = _parse_source(root, path, sheet_name)
                units.append({
                    "id": uuid.uuid4().hex,
                    "path": path,
                    "relative_path": relative_path,
                    "path_sha256": sha256(relative_path.encode("utf-8")).hexdigest(),
                    "file_sha256": digests[path],
                    "sheet_name": sheet_name,
                    "columns": columns,
                    "period": period,
                    "entity": entity,
                    "category": category,
                    "category_warning": warning,
                    "scanned_rows": 0,
                    "target_rows": 0,
                    "included_rows": 0,
                })
        finally:
            workbook.close()
    units.sort(key=lambda unit: (unit["period"], unit["relative_path"], unit["sheet_name"]))
    return units, auxiliary


def scan_source_tree(
    root: str | Path,
    *,
    rules: Iterable[dict[str, Any]] | None = None,
    on_unit: Callable[[dict[str, Any], bool], None] | None = None,
    on_fact_batch: Callable[[list[dict[str, Any]]], None] | None = None,
    batch_size: int = 1000,
    run_id: str | None = None,
) -> dict[str, Any]:
    """Scan and reconcile a full tree; without sinks this is a read-only dry run.

    The multiset of same-key rows in each source unit is compared with the
    largest prior source multiset. This retains true duplicates within a
    source while excluding one-to-one cross-source copies.
    """
    root_path = Path(root).resolve()
    paths = _source_files(root_path)
    if not paths:
        raise ValueError("导入目录中没有 .xlsx 工作簿")
    manifest, digests = _manifest(root_path, paths)
    units, auxiliary = _discover_units(root_path, paths, digests)
    normalized_rules = _normalize_rules(rules if rules is not None else DEFAULT_BRAND_RULES)
    report: dict[str, Any] = {
        "manifest_sha256": manifest,
        "file_count": len(paths),
        "source_units": len(units),
        "auxiliary_sheets": len(auxiliary),
        "scanned_rows": 0,
        "target_rows": 0,
        "isolated_rows": 0,
        "duplicate_rows": 0,
        "included_rows": 0,
        "by_document_type": Counter(),
        "by_period_entity": Counter(),
        "quality": Counter(),
        "unknown_categories": [],
        "new_category_combinations": [],
        "isolated_samples": [],
        "duplicate_samples": [],
        "warnings": [],
        "coverage": [],
    }
    for unit in units:
        if unit["category"] == "未分类" and unit["entity"] != "示例户外":
            report["unknown_categories"].append({"path": unit["relative_path"], "token": unit["category_warning"]})
        elif unit["category_warning"] == "new_entity_category":
            report["new_category_combinations"].append({"path": unit["relative_path"], "entity": unit["entity"], "category": unit["category"]})

    # Key: full SHA-256 bytes; value: earliest retained occurrence per ordinal.
    first_by_key: dict[bytes, list[tuple[str, int, str, str]]] = {}
    batch: list[dict[str, Any]] = []
    for unit in units:
        if on_unit:
            on_unit(unit, False)
        workbook = load_workbook(unit["path"], read_only=True, data_only=True)
        local_occurrences: dict[bytes, int] = defaultdict(int)
        try:
            worksheet = workbook[unit["sheet_name"]]
            rows = worksheet.iter_rows(min_row=2, values_only=False)
            for source_row, cells in enumerate(rows, start=2):
                if not any(cell.value is not None and _clean(cell.value) for cell in cells):
                    continue
                report["scanned_rows"] += 1
                unit["scanned_rows"] += 1
                columns = unit["columns"]

                def cell(field: str) -> Any:
                    index = columns[field]
                    return cells[index] if index < len(cells) else None

                document_type = _clean(cell("document_type").value if cell("document_type") else None)
                if document_type not in TARGET_TYPES:
                    continue
                report["target_rows"] += 1
                unit["target_rows"] += 1
                report["by_document_type"][document_type] += 1
                supplier_raw = _clean(cell("supplier_name").value if cell("supplier_name") else None)
                sku_name_raw = _clean(cell("sku_name").value if cell("sku_name") else None)
                supplier = supplier_raw or "未提供"
                sku_name = sku_name_raw or "未提供名称"
                sku_cell = cell("sku_code")
                sku_code, sku_warning = _identifier(sku_cell.value if sku_cell else None, sku_cell.number_format if sku_cell else None)
                doc_id, _ = _identifier(cell("technical_document_id").value if cell("technical_document_id") else None)
                doc_no, _ = _identifier(cell("technical_document_no").value if cell("technical_document_no") else None)
                if not supplier_raw:
                    report["quality"]["empty_supplier"] += 1
                if not sku_name_raw:
                    report["quality"]["empty_sku_name"] += 1
                if sku_warning:
                    report["quality"][sku_warning] += 1
                if not doc_id or not doc_no:
                    report["quality"]["missing_technical_key"] += 1

                reason = ""
                if not doc_id or not doc_no:
                    reason = "missing_technical_key"
                quantity: Decimal | None = None
                try:
                    quantity_cell = cell("quantity")
                    quantity = _decimal(quantity_cell.value if quantity_cell else None)
                except ValueError as exc:
                    reason = str(exc)
                    report["quality"]["invalid_quantity"] += 1
                if not sku_code:
                    reason = reason or "SKU编码为空"
                    report["quality"]["empty_sku_code"] += 1
                if len(sku_code) > 191 or len(sku_name) > 512 or len(supplier) > 255 or len(doc_id) > 191 or len(doc_no) > 191:
                    reason = reason or "源文本超过数据库字段长度"
                    report["quality"]["oversized_text"] += 1
                if quantity is not None and TARGET_TYPES[document_type] == "sale" and quantity < 0:
                    report["quality"]["negative_sales"] += 1

                brand, brand_rule_id = match_brand_for_entity(sku_name, unit["entity"], normalized_rules)
                if brand == "待识别":
                    report["quality"]["unknown_brand_rows"] += 1
                sales = ZERO
                return_in = ZERO
                after_sales = ZERO
                if quantity is not None:
                    kind = TARGET_TYPES[document_type]
                    if kind == "sale":
                        sales = quantity
                    elif kind == "return_in":
                        return_in = abs(quantity)
                    else:
                        after_sales = abs(quantity)

                dedupe_hash: str | None = None
                duplicate_ref: tuple[str, int, str, str] | None = None
                if reason:
                    report["isolated_rows"] += 1
                    if len(report["isolated_samples"]) < 1000:
                        report["isolated_samples"].append({"path": unit["relative_path"], "sheet": unit["sheet_name"], "row": source_row, "reason": reason})
                elif doc_id and doc_no:
                    payload = [doc_id, doc_no, unit["entity"], document_type, supplier, sku_code, sku_name, _decimal_str(quantity)]
                    digest = sha256(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).digest()
                    dedupe_hash = digest.hex()
                    occurrence = local_occurrences[digest]
                    local_occurrences[digest] += 1
                    origins = first_by_key.setdefault(digest, [])
                    if occurrence < len(origins):
                        duplicate_ref = origins[occurrence]
                        reason = "cross_source_duplicate"
                        report["duplicate_rows"] += 1
                        if unit["category"] != duplicate_ref[2]:
                            report["quality"]["duplicate_category_mismatch"] += 1
                        if len(report["duplicate_samples"]) < 1000:
                            report["duplicate_samples"].append({
                                "path": unit["relative_path"], "sheet": unit["sheet_name"], "row": source_row,
                                "first_path": duplicate_ref[3], "first_row": duplicate_ref[1],
                                "first_category": duplicate_ref[2], "later_category": unit["category"],
                            })
                    else:
                        origins.append((unit["id"], source_row, unit["category"], unit["relative_path"]))
                if not reason:
                    report["included_rows"] += 1
                    unit["included_rows"] += 1
                    report["by_period_entity"][(unit["period"], unit["entity"])] += 1
                if on_fact_batch:
                    batch.append({
                        "run_id": run_id,
                        "source_unit_id": unit["id"], "source_row": source_row,
                        "document_type": document_type,
                        "supplier_name": supplier, "supplier_raw": supplier_raw,
                        "sku_code": sku_code[:191], "sku_name": sku_name[:512], "sku_name_raw": sku_name_raw[:512],
                        "brand": brand, "brand_rule_id": brand_rule_id,
                        "raw_quantity": quantity, "sales_quantity": sales,
                        "return_in_quantity": return_in, "after_sales_quantity": after_sales,
                        "technical_document_id": doc_id[:191], "technical_document_no": doc_no[:191],
                        "dedupe_sha256": dedupe_hash,
                        "included": 0 if reason else 1,
                        "excluded_reason": reason,
                        "duplicate_of_source_unit_id": duplicate_ref[0] if duplicate_ref else None,
                        "duplicate_of_row": duplicate_ref[1] if duplicate_ref else None,
                    })
                    if len(batch) >= batch_size:
                        on_fact_batch(batch)
                        batch = []
        finally:
            workbook.close()
        if batch and on_fact_batch:
            on_fact_batch(batch)
            batch = []
        if on_unit:
            on_unit(unit, True)
        report["coverage"].append({
            "period": unit["period"], "entity": unit["entity"], "category": unit["category"],
            "path": unit["relative_path"], "sheet": unit["sheet_name"],
            "scanned_rows": unit["scanned_rows"], "target_rows": unit["target_rows"],
            "included_rows": unit["included_rows"],
        })
    if report["target_rows"] != report["isolated_rows"] + report["duplicate_rows"] + report["included_rows"]:
        raise AssertionError("目标行与隔离、去重、计入行不平衡")
    report["by_document_type"] = dict(report["by_document_type"])
    report["by_period_entity"] = [
        {"period": period, "entity": entity, "included_rows": count}
        for (period, entity), count in sorted(report["by_period_entity"].items())
    ]
    report["quality"] = dict(report["quality"])
    return report


def ensure_schema(engine: Engine | None = None) -> None:
    """Create the MySQL tables and first immutable brand rule set if missing."""
    db_engine = _engine(engine)
    script = re.sub(r"(?m)^--.*$", "", MIGRATION.read_text(encoding="utf-8"))
    with db_engine.begin() as conn:
        for statement in script.split(";"):
            if statement.strip():
                conn.execute(text(statement))
        conn.execute(text("INSERT IGNORE INTO sku_return_state (id, active_run_id, rule_version, updated_at) VALUES (1, NULL, 1, CURRENT_TIMESTAMP)"))
        existing = conn.execute(text("SELECT COUNT(*) FROM sku_return_brand_rule WHERE rule_version=1")).scalar_one()
        if not existing:
            params = [{**rule, "rule_version": 1, "updated_by": "system"} for rule in _normalize_rules(DEFAULT_BRAND_RULES)]
            conn.execute(text("""
                INSERT INTO sku_return_brand_rule
                  (rule_version, match_type, pattern, pattern_key, canonical_brand, priority, enabled, updated_at, updated_by)
                VALUES (:rule_version, :match_type, :pattern, :pattern_key, :canonical_brand, :priority, :enabled, CURRENT_TIMESTAMP, :updated_by)
            """), params)


def _state(conn: Any) -> dict[str, Any]:
    result = conn.execute(text("SELECT active_run_id, rule_version, updated_at FROM sku_return_state WHERE id=1")).mappings().first()
    return dict(result) if result else {"active_run_id": None, "rule_version": 1, "updated_at": None}


def _selected_run(conn: Any, run_id: str | None = None) -> dict[str, Any] | None:
    state = _state(conn)
    selected = run_id or state["active_run_id"]
    if not selected:
        return None
    result = conn.execute(text("""
        SELECT id, manifest_sha256, rule_version, status, started_at, completed_at,
               source_units, scanned_rows, target_rows, isolated_rows, duplicate_rows,
               included_rows, report_json
        FROM sku_return_import_run WHERE id=:run_id
    """), {"run_id": selected}).mappings().first()
    if not result or result["status"] != "complete":
        if run_id:
            raise StaleRunError("该数据版本已不可用，请刷新页面")
        return None
    return dict(result)


def _current_rules(conn: Any, version: int) -> list[dict[str, Any]]:
    rows = conn.execute(text("""
        SELECT id, match_type, pattern, pattern_key, canonical_brand, priority, enabled
        FROM sku_return_brand_rule WHERE rule_version=:version ORDER BY id
    """), {"version": version}).mappings().all()
    return [dict(row) for row in rows]


_FACT_INSERT = text("""
    INSERT INTO sku_return_fact (
      run_id, source_unit_id, source_row, document_type, supplier_name, supplier_raw,
      sku_code, sku_name, sku_name_raw, brand, brand_rule_id, raw_quantity,
      sales_quantity, return_in_quantity, after_sales_quantity,
      technical_document_id, technical_document_no, dedupe_sha256, included,
      excluded_reason, duplicate_of_source_unit_id, duplicate_of_row
    ) VALUES (
      :run_id, :source_unit_id, :source_row, :document_type, :supplier_name, :supplier_raw,
      :sku_code, :sku_name, :sku_name_raw, :brand, :brand_rule_id, :raw_quantity,
      :sales_quantity, :return_in_quantity, :after_sales_quantity,
      :technical_document_id, :technical_document_no, :dedupe_sha256, :included,
      :excluded_reason, :duplicate_of_source_unit_id, :duplicate_of_row
    )
""")


def _build_aggregate(conn: Any, run_id: str) -> None:
    conn.execute(text("""
        INSERT INTO sku_return_aggregate (
          run_id, period, entity, category, supplier_name, sku_code, sku_name, brand,
          sales_quantity, return_in_quantity, after_sales_quantity, fact_count
        )
        SELECT f.run_id, u.period, u.entity, u.category, f.supplier_name,
               f.sku_code, f.sku_name, f.brand,
               SUM(f.sales_quantity), SUM(f.return_in_quantity), SUM(f.after_sales_quantity), COUNT(*)
        FROM sku_return_fact f JOIN sku_return_source_unit u ON u.id=f.source_unit_id
        WHERE f.run_id=:run_id AND f.included=1
        GROUP BY f.run_id, u.period, u.entity, u.category, f.supplier_name,
                 f.sku_code, f.sku_name, f.brand
    """), {"run_id": run_id})


def import_source_tree(root: str | Path, *, apply: bool = False, engine: Engine | None = None) -> dict[str, Any]:
    """Preview or publish the complete source tree.

    Dry run does not request an engine and cannot mutate any database state.
    Failed applied runs remain inspectable, while the single active pointer
    remains unchanged. Repeating the same manifest and rule version is a no-op.
    """
    if not apply:
        return scan_source_tree(root)
    db_engine = _engine(engine)
    ensure_schema(db_engine)
    root_path = Path(root).resolve()
    paths = _source_files(root_path)
    if not paths:
        raise ValueError("导入目录中没有 .xlsx 工作簿")
    manifest, _ = _manifest(root_path, paths)
    with db_engine.connect() as conn:
        state = _state(conn)
        active = _selected_run(conn)
        rules = _current_rules(conn, int(state["rule_version"]))
        if active and active["manifest_sha256"] == manifest and active["rule_version"] == state["rule_version"]:
            return {"status": "unchanged", "run_id": active["id"], "manifest_sha256": manifest,
                    "report": json.loads(active["report_json"] or "{}")}
    run_id = uuid.uuid4().hex
    with db_engine.begin() as conn:
        conn.execute(text("""
            INSERT INTO sku_return_import_run (id, manifest_sha256, rule_version, status, started_at)
            VALUES (:id, :manifest, :version, 'building', CURRENT_TIMESTAMP)
        """), {"id": run_id, "manifest": manifest, "version": state["rule_version"]})

    def write_unit(unit: dict[str, Any], finished: bool) -> None:
        with db_engine.begin() as conn:
            if not finished:
                conn.execute(text("""
                    INSERT INTO sku_return_source_unit (
                      id, run_id, relative_path, path_sha256, file_sha256, sheet_name,
                      period, entity, category
                    ) VALUES (:id, :run_id, :relative_path, :path_sha256, :file_sha256,
                              :sheet_name, :period, :entity, :category)
                """), {**unit, "run_id": run_id})
            else:
                conn.execute(text("""
                    UPDATE sku_return_source_unit SET scanned_rows=:scanned_rows,
                      target_rows=:target_rows, included_rows=:included_rows WHERE id=:id
                """), unit)

    def write_facts(batch: list[dict[str, Any]]) -> None:
        with db_engine.begin() as conn:
            conn.execute(_FACT_INSERT, batch)

    try:
        report = scan_source_tree(root_path, rules=rules, on_unit=write_unit, on_fact_batch=write_facts, run_id=run_id)
        if report["manifest_sha256"] != manifest:
            raise RuntimeError("源文件在导入期间变化，请重新执行")
        if report["source_units"] < 1 or report["included_rows"] < 1:
            raise ValueError("没有可发布的有效数据")
        with db_engine.begin() as conn:
            source_count = conn.execute(text("SELECT COUNT(*) FROM sku_return_source_unit WHERE run_id=:run_id"), {"run_id": run_id}).scalar_one()
            fact_count = conn.execute(text("SELECT COUNT(*) FROM sku_return_fact WHERE run_id=:run_id"), {"run_id": run_id}).scalar_one()
            included_count = conn.execute(text("SELECT COUNT(*) FROM sku_return_fact WHERE run_id=:run_id AND included=1"), {"run_id": run_id}).scalar_one()
            if (source_count, fact_count, included_count) != (report["source_units"], report["target_rows"], report["included_rows"]):
                raise AssertionError("数据库记录数与扫描质量报告不一致")
            _build_aggregate(conn, run_id)
            aggregate_count = conn.execute(text("SELECT COALESCE(SUM(fact_count),0) FROM sku_return_aggregate WHERE run_id=:run_id"), {"run_id": run_id}).scalar_one()
            if aggregate_count != report["included_rows"]:
                raise AssertionError("预聚合事实数与计入行数不一致")
            conn.execute(text("""
                UPDATE sku_return_import_run SET status='complete', completed_at=CURRENT_TIMESTAMP,
                  source_units=:source_units, scanned_rows=:scanned_rows, target_rows=:target_rows,
                  isolated_rows=:isolated_rows, duplicate_rows=:duplicate_rows,
                  included_rows=:included_rows, report_json=:report_json WHERE id=:run_id
            """), {**{name: report[name] for name in ("source_units", "scanned_rows", "target_rows", "isolated_rows", "duplicate_rows", "included_rows")},
                    "report_json": json.dumps(report, ensure_ascii=False), "run_id": run_id})
            conn.execute(text("SELECT active_run_id FROM sku_return_state WHERE id=1 FOR UPDATE"))
            current_version = conn.execute(text("SELECT rule_version FROM sku_return_state WHERE id=1")).scalar_one()
            if current_version != state["rule_version"]:
                raise RuleConflictError("导入期间品牌规则已变化，请重新执行")
            conn.execute(text("UPDATE sku_return_state SET active_run_id=:run_id, updated_at=CURRENT_TIMESTAMP WHERE id=1"), {"run_id": run_id})
        return {"status": "published", "run_id": run_id, "manifest_sha256": manifest, "report": report}
    except Exception as exc:
        with db_engine.begin() as conn:
            conn.execute(text("""
                UPDATE sku_return_import_run SET status='failed', completed_at=CURRENT_TIMESTAMP,
                  error_message=:error WHERE id=:run_id
            """), {"run_id": run_id, "error": str(exc)[:4000]})
        raise


def _empty_metadata() -> dict[str, Any]:
    return {"run_id": None, "manifest_sha256": None, "rule_version": None,
            "imported_at": None, "periods": [], "entities": [], "categories": [],
            "brands": [], "coverage": [], "quality": {}}


def _iso(value: Any) -> str | None:
    return value.isoformat(timespec="seconds") if isinstance(value, datetime) else None


def get_metadata(run_id: str | None = None, *, engine: Engine | None = None) -> dict[str, Any]:
    try:
        with _engine(engine).connect() as conn:
            run = _selected_run(conn, run_id)
            if not run:
                return _empty_metadata()
            units = conn.execute(text("""
                SELECT period, entity, COUNT(*) AS source_units, SUM(included_rows) AS included_rows
                FROM sku_return_source_unit WHERE run_id=:run_id
                GROUP BY period, entity ORDER BY period, entity
            """), {"run_id": run["id"]}).mappings().all()
            dimensions = {}
            for column in ("entity", "category", "brand"):
                dimensions[column] = [row[0] for row in conn.execute(
                    text(f"SELECT DISTINCT {column} FROM sku_return_aggregate WHERE run_id=:run_id ORDER BY {column}"),
                    {"run_id": run["id"]}).all()]
            report = json.loads(run["report_json"] or "{}")
            quality = {**report.get("quality", {}),
                       "isolated_rows": int(run["isolated_rows"]),
                       "duplicate_rows": int(run["duplicate_rows"]),
                       "included_rows": int(run["included_rows"]),
                       "target_rows": int(run["target_rows"]),
                       "unknown_category_files": len(report.get("unknown_categories", [])),
                       "new_category_combinations": len(report.get("new_category_combinations", []))}
            return {
                "run_id": run["id"], "manifest_sha256": run["manifest_sha256"],
                "rule_version": run["rule_version"], "imported_at": _iso(run["completed_at"]),
                "periods": sorted({row["period"] for row in units}),
                "entities": dimensions["entity"], "categories": dimensions["category"],
                "brands": dimensions["brand"],
                "coverage": [{"period": row["period"], "entity": row["entity"],
                              "source_units": int(row["source_units"]), "included_rows": int(row["included_rows"] or 0)} for row in units],
                "quality": quality,
            }
    except SQLAlchemyError as exc:
        if run_id:
            raise
        code = getattr(getattr(exc, "orig", None), "args", (None,))[0]
        message = str(exc).lower()
        if code == 1146 or "no such table: sku_return_state" in message:
            # A fresh installation before its migration has no active data.
            return _empty_metadata()
        raise


def _filters(filters: dict[str, Any] | None, *, alias: str = "a") -> tuple[str, dict[str, Any]]:
    filters = filters or {}
    clauses: list[str] = []
    params: dict[str, Any] = {}
    for input_name, column in (("period", "period"), ("entity", "entity"),
                               ("category", "category"), ("brand", "brand"),
                               ("supplier", "supplier_name")):
        raw = filters.get(input_name) or []
        values = [raw] if isinstance(raw, str) else list(raw)
        values = [_clean(value) for value in values if _clean(value)]
        if len(values) > 100:
            raise ValueError(f"{input_name} 筛选值过多")
        if input_name == "period" and any(not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", value) for value in values):
            raise ValueError("月份格式应为 YYYY-MM")
        if any(len(value) > 255 for value in values):
            raise ValueError(f"{input_name} 筛选值过长")
        if values:
            names = []
            for index, value in enumerate(values):
                param_name = f"{input_name}_{index}"
                params[param_name] = value
                names.append(f":{param_name}")
            clauses.append(f"{alias}.{column} IN ({', '.join(names)})")
    keyword = _clean(filters.get("sku_keyword") or filters.get("sku"))
    if len(keyword) > 191:
        raise ValueError("SKU 关键字过长")
    if keyword:
        escaped = keyword.replace("!", "!!").replace("%", "!%").replace("_", "!_")
        params["sku_keyword"] = f"%{escaped}%"
        clauses.append(f"({alias}.sku_code LIKE :sku_keyword ESCAPE '!' OR {alias}.sku_name LIKE :sku_keyword ESCAPE '!')")
    return (" AND " + " AND ".join(clauses)) if clauses else "", params


def _scope(filters: dict[str, Any] | None) -> dict[str, Any]:
    filters = filters or {}
    return {key: filters.get(key) or ([] if key != "sku_keyword" else "")
            for key in ("period", "entity", "category", "brand", "supplier", "sku_keyword")}


def get_filter_options(
    filters: dict[str, Any] | None = None,
    run_id: str | None = None,
    *,
    engine: Engine | None = None,
) -> dict[str, Any]:
    """Return each filter's choices in the context of the other current filters."""
    columns = {
        "period": "period",
        "entity": "entity",
        "category": "category",
        "brand": "brand",
    }
    selected = filters or {}
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        if not run:
            return {"run_id": None, **{key: [] for key in columns}}

        result: dict[str, Any] = {"run_id": run["id"]}
        for key, column in columns.items():
            facet_filters = {**selected, key: []}
            where, params = _filters(facet_filters, alias="a")
            rows = conn.execute(text(f"""
                SELECT DISTINCT a.{column}
                FROM sku_return_aggregate a
                WHERE a.run_id=:run_id {where}
                ORDER BY a.{column}
            """), {"run_id": run["id"], **params}).all()
            result[key] = [row[0] for row in rows]
        return result


def list_suppliers(
    keyword: str = "",
    limit: int = 50,
    run_id: str | None = None,
    filters: dict[str, Any] | None = None,
    *,
    engine: Engine | None = None,
) -> dict[str, Any]:
    if not 1 <= limit <= 200:
        raise ValueError("limit 必须在 1–200 之间")
    keyword = _clean(keyword)
    if len(keyword) > 191:
        raise ValueError("供应商关键字过长")
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        if not run:
            return {"run_id": None, "items": []}
        filter_scope = {**(filters or {}), "supplier": []}
        where, params = _filters(filter_scope, alias="a")
        statement = text(f"""
            SELECT DISTINCT a.supplier_name FROM sku_return_aggregate a
            WHERE a.run_id=:run_id {where}
              AND a.supplier_name LIKE :keyword ESCAPE '!'
              AND LOWER(TRIM(a.supplier_name)) NOT IN ('null', 'none', 'nan')
            ORDER BY a.supplier_name LIMIT :limit
        """)
        escaped = keyword.replace("!", "!!").replace("%", "!%").replace("_", "!_")
        items = [row[0] for row in conn.execute(
            statement,
            {"run_id": run["id"], "keyword": f"%{escaped}%", "limit": limit, **params},
        ).all()]
        return {"run_id": run["id"], "items": items}


def get_summary(filters: dict[str, Any] | None = None, run_id: str | None = None, *, engine: Engine | None = None) -> dict[str, Any]:
    where, params = _filters(filters)
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        result = {"run_id": run["id"] if run else None, "sales_quantity": "0",
                  "return_in_quantity": "0", "after_sales_quantity": "0",
                  "return_quantity": "0", "return_rate": None,
                  "sku_count": 0, "no_sales_sku_count": 0, "scope": _scope(filters)}
        if not run:
            return result
        args = {"run_id": run["id"], **params}
        row = conn.execute(text(f"""
            SELECT COALESCE(SUM(a.sales_quantity),0) AS sales_quantity,
                   COALESCE(SUM(a.return_in_quantity),0) AS return_in_quantity,
                   COALESCE(SUM(a.after_sales_quantity),0) AS after_sales_quantity
            FROM sku_return_aggregate a WHERE a.run_id=:run_id {where}
        """), args).mappings().one()
        counts = conn.execute(text(f"""
            SELECT COUNT(*) AS sku_count,
                   COALESCE(SUM(CASE WHEN sales_quantity<=0 THEN 1 ELSE 0 END),0) AS no_sales_sku_count
            FROM (
                SELECT a.sku_code, a.sku_name, SUM(a.sales_quantity) AS sales_quantity
                FROM sku_return_aggregate a WHERE a.run_id=:run_id {where}
                GROUP BY a.sku_code, a.sku_name
            ) grouped_skus
        """), args).mappings().one()
        returned = Decimal(str(row["return_in_quantity"])) + Decimal(str(row["after_sales_quantity"]))
        result.update({
            "sales_quantity": _decimal_str(row["sales_quantity"]),
            "return_in_quantity": _decimal_str(row["return_in_quantity"]),
            "after_sales_quantity": _decimal_str(row["after_sales_quantity"]),
            "return_quantity": _decimal_str(returned),
            "return_rate": _rate(row["sales_quantity"], returned),
            "sku_count": int(counts["sku_count"]),
            "no_sales_sku_count": int(counts["no_sales_sku_count"]),
        })
        return result


def get_monthly_trend(filters: dict[str, Any] | None = None, run_id: str | None = None, *, engine: Engine | None = None) -> dict[str, Any]:
    """按结算归属月汇总当前筛选结果；比率始终由月度合计重新计算。"""
    where, params = _filters(filters)
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        result = {"run_id": run["id"] if run else None, "scope": _scope(filters), "items": []}
        if not run:
            return result
        rows = conn.execute(text(f"""
            SELECT a.period, SUM(a.sales_quantity) AS sales_quantity,
                   SUM(a.return_in_quantity) AS return_in_quantity,
                   SUM(a.after_sales_quantity) AS after_sales_quantity
            FROM sku_return_aggregate a
            WHERE a.run_id=:run_id {where}
            GROUP BY a.period ORDER BY a.period ASC
        """), {"run_id": run["id"], **params}).mappings().all()
        result["items"] = []
        for row in rows:
            returned = Decimal(str(row["return_in_quantity"])) + Decimal(str(row["after_sales_quantity"]))
            result["items"].append({
                "period": row["period"],
                "sales_quantity": _decimal_str(row["sales_quantity"]),
                "return_in_quantity": _decimal_str(row["return_in_quantity"]),
                "after_sales_quantity": _decimal_str(row["after_sales_quantity"]),
                "return_quantity": _decimal_str(returned),
                "return_rate": _rate(row["sales_quantity"], returned),
            })
        return result


_SORT_COLUMNS = {
    "sales_quantity": "sales_quantity", "return_in_quantity": "return_in_quantity",
    "after_sales_quantity": "after_sales_quantity", "return_quantity": "return_quantity",
    "return_rate": "return_rate", "sku_code": "sku_code", "sku_name": "sku_name",
    "period": "period",
}


def _items_sql(filters: dict[str, Any] | None, view: str, sort_by: str, sort_dir: str) -> tuple[str, str, dict[str, Any]]:
    if view not in {"sku", "detail"}:
        raise ValueError("view 必须是 sku 或 detail")
    if sort_by not in _SORT_COLUMNS or (sort_by == "period" and view != "detail"):
        raise ValueError("不支持的排序字段")
    if sort_dir not in {"asc", "desc"}:
        raise ValueError("sort_dir 必须是 asc 或 desc")
    where, params = _filters(filters)
    detail = "a.period, a.entity, a.category, a.supplier_name, " if view == "detail" else ""
    group = f"{detail}a.sku_code, a.sku_name, a.brand"
    base = f"""
        SELECT {group}, SUM(a.sales_quantity) AS sales_quantity,
               SUM(a.return_in_quantity) AS return_in_quantity,
               SUM(a.after_sales_quantity) AS after_sales_quantity,
               SUM(a.return_in_quantity)+SUM(a.after_sales_quantity) AS return_quantity,
               CASE WHEN SUM(a.sales_quantity)>0 THEN
                 (SUM(a.return_in_quantity)+SUM(a.after_sales_quantity))/SUM(a.sales_quantity)
               ELSE NULL END AS return_rate,
               COUNT(DISTINCT a.entity) AS entity_count,
               COUNT(DISTINCT a.supplier_name) AS supplier_count
        FROM sku_return_aggregate a WHERE a.run_id=:run_id {where}
        GROUP BY {group}
    """
    sort_column = _SORT_COLUMNS[sort_by]
    null_sort = "CASE WHEN result.return_rate IS NULL THEN 1 ELSE 0 END ASC, " if sort_by == "return_rate" else ""
    tie_breakers = "result.sku_code ASC, result.sku_name ASC, result.brand ASC"
    if view == "detail":
        tie_breakers += ", result.period ASC, result.entity ASC, result.category ASC, result.supplier_name ASC"
    ordered = f"SELECT * FROM ({base}) result ORDER BY {null_sort}result.{sort_column} {sort_dir.upper()}, {tie_breakers}"
    count = f"SELECT COUNT(*) FROM ({base}) result"
    return ordered, count, params


def _item_dict(row: Any) -> dict[str, Any]:
    item = dict(row)
    for key in ("sales_quantity", "return_in_quantity", "after_sales_quantity", "return_quantity"):
        item[key] = _decimal_str(item[key])
    item["return_rate"] = _decimal_str(item["return_rate"]) if item["return_rate"] is not None else None
    item["entity_count"] = int(item["entity_count"])
    item["supplier_count"] = int(item["supplier_count"])
    return item


def list_items(
    filters: dict[str, Any] | None = None, view: str = "sku", page: int = 1,
    page_size: int = 50, sort_by: str = "return_quantity", sort_dir: str = "desc",
    run_id: str | None = None, *, engine: Engine | None = None,
) -> dict[str, Any]:
    if page < 1 or not 1 <= page_size <= 200:
        raise ValueError("page 必须大于 0，page_size 必须在 1–200 之间")
    ordered, count_sql, params = _items_sql(filters, view, sort_by, sort_dir)
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        if not run:
            return {"run_id": None, "total": 0, "items": [], "page": page, "page_size": page_size, "scope": _scope(filters)}
        args = {"run_id": run["id"], **params}
        total = int(conn.execute(text(count_sql), args).scalar_one())
        rows = conn.execute(text(ordered + " LIMIT :limit OFFSET :offset"),
                            {**args, "limit": page_size, "offset": (page - 1) * page_size}).mappings().all()
        return {"run_id": run["id"], "total": total, "items": [_item_dict(row) for row in rows],
                "page": page, "page_size": page_size, "scope": _scope(filters)}


def list_sources(
    sku_code: str, sku_name: str, filters: dict[str, Any] | None = None,
    period: str | None = None, entity: str | None = None, category: str | None = None,
    supplier_name: str | None = None, limit: int = 100, run_id: str | None = None,
    *, engine: Engine | None = None,
) -> dict[str, Any]:
    if not _clean(sku_code) or len(sku_code) > 191 or not 1 <= limit <= 500:
        raise ValueError("请提供 SKU 编码，limit 须为 1–500")
    name = _clean(sku_name) or "未提供名称"
    if len(name) > 512:
        raise ValueError("SKU 名称过长")
    where, params = _filters(filters, alias="f")
    # The generic filter builder references dimensions on the aggregate row;
    # source lookup has those dimensions on the source unit instead.
    where = where.replace("f.period", "u.period").replace("f.entity", "u.entity").replace("f.category", "u.category")
    for column, value in (("period", period), ("entity", entity), ("category", category),
                          ("supplier_name", supplier_name)):
        if value is not None and _clean(value):
            if len(value) > 255:
                raise ValueError("来源维度筛选值过长")
            alias = "f" if column == "supplier_name" else "u"
            where += f" AND {alias}.{column}=:detail_{column}"
            params[f"detail_{column}"] = _clean(value)
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        if not run:
            return {"run_id": None, "items": []}
        rows = conn.execute(text(f"""
            SELECT u.relative_path, u.sheet_name, u.period, u.entity, u.category,
                   f.source_row AS `row_number`, f.document_type, f.supplier_name, f.raw_quantity,
                   f.included, f.excluded_reason
            FROM sku_return_fact f JOIN sku_return_source_unit u ON u.id=f.source_unit_id
            WHERE f.run_id=:run_id AND f.sku_code=:sku_code AND f.sku_name=:sku_name {where}
            ORDER BY u.period, u.relative_path, u.sheet_name, f.source_row LIMIT :limit
        """), {"run_id": run["id"], "sku_code": _clean(sku_code), "sku_name": name,
                "limit": limit, **params}).mappings().all()
        items = []
        for row in rows:
            item = dict(row)
            item["raw_quantity"] = _decimal_str(item["raw_quantity"]) if item["raw_quantity"] is not None else None
            item["included"] = bool(item["included"])
            items.append(item)
        return {"run_id": run["id"], "items": items}


def _text_cell(worksheet: Any, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value="" if value is None else str(value))
    cell.data_type = "s"  # Never treat untrusted SKU/supplier strings as formulas.
    return cell


def export_items(
    filters: dict[str, Any] | None = None, view: str = "sku",
    sort_by: str = "return_quantity", sort_dir: str = "desc",
    run_id: str | None = None, *, engine: Engine | None = None,
) -> tuple[str, bytes]:
    ordered, count_sql, params = _items_sql(filters, view, sort_by, sort_dir)
    with _engine(engine).connect() as conn:
        run = _selected_run(conn, run_id)
        rows: list[dict[str, Any]] = []
        if run:
            args = {"run_id": run["id"], **params}
            total = int(conn.execute(text(count_sql), args).scalar_one())
            if total > MAX_EXPORT_ROWS:
                raise ValueError(f"导出匹配 {total} 行，超过 {MAX_EXPORT_ROWS} 行上限；请缩小筛选范围")
            rows = [_item_dict(row) for row in conn.execute(text(ordered), args).mappings().all()]
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("SKU退货数据")
    dimension_fields = (["period", "entity", "category", "supplier_name"] if view == "detail" else [])
    columns = dimension_fields + ["sku_code", "sku_name", "brand", "sales_quantity", "return_in_quantity",
                                  "after_sales_quantity", "return_quantity", "return_rate", "entity_count", "supplier_count"]
    titles = {"period": "结算归属月", "entity": "主体", "category": "类目", "supplier_name": "供应商名称",
              "sku_code": "SKU编码", "sku_name": "SKU名称", "brand": "品牌", "sales_quantity": "销售数量",
              "return_in_quantity": "退货入库数量", "after_sales_quantity": "售后退货数量",
              "return_quantity": "退货数量", "return_rate": "退货率", "entity_count": "覆盖主体数",
              "supplier_count": "覆盖供应商数"}
    sheet.append([titles[column] for column in columns])
    string_columns = {"period", "entity", "category", "supplier_name", "sku_code", "sku_name", "brand"}
    for item in rows:
        cells = []
        for column in columns:
            value = item.get(column)
            if column in string_columns:
                cells.append(_text_cell(sheet, value))
            elif column == "return_rate":
                if value is None:
                    cells.append("-")
                else:
                    rate_cell = WriteOnlyCell(sheet, value=Decimal(value))
                    rate_cell.number_format = "0.00%"
                    cells.append(rate_cell)
            elif column in {"sales_quantity", "return_in_quantity", "after_sales_quantity", "return_quantity"}:
                cells.append(Decimal(value))
            else:
                cells.append(value)
        sheet.append(cells)
    notes = workbook.create_sheet("口径与版本")
    notes.append(["数据版本", run["id"] if run else "未导入"])
    notes.append(["导入时间", _iso(run["completed_at"]) if run else ""])
    notes.append(["结算月份", "主目录按月份文件夹；示例户外按月份工作表；与业务发生时间可能不同"])
    notes.append(["指标", "退货数量 = |实销实结退货入库单数量| + |售后退货数量|；销售数量 = 实销实结销售单带符号数量"])
    notes.append(["退货率", "筛选后退货数量合计 / 筛选后销售数量合计；销售数量小于等于 0 时为空"])
    notes.append(["筛选条件", json.dumps(_scope(filters), ensure_ascii=False)])
    if run:
        report = json.loads(run["report_json"] or "{}")
        notes.append(["质量说明", f"隔离 {run['isolated_rows']} 行；跨来源重复排除 {run['duplicate_rows']} 行；未分类来源 {len(report.get('unknown_categories', []))} 个"])
    output = BytesIO()
    workbook.save(output)
    filename = f"SKU退货数据-{view}-{datetime.now(timezone.utc):%Y%m%d%H%M%S}.xlsx"
    return filename, output.getvalue()


def list_brand_rules(*, engine: Engine | None = None) -> dict[str, Any]:
    with _engine(engine).connect() as conn:
        state = _state(conn)
        rules = _current_rules(conn, int(state["rule_version"]))
        return {"version": int(state["rule_version"]), "items": [
            {"id": rule["id"], "match_type": rule["match_type"], "pattern": rule["pattern"],
             "canonical_brand": rule["canonical_brand"], "priority": int(rule["priority"]),
             "enabled": bool(rule["enabled"])} for rule in rules
        ]}


def preview_brand_rules(rules: list[dict[str, Any]], *, engine: Engine | None = None) -> dict[str, Any]:
    normalized = _normalize_rules(rules)
    if len(normalized) > 1000:
        raise ValueError("品牌规则最多 1000 条")
    with _engine(engine).connect() as conn:
        state = _state(conn)
        run = _selected_run(conn)
        if not run:
            return {"version": int(state["rule_version"]), "run_id": None,
                    "affected_sku_count": 0, "affected_row_count": 0,
                    "pending_before": 0, "pending_after": 0, "samples": []}
        rows = conn.execute(text("""
            SELECT f.sku_name, u.entity, f.brand AS current_brand, COUNT(*) AS row_count
            FROM sku_return_fact f JOIN sku_return_source_unit u ON u.id=f.source_unit_id
            WHERE f.run_id=:run_id GROUP BY f.sku_name, u.entity, f.brand
        """), {"run_id": run["id"]}).mappings()
        affected_sku_names: set[str] = set()
        affected_rows = 0
        pending_before = 0
        pending_after = 0
        samples = []
        for row in rows:
            old_brand = row["current_brand"]
            new_brand, _ = match_brand_for_entity(row["sku_name"], row["entity"], normalized)
            count = int(row["row_count"])
            if old_brand == "待识别":
                pending_before += count
            if new_brand == "待识别":
                pending_after += count
            if old_brand != new_brand:
                affected_sku_names.add(row["sku_name"])
                affected_rows += count
                if len(samples) < 50:
                    samples.append({"sku_name": row["sku_name"], "entity": row["entity"], "before": old_brand,
                                    "after": new_brand, "row_count": count})
        return {"version": int(state["rule_version"]), "run_id": run["id"],
                "affected_sku_count": len(affected_sku_names), "affected_row_count": affected_rows,
                "pending_before": pending_before, "pending_after": pending_after,
                "samples": samples}


def _copy_brand_run(
    db_engine: Engine, previous: dict[str, Any], previous_version: int,
    new_version: int, rules: list[dict[str, Any]],
) -> str:
    """Clone facts from the active database version and atomically publish."""
    old_run_id = previous["id"]
    new_run_id = uuid.uuid4().hex
    with db_engine.begin() as conn:
        conn.execute(text("""
            INSERT INTO sku_return_import_run (id, manifest_sha256, rule_version, status, started_at)
            VALUES (:id, :manifest, :version, 'building', CURRENT_TIMESTAMP)
        """), {"id": new_run_id, "manifest": previous["manifest_sha256"], "version": new_version})
    try:
        with db_engine.connect() as conn:
            units = [dict(row) for row in conn.execute(text("""
                SELECT id, relative_path, path_sha256, file_sha256, sheet_name,
                       period, entity, category, scanned_rows, target_rows, included_rows
                FROM sku_return_source_unit WHERE run_id=:run_id
            """), {"run_id": old_run_id}).mappings().all()]
        unit_ids = {unit["id"]: uuid.uuid4().hex for unit in units}
        unit_entities = {unit["id"]: unit["entity"] for unit in units}
        with db_engine.begin() as conn:
            conn.execute(text("""
                INSERT INTO sku_return_source_unit (
                  id, run_id, relative_path, path_sha256, file_sha256, sheet_name,
                  period, entity, category, scanned_rows, target_rows, included_rows
                ) VALUES (:id, :run_id, :relative_path, :path_sha256, :file_sha256,
                          :sheet_name, :period, :entity, :category,
                          :scanned_rows, :target_rows, :included_rows)
            """), [{**unit, "id": unit_ids[unit["id"]], "run_id": new_run_id} for unit in units])
        select_columns = """
            source_unit_id, source_row, document_type, supplier_name, supplier_raw,
            sku_code, sku_name, sku_name_raw, raw_quantity, sales_quantity,
            return_in_quantity, after_sales_quantity, technical_document_id,
            technical_document_no, dedupe_sha256, included, excluded_reason,
            duplicate_of_source_unit_id, duplicate_of_row
        """
        unknown_brand_rows = 0
        copied = 0
        with db_engine.connect() as reader:
            cursor = reader.execution_options(stream_results=True).execute(
                text(f"SELECT {select_columns} FROM sku_return_fact WHERE run_id=:run_id ORDER BY id"),
                {"run_id": old_run_id}).mappings()
            while True:
                rows = cursor.fetchmany(1000)
                if not rows:
                    break
                batch = []
                for record in rows:
                    item = dict(record)
                    brand, rule_id = match_brand_for_entity(item["sku_name"], unit_entities[item["source_unit_id"]], rules)
                    item.update({"run_id": new_run_id, "source_unit_id": unit_ids[item["source_unit_id"]],
                                 "duplicate_of_source_unit_id": unit_ids.get(item["duplicate_of_source_unit_id"]),
                                 "brand": brand, "brand_rule_id": rule_id})
                    if brand == "待识别":
                        unknown_brand_rows += 1
                    batch.append(item)
                with db_engine.begin() as writer:
                    writer.execute(_FACT_INSERT, batch)
                copied += len(batch)
        if copied != int(previous["target_rows"]):
            raise AssertionError("品牌重算复制行数与原版本不一致")
        report = json.loads(previous["report_json"] or "{}")
        report.setdefault("quality", {})["unknown_brand_rows"] = unknown_brand_rows
        with db_engine.begin() as conn:
            _build_aggregate(conn, new_run_id)
            aggregate_count = conn.execute(text("SELECT COALESCE(SUM(fact_count),0) FROM sku_return_aggregate WHERE run_id=:run_id"), {"run_id": new_run_id}).scalar_one()
            if aggregate_count != int(previous["included_rows"]):
                raise AssertionError("品牌重算的聚合计数与原版本不一致")
            conn.execute(text("""
                UPDATE sku_return_import_run SET status='complete', completed_at=CURRENT_TIMESTAMP,
                  source_units=:source_units, scanned_rows=:scanned_rows, target_rows=:target_rows,
                  isolated_rows=:isolated_rows, duplicate_rows=:duplicate_rows,
                  included_rows=:included_rows, report_json=:report_json WHERE id=:run_id
            """), {**{name: previous[name] for name in ("source_units", "scanned_rows", "target_rows", "isolated_rows", "duplicate_rows", "included_rows")},
                    "report_json": json.dumps(report, ensure_ascii=False), "run_id": new_run_id})
            conn.execute(text("SELECT active_run_id FROM sku_return_state WHERE id=1 FOR UPDATE"))
            state = _state(conn)
            if state["active_run_id"] != old_run_id or int(state["rule_version"]) != previous_version:
                raise RuleConflictError("品牌重算期间活动版本已变化，请重新预览并保存")
            conn.execute(text("""
                UPDATE sku_return_state SET active_run_id=:run_id, rule_version=:version,
                  updated_at=CURRENT_TIMESTAMP WHERE id=1
            """), {"run_id": new_run_id, "version": new_version})
        return new_run_id
    except Exception as exc:
        with db_engine.begin() as conn:
            conn.execute(text("""
                UPDATE sku_return_import_run SET status='failed', completed_at=CURRENT_TIMESTAMP,
                  error_message=:error WHERE id=:run_id
            """), {"run_id": new_run_id, "error": str(exc)[:4000]})
        raise


def replace_brand_rules(
    rules: list[dict[str, Any]], expected_version: int | str, updated_by: str,
    *, engine: Engine | None = None,
) -> dict[str, Any]:
    normalized = _normalize_rules(rules)
    if len(normalized) > 1000:
        raise ValueError("品牌规则最多 1000 条")
    try:
        expected = int(expected_version)
    except (TypeError, ValueError) as exc:
        raise ValueError("expected_version 必须是整数") from exc
    db_engine = _engine(engine)
    with db_engine.begin() as conn:
        conn.execute(text("SELECT active_run_id FROM sku_return_state WHERE id=1 FOR UPDATE"))
        state = _state(conn)
        previous_version = int(state["rule_version"])
        if previous_version != expected:
            raise RuleConflictError("品牌规则已被其他用户修改，请刷新后重试")
        previous = _selected_run(conn)
        new_version = int(conn.execute(text("SELECT COALESCE(MAX(rule_version),0)+1 FROM sku_return_brand_rule")).scalar_one())
        if new_version <= previous_version:
            new_version = previous_version + 1
        if normalized:
            conn.execute(text("""
                INSERT INTO sku_return_brand_rule
                  (rule_version, match_type, pattern, pattern_key, canonical_brand,
                   priority, enabled, updated_at, updated_by)
                VALUES (:rule_version, :match_type, :pattern, :pattern_key,
                        :canonical_brand, :priority, :enabled, CURRENT_TIMESTAMP, :updated_by)
            """), [{**rule, "rule_version": new_version, "updated_by": _clean(updated_by)[:191]} for rule in normalized])
        if not previous:
            conn.execute(text("UPDATE sku_return_state SET rule_version=:version, updated_at=CURRENT_TIMESTAMP WHERE id=1"), {"version": new_version})
            return {"version": new_version, "run_id": None, "status": "saved"}
    with db_engine.connect() as conn:
        persisted = _current_rules(conn, new_version)
    new_run_id = _copy_brand_run(db_engine, previous, previous_version, new_version, persisted)
    return {"version": new_version, "run_id": new_run_id, "status": "published"}
