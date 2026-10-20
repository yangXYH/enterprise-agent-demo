"""Import and query the warehouse Jushuitan product-stock dashboard."""

from __future__ import annotations

import hashlib
import re
import shutil
import uuid
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from sqlalchemy import Engine, bindparam, text

from utils import import_service


UPLOAD_ROOT = Path(__file__).resolve().parent.parent / "uploads" / "jushuitan"
_DRAWING_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

_BATCH_DDL = """
CREATE TABLE IF NOT EXISTS jst_product_import_batch (
    batch_id VARCHAR(32) NOT NULL,
    source_name VARCHAR(255) NOT NULL DEFAULT '',
    row_count INT NOT NULL DEFAULT 0,
    image_count INT NOT NULL DEFAULT 0,
    imported_by VARCHAR(64) NOT NULL DEFAULT '',
    imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (batch_id)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='聚水潭商品资料导入批次'
"""

_PRODUCT_DDL = """
CREATE TABLE IF NOT EXISTS jst_product_stock (
    id BIGINT NOT NULL AUTO_INCREMENT,
    batch_id VARCHAR(32) NOT NULL,
    source_row INT NOT NULL,
    style_code VARCHAR(128) NOT NULL DEFAULT '',
    product_code VARCHAR(128) NOT NULL,
    product_name TEXT NOT NULL,
    brand VARCHAR(128) NOT NULL DEFAULT '',
    color_spec VARCHAR(255) NOT NULL DEFAULT '',
    color VARCHAR(128) NOT NULL DEFAULT '',
    spec VARCHAR(255) NOT NULL DEFAULT '',
    actual_stock DECIMAL(18,2) NOT NULL DEFAULT 0,
    order_occupied DECIMAL(18,2) NOT NULL DEFAULT 0,
    available_stock DECIMAL(18,2) NOT NULL DEFAULT 0,
    stock_updated_at DATETIME NULL,
    image_key VARCHAR(96) NOT NULL DEFAULT '',
    imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (id),
    UNIQUE KEY uk_jst_product_code (product_code),
    KEY idx_jst_product_brand (brand),
    KEY idx_jst_product_available (available_stock),
    KEY idx_jst_product_updated (stock_updated_at)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='聚水潭商品库存看板当前快照'
"""

EXPECTED_HEADERS = (
    "图片",
    "款式编码",
    "商品编码",
    "商品名称",
    "品牌",
    "颜色及规格",
    "颜色",
    "规格",
    "实际库存数",
    "订单占有数",
    "可用数",
    "库存更新时间",
)


def get_engine(engine: Engine | None = None) -> Engine:
    return engine or import_service.get_engine()


def ensure_schema(engine: Engine | None = None) -> None:
    db_engine = get_engine(engine)
    with db_engine.begin() as conn:
        conn.execute(text(_BATCH_DDL))
        conn.execute(text(_PRODUCT_DDL))


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    raw = _clean(value)
    if not raw:
        return None
    for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _normalise_headers(values: tuple[Any, ...]) -> tuple[str, ...]:
    return tuple(_clean(value).replace(" ", "") for value in values)


def _read_products(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        # Several Jushuitan exports incorrectly declare A1 as the used dimension.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        headers = _normalise_headers(next(rows, ()))
        if headers[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
            missing = [header for header in EXPECTED_HEADERS if header not in headers]
            detail = f"，缺少字段：{'、'.join(missing)}" if missing else ""
            raise ValueError(f"不是可识别的聚水潭商品资料表{detail}")

        products: list[dict[str, Any]] = []
        seen_codes: set[str] = set()
        for source_row, values in enumerate(rows, start=2):
            padded = list(values[:12]) + [None] * max(0, 12 - len(values))
            product_code = _clean(padded[2])
            if not product_code or product_code in seen_codes:
                continue
            seen_codes.add(product_code)
            products.append(
                {
                    "source_row": source_row,
                    "style_code": _clean(padded[1]),
                    "product_code": product_code,
                    "product_name": _clean(padded[3]),
                    "brand": _clean(padded[4]) or "未标品牌",
                    "color_spec": _clean(padded[5]),
                    "color": _clean(padded[6]),
                    "spec": _clean(padded[7]),
                    "actual_stock": _number(padded[8]),
                    "order_occupied": _number(padded[9]),
                    "available_stock": _number(padded[10]),
                    "stock_updated_at": _parse_datetime(padded[11]),
                    "image_key": "",
                }
            )
    finally:
        workbook.close()

    if not products:
        raise ValueError("聚水潭商品资料表未读取到有效商品")
    return products


def _drawing_image_map(archive: ZipFile) -> dict[int, str]:
    drawing_name = "xl/drawings/drawing1.xml"
    rels_name = "xl/drawings/_rels/drawing1.xml.rels"
    if drawing_name not in archive.namelist() or rels_name not in archive.namelist():
        return {}

    rels_root = ET.fromstring(archive.read(rels_name))
    targets = {
        rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
        for rel in rels_root.findall("pr:Relationship", _DRAWING_NS)
    }
    drawing_root = ET.fromstring(archive.read(drawing_name))
    result: dict[int, str] = {}
    embed_attr = f"{{{_DRAWING_NS['r']}}}embed"
    for anchor in drawing_root.findall("xdr:twoCellAnchor", _DRAWING_NS):
        row_text = anchor.findtext("xdr:from/xdr:row", namespaces=_DRAWING_NS)
        blip = anchor.find("xdr:pic/xdr:blipFill/a:blip", _DRAWING_NS)
        if row_text is None or blip is None:
            continue
        target = targets.get(blip.attrib.get(embed_attr, ""), "")
        if not target:
            continue
        media_name = str(PurePosixPath("xl/drawings") / target)
        media_name = str(PurePosixPath(media_name))
        while "/../" in media_name:
            left, right = media_name.split("/../", 1)
            media_name = f"{left.rsplit('/', 1)[0]}/{right}"
        result[int(row_text) + 1] = media_name
    return result


def _extract_images(path: Path, products: list[dict[str, Any]], target_dir: Path) -> int:
    target_dir.mkdir(parents=True, exist_ok=False)
    written: set[str] = set()
    with ZipFile(path) as archive:
        image_map = _drawing_image_map(archive)
        for product in products:
            media_name = image_map.get(product["source_row"])
            if not media_name or media_name not in archive.namelist():
                continue
            payload = archive.read(media_name)
            suffix = Path(media_name).suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
                suffix = ".bin"
            image_key = f"{hashlib.sha256(payload).hexdigest()}{suffix}"
            product["image_key"] = image_key
            if image_key in written:
                continue
            (target_dir / image_key).write_bytes(payload)
            written.add(image_key)
    return sum(1 for product in products if product["image_key"])


def import_excel(
    path: str | Path,
    source_name: str = "",
    imported_by: str = "",
    engine: Engine | None = None,
) -> dict[str, Any]:
    source_path = Path(path)
    products = _read_products(source_path)
    batch_id = uuid.uuid4().hex
    batch_dir = UPLOAD_ROOT / batch_id
    image_count = 0
    try:
        image_count = _extract_images(source_path, products, batch_dir)
        db_engine = get_engine(engine)
        ensure_schema(db_engine)
        imported_at = datetime.now()
        payload = [
            {
                **product,
                "batch_id": batch_id,
                "imported_at": imported_at,
            }
            for product in products
        ]
        with db_engine.begin() as conn:
            conn.execute(text("DELETE FROM jst_product_stock"))
            conn.execute(text("DELETE FROM jst_product_import_batch"))
            conn.execute(
                text(
                    """
                    INSERT INTO jst_product_import_batch
                        (batch_id, source_name, row_count, image_count, imported_by, imported_at)
                    VALUES
                        (:batch_id, :source_name, :row_count, :image_count, :imported_by, :imported_at)
                    """
                ),
                {
                    "batch_id": batch_id,
                    "source_name": (source_name or source_path.name)[:255],
                    "row_count": len(products),
                    "image_count": image_count,
                    "imported_by": _clean(imported_by)[:64],
                    "imported_at": imported_at,
                },
            )
            for start in range(0, len(payload), 500):
                conn.execute(
                    text(
                        """
                        INSERT INTO jst_product_stock (
                            batch_id, source_row, style_code, product_code, product_name,
                            brand, color_spec, color, spec, actual_stock, order_occupied,
                            available_stock, stock_updated_at, image_key, imported_at
                        ) VALUES (
                            :batch_id, :source_row, :style_code, :product_code, :product_name,
                            :brand, :color_spec, :color, :spec, :actual_stock, :order_occupied,
                            :available_stock, :stock_updated_at, :image_key, :imported_at
                        )
                        """
                    ),
                    payload[start : start + 500],
                )
    except Exception:
        shutil.rmtree(batch_dir, ignore_errors=True)
        raise

    for child in UPLOAD_ROOT.iterdir():
        if child.is_dir() and child.name != batch_id:
            shutil.rmtree(child, ignore_errors=True)

    return {
        "batch_id": batch_id,
        "source_name": source_name or source_path.name,
        "row_count": len(products),
        "image_count": image_count,
        "imported_at": imported_at.strftime("%Y-%m-%d %H:%M:%S"),
    }


def _where_clause(
    keyword: str,
    brands: list[str] | None,
    stock_status: str,
) -> tuple[str, dict[str, Any], list[Any]]:
    clauses: list[str] = []
    params: dict[str, Any] = {}
    bind_params: list[Any] = []
    clean_keyword = _clean(keyword)
    if clean_keyword:
        clauses.append(
            "(product_code LIKE :keyword OR style_code LIKE :keyword "
            "OR product_name LIKE :keyword OR color_spec LIKE :keyword)"
        )
        params["keyword"] = f"%{clean_keyword}%"
    clean_brands = [item.strip() for item in (brands or []) if item and item.strip()]
    if clean_brands:
        clauses.append("brand IN :brands")
        params["brands"] = clean_brands
        bind_params.append(bindparam("brands", expanding=True))
    status_sql = {
        "negative": "available_stock < 0",
        "zero": "available_stock = 0",
        "low": "available_stock > 0 AND available_stock <= 5",
        "healthy": "available_stock > 5",
        "occupied": "order_occupied > 0",
    }.get(stock_status)
    if status_sql:
        clauses.append(status_sql)
    return (" WHERE " + " AND ".join(clauses)) if clauses else "", params, bind_params


def list_dashboard(
    keyword: str = "",
    brands: list[str] | None = None,
    stock_status: str = "",
    page: int = 1,
    page_size: int = 50,
    engine: Engine | None = None,
) -> dict[str, Any]:
    db_engine = get_engine(engine)
    ensure_schema(db_engine)
    page = max(1, int(page))
    page_size = min(200, max(10, int(page_size)))
    where_sql, params, bind_params = _where_clause(keyword, brands, stock_status)
    offset = (page - 1) * page_size

    def stmt(sql: str):
        statement = text(sql)
        return statement.bindparams(*bind_params) if bind_params else statement

    with db_engine.connect() as conn:
        batch = conn.execute(
            text(
                "SELECT batch_id, source_name, row_count, image_count, imported_by, imported_at "
                "FROM jst_product_import_batch ORDER BY imported_at DESC LIMIT 1"
            )
        ).mappings().fetchone()
        summary = conn.execute(
            stmt(
                """
                SELECT COUNT(*) AS sku_count,
                       COUNT(DISTINCT brand) AS brand_count,
                       COALESCE(SUM(actual_stock), 0) AS actual_stock,
                       COALESCE(SUM(order_occupied), 0) AS order_occupied,
                       COALESCE(SUM(available_stock), 0) AS available_stock,
                       COALESCE(SUM(CASE WHEN available_stock < 0 THEN 1 ELSE 0 END), 0) AS negative_count,
                       COALESCE(SUM(CASE WHEN available_stock = 0 THEN 1 ELSE 0 END), 0) AS zero_count,
                       COALESCE(SUM(CASE WHEN available_stock > 0 AND available_stock <= 5 THEN 1 ELSE 0 END), 0) AS low_count,
                       COALESCE(SUM(CASE WHEN available_stock > 5 THEN 1 ELSE 0 END), 0) AS healthy_count,
                       MIN(stock_updated_at) AS stock_updated_min,
                       MAX(stock_updated_at) AS stock_updated_max
                FROM jst_product_stock
                """ + where_sql
            ),
            params,
        ).mappings().one()
        brand_ranking = conn.execute(
            stmt(
                """
                SELECT brand, COUNT(*) AS sku_count,
                       COALESCE(SUM(actual_stock), 0) AS actual_stock,
                       COALESCE(SUM(available_stock), 0) AS available_stock
                FROM jst_product_stock
                """ + where_sql + " GROUP BY brand ORDER BY actual_stock DESC, sku_count DESC LIMIT 12"
            ),
            params,
        ).mappings().all()
        total = int(summary["sku_count"] or 0)
        items = conn.execute(
            stmt(
                """
                SELECT id, batch_id, style_code, product_code, product_name, brand,
                       color_spec, color, spec, actual_stock, order_occupied,
                       available_stock, stock_updated_at, image_key
                FROM jst_product_stock
                """ + where_sql + " ORDER BY available_stock DESC, actual_stock DESC, product_code LIMIT :limit OFFSET :offset"
            ),
            {**params, "limit": page_size, "offset": offset},
        ).mappings().all()
        brand_options = [
            str(value)
            for value in conn.execute(text("SELECT DISTINCT brand FROM jst_product_stock ORDER BY brand")).scalars()
        ]

    def serialise_row(row: Any) -> dict[str, Any]:
        item = dict(row)
        for key in ("actual_stock", "order_occupied", "available_stock"):
            item[key] = float(item[key] or 0)
        value = item.get("stock_updated_at")
        item["stock_updated_at"] = value.strftime("%Y-%m-%d %H:%M:%S") if value else ""
        image_key = item.pop("image_key", "")
        item["image_url"] = (
            f"/api/jushuitan-dashboard/images/{item['batch_id']}/{image_key}" if image_key else ""
        )
        return item

    serialised_summary = dict(summary)
    for key, value in list(serialised_summary.items()):
        if isinstance(value, datetime):
            serialised_summary[key] = value.strftime("%Y-%m-%d %H:%M:%S")
        elif key not in {"sku_count", "brand_count", "negative_count", "zero_count", "low_count", "healthy_count"}:
            serialised_summary[key] = float(value or 0) if value is not None else 0
        else:
            serialised_summary[key] = int(value or 0)

    return {
        "batch": {
            **(dict(batch) if batch else {}),
            "imported_at": batch["imported_at"].strftime("%Y-%m-%d %H:%M:%S") if batch and batch["imported_at"] else "",
        },
        "summary": serialised_summary,
        "brand_ranking": [
            {
                "brand": row["brand"],
                "sku_count": int(row["sku_count"] or 0),
                "actual_stock": float(row["actual_stock"] or 0),
                "available_stock": float(row["available_stock"] or 0),
            }
            for row in brand_ranking
        ],
        "filter_options": {"brands": brand_options},
        "items": [serialise_row(row) for row in items],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def resolve_image(batch_id: str, image_key: str, engine: Engine | None = None) -> Path | None:
    if not re.fullmatch(r"[0-9a-f]{32}", batch_id.lower()) or not re.fullmatch(
        r"[0-9a-f]{64}\.(?:png|jpe?g|webp|gif)", image_key.lower()
    ):
        return None
    db_engine = get_engine(engine)
    ensure_schema(db_engine)
    with db_engine.connect() as conn:
        exists = conn.execute(
            text(
                "SELECT 1 FROM jst_product_stock "
                "WHERE batch_id = :batch_id AND image_key = :image_key LIMIT 1"
            ),
            {"batch_id": batch_id, "image_key": image_key},
        ).scalar()
    if not exists:
        return None
    candidate = UPLOAD_ROOT / batch_id / image_key
    return candidate if candidate.is_file() else None
