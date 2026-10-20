"""Daily operating data: preserve source cells, parameterized retrieval and evidence."""
from __future__ import annotations

import calendar
import hashlib
import json
import os
import re
from threading import Lock
from io import BytesIO
from contextlib import contextmanager
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from utils.operating_db import connect as database_connect

ROOT = Path(__file__).resolve().parents[1]
MONEY_SCALE = 1_000_000
METRICS = {'revenue': '成交金额', 'quantity': '成交商品件数', 'orders': '成交单量',
           'customers': '成交客户数', 'views': '商品浏览量', 'visitors': '商品访客数',
           'cart_quantity': '加购商品件数', 'cart_customers': '加购客户数', 'cart_amount': '加购金额'}
NOTE = '商品访客、成交客户、单量按商品日记录相加，跨商品/跨日期未去重，不等于店铺独立人数或订单数；转化率=成交客户数合计/商品访客数合计。金额按百万分之一元存储，汇总后保留两位小数，原始负值保留。缺失日期不补零。'
WEIGE_SOURCE_MARKER = '示例户外'
WEIGE_BRAND_LABELS = {'测试旅行': '测试旅行企业购', '美旅': '美旅企业购'}
WEIGE_BRAND_BASES = {label: brand for brand, label in WEIGE_BRAND_LABELS.items()}
_source_schema_lock = Lock()
_source_schema_ready = False

BRAND_PREFIXES = [
    ('测试旅行', ('测试旅行', 'samsonite')), ('金宇（JInyu）', ('金宇', 'jinyu')),
    ('DEMO_BAGS', ('demo_bags', '测试箱包', '法国大使')),
    ('测试运动', ('测试运动', 'aonijie')), ('阿迪达斯', ('阿迪达斯', 'adidas')),
    ('高山', ('高山', 'high sierra')), ('测试家居', ('campingmoon', '测试家居')),
    ('Jordan', ('jordan',)), ('测试服饰', ('demo_apparel',)), ('联想', ('联想', 'lenovo')),
    ('魅族', ('魅族', 'meizu')), ('Rains', ('demo_rain',)),
    ('GRANITEGEAR', ('demo_travel', '花岗岩')), ('iGear', ('demo_life',)),
    ('VENQUE', ('venque', '测试户外')), ('迪卡侬', ('迪卡侬', 'decathlon')),
    ('美旅', ('美旅', 'american tourister')),
    ('京塑', ('京塑',)), ('米小圈', ('米小圈',)),
]

def infer_brand(name):
    name=str(name or '').strip()
    for brand,prefixes in BRAND_PREFIXES:
        if any(name.casefold().startswith(p) for p in prefixes):return brand
    # Preserve an explicit Chinese brand with its parenthesized English name.
    match=re.match(r'^([\u4e00-\u9fff]{2,8}[（(][A-Za-z][A-Za-z\s.-]*[）)])',name)
    if match:return match[1].replace('(','（').replace(')','）')
    match=re.match(r'^([A-Za-z][A-Za-z&.-]{1,24})(?=[\u4e00-\u9fff\s（(]|$)',name)
    return match[1] if match else '待识别'


def source_key(filename=''):
    """Use the uploaded report filename to distinguish the two data sources."""
    return '示例户外' if WEIGE_SOURCE_MARKER in str(filename or '') else '示例科技'


def brand_for_record(name, filename='', brand_mapping=None):
    base = (brand_mapping or {}).get(name, infer_brand(name))
    if source_key(filename) == '示例户外':
        return WEIGE_BRAND_LABELS.get(base, base)
    return base


def brand_sql_expression(record_alias='jy_records', brand_alias=None):
    """Apply source-specific display names only where brand grouping is requested."""
    base = (f'{brand_alias}.brand' if brand_alias else
            f'(SELECT brand FROM jy_operating_brands b WHERE b.name={record_alias}.name)')
    filename = f'{record_alias}.filename'
    return (f"CASE WHEN instr(COALESCE({filename},''),'{WEIGE_SOURCE_MARKER}')>0 AND {base}='测试旅行' "
            "THEN '测试旅行企业购' "
            f"WHEN instr(COALESCE({filename},''),'{WEIGE_SOURCE_MARKER}')>0 AND {base}='美旅' "
            "THEN '美旅企业购' ELSE COALESCE("+base+",'待识别') END")


def _require_multi_source_schema(c):
    """Check the schema only on a Weige import, never on normal page reads."""
    global _source_schema_ready
    if c.dialect != 'mysql' or _source_schema_ready:
        return
    with _source_schema_lock:
        if _source_schema_ready:
            return
        row = c.execute(
            "SELECT GROUP_CONCAT(COLUMN_NAME ORDER BY SEQ_IN_INDEX SEPARATOR ',') "
            "FROM information_schema.STATISTICS "
            "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='jy_records' AND INDEX_NAME='PRIMARY'"
        ).fetchone()
        if not row or row[0] != 'day,sku,filename':
            raise ValueError(
                '数据库尚未启用多数据源，请先执行 database/migrations/20260927_allow_operating_sku_across_sources.sql'
            )
        _source_schema_ready = True

@contextmanager
def connect():
    with database_connect() as c:
        if not c.execute("SELECT 1 FROM jy_operating_settings WHERE setting_key='money_scale'").fetchone():
            # This path only applies to a pre-scale database. New databases are
            # empty, while the SQLite migration copies the setting first.
            for row in c.execute('SELECT day,sku,raw FROM jy_records'):
                raw=json.loads(row['raw'] or '{}')
                c.execute('UPDATE jy_records SET revenue=?,cart_amount=? WHERE day=? AND sku=?',
                          (amount(raw.get('成交金额')),amount(raw.get('加购金额')),row['day'],row['sku']))
            c.upsert_many('jy_operating_settings',[{'setting_key':'money_scale','value':str(MONEY_SCALE)}],('setting_key',))
        if not c.execute("SELECT 1 FROM jy_operating_settings WHERE setting_key='brand_rules_v2'").fetchone():
            names=[r[0] for r in c.execute('SELECT DISTINCT name FROM jy_records')]
            c.upsert_many('jy_operating_brands',[{'name':n,'brand':infer_brand(n)} for n in names],('name',))
            c.upsert_many('jy_operating_settings',[{'setting_key':'brand_rules_v2','value':'1'}],('setting_key',))
        yield c

def amount(v):
    return int((Decimal(str(v or 0)) * MONEY_SCALE).quantize(Decimal('1'), rounding=ROUND_HALF_UP))

def money_value(v):
    return float((Decimal(v or 0)/MONEY_SCALE).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))


def _cell_text(value):
    """Read Excel identifiers as text without turning 123.0 into a new SKU."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _normalise_header(value):
    return re.sub(r'[\s_]+', '', _cell_text(value)).casefold()

def import_excel(path, filename=None):
    path = Path(path)
    filename = Path(filename or path.name).name
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    with connect() as c:
        if source_key(filename) == '示例户外':
            _require_multi_source_schema(c)
        if c.execute('SELECT 1 FROM jy_batches WHERE filename=? AND digest=?', (filename,digest)).fetchone():
            return {'filename':filename,'unchanged':True,'updated_rows':0}
        duplicate = next(
            (row['filename'] for row in c.execute('SELECT filename FROM jy_batches WHERE digest=?', (digest,))
             if source_key(row['filename']) == source_key(filename)),
            None,
        )
        if duplicate:
            return {'filename':filename,'unchanged':True,'duplicate_of':duplicate,'updated_rows':0}
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        s = w.active
        rows = s.iter_rows(values_only=True)
        headers = [str(v or '') for v in next(rows)]
        if not {'时间','SKU','SKU名称',*METRICS.values()}.issubset(headers):
            raise ValueError('报表字段不匹配，请上传商智商品明细分天下载文件')
        records, seen, control = [], set(), None
        for n, row in enumerate(rows, 2):
            raw = dict(zip(headers, row))
            day, sku = str(raw.get('时间') or '')[:10], str(raw.get('SKU') or '')
            if sku == '合计':
                control = raw
                continue
            if not sku and not raw.get('时间'):
                continue
            try:
                day = date.fromisoformat(str(raw['时间'])[:10]).isoformat()
                if '~' in str(raw['时间']) or not sku:
                    raise ValueError()
            except ValueError:
                raise ValueError(f'{filename} 第 {n} 行不是有效日明细')
            if (day,sku) in seen:
                raise ValueError(f'{filename} 第 {n} 行日期和 SKU 重复')
            seen.add((day,sku))
            values = [amount(raw[k]) if m in ('revenue','cart_amount') else float(raw[k] or 0) for m,k in METRICS.items()]
            records.append((day,sku,str(raw.get('SKU名称') or ''),str(raw.get('货号') or ''),str(raw.get('三级类目') or ''),filename,n,*values,json.dumps(raw,ensure_ascii=False,default=str)))
        if not records:
            raise ValueError('报表没有日明细')
        if control and abs(sum(r[7] for r in records)-amount(control.get('成交金额'))) > len(records):
            raise ValueError('日明细成交金额与文件合计不一致，导入已停止')
        with connect() as c:
            current_source = source_key(filename)
            import_days = sorted({day for day, _sku in seen})
            overlap_keys = set()
            replaced_rows = []
            superseded_filenames = set()
            for offset in range(0, len(import_days), 100):
                day_chunk = import_days[offset:offset + 100]
                placeholders = ','.join('?' for _ in day_chunk)
                for old in c.execute(
                    f'SELECT day,sku,filename FROM jy_records WHERE day IN ({placeholders})', day_chunk
                ):
                    key = (old['day'], old['sku'])
                    old_filename = old['filename'] or ''
                    if key not in seen or source_key(old_filename) != current_source:
                        continue
                    overlap_keys.add(key)
                    if old_filename != filename:
                        replaced_rows.append((old['day'], old['sku'], old_filename))
                        superseded_filenames.add(old_filename)

            # Re-importing the same filename replaces that whole batch. For a
            # different file from the same source, replace only overlapping
            # product-day rows and retain its non-overlapping rows.
            c.execute('DELETE FROM jy_records WHERE filename=?',(filename,))
            c.executemany(
                'DELETE FROM jy_records WHERE day=? AND sku=? AND filename=?', replaced_rows
            )
            for old_filename in superseded_filenames:
                remaining = c.execute(
                    'SELECT COUNT(*) row_count,MIN(day) start_day,MAX(day) end_day '
                    'FROM jy_records WHERE filename=?', (old_filename,)
                ).fetchone()
                if remaining['row_count']:
                    # The old file no longer fully represents its stored rows;
                    # allow it to be uploaded again and refresh its batch stats.
                    c.execute(
                        'UPDATE jy_batches SET digest=NULL,row_count=?,`start`=?,`end`=? WHERE filename=?',
                        (remaining['row_count'], remaining['start_day'], remaining['end_day'], old_filename),
                    )
                else:
                    c.execute('DELETE FROM jy_batches WHERE filename=?', (old_filename,))
            record_columns = ('day','sku','name','code','category','filename','source_row',*METRICS,'raw')
            c.executemany(
                'INSERT INTO jy_records ('+','.join(f'`{column}`' for column in record_columns)+') VALUES ('+
                ','.join('?' for _ in record_columns)+')', records
            )
            names={r[2] for r in records}
            c.upsert_many('jy_operating_brands',[{'name':n,'brand':infer_brand(n)} for n in names],('name',))
            c.upsert_many('jy_batches',[{'filename':filename,'digest':digest,'row_count':len(records),
                'start':min(r[0] for r in records),'end':max(r[0] for r in records),
                'control':json.dumps(control,ensure_ascii=False,default=str),
                'headers':json.dumps(headers,ensure_ascii=False)}],('filename',))
        return {'filename':filename,'rows':len(records),'updated_rows':len(overlap_keys)}
    finally:
        w.close()

def metadata():
    with connect() as c:
        dates = [r[0] for r in c.execute('SELECT DISTINCT day FROM jy_records ORDER BY day')]
        batches = [dict(r) for r in c.execute('SELECT filename,row_count,start,end,imported_at FROM jy_batches ORDER BY start')]
        source_keys = {source_key(batch['filename']) for batch in batches}
        base_brands = {r[0] for r in c.execute(
            'SELECT DISTINCT brand FROM jy_operating_brands WHERE brand IS NOT NULL')}
        brand_labels = set()
        for brand in base_brands:
            if brand in WEIGE_BRAND_LABELS:
                if not source_keys or '示例科技' in source_keys:
                    brand_labels.add(brand)
                if '示例户外' in source_keys:
                    brand_labels.add(WEIGE_BRAND_LABELS[brand])
            else:
                brand_labels.add(brand)
        return {'dates':dates,'batches':batches,
                'categories':[r[0] for r in c.execute('SELECT DISTINCT category FROM jy_records ORDER BY category')],
                'brands':sorted(brand_labels),
                'brand_note':'品牌根据商品名称前缀识别并统一常见写法；示例户外来源中的测试旅行、美旅分别标为测试旅行企业购、美旅企业购；无法明确识别的商品标为待识别。',
                'metrics':METRICS,'note':NOTE,'ai_configured':bool(os.getenv('OPERATING_AI_KEY')),
                'overall_coverage':coverage(c,dates[0],dates[-1]) if dates else None}


def build_model_mapping_template():
    """Build the SKU-to-model mapping workbook used by the model ranking page."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '型号匹配'
    sheet.append(['sku', '货号'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:B1'
    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['B'].width = 28
    for cell in sheet['A']:
        cell.number_format = '@'
    for cell in sheet['B']:
        cell.number_format = '@'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def import_model_mapping(path):
    """Import SKU/model mappings and apply each non-empty model to every SKU row."""
    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration as exc:
            raise ValueError('型号匹配模板为空') from exc

        header_indexes = {_normalise_header(value): index for index, value in enumerate(headers)}
        sku_index = next((header_indexes[key] for key in ('sku', 'sku编码', '商品sku', '商品编码') if key in header_indexes), None)
        code_index = next((header_indexes[key] for key in ('货号', '型号', '商品货号') if key in header_indexes), None)
        if sku_index is None or code_index is None:
            raise ValueError('模板字段不匹配，需要包含 sku、货号两列')

        imported_skus = {}
        for row_number, row in enumerate(rows, 2):
            sku = _cell_text(row[sku_index] if sku_index < len(row) else None)
            code = _cell_text(row[code_index] if code_index < len(row) else None)
            if not sku and not code:
                continue
            if not sku:
                raise ValueError(f'第 {row_number} 行 SKU 为空')
            previous = imported_skus.get(sku, '')
            if previous and code and previous != code:
                raise ValueError(f'第 {row_number} 行 SKU {sku} 存在多个不同货号')
            imported_skus[sku] = code or previous
        if not imported_skus:
            raise ValueError('模板中没有可导入的 SKU 数据')

        with connect() as c:
            imported_values = list(imported_skus)
            imported_placeholders = ','.join('?' for _ in imported_values)
            # Keep this query deliberately index-covering.  The old aggregate
            # (MAX(code) + COUNT(*) + GROUP BY) made MySQL scan a large range
            # and could keep the import request open for minutes.  The
            # (sku, code) index lets MySQL return just the indexed columns;
            # aggregation is cheap for this bounded import batch in Python.
            matched_by_sku = {}
            for row in c.execute(
                f"SELECT sku, code FROM jy_records WHERE sku IN ({imported_placeholders})",
                imported_values,
            ):
                sku = _cell_text(row[0])
                entry = matched_by_sku.setdefault(sku, {'previous_code': '', 'record_count': 0})
                entry['record_count'] += 1
                code = _cell_text(row[1])
                if code and not entry['previous_code']:
                    entry['previous_code'] = code
            matched = [sku for sku in imported_values if sku in matched_by_sku]
            unknown = [sku for sku in imported_values if sku not in matched_by_sku]
            explicit_rows = c.execute(
                f"SELECT sku, code FROM jy_operating_model_mapping WHERE sku IN ({imported_placeholders})",
                imported_values,
            ).fetchall()
            explicit_codes = {
                _cell_text(row[0]): _cell_text(row[1])
                for row in explicit_rows if _cell_text(row[1])
            }
            for sku, code in explicit_codes.items():
                if sku in matched_by_sku:
                    matched_by_sku[sku]['previous_code'] = code
            mappings = {sku: imported_skus[sku] for sku in matched if imported_skus[sku]}

            new_models = sum(
                1 for sku in mappings
                if not matched_by_sku[sku]['previous_code']
            )
            updated_records = sum(
                int(matched_by_sku[sku]['record_count'] or 0)
                for sku in mappings
            )
            if mappings:
                # Do not rewrite every historical jy_records row. A single
                # SKU can have hundreds of daily records, so that approach
                # created large locks and undo logs. The compact mapping table
                # is the fast source of truth for imported model codes.
                c.upsert_many(
                    'jy_operating_model_mapping',
                    [{'sku': sku, 'code': code} for sku, code in mappings.items()],
                    ('sku',),
                )
        total_skus = len(imported_skus)
        matched_skus = len(matched)
        message = f'本次导入{total_skus}个sku，{matched_skus}个匹配，新增{new_models}个型号'
        if unknown:
            message += f'，{len(unknown)}个sku未匹配，已跳过'
        return {
            'rows': total_skus,
            'total_skus': total_skus,
            'matched_skus': matched_skus,
            'updated_skus': len(mappings),
            'updated_records': updated_records,
            'new_models': new_models,
            'unknown_skus': unknown,
            'unknown_count': len(unknown),
            'message': message,
        }
    finally:
        workbook.close()

def validate_range(start,end):
    a,b = date.fromisoformat(start),date.fromisoformat(end)
    if a>b or (b-a).days>1096:
        raise ValueError('日期区间必须正序且不超过三年')
    return a,b

def where(start,end,keyword='',category='',brand=''):
    validate_range(start,end)
    sql,args = 'day BETWEEN ? AND ?', [start,end]
    if keyword:
        sql += (' AND (instr(sku,?)>0 OR instr(name,?)>0 OR instr(code,?)>0 '
                'OR EXISTS (SELECT 1 FROM jy_operating_model_mapping model_mapping '
                'WHERE model_mapping.sku=jy_records.sku AND instr(model_mapping.code,?)>0))')
        args += [keyword]*4
    if category:
        sql += ' AND category=?'
        args += [category]
    if brand:
        if brand in WEIGE_BRAND_BASES:
            sql += (" AND instr(COALESCE(filename,''),?)>0 "
                    'AND name IN (SELECT name FROM jy_operating_brands WHERE brand=?)')
            args += [WEIGE_SOURCE_MARKER, WEIGE_BRAND_BASES[brand]]
        elif brand in WEIGE_BRAND_LABELS:
            sql += (" AND instr(COALESCE(filename,''),?)=0 "
                    'AND name IN (SELECT name FROM jy_operating_brands WHERE brand=?)')
            args += [WEIGE_SOURCE_MARKER, brand]
        else:
            sql += ' AND name IN (SELECT name FROM jy_operating_brands WHERE brand=?)'
            args += [brand]
    return sql,args

def totals(c,sql,args):
    r=dict(c.execute('SELECT COUNT(*) row_count,COUNT(DISTINCT sku) sku_count,'+','.join(f'SUM({k}) {k}' for k in METRICS)+' FROM jy_records WHERE '+sql,args).fetchone())
    for k in METRICS:
        r[k] = money_value(r[k]) if k in ('revenue','cart_amount') else r[k] or 0
    r['conversion'] = r['customers']/r['visitors'] if r['visitors'] else None
    r['unit_price'] = r['revenue']/r['quantity'] if r['quantity'] else None
    return r

def coverage(c,start,end):
    a,b=validate_range(start,end)
    present={r[0] for r in c.execute('SELECT DISTINCT day FROM jy_records WHERE day BETWEEN ? AND ?',(start,end))}
    missing=[(a+timedelta(days=i)).isoformat() for i in range((b-a).days+1) if (a+timedelta(days=i)).isoformat() not in present]
    return {'start':start,'end':end,'available_days':len(present),'expected_days':(b-a).days+1,'missing_dates':missing,'complete':not missing}

def shift_month(d,n):
    m=d.year*12+d.month-1+n
    y,mo=divmod(m,12)
    return date(y,mo+1,min(d.day,calendar.monthrange(y,mo+1)[1]))

def dashboard(start,end,keyword='',category='',page=1,page_size=50,metric='revenue',group='sku',brand=''):
    if metric not in METRICS or group not in ('sku','category','model'):
        raise ValueError('未知指标或分组')
    a,b=validate_range(start,end)
    sql,args=where(start,end,keyword,category,brand)
    with connect() as c:
        summary=totals(c,sql,args)
        cov=coverage(c,start,end)
        comparisons={}
        for label,x,y in [('yoy',shift_month(a,-12),shift_month(b,-12)),('mom',shift_month(a,-1),shift_month(b,-1))]:
            cs,ca=where(x.isoformat(),y.isoformat(),keyword,category,brand)
            cc=coverage(c,x.isoformat(),y.isoformat())
            t=totals(c,cs,ca)
            comparisons[label]={'coverage':cc,'summary':t,'rates':{k:(summary[k]-t[k])/abs(t[k]) if t[k] and cc['complete'] and cov['complete'] and cc['expected_days']==cov['expected_days'] else None for k in METRICS}}
        trend=[]
        for r in c.execute('SELECT day,'+','.join(f'SUM({k}) {k}' for k in METRICS)+' FROM jy_records WHERE '+sql+' GROUP BY day ORDER BY day',args):
            item=dict(r)
            for k in ('revenue','cart_amount'): item[k]=money_value(item[k])
            trend.append(item)
        ranking_metrics = ','.join(f'SUM({key}) {key}' for key in METRICS)
        if group == 'model':
            # Group by SKU once, then merge the small result set in Python.
            # Explicitly imported model codes are read from the compact mapping
            # table and override the raw report code without rewriting history.
            sku_rows=[dict(r) for r in c.execute(
                f'SELECT sku label,MAX(name) name,MAX(code) code,'
                f'COUNT(DISTINCT sku) sku_count,SUM({metric}) value,{ranking_metrics} '
                f'FROM jy_records WHERE '+sql+
                f' GROUP BY sku', args)]
            sku_values = [r['label'] for r in sku_rows]
            mapping = {}
            if sku_values:
                mapping_placeholders = ','.join('?' for _ in sku_values)
                mapping = {r[0]: r[1] for r in c.execute(
                    f'SELECT sku,code FROM jy_operating_model_mapping '
                    f'WHERE sku IN ({mapping_placeholders})', sku_values) if r[1]}
            merged = {}
            for r in sku_rows:
                label = mapping.get(r['label']) or _cell_text(r['code']) or r['label']
                target = merged.get(label)
                if target is None:
                    item = dict(r); item['label'] = label
                    merged[label] = item
                else:
                    target['value'] += r['value']
                    target['sku_count'] += r['sku_count']
                    for key in METRICS: target[key] += r[key]
                    target['name'] = max(target['name'] or '', r['name'] or '')
                    target['code'] = max(target['code'] or '', r['code'] or '')
            ranks = sorted(merged.values(), key=lambda x: (-x['value'], x['label']))[:20]
        else:
            group_expression = {'sku': 'sku', 'category': 'category'}[group]
            ranks=[dict(r) for r in c.execute(
                f'SELECT {group_expression} label,MAX(name) name,MAX(code) code,'
                f'COUNT(DISTINCT sku) sku_count,SUM({metric}) value,{ranking_metrics} '
                f'FROM jy_records WHERE '+sql+
                f' GROUP BY {group_expression} ORDER BY value DESC,{group_expression} LIMIT 20', args)]
        for r in ranks:
            for key in ('revenue','cart_amount'):
                r[key]=money_value(r[key])
            r['value']=r[metric]
        detail=[dict(r) for r in c.execute('SELECT * FROM jy_records WHERE '+sql+' ORDER BY day DESC,revenue DESC,sku LIMIT ? OFFSET ?',args+[page_size,(page-1)*page_size])]
        model_overrides = {r[0]: r[1] for r in c.execute('SELECT sku,code FROM jy_operating_model_mapping') if r[1]}
        brands=dict(c.execute('SELECT name,brand FROM jy_operating_brands'))
        for r in detail:
            if r['sku'] in model_overrides:
                r['code'] = model_overrides[r['sku']]
            r['brand']=brand_for_record(r['name'], r['filename'], brands)
            r['raw']=json.loads(r['raw'])
            for k in ('revenue','cart_amount'):r[k]=(r[k] or 0)/MONEY_SCALE
        sources=[dict(r) for r in c.execute('SELECT filename,COUNT(*) AS `rows`,MIN(source_row) first_row,MAX(source_row) last_row FROM jy_records WHERE '+sql+' GROUP BY filename',args)]
        distribution=[dict(r) for r in c.execute(f'SELECT category label,COUNT(DISTINCT sku) sku_count,SUM({metric}) value FROM jy_records WHERE '+sql+' GROUP BY category ORDER BY value DESC,category',args)]
        if metric in ('revenue','cart_amount'):
            for r in distribution:r['value']=money_value(r['value'])
        return {'summary':summary,'coverage':cov,'comparisons':comparisons,'trend':trend,'ranking':ranks,'items':detail,'total':summary['row_count'],'sources':sources,'note':NOTE,
                'category_distribution':distribution,
                'filters':{'start':start,'end':end,'keyword':keyword,'category':category,'brand':brand,'metric':metric,'group':group},'comparison_note':'同比按上年同月同日起止，环比按上月同日起止；月末日期截至该月最后一天，天数不等或缺数不输出增长率。'}


def executive_insights(start, end, brand='', category=''):
    """Read-only executive view over the same product-day facts and comparison rules."""
    a, b = validate_range(start, end)
    aggregate = ','.join(f'SUM({key}) {key}' for key in METRICS)

    def normalized(row):
        item = dict(row)
        for key in ('revenue', 'cart_amount'):
            if key in item:
                item[key] = money_value(item[key])
        return item

    def dimensions(c, first, last):
        clause, args = where(first, last, category=category, brand=brand)
        categories = [normalized(r) for r in c.execute(
            f'SELECT category label,COUNT(DISTINCT sku) sku_count,{aggregate} '
            f'FROM jy_records WHERE {clause} GROUP BY category ORDER BY SUM(revenue) DESC,category', args)]
        # The derived table keeps the common filter unambiguous (both tables have `name`).
        branded = (f'(SELECT jy_records.*,{brand_sql_expression(brand_alias="brands")} brand_label '
                   'FROM jy_records LEFT JOIN jy_operating_brands brands ON jy_records.name=brands.name) facts')
        brands = [normalized(r) for r in c.execute(
            f'SELECT brand_label label,COUNT(DISTINCT sku) sku_count,{aggregate} '
            f'FROM {branded} WHERE {clause} GROUP BY brand_label ORDER BY SUM(revenue) DESC,brand_label', args)]
        return categories, brands

    with connect() as c:
        clause, args = where(start, end, category=category, brand=brand)
        summary = totals(c, clause, args)
        current_coverage = coverage(c, start, end)
        comparisons = {}
        previous_dimensions = {}
        for key, first, last in (('yoy', shift_month(a, -12), shift_month(b, -12)),
                                 ('mom', shift_month(a, -1), shift_month(b, -1))):
            ps, pe = first.isoformat(), last.isoformat()
            old_clause, old_args = where(ps, pe, category=category, brand=brand)
            old_coverage = coverage(c, ps, pe)
            old_summary = totals(c, old_clause, old_args)
            comparable = (current_coverage['complete'] and old_coverage['complete']
                          and current_coverage['expected_days'] == old_coverage['expected_days'])
            comparisons[key] = {
                'summary': old_summary, 'coverage': old_coverage, 'comparable': comparable,
                'rates': {metric: (summary[metric] - old_summary[metric]) / abs(old_summary[metric])
                          if comparable and old_summary[metric] else None for metric in METRICS},
            }
            if key == 'mom':
                previous_dimensions = dict(zip(('categories', 'brands'), dimensions(c, ps, pe)))
        trend = [normalized(r) for r in c.execute(
            f'SELECT day,{aggregate} FROM jy_records WHERE {clause} GROUP BY day ORDER BY day', args)]
        categories, brands = dimensions(c, start, end)
        products = [normalized(r) for r in c.execute(
            f'SELECT sku label,MAX(name) name,MAX(code) code,MAX(category) category,{aggregate} '
            f'FROM jy_records WHERE {clause} GROUP BY sku ORDER BY SUM(revenue) DESC,sku LIMIT 20', args)]
        sources = [dict(r) for r in c.execute(
            f'SELECT filename,COUNT(*) AS `rows`,MIN(source_row) first_row,MAX(source_row) last_row '
            f'FROM jy_records WHERE {clause} GROUP BY filename', args)]
        return {'summary': summary, 'coverage': current_coverage, 'comparisons': comparisons,
                'trend': trend, 'categories': categories, 'brands': brands, 'products': products,
                'previous_dimensions': previous_dimensions, 'sources': sources,
                'filters': {'start': start, 'end': end, 'brand': brand, 'category': category}, 'note': NOTE}
