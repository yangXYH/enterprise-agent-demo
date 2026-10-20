"""使用临时 Excel 验证 SKU 退货扫描口径；完全不连接数据库。"""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from utils import sku_return_service as service


HEADERS = ["供应商名称", "应付帐单据类型", "SKU编码", "SKU名称", "数量", "单据ID", "单据编号"]


def make_workbook(path: Path, rows: list[list], sheet_name: str = "数据") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


class SkuReturnScanTests(unittest.TestCase):
    def test_weige_samsonite_brand_is_separate_from_yiwangou(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            make_workbook(
                root / "2026年示例科技结算单" / "1月" / "示例科技" / "1 家电 2026-01-01.xlsx",
                [["供应商一", "实销实结销售单", "001", "测试旅行行李箱", 5, "D1", "N1"]],
            )
            make_workbook(
                root / "2026年示例户外结算单" / "示例户外2026年1-7月.xlsx",
                [["供应商二", "实销实结销售单", "001", "测试旅行行李箱", 7, "D2", "N2"],
                 ["供应商二", "售后退货", "002", "Samsonite背包", -1, "D3", "N3"]],
                sheet_name="1月",
            )
            unit_entities: dict[str, str] = {}
            facts: list[dict] = []
            report = service.scan_source_tree(
                root, rules=service.DEFAULT_BRAND_RULES,
                on_unit=lambda unit, _: unit_entities.update({unit["id"]: unit["entity"]}),
                on_fact_batch=lambda batch: facts.extend(batch),
            )
        self.assertEqual(report["included_rows"], 3)
        self.assertEqual(
            {(unit_entities[fact["source_unit_id"]], fact["sku_code"]): fact["brand"] for fact in facts},
            {("示例科技", "001"): "测试旅行", ("示例户外", "001"): "测试旅行企业购",
             ("示例户外", "002"): "测试旅行企业购"},
        )

    def test_fly_globe_name_variants_share_feiqiu_brand(self):
        rules = service._normalize_rules(service.DEFAULT_BRAND_RULES)
        for name in (
            "测试五金室内门锁单舌不锈钢款",
            "测试五金( DEMO_TOOLS) 球形锁",
            "测试五金（DEMO_TOOLS）门锁",
            "DEMO_TOOLS门锁",
        ):
            with self.subTest(name=name):
                self.assertEqual(service.match_brand(name, rules)[0], "测试五金")

    def test_demo_home_prefix_is_recognized_case_insensitively(self):
        rules = service._normalize_rules(service.DEFAULT_BRAND_RULES)
        for name in ("demo_home克鲁伯咖啡机", "KRUPS德国克鲁伯咖啡机"):
            with self.subTest(name=name):
                self.assertEqual(service.match_brand(name, rules)[0], "demo_home")

    def test_source_multiset_keeps_first_month_and_isolates_missing_quantity(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            base = root / "2026年示例科技结算单"
            sale = ["供应商一", "实销实结销售单", "001", "商品甲", 10, "D1", "N1"]
            make_workbook(
                base / "6月" / "B2B" / "001 B2B箱包 2026-06-01.xlsx",
                [
                    sale, sale,  # 同一来源表里的两个真实出现次数都计入
                    ["供应商一", "实销实结退货入库单", "001", "商品甲", -2, "D2", "N2"],
                    ["供应商一", "售后退货", "001", "商品甲", -1, "D3", "N3"],
                    ["供应商一", "售后退货", "002", "商品乙", -2, "D4", "N4"],
                    ["供应商一", "实销实结销售单", "003", "商品丙", None, "D5", "N5"],
                ],
            )
            make_workbook(
                base / "7月" / "B2B" / "002 B2B箱包 2026-07-01.xlsx",
                [sale, sale],  # 跨来源重复，仅 6 月的两个出现次数保留
            )
            facts: list[dict] = []
            report = service.scan_source_tree(
                root, rules=[], on_fact_batch=lambda batch: facts.extend(batch)
            )

        self.assertEqual(report["source_units"], 2)
        self.assertEqual(report["target_rows"], 8)
        self.assertEqual(report["isolated_rows"], 1)
        self.assertEqual(report["duplicate_rows"], 2)
        self.assertEqual(report["included_rows"], 5)
        self.assertEqual(
            {(row["period"], row["included_rows"]) for row in report["coverage"]},
            {("2026-06", 5), ("2026-07", 0)},
        )
        self.assertEqual(len([fact for fact in facts if fact["excluded_reason"] == "cross_source_duplicate"]), 2)
        self.assertEqual(len([fact for fact in facts if fact["included"] and fact["technical_document_id"] == "D1"]), 2)
        self.assertEqual(len([fact for fact in facts if not fact["included"] and fact["raw_quantity"] is None]), 1)

        totals: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
        for fact in facts:
            if fact["included"]:
                for key in ("sales_quantity", "return_in_quantity", "after_sales_quantity"):
                    totals[fact["sku_code"]][key] += fact[key]
        self.assertEqual(totals["001"]["sales_quantity"], Decimal("20"))
        self.assertEqual(totals["001"]["return_in_quantity"], Decimal("2"))
        self.assertEqual(totals["001"]["after_sales_quantity"], Decimal("1"))
        self.assertEqual(service._rate(Decimal("20"), Decimal("3")), "0.15")
        self.assertIsNone(service._rate(totals["002"]["sales_quantity"], totals["002"]["after_sales_quantity"]))

    def test_export_text_cell_does_not_execute_excel_formula(self):
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("导出")
        worksheet.append([service._text_cell(worksheet, "=1+1")])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        reopened = load_workbook(stream, read_only=True, data_only=False)
        try:
            cell = reopened.active["A1"]
            self.assertEqual(cell.value, "=1+1")
            self.assertEqual(cell.data_type, "s")
        finally:
            reopened.close()


if __name__ == "__main__":
    unittest.main()
