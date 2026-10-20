import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
import openpyxl
from utils import operating_analysis_service as s

class OperatingAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        database=(Path(self.tmp.name)/'data.sqlite3').resolve().as_posix()
        self.env=patch.dict(os.environ,{'DB_URL':'sqlite+pysqlite:///'+database})
        self.env.start()
    def tearDown(self):
        self.env.stop(); self.tmp.cleanup()
    def report(self,name,rows,control=None,product_name='电脑包'):
        w=openpyxl.Workbook(); sh=w.active
        headers=['时间','SKU','SKU名称','货号','三级类目',*s.METRICS.values()]
        sh.append(headers)
        if control is not None: sh.append(['2026-01-01~2026-01-31','合计','','','',control,*([0]*8)])
        for day,sku,money in rows:sh.append([day,sku,product_name,'CODE','电脑包',money,1,1,1,10,5,1,1,money])
        p=Path(self.tmp.name)/name;w.save(p);w.close();return p
    def test_import_controls_raw_precision_and_idempotence(self):
        p=self.report('a.xlsx',[('2026-01-01','1',10.12),('2026-01-02','1',-2.01)],8.11)
        self.assertEqual(s.import_excel(p)['rows'],2)
        self.assertTrue(s.import_excel(p)['unchanged'])
        d=s.dashboard('2026-01-01','2026-01-02')
        self.assertEqual(d['summary']['revenue'],8.11)
        self.assertEqual(d['items'][0]['source_row'],4)
        self.assertEqual(d['items'][0]['raw']['成交金额'],-2.01)
        self.assertEqual(d['summary']['conversion'],.2)
    def test_overlapping_import_replaces_existing_rows(self):
        p=self.report('a.xlsx',[('2026-01-01','1',10)])
        s.import_excel(p)
        q=self.report('b.xlsx',[('2026-01-02','1',2),('2026-01-01','1',3)])
        result=s.import_excel(q)
        self.assertEqual(result['updated_rows'],1)
        self.assertEqual(s.metadata()['dates'],['2026-01-01','2026-01-02'])
        self.assertEqual(len(s.metadata()['batches']),1)
        dashboard=s.dashboard('2026-01-01','2026-01-02')
        self.assertEqual(dashboard['summary']['revenue'],5)
        self.assertEqual(dashboard['items'][1]['revenue'],3)
    def test_missing_days_and_parameterized_filters(self):
        s.import_excel(self.report('a.xlsx',[('2026-01-01','1',10)]))
        d=s.dashboard('2026-01-01','2026-01-02')
        self.assertEqual(d['coverage']['missing_dates'],['2026-01-02'])
        self.assertIsNone(d['comparisons']['yoy']['rates']['revenue'])
        self.assertEqual(s.dashboard('2026-01-01','2026-01-02',keyword="' OR 1=1 --")['total'],0)
    def test_comparison_and_zero_baseline(self):
        for name,day,m in [('a','2026-02-01',20),('b','2026-01-01',10),('c','2025-02-01',0)]:
            s.import_excel(self.report(name+'.xlsx',[(day,'1',m)]))
        d=s.dashboard('2026-02-01','2026-02-01')
        self.assertEqual(d['comparisons']['mom']['rates']['revenue'],1)
        self.assertIsNone(d['comparisons']['yoy']['rates']['revenue'])
        self.assertEqual(s.shift_month(date(2024,2,29),-12),date(2023,2,28))
    def test_invalid_report_preserves_data(self):
        with self.assertRaises(ValueError):s.import_excel(self.report('a.xlsx',[('2026-01-01','1',10)],500))
        self.assertEqual(s.metadata()['dates'],[])

    def test_fractional_cents_are_rounded_after_summing(self):
        s.import_excel(self.report('a.xlsx',[('2026-01-01','1',1.004),('2026-01-01','2',1.004)],2.008))
        d=s.dashboard('2026-01-01','2026-01-01')
        self.assertEqual(d['summary']['revenue'],2.01)
        self.assertEqual(d['items'][0]['revenue'],1.004)

    def test_brand_names_and_comparison_filters(self):
        self.assertEqual(s.infer_brand('测试旅行双肩包电脑包16英寸笔记本男女商务通勤背包旅行包大容量书包礼物'),'测试旅行')
        self.assertEqual(s.infer_brand('金宇（JInyu）卡式炉便携式户外防风防爆露营家用卡磁炉'),'金宇（JInyu）')
        self.assertEqual(s.infer_brand('DEMO_BAGS测试箱包拉杆箱法国大使行李箱28英寸密码旅行箱万向轮行李箱大容量'),'DEMO_BAGS')
        self.assertEqual(s.infer_brand('金宇（jinyu）户外炉'),'金宇（JInyu）')
        self.assertEqual(s.infer_brand('无明确品牌商品'),'待识别')
        s.import_excel(self.report('a.xlsx',[('2026-08-01','1',30)],product_name='测试旅行双肩包'))
        s.import_excel(self.report('b.xlsx',[('2025-08-01','1',10)],product_name='测试旅行（Samsonite）双肩包'))
        s.import_excel(self.report('c.xlsx',[('2026-08-01','2',100)],product_name='金宇（jinyu）户外炉'))
        d=s.dashboard('2026-08-01','2026-08-01',brand='测试旅行')
        self.assertEqual(d['total'],1)
        self.assertEqual(d['items'][0]['brand'],'测试旅行')
        self.assertEqual(d['summary']['revenue'],30)
        self.assertEqual(d['comparisons']['yoy']['rates']['revenue'],2)
        self.assertEqual(d['category_distribution'][0]['value'],30)
        self.assertEqual(s.dashboard('2026-08-01','2026-08-01',brand="' OR 1=1 --")['total'],0)

    def test_executive_insights_are_sourced_and_comparable(self):
        s.import_excel(self.report('current-a.xlsx',[('2026-08-01','1',30)],product_name='测试旅行双肩包'))
        s.import_excel(self.report('current-b.xlsx',[('2026-08-01','2',-5)],product_name='金宇（jinyu）户外炉'))
        s.import_excel(self.report('previous.xlsx',[('2026-07-01','1',10)],product_name='测试旅行双肩包'))
        with s.connect() as c:
            c.execute('UPDATE jy_records SET category=? WHERE sku=?', ('户外装备', '2'))
        result=s.executive_insights('2026-08-01','2026-08-01')
        self.assertEqual(result['summary']['revenue'],25)
        self.assertEqual(result['comparisons']['mom']['rates']['revenue'],1.5)
        self.assertEqual({r['label']:r['revenue'] for r in result['categories']}, {'电脑包':30,'户外装备':-5})
        self.assertEqual({r['label']:r['revenue'] for r in result['brands']}, {'测试旅行':30,'金宇（JInyu）':-5})
        self.assertEqual(len(result['sources']),2)
        filtered=s.executive_insights('2026-08-01','2026-08-01',brand='测试旅行')
        self.assertEqual(filtered['summary']['revenue'],30)
        self.assertEqual(filtered['comparisons']['mom']['rates']['revenue'],2)
        self.assertEqual(len(filtered['brands']),1)
        self.assertEqual(s.executive_insights('2026-08-01','2026-08-01',brand="' OR 1=1 --")['summary']['row_count'],0)
        incomplete=s.executive_insights('2026-08-01','2026-08-02')
        self.assertFalse(incomplete['coverage']['complete'])
        self.assertIsNone(incomplete['comparisons']['mom']['rates']['revenue'])

    def test_executive_insights_http_route_is_registered(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from operating_analysis_routes import router
        s.import_excel(self.report('current.xlsx',[('2026-08-01','1',30)]))
        app=FastAPI()
        app.include_router(router)
        with patch('operating_analysis_routes.authorize', return_value={'role':'boss'}):
            response=TestClient(app).get('/api/operating-analysis/insights',params={'start':'2026-08-01','end':'2026-08-01'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['summary']['revenue'],30)

    def test_model_ranking_falls_back_to_sku_and_mapping_updates_all_dates(self):
        s.import_excel(self.report('a.xlsx', [('2026-08-01', '1', 30), ('2026-08-02', '1', 20)]))
        s.import_excel(self.report('b.xlsx', [('2026-08-01', '2', 100)]))
        with s.connect() as c:
            c.execute('UPDATE jy_records SET code=? WHERE sku=?', ('', '2'))
            c.execute('UPDATE jy_records SET code=? WHERE sku=? AND day=?', ('', '1', '2026-08-02'))

        ranking = s.dashboard('2026-08-01', '2026-08-02', group='model')['ranking']
        self.assertEqual(ranking[0]['label'], '2')
        self.assertEqual(ranking[1]['label'], 'CODE')
        self.assertEqual(ranking[1]['value'], 50)

        template = Path(self.tmp.name) / 'mapping.xlsx'
        template.write_bytes(s.build_model_mapping_template())
        workbook = openpyxl.load_workbook(template)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual([sheet['A1'].value, sheet['B1'].value], ['sku', '货号'])
        sheet['A2'] = '1'
        sheet['B2'] = 'MODEL-1'
        sheet['A3'] = '2'
        sheet['B3'] = 'MODEL-2'
        workbook.save(template)
        workbook.close()
        result = s.import_model_mapping(template)
        self.assertEqual(result['updated_skus'], 2)
        with s.connect() as c:
            self.assertEqual(c.execute('SELECT code FROM jy_operating_model_mapping WHERE sku=?', ('1',)).fetchone()[0], 'MODEL-1')
            self.assertEqual(c.execute('SELECT code FROM jy_operating_model_mapping WHERE sku=?', ('2',)).fetchone()[0], 'MODEL-2')
        ranking = s.dashboard('2026-08-01', '2026-08-02', group='model')['ranking']
        self.assertEqual({row['label'] for row in ranking}, {'MODEL-1', 'MODEL-2'})

    def test_model_mapping_skips_unknown_sku_and_reports_import_summary(self):
        s.import_excel(self.report('a.xlsx', [('2026-08-01', '1', 30)]))
        with s.connect() as c:
            c.execute('UPDATE jy_records SET code=? WHERE sku=?', ('', '1'))
        template = Path(self.tmp.name) / 'mapping-invalid.xlsx'
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['sku', '货号'])
        sheet.append(['1', 'MODEL-1'])
        sheet.append(['missing', 'MODEL-X'])
        workbook.save(template)
        workbook.close()
        result = s.import_model_mapping(template)
        self.assertEqual(result['total_skus'], 2)
        self.assertEqual(result['matched_skus'], 1)
        self.assertEqual(result['new_models'], 1)
        self.assertEqual(result['unknown_skus'], ['missing'])
        self.assertEqual(result['message'], '本次导入2个sku，1个匹配，新增1个型号，1个sku未匹配，已跳过')
        with s.connect() as c:
            self.assertEqual(c.execute('SELECT code FROM jy_operating_model_mapping WHERE sku=?', ('1',)).fetchone()[0], 'MODEL-1')

if __name__=='__main__':unittest.main()
