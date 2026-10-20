"""Optional isolated MySQL integration test; never touches the configured database.

Run explicitly with SKU_RETURN_MYSQL_TEST=1. A disposable schema is created,
used for the test, and removed in tearDownClass.
"""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
import uuid

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text

from utils import import_service
from utils import sku_return_service as service


@unittest.skipUnless(os.getenv("SKU_RETURN_MYSQL_TEST") == "1", "isolated MySQL test is opt-in")
class SkuReturnMysqlServiceTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from dotenv import load_dotenv
        load_dotenv(Path(__file__).resolve().parents[1] / ".env", override=False)
        configured = import_service.get_engine()
        if configured.dialect.name != "mysql":
            raise unittest.SkipTest("requires a MySQL server")
        cls.schema = "sku_return_test_" + uuid.uuid4().hex[:12]
        with configured.begin() as conn:
            conn.execute(text(f"CREATE DATABASE `{cls.schema}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci"))
        cls.engine = create_engine(configured.url.set(database=cls.schema), pool_pre_ping=True)

    @classmethod
    def tearDownClass(cls) -> None:
        if not hasattr(cls, "schema"):
            return
        cls.engine.dispose()
        configured = import_service.get_engine()
        with configured.begin() as conn:
            conn.execute(text(f"DROP DATABASE `{cls.schema}`"))

    @staticmethod
    def _book(path: Path, rows: list[list[object]]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "业务表"
        sheet.append(["供应商名称", "应付帐单据类型", "SKU编码", "SKU名称", "数量", "单据ID", "单据编号"])
        for row in rows:
            sheet.append(row)
        workbook.save(path)

    def test_publish_query_export_and_rule_rebuild(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            june = root / "2026年示例科技结算单" / "6月" / "B2B" / "1 B2B户外 2026-06-01.xlsx"
            july = root / "2026年示例科技结算单" / "7月" / "B2B" / "2 B2B 2026-07-01.xlsx"
            sale = ["甲供应商", "实销实结销售单", "0001", "测试商品", 10, "D1", "N1"]
            returned = ["甲供应商", "售后退货", "0001", "测试商品", -2, "D2", "N2"]
            self._book(june, [sale, returned])
            self._book(july, [sale, ["甲供应商", "实销实结退货入库单", "0001", "测试商品", -1, "D3", "N3"]])
            preview = service.import_source_tree(root)
            self.assertEqual((preview["target_rows"], preview["duplicate_rows"], preview["included_rows"]), (4, 1, 3))
            published = service.import_source_tree(root, apply=True, engine=self.engine)
            run_id = published["run_id"]
            self.assertEqual(service.import_source_tree(root, apply=True, engine=self.engine)["status"], "unchanged")
            meta = service.get_metadata(engine=self.engine)
            self.assertEqual(meta["run_id"], run_id)
            self.assertEqual(meta["periods"], ["2026-06", "2026-07"])
            summary = service.get_summary(engine=self.engine)
            self.assertEqual((summary["sales_quantity"], summary["return_quantity"], summary["return_rate"]), ("10", "3", "0.3"))
            filtered = service.get_summary({"period": ["2026-07"]}, engine=self.engine)
            self.assertEqual((filtered["sales_quantity"], filtered["return_quantity"], filtered["return_rate"]), ("0", "1", None))
            trend = service.get_monthly_trend(engine=self.engine)
            self.assertEqual([row["period"] for row in trend["items"]], ["2026-06", "2026-07"])
            self.assertEqual((trend["items"][0]["sales_quantity"], trend["items"][0]["return_quantity"], trend["items"][0]["return_rate"]), ("10", "2", "0.2"))
            self.assertEqual((trend["items"][1]["sales_quantity"], trend["items"][1]["return_quantity"], trend["items"][1]["return_rate"]), ("0", "1", None))
            self.assertEqual([row["period"] for row in service.get_monthly_trend({"period": ["2026-07"]}, engine=self.engine)["items"]], ["2026-07"])
            item = service.list_items(engine=self.engine)["items"][0]
            self.assertEqual(item["sku_code"], "0001")
            self.assertEqual(item["return_quantity"], "3")
            sources = service.list_sources("0001", "测试商品", engine=self.engine)
            self.assertEqual(len(sources["items"]), 4)
            self.assertIn("row_number", sources["items"][0])
            _, exported = service.export_items(engine=self.engine)
            workbook = load_workbook(BytesIO(exported), read_only=True, data_only=False)
            try:
                cells = list(workbook.worksheets[0].iter_rows(min_row=2, max_row=2))[0]
                self.assertEqual(cells[0].data_type, "s")
                self.assertEqual(cells[3].data_type, "n")
                self.assertEqual(cells[7].number_format, "0.00%")
            finally:
                workbook.close()
            current_rules = service.list_brand_rules(engine=self.engine)
            rules = [{key: rule[key] for key in ("match_type", "pattern", "canonical_brand", "priority", "enabled")}
                     for rule in current_rules["items"]]
            rules.append({"match_type": "exact", "pattern": "测试商品", "canonical_brand": "测试牌", "priority": 0, "enabled": True})
            updated = service.replace_brand_rules(rules, current_rules["version"], "unit-test", engine=self.engine)
            self.assertNotEqual(updated["run_id"], run_id)
            self.assertEqual(service.list_items(engine=self.engine)["items"][0]["brand"], "测试牌")
            self.assertEqual(service.get_summary(engine=self.engine)["return_rate"], "0.3")
            self.assertEqual(service.list_items(run_id=run_id, engine=self.engine)["items"][0]["brand"], "待识别")


if __name__ == "__main__":
    unittest.main()
