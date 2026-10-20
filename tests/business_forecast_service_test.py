import json
import unittest
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from utils import business_forecast_service as service
from utils import inventory_service


class BusinessForecastBrandNormalizationTest(unittest.TestCase):
    def test_business_import_keeps_source_brand_name(self) -> None:
        self.assertEqual(
            service._normalize_business_brand("测试旅行", "测试旅行（Samsonite）企业购京东自营旗舰店"),
            "测试旅行",
        )
        self.assertEqual(
            service._normalize_business_brand("美旅（AMERICAN TOURISTER）", "美旅（AMERICAN TOURISTER）企业购京东自营旗舰店"),
            "美旅（AMERICAN TOURISTER）",
        )

    def test_regular_shop_keeps_original_brand_name(self) -> None:
        self.assertEqual(
            service._normalize_business_brand("测试旅行", "测试旅行数码京东自营旗舰店"),
            "测试旅行",
        )
        self.assertEqual(
            service._normalize_business_brand("美旅（AMERICAN TOURISTER）", "美旅京东自营旗舰店"),
            "美旅（AMERICAN TOURISTER）",
        )


    def test_inventory_brand_controls_forecast_brand(self) -> None:
        self.assertEqual(
            service._forecast_brand("测试旅行企业购", "测试旅行"),
            "测试旅行企业购",
        )
        self.assertEqual(
            service._forecast_brand("美旅企业购", "美旅（AMERICAN TOURISTER）"),
            "美旅企业购",
        )
        self.assertEqual(
            service._forecast_brand("测试旅行", "测试旅行企业购"),
            "测试旅行",
        )

    def test_regular_business_brand_still_prefers_inventory_brand(self) -> None:
        self.assertEqual(
            service._forecast_brand("测试旅行", "测试旅行"),
            "测试旅行",
        )


class InventoryBrandNormalizationTest(unittest.TestCase):
    def test_inventory_supplier_header_uses_readable_column_name(self) -> None:
        self.assertEqual(inventory_service.BASE_COLUMNS["supplier_name"], "供应商名称")
        self.assertIn("供应商", inventory_service.SUPPLIER_NAME_COLUMN_ALIASES)

    def test_enterprise_constants_use_real_chinese_text(self) -> None:
        self.assertEqual(inventory_service.ENTERPRISE_PURCHASE_SUPPLIER_NAME, "广州示例户外网络科技有限公司")
        self.assertEqual(inventory_service.SAMSONITE_ENTERPRISE_BRAND, "测试旅行企业购")
        self.assertEqual(inventory_service.AMERICAN_TOURISTER_ENTERPRISE_BRAND, "美旅企业购")

    def test_enterprise_supplier_converts_supported_brands(self) -> None:
        supplier = "广州示例户外网络科技有限公司"
        self.assertEqual(
            inventory_service._normalize_inventory_brand_by_supplier("测试旅行", supplier),
            "测试旅行企业购",
        )
        self.assertEqual(
            inventory_service._normalize_inventory_brand_by_supplier("美旅（AMERICAN TOURISTER）", supplier),
            "美旅企业购",
        )

    def test_other_supplier_keeps_original_brand(self) -> None:
        self.assertEqual(
            inventory_service._normalize_inventory_brand_by_supplier("测试旅行", "其他供应商"),
            "测试旅行",
        )

    def test_enterprise_purchase_brands_are_box_category(self) -> None:
        self.assertEqual(inventory_service._brand_category("测试旅行企业购"), "箱包")
        self.assertEqual(inventory_service._brand_category("美旅企业购"), "箱包")
        self.assertEqual(inventory_service._brand_category("美旅（AMERICAN TOURISTER）"), "箱包")


class BusinessForecastDateNormalizationTest(unittest.TestCase):
    def test_database_date_is_restored_before_forecast_calculation(self) -> None:
        forecast_date = date(2026, 7, 13)
        rows = []
        for offset in range(30):
            biz_date = forecast_date - timedelta(days=29 - offset)
            rows.append(
                service._normalize_business_daily_row(
                    {
                        "biz_date": biz_date,
                        "sku_code": "SKU-001",
                        "sold_units": 10,
                        "visitors": 100,
                        "cart_users": 20,
                        "conversion_rate": 0.1,
                    }
                )
            )

        self.assertIsInstance(rows[0]["biz_date"], date)

        forecast, cities = service._build_sku_forecast(
            forecast_date=forecast_date,
            sku_code="SKU-001",
            rows=rows,
            inventory_meta=None,
            snapshots=[],
            policies={},
            supply_stocks={},
        )

        self.assertEqual(forecast["days_covered"], 30)
        self.assertEqual(forecast["base_pred_7d"], 70.0)
        self.assertEqual(forecast["adjusted_pred_7d"], 70.0)
        self.assertEqual(cities, [])

    def test_business_signals_adjust_forecast_even_when_units_are_flat(self) -> None:
        forecast_date = date(2026, 7, 13)
        rows = []
        for offset in range(30):
            biz_date = forecast_date - timedelta(days=29 - offset)
            recent_boost = offset >= 27
            rows.append(
                service._normalize_business_daily_row(
                    {
                        "biz_date": biz_date,
                        "sku_code": "SKU-001",
                        "page_views": 300 if recent_boost else 120,
                        "visitors": 160 if recent_boost else 100,
                        "views_per_visitor": 1.9 if recent_boost else 1.2,
                        "avg_stay_seconds": 95 if recent_boost else 60,
                        "buyers": 18 if recent_boost else 10,
                        "conversion_rate": 0.14 if recent_boost else 0.1,
                        "order_count": 16 if recent_boost else 10,
                        "sold_units": 10,
                        "gmv_amount": 1300 if recent_boost else 800,
                        "avg_order_value": 90 if recent_boost else 80,
                        "cart_users": 36 if recent_boost else 20,
                        "cart_conversion_rate": 0.22 if recent_boost else 0.15,
                        "cart_units": 45 if recent_boost else 24,
                    }
                )
            )

        forecast, _ = service._build_sku_forecast(
            forecast_date=forecast_date,
            sku_code="SKU-001",
            rows=rows,
            inventory_meta=None,
            snapshots=[],
            policies={},
            supply_stocks={},
        )

        self.assertEqual(forecast["base_pred_7d"], 70.0)
        self.assertGreater(forecast["adjusted_pred_7d"], 70.0)
        self.assertLessEqual(forecast["adjustment_factor"], 1.25)
        self.assertEqual(forecast["trend_level"], service.TREND_UP)
        reason = json.loads(forecast["reason_json"])
        self.assertIn("signal_scores", reason)
        self.assertGreater(reason["signal_scores"]["traffic"], 0)
        self.assertGreater(reason["signal_scores"]["interest"], 0)
        self.assertGreater(reason["signal_scores"]["conversion"], 0)
        self.assertGreater(reason["signal_scores"]["quality"], 0)
        self.assertEqual([point["name"] for point in reason["signal_windows"]["traffic"]], ["近30", "近14", "近7", "近3", "未来"])
        self.assertEqual(reason["signal_windows"]["traffic"][0]["label"], "100.0")
        self.assertGreater(reason["signal_windows"]["traffic"][3]["value"], 100)
        self.assertIn("近30日=100", reason["signal_window_unit"])

    def test_business_signals_lower_forecast_with_guardrail_when_demand_weakens(self) -> None:
        forecast_date = date(2026, 7, 13)
        rows = []
        for offset in range(30):
            biz_date = forecast_date - timedelta(days=29 - offset)
            recent_drop = offset >= 27
            rows.append(
                service._normalize_business_daily_row(
                    {
                        "biz_date": biz_date,
                        "sku_code": "SKU-002",
                        "page_views": 80 if recent_drop else 180,
                        "visitors": 55 if recent_drop else 110,
                        "views_per_visitor": 0.9 if recent_drop else 1.5,
                        "avg_stay_seconds": 30 if recent_drop else 70,
                        "buyers": 4 if recent_drop else 12,
                        "conversion_rate": 0.04 if recent_drop else 0.11,
                        "order_count": 4 if recent_drop else 12,
                        "sold_units": 6 if recent_drop else 10,
                        "gmv_amount": 280 if recent_drop else 850,
                        "avg_order_value": 70 if recent_drop else 78,
                        "cart_users": 8 if recent_drop else 24,
                        "cart_conversion_rate": 0.08 if recent_drop else 0.16,
                        "cart_units": 9 if recent_drop else 26,
                    }
                )
            )

        forecast, _ = service._build_sku_forecast(
            forecast_date=forecast_date,
            sku_code="SKU-002",
            rows=rows,
            inventory_meta=None,
            snapshots=[],
            policies={},
            supply_stocks={},
        )

        self.assertLess(forecast["adjusted_pred_7d"], forecast["base_pred_7d"])
        self.assertGreaterEqual(forecast["adjustment_factor"], 0.8)
        self.assertEqual(forecast["trend_level"], service.TREND_DOWN)
        reason = json.loads(forecast["reason_json"])
        self.assertLess(reason["signal_scores"]["traffic"], 0)
        self.assertLess(reason["signal_scores"]["interest"], 0)
        self.assertLess(reason["signal_scores"]["conversion"], 0)

    def test_apply_trend_factor_switch_disables_national_trend_amplification(self) -> None:
        """关闭趋势因子开关后：调整因子按 1.0 计算，预测与城市补货建议不再被全国趋势放大。"""

        forecast_date = date(2026, 7, 13)
        rows = []
        for offset in range(30):
            biz_date = forecast_date - timedelta(days=29 - offset)
            boost = 10 if offset >= 20 else 2  # 近期销量上涨 → 趋势上调因子 > 1
            rows.append(
                service._normalize_business_daily_row(
                    {
                        "biz_date": biz_date,
                        "sku_code": "SKU-001",
                        "sold_units": boost,
                        "visitors": 100,
                        "cart_users": 20,
                        "conversion_rate": 0.1,
                    }
                )
            )
        snapshots = [
            {
                "warehouse_code": "成都仓",
                "stock_qty": 4,
                "inbound_qty": 0,
                "sales_1d": 1,
                "sales_7d": 7,
                "sales_30d": 30,
            }
        ]
        common = dict(
            forecast_date=forecast_date,
            sku_code="SKU-001",
            rows=rows,
            inventory_meta={"sku_code": "SKU-001", "product_name": "P", "brand": "Test Brand", "category": "箱包"},
            snapshots=snapshots,
            policies={},
            supply_stocks={},
        )
        on_forecast, on_cities = service._build_sku_forecast(**common, apply_trend_factor=True)
        off_forecast, off_cities = service._build_sku_forecast(**common, apply_trend_factor=False)

        self.assertGreater(on_forecast["adjustment_factor"], 1.0)
        self.assertEqual(off_forecast["adjustment_factor"], 1.0)
        # 开关关闭后预测不再放大；基线保持不变
        self.assertGreater(on_forecast["adjusted_pred_7d"], off_forecast["adjusted_pred_7d"])
        self.assertEqual(on_forecast["base_pred_7d"], off_forecast["base_pred_7d"])
        # 城市补货建议：关闭后不放大，建议量 ≤ 开启时
        on_total = sum(row["suggested_qty"] for row in on_cities)
        off_total = sum(row["suggested_qty"] for row in off_cities)
        self.assertGreaterEqual(on_total, off_total)
        self.assertGreater(on_total, 0)


class BusinessForecastCityDemandTest(unittest.TestCase):
    FORECAST_CITIES = {"广州", "上海", "北京", "成都", "武汉", "西安", "德州", "沈阳"}

    @staticmethod
    def _policy() -> SimpleNamespace:
        return SimpleNamespace(
            target_days=30,
            reserve_days=30,
            allow_jd_transfer=False,
            preferred_source_warehouse="",
            supplier_eta_days=0,
        )

    def test_city_forecast_uses_city_baseline_times_national_trend(self) -> None:
        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 7, 13),
            sku_code="SKU-001",
            brand="Test Brand",
            category="Test Category",
            adjustment_factor=1.2,
            snapshots=[
                {
                    "warehouse_code": "广州仓",
                    "stock_qty": 100,
                    "inbound_qty": 0,
                    "sales_1d": 10,
                    "sales_7d": 70,
                    "sales_30d": 300,
                },
                {
                    "warehouse_code": "上海RDC",
                    "stock_qty": 100,
                    "inbound_qty": 0,
                    "sales_1d": 5,
                    "sales_7d": 35,
                    "sales_30d": 150,
                },
            ],
            policy=self._policy(),
            supply_stocks={},
        )

        self.assertFalse(pending)
        self.assertEqual(
            [service.inventory_service._base_city(row["warehouse_code"]) for row in cities],
            ["广州", "上海"],
        )
        self.assertTrue(
            all(
                service.inventory_service._base_city(row["warehouse_code"]) in self.FORECAST_CITIES
                for row in cities
            )
        )
        self.assertEqual(cities[0]["daily_city_forecast"], 12.0)
        self.assertEqual(cities[0]["predicted_city_7d"], 84.0)
        self.assertEqual(cities[1]["daily_city_forecast"], 6.0)
        self.assertEqual(cities[1]["predicted_city_7d"], 42.0)

    def test_pending_link_sku_never_generates_city_forecasts(self) -> None:
        forecast_date = date(2026, 7, 13)
        rows = [
            service._normalize_business_daily_row(
                {
                    "biz_date": forecast_date - timedelta(days=offset),
                    "sku_code": "UNMATCHED-SKU",
                    "sold_units": 10,
                    "visitors": 100,
                    "cart_users": 20,
                    "conversion_rate": 0.1,
                }
            )
            for offset in range(30)
        ]

        forecast, cities = service._build_sku_forecast(
            forecast_date=forecast_date,
            sku_code="UNMATCHED-SKU",
            rows=rows,
            inventory_meta=None,
            snapshots=[
                {
                    "warehouse_code": "CITY-A",
                    "stock_qty": 100,
                    "inbound_qty": 0,
                    "sales_1d": 10,
                    "sales_7d": 70,
                    "sales_30d": 300,
                }
            ],
            policies={},
            supply_stocks={},
        )

        self.assertEqual(forecast["sku_link_status"], service.PENDING_LINK)
        self.assertEqual(forecast["affected_warehouse_count"], 0)
        self.assertEqual(forecast["suggestion_count"], 0)
        self.assertEqual(cities, [])

    def test_city_without_outbound_history_is_excluded(self) -> None:
        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 7, 13),
            sku_code="SKU-001",
            brand="Test Brand",
            category="Test Category",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "北京仓",
                    "stock_qty": 0,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 7,
                    "sales_30d": 30,
                },
                {
                    "warehouse_code": "广州仓",
                    "stock_qty": 500,
                    "inbound_qty": 200,
                    "sales_1d": 0,
                    "sales_7d": 0,
                    "sales_30d": 0,
                },
            ],
            policy=self._policy(),
            supply_stocks={},
        )

        self.assertFalse(pending)
        self.assertEqual(
            [service.inventory_service._base_city(row["warehouse_code"]) for row in cities],
            ["北京"],
        )

    def test_demo_life_purchase_uses_conservative_three_period_shortage_without_city_case_pack(self) -> None:
        """iGear货盘口径取30/14/7日缺口中的最小值，城市层不按整箱放大。"""

        policy = self._policy()
        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 9, 2),
            sku_code="IGEAR-001",
            brand="iGear工业品",
            category="工业品",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "北京仓",
                    "stock_qty": 10,
                    "inbound_qty": 0,
                    "sales_1d": 2,
                    "sales_7d": 14,
                    "sales_14d": 28,
                    "sales_30d": 60,
                }
            ],
            policy=policy,
            supply_stocks={},
            purchase_formula="excel",
            case_pack=100,
        )

        self.assertFalse(pending)
        self.assertEqual(cities[0]["suggested_qty"], 40)
        self.assertEqual(cities[0]["shortage_qty"], 40)
        self.assertIn("iGear货盘表保守补货口径", cities[0]["reason"])
        self.assertNotIn("按箱规", cities[0]["reason"])

    def test_excel_formula_uses_orderable_qty_field(self) -> None:
        """Excel 运营表算法：建议 = 近30日出库 × (目标周转/30) − 可订购（orderable_qty 字段）。"""

        policy = self._policy()
        policy.target_days = 40
        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 9, 2),
            sku_code="SKU-ORDERABLE",
            brand="测试服饰",
            category="箱包",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "北京仓",
                    "stock_qty": 8,
                    "inbound_qty": 0,
                    "available_qty": 8,
                    "orderable_qty": 46,
                    "sales_1d": 1,
                    "sales_7d": 10,
                    "sales_14d": 20,
                    "sales_30d": 32,
                }
            ],
            policy=policy,
            supply_stocks={},
            purchase_formula="excel",
        )

        self.assertFalse(pending)
        # 32 × 40/30 − 46 = -3.333 → ceil = -3（负数=已超储，保留展示）
        self.assertEqual(cities[0]["suggested_qty"], -3)
        self.assertEqual(cities[0]["orderable_qty"], 46)
        self.assertIn("减可订购46件", cities[0]["reason"])
        self.assertIn("按Excel 40周转口径", cities[0]["reason"])

    def test_excel_formula_falls_back_to_stock_plus_inbound_without_orderable(self) -> None:
        """旧快照未导入“可订购”字段时，退回 库存件数+采购未到货，保持历史口径。"""

        policy = self._policy()
        policy.target_days = 40
        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 9, 2),
            sku_code="SKU-ORDERABLE",
            brand="测试服饰",
            category="箱包",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "北京仓",
                    "stock_qty": 8,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 10,
                    "sales_14d": 20,
                    "sales_30d": 32,
                }
            ],
            policy=policy,
            supply_stocks={},
            purchase_formula="excel",
        )

        self.assertFalse(pending)
        # 无 orderable_qty → 回退 available = 库存 8；32 × 40/30 − 8 = 34.667 → ceil = 35
        self.assertEqual(cities[0]["suggested_qty"], 35)
        self.assertEqual(cities[0]["orderable_qty"], 8)

    def test_supply_stock_is_consumed_across_city_suggestions(self) -> None:
        policy = self._policy()
        policy.target_days = 10
        supply_warehouse = service.inventory_service.SUPPLY_HUB_WAREHOUSE

        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 7, 13),
            sku_code="SKU-001",
            brand="Test Brand",
            category="Test Category",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "成都仓",
                    "stock_qty": 2,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 7,
                    "sales_30d": 30,
                },
                {
                    "warehouse_code": "武汉配送中心",
                    "stock_qty": 2,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 7,
                    "sales_30d": 30,
                },
            ],
            policy=policy,
            supply_stocks={
                (supply_warehouse, "SKU-001"): {
                    "available_qty": 10,
                    "stock_qty": 10,
                }
            },
        )

        self.assertFalse(pending)
        self.assertEqual(
            [service.inventory_service._base_city(row["warehouse_code"]) for row in cities],
            ["成都", "武汉"],
        )
        self.assertEqual([row["shortage_qty"] for row in cities], [8.0, 8.0])
        self.assertEqual([row["suggested_qty"] for row in cities], [8, 2])
        self.assertEqual(sum(row["suggested_qty"] for row in cities), 10)
        self.assertEqual([row["suggested_source"] for row in cities], [supply_warehouse, supply_warehouse])
        self.assertTrue(all(row["status"] == "OBSERVE" for row in cities))

    def test_missing_supply_stock_still_offers_shortage_suggestion(self) -> None:
        """无供应仓（花东/南沙·聚水潭）库存数据时，只能靠供应仓补货的 SKU
        仍按缺货量给出补货建议，避免建议量被压成 0 而消失在补货导出明细。"""

        policy = self._policy()
        policy.target_days = 10
        supply_warehouse = service.inventory_service.SUPPLY_HUB_WAREHOUSE

        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 7, 13),
            sku_code="SKU-001",
            brand="Test Brand",
            category="Test Category",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "成都仓",
                    "stock_qty": 2,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 7,
                    "sales_30d": 30,
                },
                {
                    "warehouse_code": "武汉配送中心",
                    "stock_qty": 2,
                    "inbound_qty": 0,
                    "sales_1d": 1,
                    "sales_7d": 7,
                    "sales_30d": 30,
                },
            ],
            policy=policy,
            supply_stocks={},
        )

        self.assertFalse(pending)
        self.assertEqual([row["suggested_qty"] for row in cities], [8, 8])
        self.assertEqual([row["suggested_source"] for row in cities], [supply_warehouse, supply_warehouse])
        self.assertTrue(all(row["status"] == "OBSERVE" for row in cities))

    def test_days_supply_counts_inbound_for_business_replenishment(self) -> None:
        policy = self._policy()
        policy.target_days = 30

        cities, pending = service._build_city_rows(
            forecast_date=date(2026, 7, 13),
            sku_code="SKU-INBOUND",
            brand="Test Brand",
            category="Test Category",
            adjustment_factor=1.0,
            snapshots=[
                {
                    "warehouse_code": "上海仓",
                    "stock_qty": 15,
                    "inbound_qty": 6,
                    "sales_1d": 1,
                    "sales_7d": 5,
                    "sales_30d": 21,
                },
            ],
            policy=policy,
            supply_stocks={},
        )

        self.assertFalse(pending)
        self.assertEqual(len(cities), 1)
        self.assertAlmostEqual(
            cities[0]["days_of_supply"],
            cities[0]["available_qty"] / cities[0]["daily_city_forecast"],
            places=2,
        )
        self.assertGreater(cities[0]["days_of_supply"], cities[0]["stock_qty"] / cities[0]["daily_city_forecast"])

    def test_only_eight_forecast_cities_are_included_for_every_brand(self) -> None:
        policy = self._policy()
        policy.target_days = 10
        supply_warehouse = service.inventory_service.SUPPLY_HUB_WAREHOUSE
        snapshots = [
            {
                "warehouse_code": "西安仓",
                "stock_qty": 2,
                "inbound_qty": 0,
                "sales_1d": 1,
                "sales_7d": 7,
                "sales_30d": 30,
            },
            {
                "warehouse_code": "山东德州RDC",
                "stock_qty": 2,
                "inbound_qty": 0,
                "sales_1d": 1,
                "sales_7d": 7,
                "sales_30d": 30,
            },
            {
                "warehouse_code": "杭州仓",
                "stock_qty": 2,
                "inbound_qty": 0,
                "sales_1d": 1,
                "sales_7d": 7,
                "sales_30d": 30,
            },
        ]

        for brand in ("Test Brand", "测试五金", "iGear"):
            with self.subTest(brand=brand):
                cities, pending = service._build_city_rows(
                    forecast_date=date(2026, 7, 13),
                    sku_code="SKU-001",
                    brand=brand,
                    category="Test Category",
                    adjustment_factor=1.0,
                    snapshots=snapshots,
                    policy=policy,
                    supply_stocks={
                        (supply_warehouse, "SKU-001"): {
                            "available_qty": 100,
                            "stock_qty": 100,
                        }
                    },
                )

                normalized_cities = [
                    service.inventory_service._base_city(row["warehouse_code"])
                    for row in cities
                ]
                self.assertFalse(pending)
                self.assertEqual(normalized_cities, ["西安", "德州"])
                self.assertTrue(set(normalized_cities).issubset(self.FORECAST_CITIES))
                self.assertEqual(len(cities), 2)
                self.assertEqual(sum(row["suggested_qty"] > 0 for row in cities), 2)


class BusinessForecastVisibleCityCountTest(unittest.TestCase):
    def test_counts_only_visible_eight_warehouse_rows(self) -> None:
        rows = [
            {
                "warehouse_code": "华南广州仓",
                "risk_level": "HIGH",
                "suggested_qty": 6,
            },
            {
                "warehouse_code": "上海RDC",
                "risk_level": "LOW",
                "suggested_qty": 3,
            },
            {
                "warehouse_code": "杭州仓",
                "risk_level": "HIGH",
                "suggested_qty": 9,
            },
        ]

        affected_count, suggestion_count = service._visible_city_counts(rows)

        self.assertEqual(affected_count, 1)
        self.assertEqual(suggestion_count, 2)

    def test_empty_rows_have_no_visible_counts(self) -> None:
        self.assertEqual(service._visible_city_counts([]), (0, 0))


class BusinessForecastPurchaseExportTest(unittest.TestCase):
    def test_purchase_export_workbook_matches_operator_template_columns(self) -> None:
        content = service.build_business_purchase_export_workbook(
            [
                {
                    "purchaser_erp": "fengjingwei",
                    "sku_code": "100005166204",
                    "product_name": "测试旅行电脑包",
                    "purchase_reason": "经营趋势显示未来7天有补货需求，北京等仓建议补货。",
                    "北京": 2,
                    "上海": 0,
                    "广州": 4,
                    "成都": 0,
                    "武汉": 1,
                    "沈阳": 3,
                    "西安": 0,
                    "德州": 5,
                }
            ]
        )

        wb = load_workbook(BytesIO(content))
        ws = wb["Horizontal"]

        self.assertEqual(
            [ws.cell(row=1, column=col).value for col in range(1, 16)],
            list(service.PURCHASE_EXPORT_HEADERS),
        )
        self.assertEqual(ws["A2"].value, "fengjingwei")
        self.assertEqual(ws["B2"].value, "100005166204")
        self.assertEqual(ws["C2"].value, "测试旅行电脑包")
        self.assertFalse(any((2, col) in ws._cells for col in range(4, 8)))
        self.assertEqual([ws.cell(row=2, column=col).value for col in range(4, 7)], [None, None, None])
        self.assertIsNone(ws["G2"].value)
        self.assertEqual([ws.cell(row=2, column=col).value for col in range(8, 16)], [2, 0, 4, 0, 1, 3, 0, 5])
        self.assertEqual(ws.column_dimensions["A"].width, service.PURCHASE_EXPORT_WIDTHS["A"])
        self.assertEqual(ws.row_dimensions[1].height, 36)
        self.assertEqual(ws.row_dimensions[2].height, 16.5)

    def test_purchase_export_filename_and_purchaser_erp_follow_brand_rules(self) -> None:
        self.assertEqual(service._purchase_export_filename("测试旅行"), "测试旅行补货.xlsx")
        self.assertEqual(service._purchase_export_filename(""), "全部导出.xlsx")
        self.assertEqual(service._purchase_export_erp("测试旅行"), "fengjingwei")
        self.assertEqual(service._purchase_export_erp("Samsonite"), "fengjingwei")
        self.assertEqual(service._purchase_export_erp("测试运动"), "")

    def test_purchase_filter_options_use_full_result_set(self) -> None:
        options = service._purchase_filter_options(
            [
                {"brand": "Brand-C", "category": "箱包"},
                {"brand": "Brand-A", "category": "户外"},
                {"brand": "Brand-B", "category": "箱包"},
                {"brand": "Brand-A", "category": "户外"},
            ]
        )

        self.assertEqual(options["brands"], ["Brand-A", "Brand-B", "Brand-C"])
        self.assertEqual(options["categories"], ["户外", "箱包"])

    def test_purchase_rows_keep_linked_sku_even_when_city_quantities_are_zero(self) -> None:
        rows = service._build_purchase_rows_from_forecasts(
            [
                {
                    "forecast_date": date(2026, 7, 15),
                    "target_start_date": date(2026, 7, 16),
                    "target_end_date": date(2026, 7, 22),
                    "sku_code": "SKU-ZERO",
                    "product_name": "零补货测试商品",
                    "brand": "测试旅行",
                    "sku_link_status": service.LINKED,
                    "trend_level": service.TREND_UP,
                    "confidence_level": service.CONFIDENCE_HIGH,
                    "adjusted_pred_7d": 18,
                }
            ],
            [
                {"sku_code": "SKU-ZERO", "warehouse_code": "北京", "suggested_qty": 0},
                {"sku_code": "SKU-ZERO", "warehouse_code": "上海", "suggested_qty": 0},
            ],
            only_with_qty=False,
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["sku_code"], "SKU-ZERO")
        self.assertEqual(sum(rows[0][city] for city in service.PURCHASE_EXPORT_CITIES), 0)
        self.assertIn("暂无可执行补货数量", rows[0]["purchase_reason"])

    def test_purchase_rows_explain_eight_warehouse_replenishment(self) -> None:
        rows = service._build_purchase_rows_from_forecasts(
            [
                {
                    "forecast_date": date(2026, 7, 21),
                    "sku_code": "SKU-REASON",
                    "product_name": "原因测试商品",
                    "brand": "测试旅行",
                    "category": "",
                    "category_l1": "箱包",
                    "sku_link_status": service.LINKED,
                    "trend_level": service.TREND_UP,
                    "confidence_level": service.CONFIDENCE_HIGH,
                    "base_pred_7d": 16,
                    "adjusted_pred_7d": 20,
                }
            ],
            [
                {
                    "sku_code": "SKU-REASON",
                    "warehouse_code": "北京仓",
                    "predicted_city_7d": 7,
                    "daily_city_forecast": 1,
                    "stock_qty": 2,
                    "inbound_qty": 3,
                    "available_qty": 5,
                    "days_of_supply": 2,
                    "target_days": 30,
                    "suggested_qty": 4,
                    "route_path": "花东仓 → 北京仓",
                    "eta_days": 4,
                },
                {
                    "sku_code": "SKU-REASON",
                    "warehouse_code": "上海RDC",
                    "predicted_city_7d": 3.1,
                    "daily_city_forecast": 0.5,
                    "stock_qty": 3,
                    "inbound_qty": 0,
                    "available_qty": 3,
                    "days_of_supply": 6,
                    "target_days": 30,
                    "suggested_qty": 3,
                    "route_path": "花东仓 → 上海RDC",
                    "eta_days": 2,
                },
            ],
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["brand"], "测试旅行")
        self.assertEqual(rows[0]["category"], "箱包")
        self.assertEqual(rows[0]["category_l1"], "箱包")
        self.assertIn("八大仓里北京4件、上海3件需要补货，合计7件", rows[0]["purchase_reason"])
        self.assertIn("北京补4件：未来7天预计卖7件", rows[0]["purchase_reason"])
        self.assertIn("库存2件、在途3件，现在约能卖5天", rows[0]["purchase_reason"])
        self.assertIn("没达到30天安全线", rows[0]["purchase_reason"])
        self.assertIn("参考路径花东仓 → 北京仓", rows[0]["purchase_reason"])
        self.assertIn("八大仓里北京4件、上海3件需要补货，合计7件", rows[0]["purchase_reason_summary"])
        self.assertEqual(
            rows[0]["purchase_reason_details"],
            [
                {
                    "city": "北京",
                    "suggested_qty": 4,
                    "predicted_city_7d": 7.0,
                    "stock_qty": 2.0,
                    "inbound_qty": 3.0,
                    "huadong_available_qty": 0.0,
                    "days_of_supply": 5.0,
                    "target_days": 30,
                    "route_path": "花东仓 → 北京仓",
                    "eta_days": 4,
                },
                {
                    "city": "上海",
                    "suggested_qty": 3,
                    "predicted_city_7d": 3.1,
                    "stock_qty": 3.0,
                    "inbound_qty": 0.0,
                    "huadong_available_qty": 0.0,
                    "days_of_supply": 6.0,
                    "target_days": 30,
                    "route_path": "花东仓 → 上海RDC",
                    "eta_days": 2,
                },
            ],
        )

    def test_purchase_rows_use_inventory_brand_and_huadong_available_qty(self) -> None:
        rows = service._build_purchase_rows_from_forecasts(
            [
                {
                    "forecast_date": date(2026, 7, 31),
                    "sku_code": "SKU-HUB",
                    "product_name": "旧商品名",
                    "inventory_product_name": "库存商品名",
                    "brand": "测试旅行",
                    "inventory_brand": "测试旅行企业购",
                    "category": "旧类目",
                    "inventory_category": "箱包",
                    "sku_link_status": service.LINKED,
                    "trend_level": service.TREND_UP,
                    "confidence_level": service.CONFIDENCE_HIGH,
                    "base_pred_7d": 10,
                    "adjusted_pred_7d": 12,
                }
            ],
            [
                {
                    "sku_code": "SKU-HUB",
                    "warehouse_code": "北京仓",
                    "predicted_city_7d": 8,
                    "daily_city_forecast": 1,
                    "stock_qty": 1,
                    "inbound_qty": 0,
                    "available_qty": 1,
                    "days_of_supply": 1,
                    "target_days": 30,
                    "suggested_qty": 5,
                    "route_path": "花东仓 → 北京仓",
                    "eta_days": 4,
                },
            ],
            hub_stock_by_sku={
                "SKU-HUB": {
                    "available_qty": 26,
                    "stock_qty": 20,
                    "inbound_qty": 6,
                }
            },
        )

        self.assertEqual(rows[0]["product_name"], "库存商品名")
        self.assertEqual(rows[0]["brand"], "测试旅行企业购")
        self.assertEqual(rows[0]["category"], "箱包")
        self.assertEqual(rows[0]["huadong_available_qty"], 26.0)
        self.assertEqual(rows[0]["purchase_reason_details"][0]["huadong_available_qty"], 26.0)

    def test_purchase_rows_fallback_category_from_enterprise_brand(self) -> None:
        rows = service._build_purchase_rows_from_forecasts(
            [
                {
                    "forecast_date": date(2026, 7, 31),
                    "sku_code": "SKU-AMT",
                    "product_name": "美旅箱包",
                    "inventory_brand": "美旅企业购",
                    "inventory_category": "",
                    "category": "",
                    "sku_link_status": service.LINKED,
                    "trend_level": service.TREND_STABLE,
                    "confidence_level": service.CONFIDENCE_MID,
                    "base_pred_7d": 7,
                    "adjusted_pred_7d": 7,
                }
            ],
            [
                {
                    "sku_code": "SKU-AMT",
                    "warehouse_code": "北京仓",
                    "suggested_qty": 1,
                }
            ],
        )

        self.assertEqual(rows[0]["category"], "箱包")

    def test_purchase_rows_include_transfer_plan_fallback_sku(self) -> None:
        forecast_rows, city_rows = service._merge_transfer_purchase_fallback(
            forecast_rows=[],
            city_rows=[],
            transfer_rows=[
                {
                    "forecast_date": date(2026, 7, 31),
                    "sku_code": "SKU-ENTERPRISE",
                    "warehouse_code": "北京仓",
                    "inventory_product_name": "企业购电脑包",
                    "inventory_brand": "测试旅行企业购",
                    "inventory_category": "箱包",
                    "daily_city_forecast": 2,
                    "predicted_city_7d": 14,
                    "stock_qty": 3,
                    "inbound_qty": 1,
                    "available_qty": 4,
                    "days_of_supply": 2,
                    "target_days": 30,
                    "suggested_qty": 6,
                    "route_path": "花东仓 → 北京仓",
                    "eta_days": 4,
                }
            ],
        )

        rows = service._build_purchase_rows_from_forecasts(forecast_rows, city_rows)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["sku_code"], "SKU-ENTERPRISE")
        self.assertEqual(rows[0]["brand"], "测试旅行企业购")
        self.assertEqual(rows[0]["category"], "箱包")
        self.assertEqual(rows[0]["北京"], 6)
        self.assertEqual(rows[0]["adjusted_pred_7d"], 14)

    def test_purchase_rows_transfer_fallback_does_not_double_city_inventory(self) -> None:
        forecast_rows = [
            {
                "forecast_date": date(2026, 7, 31),
                "sku_code": "SKU-DUP",
                "product_name": "翻倍测试商品",
                "inventory_product_name": "翻倍测试商品",
                "brand": "测试旅行",
                "inventory_brand": "测试旅行",
                "inventory_category": "箱包",
                "sku_link_status": service.LINKED,
                "trend_level": service.TREND_UP,
                "confidence_level": service.CONFIDENCE_HIGH,
                "base_pred_7d": 10,
                "adjusted_pred_7d": 12,
            }
        ]
        business_city_rows = [
            {
                "forecast_date": date(2026, 7, 31),
                "sku_code": "SKU-DUP",
                "warehouse_code": "北京仓",
                "predicted_city_7d": 25.2,
                "daily_city_forecast": 3.6,
                "stock_qty": 0,
                "inbound_qty": 44,
                "available_qty": 44,
                "days_of_supply": 12.2,
                "target_days": 30,
                "suggested_qty": 0,
                "route_path": "",
                "eta_days": 0,
            }
        ]
        transfer_rows = [
            {
                "forecast_date": date(2026, 7, 31),
                "sku_code": "SKU-DUP",
                "warehouse_code": "北京仓",
                "inventory_product_name": "翻倍测试商品",
                "inventory_brand": "测试旅行",
                "inventory_category": "箱包",
                "daily_city_forecast": 3.6,
                "predicted_city_7d": 25.2,
                "stock_qty": 0,
                "inbound_qty": 44,
                "available_qty": 44,
                "days_of_supply": 12.2,
                "target_days": 30,
                "suggested_qty": 67,
                "route_path": "供应商 → 花东仓 → 北京仓",
                "eta_days": 11,
            }
        ]

        merged_forecasts, merged_cities = service._merge_transfer_purchase_fallback(
            forecast_rows=forecast_rows,
            city_rows=business_city_rows,
            transfer_rows=transfer_rows,
        )
        self.assertEqual(len(merged_cities), 1)

        rows = service._build_purchase_rows_from_forecasts(merged_forecasts, merged_cities)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["北京"], 67)
        detail = rows[0]["purchase_reason_details"][0]
        self.assertEqual(detail["inbound_qty"], 44.0)
        self.assertEqual(detail["stock_qty"], 0.0)
        self.assertEqual(detail["predicted_city_7d"], 25.2)
        self.assertEqual(detail["suggested_qty"], 67)
        self.assertEqual(detail["route_path"], "供应商 → 花东仓 → 北京仓")

    def test_purchase_rows_sku_filter_applies_after_transfer_fallback_merge(self) -> None:
        transfer_rows = [
            {
                "forecast_date": date(2026, 7, 31),
                "sku_code": "SKU-ENTERPRISE",
                "warehouse_code": "北京仓",
                "inventory_product_name": "企业购电脑包",
                "inventory_brand": "测试旅行企业购",
                "inventory_category": "箱包",
                "daily_city_forecast": 2,
                "predicted_city_7d": 14,
                "stock_qty": 3,
                "inbound_qty": 1,
                "available_qty": 4,
                "days_of_supply": 2,
                "target_days": 30,
                "suggested_qty": 6,
                "route_path": "花东仓 → 北京仓",
                "eta_days": 4,
            },
            {
                "forecast_date": date(2026, 7, 31),
                "sku_code": "SKU-OTHER",
                "warehouse_code": "上海仓",
                "inventory_product_name": "其他商品",
                "inventory_brand": "测试旅行",
                "inventory_category": "箱包",
                "daily_city_forecast": 1,
                "predicted_city_7d": 7,
                "stock_qty": 1,
                "inbound_qty": 0,
                "available_qty": 1,
                "days_of_supply": 1,
                "target_days": 30,
                "suggested_qty": 3,
                "route_path": "花东仓 → 上海仓",
                "eta_days": 2,
            },
        ]
        forecast_rows, city_rows = service._merge_transfer_purchase_fallback(
            forecast_rows=[],
            city_rows=[],
            transfer_rows=transfer_rows,
        )
        rows = service._build_purchase_rows_from_forecasts(forecast_rows, city_rows)

        matched = service._filter_purchase_rows_by_effective_fields(
            rows,
            sku_code="SKU-ENTERPRISE",
        )
        self.assertEqual([row["sku_code"] for row in matched], ["SKU-ENTERPRISE"])

        by_product = service._filter_purchase_rows_by_effective_fields(
            rows,
            product_name="企业购电脑包",
        )
        self.assertEqual([row["sku_code"] for row in by_product], ["SKU-ENTERPRISE"])

        no_match = service._filter_purchase_rows_by_effective_fields(
            rows,
            sku_code="SKU-NOT-EXISTS",
        )
        self.assertEqual(no_match, [])


if __name__ == "__main__":
    unittest.main()
