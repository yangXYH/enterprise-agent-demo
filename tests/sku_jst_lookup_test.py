import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from utils import inventory_service


class _Result:
    def __init__(self, rows):
        self._rows = rows

    def mappings(self):
        return self

    def all(self):
        return [dict(row) for row in self._rows]

    def scalar(self):
        return self._rows[0] if self._rows else None


class _Conn:
    def __init__(self, result):
        self._result = result

    def execute(self, sql, params=None):
        return self._result


class _Ctx:
    def __init__(self, value):
        self._value = value

    def __enter__(self):
        return self._value

    def __exit__(self, *args):
        return False


class _Engine:
    def __init__(self, rows):
        self._rows = rows

    def connect(self):
        return _Ctx(_Conn(_Result(self._rows)))


class SkuJstLookupTemplateTest(unittest.TestCase):
    def test_template_contains_required_headers(self) -> None:
        content = inventory_service.build_sku_jst_lookup_template()
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        headers = [ws.cell(row=1, column=col).value for col in range(1, 3)]
        self.assertEqual(headers[0], "sku")
        self.assertEqual(headers[1], "聚水潭查找编码")
        wb.close()


class SkuJstLookupImportTest(unittest.TestCase):
    def test_import_parses_mapping_rows(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "mapping.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["sku", "聚水潭查找编码"])
            ws.append([" SKU-1 ", "JST-1001"])
            ws.append(["SKU-2", "JST-1002"])
            ws.append(["SKU-3", ""])
            wb.save(path)
            wb.close()

            with patch.object(
                inventory_service,
                "upsert_sku_jst_lookup_mappings",
                return_value={"imported": 2, "skipped": 1},
            ) as upsert:
                result = inventory_service.import_sku_jst_lookup_excel(path)

            self.assertEqual(result["rows_read"], 3)
            self.assertEqual(result["imported"], 2)
            self.assertEqual(result["skipped"], 1)
            self.assertEqual(result["sku_codes"], ["SKU-1", "SKU-2"])
            upsert.assert_called_once()
            mappings = upsert.call_args.args[0]
            self.assertEqual(mappings, [
                {"sku_code": "SKU-1", "lookup_code": "JST-1001"},
                {"sku_code": "SKU-2", "lookup_code": "JST-1002"},
            ])


class SkuJstLookupSyncIntegrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.inv_sku_rows = [
            {"sku_code": "SKU-A", "barcode": "BARCODE-A"},
            {"sku_code": "SKU-B", "barcode": "BARCODE-B"},
        ]

    def test_sync_target_prefers_lookup_code(self) -> None:
        with patch.object(
            inventory_service,
            "_load_sku_jst_lookup_map",
            return_value={"SKU-A": "JST-CODE-A"},
        ):
            targets, stats = inventory_service._jushuitan_sync_target_data(
                _Engine(self.inv_sku_rows),
                codes=["SKU-A"],
            )

        self.assertEqual(len(targets), 1)
        self.assertEqual(targets[0]["query_code"], "JST-CODE-A")
        self.assertEqual(targets[0]["query_code_source"], "JST_LOOKUP")
        self.assertEqual(stats["lookup_code_links"], 1)

    def test_sync_target_falls_back_to_barcode(self) -> None:
        with patch.object(inventory_service, "_load_sku_jst_lookup_map", return_value={}):
            targets, stats = inventory_service._jushuitan_sync_target_data(
                _Engine(self.inv_sku_rows),
                codes=["SKU-A"],
            )

        self.assertEqual(targets[0]["query_code"], "BARCODE-A")
        self.assertEqual(targets[0]["query_code_source"], "BARCODE")
        self.assertEqual(stats["barcode_links"], 1)

    def test_external_code_map_includes_lookup_code(self) -> None:
        with patch.object(
            inventory_service,
            "_load_sku_jst_lookup_map",
            return_value={"SKU-A": "JST-CODE-A"},
        ):
            code_map = inventory_service._sku_code_map_for_external_codes(
                _Engine(self.inv_sku_rows),
                ["JST-CODE-A", "BARCODE-A"],
            )

        self.assertEqual(code_map["JST-CODE-A"], "SKU-A")
        self.assertEqual(code_map["BARCODE-A"], "SKU-A")


if __name__ == "__main__":
    unittest.main()
